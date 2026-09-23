"""Presentation-only simulated workbook for M5 materiality integration."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .event_materiality import (
    DECISION_ALREADY_INCORPORATED,
    DECISION_DUPLICATE,
    DECISION_NOT_MATERIAL,
    DECISION_REQUIRES_DECOMPOSITION,
    DECISION_REQUIRES_RECALCULATION,
    DECISION_RISK_MONITOR,
    DECISION_SUPPORTING,
)
from .investment_decision import ACTION_NO_ORDER
from .m5_event_run import M5EventRunReceipt, RUN_ATTENTION, RUN_HEALTHY
from .m5_event_watermark import ScanWatermark
from .m5_materiality_bridge import (
    CONFIDENCE_SILENT,
    SEVERITY_SILENT,
    MaterialityBridgeBatch,
    MaterialityDecisionPlan,
)


OVERVIEW_SHEET = "00_总览"
MATERIALITY_SHEET = "01_材料性映射"
EVENT_SHEET = "02_事件账"
INVALIDATION_SHEET = "03_依赖失效与重算"
OUTBOX_SHEET = "04_Outbox"
BOUNDARY_SHEET = "05_输入与边界"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_HEALTH_LABELS = {
    RUN_HEALTHY: "健康",
    RUN_ATTENTION: "需关注",
}
_DECISION_LABELS = {
    DECISION_NOT_MATERIAL: "非重大",
    DECISION_SUPPORTING: "支持证据",
    DECISION_ALREADY_INCORPORATED: "已纳入",
    DECISION_REQUIRES_RECALCULATION: "需重算",
    DECISION_RISK_MONITOR: "风险监控",
    DECISION_DUPLICATE: "重复/衍生",
    DECISION_REQUIRES_DECOMPOSITION: "需拆分",
}
_SEVERITY_LABELS = {
    "CRITICAL": "严重",
    "HIGH": "高",
    "MEDIUM": "中",
    "LOW": "低",
    SEVERITY_SILENT: "静默",
}
_CONFIDENCE_LABELS = {
    "HIGH": "高",
    "MEDIUM": "中",
    "LOW": "低",
    CONFIDENCE_SILENT: "静默",
}
_KIND_LABELS = {
    "financial_facts": "财务事实",
    "valuation_inputs": "估值输入",
    "distribution_history": "分红历史",
    "model_validity": "模型有效性",
    "research_thesis": "研究论点",
    "entry_consistency": "Entry一致性",
    "decision_review": "决策复核",
    "price_attractiveness": "价格吸引力",
    "current_research_status": "当前研究状态",
    "portfolio_risk": "组合风险",
    "position_guidance": "仓位边界",
    "dividend_sustainability": "股息可持续性",
    "source_health": "数据源健康",
    "valuation_result": "估值结果",
    "price_bridge": "价格桥",
}
_ALERT_LABELS = {
    "REVIEW_DUE": "待复核",
    "CRITICAL_BREAKER": "关键论点破坏",
    "MODEL_STALE": "模型过期",
    "THESIS_ALERT": "论点提醒",
    "DIVIDEND_ALERT": "股息提醒",
    "POSITION_RISK": "仓位风险",
    "SYSTEM_HEALTH": "系统健康",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_row(ws, row: int, values: list[Any], fill: str = "FFFFFF") -> None:
    for column, value in enumerate(values, 1):
        _style(ws.cell(row, column, value), fill=fill)


def _name(symbol: str, names: Mapping[str, str]) -> str:
    return names.get(symbol, symbol)


def _kind_text(kinds: tuple[str, ...]) -> str:
    if not kinds:
        return "无"
    return "、".join(_KIND_LABELS.get(item, item) for item in kinds)


def _unmapped_text(plan: MaterialityDecisionPlan) -> str:
    if plan.event is None:
        return "无"
    state = plan.event.current_state
    values = list(state.get("unmapped_domains") or [])
    values.extend(state.get("unmapped_artifacts") or [])
    return "、".join(str(item) for item in values) if values else "无"


def _overview(
    wb: Workbook,
    batch: MaterialityBridgeBatch,
    receipt: M5EventRunReceipt,
    watermark: ScanWatermark,
) -> None:
    sheet = wb.create_sheet(OVERVIEW_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [22, 26, 68])
    _title(
        sheet,
        "M5 材料性判定接入（模拟演示）",
        "人类材料性结论精确映射为事件、依赖失效和 Outbox；不产生估值、仓位或订单。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    metrics = [
        ("命名空间", batch.namespace, "公开工作簿只接受显式模拟输入。"),
        ("扫描水位", watermark.watermark_id, "覆盖范围只能前进，不允许回退。"),
        ("覆盖至", watermark.coverage_through.isoformat(), "本候选的模拟扫描覆盖终点。"),
        ("数据源健康", watermark.source_health, "失联会与无事件区分，不静默伪造覆盖。"),
        ("运行状态", _HEALTH_LABELS[receipt.health_status], "有待复核提醒时显示需关注。"),
        ("材料性判定", len(batch.plans), "包含静默和可动作判定。"),
        ("静默判定", batch.silent_count, "非重大、已纳入和重复判定不产生 M5 副作用。"),
        ("M5 事件", len(receipt.ingest_results), "本工作簿输入的事件账记录。"),
        ("依赖失效记录", len(receipt.invalidations), "每条可动作材料性事件对应一组有界失效。"),
        ("Outbox 提醒", len(receipt.alerts), "只记录状态，不执行真实通知投递。"),
        ("动作", receipt.action, "所有结论仍是 no_order，交易由人工确认。"),
    ]
    for item in metrics:
        _write_row(sheet, row, list(item))
        row += 1


def _materiality(
    wb: Workbook,
    batch: MaterialityBridgeBatch,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(MATERIALITY_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "判定ID",
        "证券",
        "人类结论",
        "是否静默",
        "事件类型",
        "严重度",
        "置信度",
        "直接失效域",
        "未注册域/工件",
    ]
    _widths(sheet, [34, 16, 18, 10, 22, 10, 10, 34, 34])
    _title(
        sheet,
        "材料性判定映射",
        "需重算进入 M5；需拆分与风险监控只失效决策复核，不把模型标记为过期。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for plan in batch.plans:
        fill = GREY if plan.silent else "FFFFFF"
        _write_row(
            sheet,
            row,
            [
                plan.decision_id,
                _name(plan.symbol, names),
                _DECISION_LABELS.get(plan.human_decision, plan.human_decision),
                "是" if plan.silent else "否",
                "无" if plan.event is None else "重要公告",
                _SEVERITY_LABELS.get(plan.severity, plan.severity),
                _CONFIDENCE_LABELS.get(plan.confidence, plan.confidence),
                _kind_text(plan.direct_kinds),
                _unmapped_text(plan),
            ],
            fill=fill,
        )
        row += 1


def _events(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(EVENT_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "来源事件ID",
        "证券",
        "事件类型",
        "严重度",
        "状态",
        "发生/可用时间",
        "抓取时间",
        "序列",
    ]
    _widths(sheet, [36, 16, 18, 10, 16, 22, 22, 8])
    _title(
        sheet,
        "M5 事件账",
        "同一材料性结论的重复输入幂等；更正必须显式引用前序事件。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for result in receipt.ingest_results:
        event = result.event
        if event is None:
            continue
        _write_row(
            sheet,
            row,
            [
                event.source_event_id,
                _name(event.symbol, names),
                "重要公告",
                _SEVERITY_LABELS.get(event.severity, event.severity),
                result.status,
                f"{event.available_at.isoformat()} / {event.effective_at.isoformat() if event.effective_at else '无'}",
                event.detected_at.isoformat(),
                event.sequence,
            ],
        )
        row += 1


def _invalidations(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(INVALIDATION_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["失效ID", "证券", "事件类型", "节点", "依赖域", "层级", "原因"]
    _widths(sheet, [34, 16, 22, 34, 22, 8, 44])
    _title(
        sheet,
        "有界依赖失效",
        "直接失效域由人类材料性结论决定；价格变化不会反向标记内在估值。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for invalidation in receipt.invalidations:
        for item in invalidation.affected_nodes:
            _write_row(
                sheet,
                row,
                [
                    invalidation.invalidation_id,
                    _name(invalidation.symbol, names),
                    invalidation.event_type,
                    item.node_id,
                    _KIND_LABELS.get(item.kind, item.kind),
                    item.depth,
                    item.reason,
                ],
            )
            row += 1


def _outbox(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(OUTBOX_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["提醒ID", "类型", "严重度", "证券", "人工复核", "状态", "尝试次数"]
    _widths(sheet, [34, 18, 10, 16, 10, 18, 9])
    _title(
        sheet,
        "Outbox 提醒账",
        "只维护投递状态；真实通知、目标和重试由后续授权后的运营层执行。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for alert in receipt.alerts:
        event = next(
            (
                result.event
                for result in receipt.ingest_results
                if result.event and result.event.event_id == alert.event_id
            ),
            None,
        )
        _write_row(
            sheet,
            row,
            [
                alert.alert_id,
                _ALERT_LABELS.get(alert.alert_type, alert.alert_type),
                _SEVERITY_LABELS.get(alert.severity, alert.severity),
                _name(event.symbol, names) if event else "系统",
                "是" if alert.requires_human_review else "否",
                alert.status,
                alert.attempts,
            ],
        )
        row += 1


def _boundaries(wb: Workbook) -> None:
    sheet = wb.create_sheet(BOUNDARY_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["边界", "状态", "说明"]
    _widths(sheet, [30, 14, 68])
    _title(
        sheet,
        "输入与运行边界",
        "本候选仅证明离线桥接；真实公告采集、生产调度与通知投递另需授权。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("材料性来源", "人类判定", "只有 human_research_lead 的结论可以进入 M5。"),
        ("静默处理", "保留", "非重大、已纳入、重复判定不创建事件或失效。"),
        ("模型有效性", "精确隔离", "需拆分与风险监控不会把模型标记为过期。"),
        ("未注册领域", "保留证据", "不静默猜测失效范围，写入 unmapped 字段。"),
        ("事件时间链", "已校验", "源发布日作为 available，人工复核日作为 detected。"),
        ("通知投递", "未启用", "只记录 outbox 状态，不发送真实通知。"),
        ("生产调度", "未授权", "不修改服务器、PTA、数据库或计划任务。"),
        ("动作", ACTION_NO_ORDER, "所有决策仍由人工完成。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def build_m5_materiality_workbook(
    batch: MaterialityBridgeBatch,
    receipt: M5EventRunReceipt,
    watermark: ScanWatermark,
    *,
    security_names: Mapping[str, str] | None = None,
) -> Workbook:
    names = dict(security_names or {})
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, batch, receipt, watermark)
    _materiality(wb, batch, names)
    _events(wb, receipt, names)
    _invalidations(wb, receipt, names)
    _outbox(wb, receipt, names)
    _boundaries(wb)
    return wb


def write_m5_materiality_workbook(
    batch: MaterialityBridgeBatch,
    receipt: M5EventRunReceipt,
    watermark: ScanWatermark,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Workbook output must remain inside the project root")
    if output.exists():
        raise ValueError("M5 materiality candidate output already exists")
    wb = build_m5_materiality_workbook(
        batch,
        receipt,
        watermark,
        security_names=security_names,
    )
    wb.save(output)
    return {
        "workbook_path": str(output),
        "workbook_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "sheet_count": len(wb.sheetnames),
        "materiality_decision_count": len(batch.plans),
        "silent_decision_count": batch.silent_count,
        "event_count": len(receipt.ingest_results),
        "invalidation_count": len(receipt.invalidations),
        "alert_count": len(receipt.alerts),
        "action": receipt.action,
    }
