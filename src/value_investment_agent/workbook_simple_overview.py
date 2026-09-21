"""Readable company progress and navigation derived from retained evidence."""
from copy import copy
from decimal import Decimal
import hashlib
import json
import math
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from .workbook_frontdoor import HOME, PRIMARY, VERSION, _rows, _band, _link, repair_internal_links

OVERVIEW = '00_公司总览'
PENDING = '00_待完成公司'
GUIDE = '00_使用说明'
DERIVED = {HOME, OVERVIEW, PENDING, GUIDE}
ROOT = Path(__file__).resolve().parents[2]
METADATA = {'初筛日期', '行情快照时间', '服务导出时间', '数据时效'}
INK, GREEN, BLUE, AMBER = '263330', '18755D', '245D83', 'FFF1D6'


def _pinned(root, pointer, filename):
    pointer_path = root / pointer
    if not pointer_path.exists():
        return None
    ref = json.loads(pointer_path.read_text(encoding='utf-8'))
    path = (root / ref['path'] / filename).resolve()
    expected = ref.get('sha256') or ref.get('summary_sha256')
    if not path.is_relative_to(root.resolve()) or hashlib.sha256(path.read_bytes()).hexdigest() != expected:
        raise ValueError('Company navigation evidence changed: ' + pointer)
    data = json.loads(path.read_text(encoding='utf-8'))
    return data, {'path': str(path), 'sha256': expected}


def current_equity_research(root):
    result = _pinned(root, 'runtime/company-research/600519-consolidated-parent-equity-residual-income-current-latest.json', 'evidence.json')
    if result is None:
        return None
    model, reference = result
    if (model.get('symbol') != '600519'
            or model.get('model_version') != 'moutai-current-parent-equity-residual-income-v1'
            or model.get('review_status') != 'current_model_and_sensitivity_ready_for_scope_review'
            or model.get('formal_fair_value') is not None
            or any(model.get(key) is not False for key in ('valuation_approved', 'simulation_eligible', 'trade_approved'))
            or not model.get('checks') or any(value is not True for value in model['checks'].values())
            or [row.get('scenario') for row in model.get('results', [])] != ['bear', 'base', 'bull']):
        raise ValueError('Current equity research display scope changed')
    for ref in [model['policy'], *model['inputs'].values()]:
        source = (root / ref['path']).resolve()
        if (not source.is_relative_to(root.resolve()) or hashlib.sha256(source.read_bytes()).hexdigest() != ref['sha256']):
            raise ValueError('Current equity research dependency changed')
    return model, reference


def company_progress(wb, root):
    """Use the union: a fixed case may be absent from this month's candidates."""
    watch = _rows(wb['01_观察名单'])
    research = _rows(wb['09_公司研究'])
    decisions = _rows(wb['21_决策验证']) if '21_决策验证' in wb else {}
    rows = {}
    for code in sorted(watch.keys() | research.keys()):
        name = (research.get(code) or watch[code])[1][1]
        row = watch.get(code)
        rows[code] = dict(
            code=code, name=name, pool='本期候选' if row else '保留研究记录',
            stage='待贯通', has_results=False, strategy_accepted=False,
            available='已有初筛和财务记录' if row else '已有公司研究记录',
            next_step='完成公司研究、适用估值和历史验证',
            tracking_state=str(decisions[code][1][3]) if code in decisions else '尚无每日决策记录',
        )
    refs = {}
    if '600519' in rows:
        closure = _pinned(root, 'runtime/strategy-validation/moutai-simulation-closure-latest.json', 'summary.json')
        case = _pinned(root, 'runtime/company-research/600519-end-to-end-case-latest.json', 'evidence.json')
        if closure and case:
            run, case_data = closure[0], case[0]
            if run.get('symbol') != '600519' or case_data.get('symbol') != '600519':
                raise ValueError('Single-company navigation identity mismatch')
            cash = run.get('cash_anchor_research_history', {})
            replayed = (run.get('execution_mechanics_verified') is True
                        and run.get('real_history', {}).get('second_run_new_journal_rows') == 0
                        and cash.get('research_model_decision_sessions', 0) > 0
                        and cash.get('second_run_new_journal_rows') == 0)
            if replayed:
                rows['600519'].update(
                    pool='固定研究案例', stage='已运行研究回放', has_results=True,
                    available='公司研究、条件估值、历史回放与虚拟账本',
                    next_step='完成估值口径与策略有效性验收',
                    closure=run, case=case_data,
                )
                refs['600519'] = [closure[1], case[1]]
                observation = _pinned(root, 'runtime/company-research/600519-current-conditional-observation-latest.json', 'evidence.json')
                pe = _pinned(root, 'runtime/strategy-validation/moutai-pe-mid-paper-contract-v2-latest.json', 'summary.json')
                if observation:
                    if observation[0].get('symbol') != '600519':
                        raise ValueError('Quote observation identity mismatch')
                    rows['600519']['observation'] = observation[0]
                    refs['600519'].append(observation[1])
                if pe:
                    if pe[0].get('symbol') != '600519':
                        raise ValueError('PE research identity mismatch')
                    rows['600519']['pe'] = pe[0]
                    refs['600519'].append(pe[1])
                contract = _pinned(root, 'runtime/company-research/600519-p1-model-contract-latest.json', 'evidence.json')
                if contract and contract[0].get('contract_version') in {'moutai-p1-model-contract-v3', 'moutai-p1-model-contract-v4'}:
                    current_contract = contract[0].get('contract_version') == 'moutai-p1-model-contract-v4'
                    model_ref = contract[0]['inputs']['current_model'] if current_contract else contract[0]['inputs']['conditional_model']
                    model_path = (root / model_ref['path']).resolve()
                    if (not model_path.is_relative_to(root.resolve())
                            or hashlib.sha256(model_path.read_bytes()).hexdigest() != model_ref['sha256']):
                        raise ValueError('Primary equity model changed after P1 contract freeze')
                    model = json.loads(model_path.read_text(encoding='utf-8'))
                    expected_model_version = 'moutai-current-parent-equity-residual-income-v1' if current_contract else 'moutai-consolidated-parent-equity-residual-income-v2'
                    if (model.get('symbol') != '600519' or model.get('valuation_approved') is not False
                            or model.get('model_version') != expected_model_version):
                        raise ValueError('Primary equity model display scope changed')
                    rows['600519']['primary_equity_model'] = model
                    rows['600519']['primary_equity_model_sha256'] = model_ref['sha256']
                    refs['600519'].extend([contract[1], {'path': str(model_path), 'sha256': model_ref['sha256']}])
                current = current_equity_research(root)
                if current:
                    rows['600519']['current_equity_research'] = current[0]
                    rows['600519']['available'] = '半年报权益/扣非TTM、当前日期衰减试算及原研究回放'
                    rows['600519']['next_step'] = '审查优势衰减假设，接通当前模型与模拟账户'
                    refs['600519'].append(current[1])
                card = _pinned(root, 'runtime/company-research/600519-research-card-evidence-latest.json', 'evidence.json')
                if card:
                    if (card[0].get('symbol') != '600519'
                            or card[0].get('research_card_version') != 'moutai-research-card-evidence-v5'
                            or card[0].get('trade_approved') is not False):
                        raise ValueError('Research-card evidence display scope changed')
                    rows['600519']['research_card'] = card[0]
                    rows['600519']['available'] = '半年报、FY2025量价、现金分配与估值假设反证已归档'
                    rows['600519']['next_step'] = '等待下一期量价/现金事实，约束优势持续期后冻结主估值模型'
                    refs['600519'].append(card[1])
                daily = _pinned(root, 'runtime/strategy-validation/moutai-daily-paper-latest.json', 'summary.json')
                if daily:
                    daily_data = daily[0]
                    account = daily_data.get('account') or {}
                    decision = daily_data.get('decision') or {}
                    if (daily_data.get('symbol') != '600519'
                            or daily_data.get('run_type') != 'current_daily_paper_simulation'
                            or decision.get('action') != 'no_order'
                            or account.get('new_journal_rows') not in {0, 1}
                            or (account.get('new_journal_rows') == 0
                                and daily_data.get('idempotent_replay') is not True)
                            or account.get('filled_orders')
                            or daily_data.get('trade_approved') is not False):
                        raise ValueError('Daily paper-account display scope changed')
                    rows['600519']['daily_paper'] = daily_data
                    rows['600519']['next_step'] = '刷新当前资本动作/估值证据后，继续每日纸面账户跟踪'
                    refs['600519'].append(daily[1])
    if '000333' in rows:
        audit = _pinned(root, 'runtime/company-research/midea-admission-audit-latest.json', 'evidence.json')
        if audit:
            if audit[0].get('symbol') != '000333':
                raise ValueError('Midea navigation identity mismatch')
            rows['000333'].update(pool='固定研究案例', stage='研究进行中',
                available='财务勾稽、股本与停复牌资料',
                next_step='匹配每股口径并完成适用估值；尚无完整回放')
            refs['000333'] = [audit[1]]
    if '601088' in rows:
        rows['601088'].update(pool='固定研究案例',
            next_step='完成周期正常化研究与估值；尚无完整回放')
    return list(rows.values()), refs


def _sheet(wb, name, widths, title, subtitle):
    if name in wb:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_view.topLeftCell = 'A1'
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, value, height in ((1, title, 40), (2, subtitle, 40)):
        _span(ws, r, 1, len(widths), value, fill=GREEN if r == 1 else 'EFF3F1', bold=r == 1)
        ws.cell(r, 1).font = Font(name='Microsoft YaHei', size=18 if r == 1 else 11,
                                  bold=r == 1, color='FFFFFF' if r == 1 else INK)
        ws.row_dimensions[r].height = height
    ws.freeze_panes = 'C4'
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def _cell(ws, row, col, value, *, fill='FFFFFF', bold=False):
    c = ws.cell(row, col, value)
    if isinstance(value, str):
        c.data_type = 's'
    c.font = Font(name='Microsoft YaHei', size=11, color=INK, bold=bold)
    c.fill = PatternFill('solid', fgColor=fill)
    c.alignment = Alignment(vertical='center', wrap_text=True)
    return c


def _span(ws, row, first, last, value, *, fill='FFFFFF', bold=False):
    if first != last:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    return _cell(ws, row, first, value, fill=fill, bold=bold)


def _links(ws, row, code, indexes, first_col=6):
    for offset, (sheet, label, column) in enumerate((
            ('09_公司研究', '看生意', 'N'), ('04_估值跟踪', '看估值', 'A'),
            ('21_决策验证', '看条件', 'T'))):
        target = indexes[sheet].get(code)
        if target:
            if column == 'N' and (len(target[1]) <= 13 or target[1][13] is None):
                column = 'A'
            if column == 'T' and (len(target[1]) <= 19 or target[1][19] is None):
                column = 'A'
            _link(ws.cell(row, first_col + offset), label, sheet, target[0], column)
        else:
            _cell(ws, row, first_col + offset, '尚未形成')


def _research_contract(ws, start_row, company):
    """Show research requirements without assigning or approving a strategy."""
    if company['code'] != '600519':
        return start_row
    _band(ws, start_row, '投资论点与财务检查 | 研究定位，不是新增买卖结论', color=BLUE, height=32)
    rows = [
        ('候选路径', '成熟优质复利 + 现金回报。以持续经营和普通股分配研究，不按烟蒂估值；标签是本轮研究定位，未获策略准入。', '候选，待验证'),
        ('回报从哪里来', '待验证论点：利润能持续且留存资本有效配置，普通股可获得现金分配和每股价值增长。估值重评单列，不重复计入回报。', '假说，不是承诺'),
        ('为什么会错价', '尚未证明市场低估。须说明市场对需求、盈利或优势期限的预期与独立研究有何差异；优质公司不等于当前价格便宜。', '错价尚未证明'),
        ('最强反证', '终端需求、渠道库存/批价与品牌溢价能否持续；留存资金能否取得合理回报。它们是待核查的反证线索，不是已发生的负面事实。', '待行业交叉核验'),
        ('利润质量 / 持续性', '已归档财报和扣非TTM只构成起点；仍需解释量价结构、分部、回款与非经常因素。扣非利润不自动等于未来正常化利润。', '不能仅凭总分放行'),
        ('现金 / 低谷生存', '继续核对受限资金、财务公司/集团往来、母公司分配能力、必要投入和压力偿债。集团现金与经营现金流不直接等于白酒可分配现金。', '现金归属须单列'),
        ('留存资本 / 治理', '比较实际分红、回购价格、投资回报和稀释；跟踪承诺兑现及关联交易。留存不自动按历史高ROE复利，管理层评价需行为证据。', '优势与资本配置待审'),
        ('市场隐含预期', '已登记逆向检验显示：归档价格在注册的利润增长、优势衰减和折现率组合之外。它至少要求其中一项更强，但变量多解，不能反推唯一市场预测。', '仅作约束，不是目标价'),
        ('下一验证事件', '沿用已验收半年报事实，有限审查优势持续期与分配假设；下一财报/经营披露及资本动作出现时更新。未设置未经论证的自动触发阈值。', '系统继续补证'),
        ('持有与退出依据', '比较更新后价值、剩余现金回报、论点是否失效及组合约束。30%与到基准/乐观价减退出仍是原候选实验，不是所有路径通则。', '不改变当前no_order'),
    ]
    for r, (label, body, state) in enumerate(rows, start_row + 1):
        _span(ws, r, 1, 2, label, fill='EDF3F8', bold=True)
        _span(ws, r, 3, 5, body)
        _span(ws, r, 6, 8, state, fill=AMBER)
        ws.row_dimensions[r].height = 78
        ws.cell(r, 3).comment = Comment(
            '目标修订：2026-09-15；value-investment-goal-framework-v2.md 第5.3、6.3、6.4、7.4及8节。'
            '\n这是研究契约/未知项说明，不新增财务事实，不改变估值、订单或准入标志。', '研究范围')
    return start_row + len(rows) + 2


def _methodology_guide(ws, start_row):
    _band(ws, start_row, '投资路径 | 可以重叠；行业口径另行匹配', color=BLUE, height=32)
    for first, label in ((1, '研究路径'), (3, '重点研究什么'), (5, '如何估值 / 实现价值'), (7, '退出复评 / 开发状态')):
        _span(ws, start_row + 1, first, first + 1, label, fill='DCEBE5', bold=True)
    ws.row_dimensions[start_row + 1].height = 34
    paths = [
        ('资产折价 / 烟蒂', '资产可变现性、完整债务索偿、现金消耗；低PB不等于可回收折价。', '回收净资产扣处置成本、税费和等待消耗；处置与持续经营情景不重复相加。', '资产兑现、折价闭合或继续减值。后续专项，未支持。'),
        ('成熟优质复利', '盈利持续性、竞争优势期限、留存资本边际回报及每股增长。', '匹配现金流/权益模型；分配和每股价值增长，不靠永久高ROE。', '论点失效或更新后回报不足。茅台首例研究中，未准入。'),
        ('成长价值', '需求/份额、量价结构、单位经济性、再投资及融资稀释。', '多阶段现金流/权益模型；增长、利润率与资本需求一致。', '单位经济性或增长被证伪、预期过高。后续复制验证。'),
        ('现金回报', '母公司可分配现金、必要投入、债务及实际分红/回购。', '可持续分配与总回报；特别股息不外推，回购与股息分开。', '分配能力下降、挤占投入或债务风险。与首例交叉研究。'),
        ('周期正常化', '供需/产能、成本曲线、库存、资源寿命和低谷生存。', '中周期现金流与低谷压力；高峰低PE不等于便宜。', '供需反转或景气充分定价。固定神华，待复制验证。'),
        ('困境反转', '经营修复、现金跑道、融资条件、索偿顺序及失败损失。', '修复/延迟/失败情景；无依据不填精确成功概率。', '里程碑失败、融资中断或过度稀释。暂不支持新增判断。'),
        ('特殊情形', '重组/分拆/处置的合同、审批、时间及失败后独立价值。', '完成及失败情景扣费用与时间；短期年化不可当可重复收益。', '条件改变、拖延或事件失败。暂不支持套利建议。'),
    ]
    for r, values in enumerate(paths, start_row + 2):
        for first, value in zip((1, 3, 5, 7), values):
            _span(ws, r, first, first + 1, value, fill=AMBER if first == 7 else 'F3F6F8', bold=first == 1)
        ws.row_dimensions[r].height = 102
    r = start_row + len(paths) + 3
    _band(ws, r, '财务五问 | 所有适用公司必答，不用单一总分替代', color=BLUE)
    questions = [
        ('利润质量', '收入、应收/合同资产、回款及税项多期对照；审计、一次性项目、资本化和关联交易。Hash与报表勾稽不等于发行人无舞弊。'),
        ('盈利持续性', '拆解量、价、产品结构、并购和周期；扣非利润不自动等于正常化利润，TTM不自动代表长期能力。'),
        ('现金可用性', '受限资金、集团/财务公司、母公司现金及必要投入；经营现金流可能受付款时点、去库存影响。'),
        ('低谷生存', '债务到期、利息、租赁、担保、必要投入和压力现金流；金融企业另看资本及流动性。'),
        ('留存资本价值', '增量回报、并购和回购价格、股份稀释与每股价值；存量高ROE不等于新资金同样高回报。'),
    ]
    for n, (label, body) in enumerate(questions, r + 1):
        _span(ws, n, 1, 2, label, fill='EDF3F8', bold=True)
        _span(ws, n, 3, 8, body)
        ws.row_dimensions[n].height = 54
    r += len(questions) + 2
    _band(ws, r, '研究交付与开发顺序 | 定义能力，不提前宣布完成', color=BLUE)
    requirements = [
        ('一家公司一张论点卡', '路径、主要回报来源、错价假说、最强反证、实现方式、持有期和下一事件。未知写未知，不批量给其他公司贴标签。'),
        ('治理与老板', '看承诺兑现、关联交易、融资摊薄、质押/减持、并购和回购结果及中小股东利益；个人形象、国有身份不替代证据。'),
        ('市场隐含预期', '先独立估值，再解释现价要求的增长/利润率/优势期限。可有多解或无解；当前尚未实现，不冒充买点。'),
        ('买卖随路径变化', '共同风险和账户门禁保留。30%安全边际、至少5年和基准/乐观价退出只属现有候选实验，不通用于烟蒂、周期或事件。'),
        ('先完成什么', '茅台适用模型与当前模拟账本、每日原表交付；随后真实历史窗口与R1，再到美的、神华和组合。研究说明完成不等于功能准入。'),
        ('以后如何扩展', '多通道筛选、6至10家研究、资产折价专项；保留失败/不入选案例及匹配基准。困境与事件暂不支持，不并行启动七套策略。'),
    ]
    for n, (label, body) in enumerate(requirements, r + 1):
        _span(ws, n, 1, 2, label, fill='EDF3F8', bold=True)
        _span(ws, n, 3, 8, body)
        ws.row_dimensions[n].height = 58
    return r + len(requirements) + 2


def _results(wb, done):
    indexes = {s: _rows(wb[s]) for s in ('09_公司研究', '04_估值跟踪', '21_决策验证')}
    ws = _sheet(wb, OVERVIEW, [12, 16, 21, 36, 40, 12, 12, 12],
        '贵州茅台 | 单公司研究卡',
        '初步事实与判断已整理；完整估值结论仍待完成。历史快照不是今日报价，研究不等于买卖指令。')
    for c, label in enumerate(['代码', '公司', '进度', '已经能看什么', '还差什么', '公司研究', '估值明细', '决策明细'], 1):
        _cell(ws, 3, c, label, fill='DCEBE5', bold=True)
    ws.row_dimensions[3].height = 28
    for r, company in enumerate(done, 4):
        for c, value in enumerate([company['code'], company['name'], company['stage'],
                                  company['available'], company['next_step']], 1):
            _cell(ws, r, c, value, fill='F3F7F5')
        _links(ws, r, company['code'], indexes)
        ws.row_dimensions[r].height = 60
    if not done:
        _span(ws, 4, 1, 8, '暂无已核验的研究回放；请从“待完成公司”查看进度。')
        ws.row_dimensions[4].height = 36
    r = max(7, len(done) + 6)
    for company_row, company in enumerate(done, 4):
        brief_start = r
        r = _company_brief(ws, r, company)
        _cell(ws, company_row, 3, '研究卡初版', fill='F3F7F5')
        _cell(ws, company_row, 4, '归档事实、盈利指标、估值分歧与条件式行动', fill='F3F7F5')
        _cell(ws, company_row, 5, '盈利持续性、可分配现金及主估值结论', fill='F3F7F5')
        for column, label, target in ((6, '研究判断', brief_start + 1), (7, '事实与公式', brief_start + 12),
                                      (8, '行动条件', brief_start + 7)):
            _link(ws.cell(company_row, column), label, OVERVIEW, target, 'A')
        archive_start = r
        _band(ws, r, company['name'] + ' | 旧实验与技术记录（保留追溯）', color=BLUE, height=32)
        case = company['case']
        observation = company.get('observation', {})
        scenarios = observation.get('scenarios', [])
        price = '尚无已核验报价记录'
        if observation.get('quote_session_verified') is True and scenarios:
            price = (observation['as_of'][:10] + ' 收盘观察价 ' +
                     f"{Decimal(scenarios[0]['observed_price_cny']):,.2f} 元；这是该日快照。")
        model_values, model_label, _ = _valuation_view(company)
        values = ' / '.join(f"{model_values[key]:.2f}" for key in ('bear', 'base', 'bull') if key in model_values)
        cash = company['closure']['cash_anchor_research_history']
        pe = company.get('pe', {})
        details = [
            ('1 公司生意', '高端白酒：品牌与渠道是主要研究线索。需求、批价与系列酒成本是需要持续跟踪的反证。',
             '09_公司研究', 'N'),
            ('2 数据时点', price + ' 财务快照截至 ' + str(case['data_snapshot']['period_end']) + '。',
             '03_财务指标', 'A'),
            ('3 条件估值', ('悲观 / 基准 / 乐观：' + values + ' 元。' + model_label + '。')
             if values else '条件值未载入；查看原研究证据。', '04_估值跟踪', 'A'),
            ('4 现在怎么做', '研究观察，未生成订单。主模型的假设与当前资本动作桥接未通过；不同日期的条件值不能直接生成当前买卖点。',
             '21_决策验证', 'T'),
            ('5 现金锚回放', f"{cash['sessions']} 个历史会话；其中 {cash['research_model_decision_sessions']} 日有研究模型判断，"
             f"{cash['proposed_orders']} 笔拟单，重复运行新增 {cash['second_run_new_journal_rows']} 条账本。",
             '21_决策验证', 'X'),
            ('6 PE独立实验', f"2015—2025 年：{pe.get('proposed_entries', '未记录')} 次研究入场、"
             f"{pe.get('proposed_exits', '未记录')} 次退出。它与现金锚、DCF属于不同实验，不能合并成一套已验证策略。",
             '21_决策验证', 'X'),
            ('7 仓位与记录', '研究虚拟账户与个人记录分开。原仓位页中的模板金额和持仓不自动视为你的实际资产。',
             '05_仓位管理', 'A'),
        ]
        current = company.get('current_equity_research')
        if current:
            facts = current['facts']
            current_values = ' / '.join(f"{Decimal(row['conditional_value_per_current_disclosed_share_cny']):.2f}" for row in current['results'])
            stress = {row['case']: Decimal(row['conditional_value_per_current_disclosed_share_cny']) for row in current['sensitivity']}
            details[3:3] = [
                ('最新事实', f"2026半年报归母权益 {Decimal(facts['parent_equity_cny']) / Decimal('1e8'):,.2f} 亿元；"
                 f"扣非TTM利润 {Decimal(current['profit_anchor_cny']) / Decimal('1e8'):,.2f} 亿元；股数 {int(facts['issued_shares']):,} 股。净利润与综合收益分列，已计分红不再扣一次。", '09_公司研究', 'N'),
                ('当前口径试算', current['valuation_at'][:10] + ' 强衰减情景：' + current_values
                 + ' 元。假设超额回报在十年内消退，已按当前日期复算；不等于中性合理价，不替代上方旧研究主模型或批准买卖。', '09_公司研究', 'N'),
                ('假设影响', f"基准盈利下，分红50%/85%对应 {stress['base_payout_0.50']:.2f}/{stress['base_payout_0.85']:.2f} 元；"
                 f"第五年后立即/十年衰减对应 {stress['base_fade_0']:.2f}/{stress['base_fade_10']:.2f} 元；"
                 f"必要回报加2个百分点为 {stress['base_higher_equity_cost']:.2f} 元。优势持续期仍须审查。", '09_公司研究', 'N'),
            ]
        daily = company.get('daily_paper')
        if daily:
            decision = daily['decision']
            account = daily['account']
            reason_text = '、'.join(decision.get('reasons') or [])
            details.insert(4, (
                '日度纸面账户',
                f"{daily['observed_at'][:10]} 已完成收盘周期：状态 {decision['state']}，"
                f"动作 {decision['action']}；新增账本 {account['new_journal_rows']} 行、"
                f"现金 {Decimal(account['ending_cash_cny']):,.2f} 元、持仓 {account['ending_shares']} 股、"
                f"成交 {len(account['filled_orders'])} 笔。阻断：{reason_text}。",
                '21_决策验证', 'T'))
        for n, (label, explanation, dest, column) in enumerate(details, r + 1):
            _span(ws, n, 1, 2, label, fill='EDF3F8', bold=True)
            _span(ws, n, 3, 7, explanation)
            target = _rows(wb[dest]).get(company['code']) if dest != '05_仓位管理' and dest in PRIMARY else None
            if target is None and dest != '05_仓位管理':
                dest, column = '09_公司研究', 'P'
                target = indexes[dest].get(company['code'])
                if target and (len(target[1]) <= 15 or target[1][15] is None):
                    column = 'A'
            _link(ws.cell(n, 8), '查看详情', dest, target[0] if target else 1, column if target else 'A')
            ws.row_dimensions[n].height = 60
            if current and label in ('最新事实', '当前口径试算', '假设影响'):
                ws.cell(n, 3).comment = Comment(
                    '半年报原文: ' + current['facts']['source_url']
                    + '\n原文件 SHA-256: ' + current['facts']['raw_file_hash']
                    + '\n当前模型: ' + current['model_version']
                    + '\n估值日期: ' + current['valuation_at']
                    + '\n政策 SHA-256: ' + current['policy']['sha256'], '研究证据')
                ws.row_dimensions[n].height = 76
        r += len(details) + 2
        r = _research_contract(ws, r, company)
        _band(ws, r, '买卖与仓位 | 条件、当前结果和下一步', color=BLUE)
        from .excel_report import read_holdings
        actual = read_holdings(wb['05_仓位管理']).get(company['code'])
        actual_text = '实际持仓未知' if actual is None else f'原表记录 {actual} 股；尚未与实际账户核对'
        rules = [
            ('买入 / 加仓', '目标规则：适用估值与必要数据通过，价格满足买入范围，现金和集中度允许；加仓还需新增证据或已登记价格条件。', '当前不生成新增拟单'),
            ('持有 / 减仓 / 退出', '先检查已证实的生意失效、治理或债务风险，再检查组合限制和估值退出。持仓未知不能直接产生卖出股数。', actual_text),
            ('建议数量与仓位', '完整适用估值、账户与执行条件尚未齐备。条件试算不能转换成正式触发价或目标仓位。', '拟变动数量 / 目标仓位：未生成'),
            ('历史模拟账户', f"现金锚研究回放期末：现金 {Decimal(cash['ending_cash_cny']):,.2f} 元，持仓 {cash['ending_shares']} 股，净值 {Decimal(cash['ending_nav_cny']):,.2f} 元。", '历史研究账户，非今日真实资产'),
            ('系统下一步', '先论证估值口径与缺口影响，冻结适用模型；接通同一规则的买卖、仓位和账本，再验收历史及每日跟踪。', '专业补证由系统继续处理'),
        ]
        for n, (label, explanation, state) in enumerate(rules, r + 1):
            _span(ws, n, 1, 2, label, fill='EDF3F8', bold=True)
            _span(ws, n, 3, 5, explanation)
            _span(ws, n, 6, 8, state, fill=AMBER)
            ws.row_dimensions[n].height = 68
        r += len(rules) + 2
        ws.row_dimensions.group(archive_start + 1, r - 1, hidden=True)
        ws.row_dimensions[archive_start].collapsed = True
        ws.sheet_properties.outlinePr.summaryBelow = False
    _link(ws.cell(r, 1), '返回首页', HOME)
    _link(ws.cell(r, 4), '查看待完成公司', PENDING)
    ws.row_dimensions[r].height = 30
    ws.print_area = f'A1:H{r}'
    return ws


def _research_quote(company):
    """Use the same retained quote for the cover and the research card."""
    daily = company.get('daily_paper')
    observation = company.get('observation', {})
    candidates = []
    if daily:
        candidates.append((daily['observed_at'], Decimal(daily['observed_close_cny'])))
    if observation.get('quote_session_verified') and observation.get('scenarios'):
        candidates.append((observation['as_of'], Decimal(observation['scenarios'][0]['observed_price_cny'])))
    if not candidates:
        return None
    at, price = max(candidates, key=lambda item: item[0])
    if not price.is_finite() or price <= 0:
        raise ValueError('Research-card quote must be positive')
    return at[:10], price


def _company_brief(ws, start, company):
    current = company.get('current_equity_research')
    research_card = company.get('research_card')
    quote = _research_quote(company)
    _band(ws, start, '研究结论 | 先理解公司，再判断价格', color=BLUE, height=32)
    dated_price = (f'{quote[0]} 归档价格 {quote[1]:,.2f} 元；未在本轮刷新。' if quote
                   else '缺少可用归档价格，不计算价格相关指标。')
    rows = [
        ('研究定位', '成熟盈利与现金回报案例。现有资料支持继续研究，不足以证明市场低估；当前不提供建议买入价。'),
        ('靠什么赚钱', '白酒产品销售是研究对象。品牌溢价、渠道与产品结构是盈利驱动假说；品牌知名度不能单独证明未来量价或长期护城河。'),
        ('财务能说明什么', '归母权益是账面权益，扣非TTM是过去四季盈利锚；两者不能直接代表可分配现金或未来利润。已披露股数用于下方参考指标，不能冒充每日流通股数。'),
        ('支持与反对', '支持：已归档扣非TTM与归母权益均为正，具备研究持续盈利的基础。反对：盈利持续性、现金归属和优势期限仍缺完整论证，不能从盈利为正跳到值得买入。'),
        ('价格如何理解', dated_price + '下方市盈率是该价格相对归档TTM盈利的描述，不是低估证明；跨日期指标不作为当时可执行信号。'),
        ('估值分歧的原因', '旧年度权益试算与强衰减试算的基点、股数和盈利路径不同，不能平均或拼成合理价区间。十年内超额回报消退是一种压力假设，不能直接当作中性预测。'),
        ('未持有时', '先比较能解释盈利与分配的主模型和报价，再讨论安全边际。目前只作研究观察，不从旧试算推导买入价，也不因价格下跌自动加仓。'),
        ('已持有时', '先核对自己的持仓、成本与现金需求；复评盈利能力、分配能力及治理变化。系统数据或执行失败不是公司恶化证据，不能直接变成卖出理由。'),
        ('三个关键缺口', '①量价/渠道反证对盈利的影响；②母公司可分配现金与必要投入；③竞争优势持续期及主估值假设。下一轮围绕这三项形成结论，不扩展公司或历史细项。'),
        ('当前交付边界', 'R0研究卡与P1研究日条件估值准入已通过；正式合理价、模拟准入和实盘指令仍未通过。'),
    ]
    source_note = None
    if current:
        facts = current['facts']
        rationale = current.get('model_policy', {}).get('rationale', {})
        if 'FY2025 parent profit fell 4.53%' in rationale.get('income', ''):
            rows[2] = ('盈利判断', '归档材料：2025年归母净利润下降4.53%，2026上半年下降1.95%。判断：公司仍有盈利，但尚不能默认恢复增长；扣非TTM不是未来利润承诺。')
        if '15,224,606,455.02' in rationale.get('payout', ''):
            rows[3] = ('分配能力与反证', '归档半年报：母公司经营现金流152.25亿元，购建长期资产支出8.29亿元，分红/利息支出350.33亿元。判断：本期经营现金流不足单独覆盖分配，需结合期初现金及子公司汇款；不能等同于无力分红。')
        source_note = ('归档财务原文：' + facts['source_url'] + '\n报告期：' + facts['period_end']
                       + '\n原文 SHA-256：' + facts['raw_file_hash']
                       + '\n模型政策依据：' + str(current.get('policy', {}))
                       + '\n归档盈利/分配摘录：' + rationale.get('income', '') + '\n' + rationale.get('payout', '')
                       + '\n这些事实不恢复已撤回的模型时点准入。')
    if research_card:
        card_facts = research_card['facts']
        conclusions = research_card['conclusions']
        actions = research_card['conditional_actions']
        rows = [
            ('研究定位', '成熟盈利与现金回报案例。第二轮事实支持继续研究，但不证明市场低估；当前不提供建议买入价。'),
            ('靠什么赚钱', conclusions['business_model']),
            ('盈利判断', conclusions['earnings']),
            ('量价与反证', conclusions['sales_realization']),
            ('现金 / 低谷韧性', conclusions['cash_distribution'] + ' ' + conclusions['resilience'] + ' ' + conclusions['liquidity_boundary'] + ' ' + conclusions['resilience_counterevidence']),
            ('量价 / 行业反证', conclusions['sales_realization'] + ' ' + conclusions['counterevidence']),
            ('资本配置 / 治理', conclusions['capital_allocation'] + ' ' + conclusions['governance'] + ' ' + conclusions['governance_counterevidence']),
            ('主模型 / 价格要求', dated_price + conclusions['model_basis']),
            ('市场对价格的要求', conclusions['market_implied_expectation']),
            ('价格如何理解', dated_price + conclusions['valuation']),
            ('未持有时', actions['not_holding']),
            ('已持有时', actions['holding']),
            ('三个关键缺口', '①' + research_card['gaps'][0] + '；②' + research_card['gaps'][1] + '；③' + research_card['gaps'][2]),
            ('下一验证事件', research_card['next_event']),
            ('当前交付边界', 'R0研究卡与P1研究日条件估值准入已通过；正式合理价、模拟准入和实盘指令仍未通过。'),
        ]
        source_note = ((source_note + '\n' if source_note else '')
                       + '第二轮研究证据：' + research_card['evidence']['fy2025_volume_constraints']['path']
                       + '\nFY2025量价：收入增速 ' + card_facts['liquor_revenue_growth_pct'] + '%；销量增速 '
                       + card_facts['liquor_sales_volume_growth_pct'] + '%；收入/吨近似变动 '
                       + f"{Decimal(card_facts['approx_revenue_per_tonne_change_pct']):.2f}" + '%。'
                       + '\n半年报现金摘录：母公司经营现金流 ' + card_facts['parent_cfo_h1_cny']
                       + '；扣资本开支后覆盖当期分红及利息约 '
                       + f"{Decimal(card_facts['parent_cfo_after_capex_coverage_of_distribution']) * Decimal('100'):.1f}" + '%。'
                       + '\n证据包 SHA-256：' + research_card['evidence']['valuation_diagnostic']['sha256']
                       + '\n该包不恢复估值或交易准入。')
    for r, (label, text) in enumerate(rows, start + 1):
        _span(ws, r, 1, 2, label, fill='EDF3F8', bold=True)
        _span(ws, r, 3, 8, text)
        ws.row_dimensions[r].height = 58
        if source_note and label in ('财务能说明什么', '支持与反对', '盈利判断', '分配能力与反证',
                                     '量价与反证', '现金与分配', '行业反证', '价格如何理解'):
            ws.cell(r, 3).comment = Comment(source_note, '研究证据')
    r = start + len(rows) + 2
    _band(ws, r, '财务与价格 | 原值、公式和含义（非合理价）', color=BLUE)
    if current:
        facts = current['facts']
        inputs = [
            ('归母权益（元）', float(Decimal(facts['parent_equity_cny'])), '#,##0.00', facts['period_end'] + '；账面权益，不等于可分配现金。'),
            ('扣非TTM（元）', float(Decimal(current['profit_anchor_cny'])), '#,##0.00', '截至' + facts['period_end'] + '；归档盈利锚，不是预测。'),
            ('已披露股数（股）', int(facts['issued_shares']), '#,##0', '与本事实包匹配的普通股数；不声明为今日股数。'),
            ('归档价格（元）', float(quote[1]) if quote else None, '#,##0.00', dated_price),
            ('每股扣非TTM（元）', f'=C{r+2}/C{r+3}', '0.00', '扣非TTM ÷ 已披露股数；非公司公告每股收益口径。'),
            ('参考市盈率（倍）', f'=IF(AND(ISNUMBER(C{r+4}),C{r+5}>0),C{r+4}/C{r+5},"")', '0.00', '归档价格 ÷ 上行每股盈利；无定价结论，不与合理PE比较。'),
            ('账面权益/股（元）', f'=C{r+1}/C{r+3}', '0.00', '账面权益 ÷ 已披露股数；非清算价值或估值下限。'),
        ]
        for n, (label, value, fmt, note) in enumerate(inputs, r + 1):
            _span(ws, n, 1, 2, label, fill='EDF3F8', bold=True)
            cell = ws.cell(n, 3, value)
            cell.number_format = fmt
            cell.font = Font(name='Microsoft YaHei', size=11, color='000000' if cell.data_type == 'f' else BLUE)
            cell.alignment = Alignment(vertical='center', horizontal='right')
            cell.comment = Comment(source_note + ('\n价格时点：' + quote[0] if quote else '')
                                   + '\n派生公式仅用于本卡描述性比较。', '研究证据')
            _span(ws, n, 4, 8, note)
            ws.row_dimensions[n].height = 42
        r += len(inputs)
    _link(ws.cell(r + 2, 1), '财报与研究原文', '09_公司研究', 1)
    _link(ws.cell(r + 2, 4), '个人持仓记录', '05_仓位管理', 1)
    ws.row_dimensions[r + 2].height = 30
    return r + 4


def _pending(wb, pending):
    indexes = {s: _rows(wb[s]) for s in ('09_公司研究', '04_估值跟踪', '21_决策验证')}
    ws = _sheet(wb, PENDING, [12, 16, 17, 21, 34, 44, 12, 12, 12],
        '待完成公司 | 按进度筛选',
        '已有财务或初筛数据，也可能尚未跑通研究、估值和历史验证。下列公司均未列入“已运行案例”。')
    for c, label in enumerate(['代码', '公司', '名单归属', '工作进度', '数据 / 跟踪状态', '系统下一步', '公司研究', '估值明细', '决策明细'], 1):
        _cell(ws, 3, c, label, fill='DCEBE5', bold=True)
    ws.row_dimensions[3].height = 30
    ordered = sorted(pending, key=lambda x: (x['pool'] != '固定研究案例', x['stage'] != '研究进行中', x['code']))
    for r, company in enumerate(ordered, 4):
        for c, key in enumerate(['code', 'name', 'pool', 'stage', 'tracking_state', 'next_step'], 1):
            _cell(ws, r, c, company[key], fill='F3F6F8' if r % 2 == 0 else 'FFFFFF')
        ws.cell(r, 4).comment = Comment(company['available'], '研究进度')
        _links(ws, r, company['code'], indexes, 7)
        ws.row_dimensions[r].height = 55
    ws.auto_filter.ref = f'A3:I{max(3, ws.max_row)}'
    _link(ws['K1'], '返回首页', HOME)
    ws.column_dimensions['K'].width = 14
    ws.print_title_rows = '1:3'
    return ws


def _home(wb, done, pending, metadata):
    ws = _sheet(wb, HOME, [16, 14, 16, 16, 16, 16, 16, 16],
        '价值投资 | 先研究一家公司',
        '短期交付：茅台研究卡。事实与初步判断已整理；完整估值结论尚未完成，不提供实盘买卖指令。')
    for first, last, label, dest in (
            (1, 3, '茅台研究卡 →', OVERVIEW),
            (4, 6, '其他公司记录（暂不扩展）', PENDING),
            (7, 8, '研究与交易边界', GUIDE)):
        _span(ws, 4, first, last, label, fill='DCEBE5' if first == 1 else 'EDF3F8', bold=True)
        _link(ws.cell(4, first), label, dest)
    ws.row_dimensions[4].height = 45
    _band(ws, 6, '当前研究观察 | 贵州茅台 600519' if done else '当前研究观察 | 尚无已核验案例', height=28)
    observation = done[0].get('observation', {}) if done else {}
    scenarios = observation.get('scenarios', [])
    model_values, model_label, _ = _valuation_view(done[0] if done else {})
    values = ' / '.join(f"{model_values[key]:.2f}" for key in ('bear', 'base', 'bull') if key in model_values)
    quote = (f"{observation['as_of'][:10]} 快照：{Decimal(scenarios[0]['observed_price_cny']):,.2f} 元"
             if scenarios and observation.get('quote_session_verified') else '尚无已核验报价快照')
    from .excel_report import read_holdings
    held = read_holdings(wb['05_仓位管理']).get('600519')
    position = '实际持仓未知；拟变动数量与目标仓位未生成' if held is None else f'原表 {held} 股，尚未与实际账户核对；拟变动与目标仓位未生成'
    summary = [
        ('市场快照', quote + ('；财报截至 ' + str(done[0]['case']['data_snapshot']['period_end']) if done else ''), OVERVIEW, 'EDF3F8'),
        ('模型 / 估值状态', (values + '；' + model_label) if values else '尚无可用条件值', '04_估值跟踪', 'EDF3F8'),
        ('买入 / 加仓提醒', '未触发：须先通过适用模型、数据时点、价格安全边际、账户和执行门禁。', '21_决策验证', 'FFF1D6'),
        ('减仓 / 卖出提醒', '未触发：须有已登记的纸面持仓，并满足估值退出或已证实基本面失效条件。', '21_决策验证', 'EDF3F8'),
        ('风险提醒', '暂无可执行交易结论；未通过的数据或模型门禁会阻断新增风险。', GUIDE, 'FDECEC'),
    ]
    if done and done[0].get('current_equity_research'):
        summary[1] = ('模型 / 估值状态', '半年报权益、扣非TTM及股数已接通；当前日期强衰减试算已完成，但尚未定为主模型。', OVERVIEW, 'FFF1D6')
    if done and done[0].get('daily_paper'):
        daily = done[0]['daily_paper']
        account = daily['account']
        decision = daily['decision']
        reasons = decision.get('reasons') or []
        blocked = decision.get('action') == 'no_order'
        summary[0] = ('市场快照',
                      f"{daily['observed_at'][:10]} 已验证收盘：{Decimal(daily['observed_close_cny']):,.2f} 元；"
                      f"纸面账户本次新增账本 {account['new_journal_rows']} 行。", OVERVIEW, 'EDF3F8')
        summary[2] = ('买入 / 加仓提醒',
                      ('未触发：系统未生成拟买单。' if blocked else '已触发纸面拟买单，须先核对次交易日开盘及限价条件。'),
                      '21_决策验证', 'FFF1D6')
        summary[3] = ('减仓 / 卖出提醒',
                      (f"未触发：纸面账户持仓 {account['ending_shares']} 股，未满足减仓或退出条件。"
                       if account['ending_shares'] == 0 else '请到决策验证页核对纸面减仓/退出条件与可卖数量。'),
                      '21_决策验证', 'EDF3F8')
        summary[4] = ('风险提醒',
                      ('高：' + '、'.join(reasons) + '。当前门禁阻断新增风险；这不是卖出指令。'
                       if blocked else '需复核次交易日成交条件、模型范围和纸面订单状态。'),
                      '21_决策验证', 'FDECEC')
    # The cover is a research brief, not a rendering of execution error codes.
    company = done[0] if done else {}
    dated = _research_quote(company)
    current = company.get('current_equity_research')
    price_text = (f'{dated[0]} 归档价格 {dated[1]:,.2f} 元；本轮未刷新，不代表今日价。' if dated
                  else '暂无已核验归档价格；不推导价格结论。')
    financial_text = '尚无本卡可引用的归档财务事实。'
    if current:
        facts = current['facts']
        eps = Decimal(current['profit_anchor_cny']) / Decimal(facts['issued_shares'])
        financial_text = (f"截至{facts['period_end']}，扣非TTM {Decimal(current['profit_anchor_cny']) / Decimal('1e8'):.2f}亿元；"
                          f'按已披露股数折算每股{eps:.2f}元。')
        if dated and eps > 0:
            financial_text += f'归档价格对应约{dated[1]/eps:.2f}倍；不代表低估。'
    summary = [
        ('事实日期', price_text, OVERVIEW, 'EDF3F8'),
        ('财务起点', financial_text, OVERVIEW, 'EDF3F8'),
        ('当前判断', '有持续盈利研究基础，但尚未证明低估。旧年度值与强衰减试算不组成合理价区间；建议买入价未形成。', OVERVIEW, 'FFF1D6'),
        ('主要反证', '量价与渠道能否支持盈利、现金能否分配、优势能持续多久。这些影响投资结论；系统执行失败不是公司高风险证据。', OVERVIEW, 'EDF3F8'),
        ('下一步行动', '先读研究卡的支持/反对理由和三个关键缺口；已持有者核对个人仓位。当前没有因本次页面更新产生的买卖指令。', OVERVIEW, 'FFF1D6'),
    ]
    for row, (label, explanation, dest, fill) in enumerate(summary, 7):
        _span(ws, row, 1, 2, label, fill=fill, bold=True)
        _span(ws, row, 3, 6, explanation, fill=fill)
        _span(ws, row, 7, 8, '查看详情', fill=fill)
        target = _rows(wb[dest]).get('600519') if dest in ('04_估值跟踪', '21_决策验证') else None
        _link(ws.cell(row, 7), '查看详情', dest, target[0] if target else 1,
              'T' if target and dest == '21_决策验证' else 'A')
        ws.row_dimensions[row].height = 68
    _band(ws, 13, '日常使用顺序 | 六个入口', height=28)
    steps = [
        ('1 读研究卡', '本期只完成茅台；其他公司记录保留，暂不扩展。', OVERVIEW, '研究卡'),
        ('2 查依据', '事实、判断和未知分开，原财报与个人笔记保留。', '09_公司研究', '公司研究'),
        ('3 查旧试算', '仅作追溯，不将不同模型区间当成当前合理价。', '04_估值跟踪', '历史试算'),
        ('4 看买卖条件', '查看买入、加减仓和退出条件，以及当前未满足的原因。', '21_决策验证', '决策验证'),
        ('5 管持仓', '核对自己的现金和实际持仓；研究模拟账户另看案例页。', '05_仓位管理', '仓位管理'),
        ('6 记交易', '人工成交后记录数量、价格、理由，并在月度复盘中回看。', '08_交易记录', '交易记录'),
    ]
    for r, (label, explanation, dest, link_text) in enumerate(steps, 14):
        _span(ws, r, 1, 2, label, fill='F0F4F2', bold=True)
        _span(ws, r, 3, 6, explanation)
        _span(ws, r, 7, 8, link_text)
        _link(ws.cell(r, 7), link_text, dest)
        ws.row_dimensions[r].height = 34
    _span(ws, 21, 1, 6, '研究卡初版 ≠ 完整研究验收 ≠ 模拟策略有效 ≠ 实盘适用。系统与公司风险分别判断。', fill=AMBER)
    _span(ws, 21, 7, 8, '使用说明')
    _link(ws['G21'], '使用说明', GUIDE)
    ws.row_dimensions[21].height = 42
    _span(ws, 23, 1, 6, '短期目标：补齐一家公司可用的研究结论；随后再接日常模拟。市场择时、多公司扩展和历史细项暂缓。', fill='EFF3F1')
    _span(ws, 23, 7, 8, '研究框架')
    _link(ws['G23'], '研究框架', GUIDE, 6)
    ws.row_dimensions[23].height = 46
    for r, (key, value) in enumerate(sorted(metadata.items()), 26):
        _cell(ws, r, 1, key)
        _cell(ws, r, 2, value)
        ws.row_dimensions[r].hidden = True
    ws.freeze_panes = 'A5'
    ws.print_area = 'A1:H23'
    return ws


def _guide(wb, refs, metadata):
    ws = _sheet(wb, GUIDE, [16, 14, 16, 16, 16, 16, 16, 16],
        '使用说明 | 投资研究框架',
        '目标修订 2026-09-20 | 先交付单公司研究卡，再接日常模拟。下列多路径及实验规则为长期参考，不是本期开发清单。')
    _link(ws['A4'], '返回首页', HOME)
    directory = [
        (OVERVIEW, '有研究回放结果的公司；首例为茅台。'),
        (PENDING, '其余公司按待完成阶段分开列示；名单不等于推荐。'),
        ('09_公司研究', '商业模式、护城河、研究假设和反证；保留个人研究。'),
        ('04_估值跟踪', '估值及其适用范围；参考值不能直接当买卖价。'),
        ('21_决策验证', '逐项核对买卖条件；数据验证与策略验证分开。'),
        ('05_仓位管理', '个人现金与持仓记录；原模板值不等于券商资产。'),
        ('08_交易记录', '实际人工成交记录；研究模拟不得写成真实交易。'),
        ('03_财务指标', '盈利、现金和资产负债数据；逐项看报告期。'),
        ('02_质量评分', '评分及其底层指标；高分不等于可以买。'),
        ('10_年报跟踪', '报告原文和更新进度。'),
        ('18_指标证据', '指标对应的原文、页码、单位及来源。'),
        ('19_金融专用指标', '银行、证券等专用口径；不套普通企业模型。'),
        ('06_月度跟踪', '历史月份快照。'),
        ('07_月度复盘', '记录决策理由、结果与改进。'),
        ('11_数据源审计', '来源、抓取与验证历史。'),
        ('20_市场覆盖', '初筛覆盖范围及排除依据。'),
        ('12_提醒', '候选公司提醒；不等于交易指令。'),
        ('01_观察名单', '月度初筛原名单，固定研究案例可在名单之外。'),
        ('12_系统设置', '原有系统配置和参数。'),
    ]
    r = _methodology_guide(ws, 6)
    _band(ws, r, '工作表目录与证据入口')
    r += 1
    for dest, description in directory:
        if dest not in wb:
            continue
        _span(ws, r, 1, 3, dest)
        if dest in PRIMARY:
            _link(ws.cell(r, 1), dest, dest)
        else:
            description = '归档资料，默认隐藏。' + description
        _span(ws, r, 4, 8, description)
        ws.row_dimensions[r].height = 32
        r += 1
    _span(ws, r + 1, 1, 8, '日常6页，参考资料4页；其余原页保留。归档页在工作表标签处取消隐藏后阅读，不设置失效跳转。', fill=AMBER)
    ws.row_dimensions[r + 1].height = 36
    _strategy_rules(ws, r + 3)
    last = ws.max_row + 2
    _band(ws, last, '方法参考 | 不是A股策略收益证明')
    sources = [
        ('增长与再投资', 'https://pages.stern.nyu.edu/adamodar/New_Home_Page/valquestions/growth.htm'),
        ('竞争优势期限', 'https://pages.stern.nyu.edu/~adamodar/New_Home_Page/valquestions/highgrowthperiod.htm'),
        ('资产折价与时间风险', 'https://www.berkshirehathaway.com/letters/1989.html'),
    ]
    for label, url in sources:
        last += 1
        _span(ws, last, 1, 2, label, fill='EDF3F8', bold=True)
        _span(ws, last, 3, 8, url)
        ws.cell(last, 3).hyperlink = url
        ws.cell(last, 3).comment = Comment('2026-09-15复核原始方法资料。七类路径是项目组织方案，非来源给出的统一分类。', '方法来源')
        ws.row_dimensions[last].height = 48
    last += 2
    _band(ws, last, '数据时点与运行证据')
    for key, value in sorted(metadata.items()):
        last += 1
        _span(ws, last, 1, 3, key)
        _span(ws, last, 4, 8, str(value))
        ws.row_dimensions[last].height = 30
    for code, sources in sorted(refs.items()):
        for source in sources:
            last += 1
            _span(ws, last, 1, 2, code + ' 运行证据')
            _span(ws, last, 3, 8, source['path'] + '\nSHA-256: ' + source['sha256'])
            ws.row_dimensions[last].height = 70
    ws.print_area = f'A1:H{last}'


def _valuation_view(company):
    model = company.get('primary_equity_model')
    if model:
        value_key = ('conditional_value_per_current_disclosed_share_cny'
                     if 'conditional_value_per_current_disclosed_share_cny' in model['results'][0]
                     else 'conditional_value_per_2025_issued_share_cny')
        values = {row['scenario']: Decimal(row[value_key])
                  for row in model['results']}
        label = ('归母剩余收益；研究日已披露股数及后续假设，未批准'
                 if value_key == 'conditional_value_per_current_disclosed_share_cny'
                 else '归母剩余收益v2；2025年末股数及后续假设，未批准')
        published = model.get('facts', {}).get('report_published_date')
        if published:
            label += '；年报披露' + published
        return values, label, company['primary_equity_model_sha256']
    observation = company.get('observation', {})
    return ({row['scenario']: Decimal(row['conditional_value_per_share_cny'])
             for row in observation.get('scenarios', [])},
            '旧条件现金流试算，未批准', observation.get('conditional_valuation_sha256'))


def _valuation_cases(wb, done):
    """Append the retained case without putting conditional values in fair value."""
    ws = wb['04_估值跟踪']
    added = []
    for company in done:
        observation = company.get('observation', {})
        scenarios = {s['scenario']: s for s in observation.get('scenarios', [])}
        model_values, model_label, model_hash = _valuation_view(company)
        if not all(s in model_values for s in ('bear', 'base', 'bull')):
            continue
        existing = _rows(ws).get(company['code'])
        row = existing[0] if existing else max(3, max((i for i, _ in _rows(ws).values()), default=3)) + 1
        values = {18: '条件悲观值(元)', 19: '条件基准值(元)', 20: '条件乐观值(元)'}
        for col, label in values.items():
            if ws.cell(3, col).value is None:
                _cell(ws, 3, col, label, bold=True)
                added.append(ws.cell(3, col).coordinate)
            elif ws.cell(3, col).value != label:
                raise ValueError('Conditional valuation column conflicts with existing data')
        if existing is None:
            cells = {1: company['code'], 2: company['name'],
                3: float(scenarios['base']['observed_price_cny']) if scenarios and observation.get('quote_session_verified') else None,
                13: '研究观察；不生成订单', 14: '条件现金流试算，未批准',
                15: observation.get('as_of'), 16: company['case']['data_snapshot']['period_end'],
                17: '右侧条件值不是正式合理价。报价为指定日期快照；模型口径、权益桥接与缺口影响尚待验收。'}
            for col, value in cells.items():
                if value is not None:
                    _cell(ws, row, col, value)
                    added.append(ws.cell(row, col).coordinate)
        managed_case = (ws.cell(row, 13).value == '研究观察；不生成订单'
                        and str(ws.cell(row, 14).value).startswith(('条件现金流试算', '归母剩余收益')))
        if managed_case:
            quote_values = {
                3: float(scenarios['base']['observed_price_cny']) if scenarios and observation.get('quote_session_verified') else None,
                15: observation.get('as_of'),
            }
            for col, value in quote_values.items():
                cell = ws.cell(row, col)
                if cell.value != value:
                    cell.value = value
                    added.append(cell.coordinate)
        for col, scenario in enumerate(('bear', 'base', 'bull'), 18):
            value = float(model_values[scenario])
            cell = ws.cell(row, col)
            legacy_value = (float(scenarios[scenario]['conditional_value_per_share_cny'])
                            if scenario in scenarios else None)
            same_value = (isinstance(cell.value, (int, float))
                          and math.isclose(cell.value, value, rel_tol=1e-15, abs_tol=0))
            can_refresh = (company.get('primary_equity_model') is not None
                           and cell.comment and cell.comment.author == '研究证据'
                           and isinstance(cell.value, (int, float)) and legacy_value is not None
                           and math.isclose(cell.value, legacy_value, rel_tol=1e-15, abs_tol=0))
            if cell.value is None or (not same_value and can_refresh):
                _cell(ws, row, col, value, fill=AMBER)
                cell.number_format = '#,##0.00'
                added.append(cell.coordinate)
            elif not same_value:
                raise ValueError('Existing conditional value changed; refresh its evidence before publishing')
            comment = model_label + '\n报价日期（不代表估值日期）：' + str(observation.get('as_of')) + '\n输入 SHA-256: ' + str(model_hash)
            if (cell.comment is None or cell.comment.author == '研究证据') and (cell.comment is None or cell.comment.text != comment):
                cell.comment = Comment(comment, '研究证据')
                if cell.coordinate not in added:
                    added.append(cell.coordinate)
        if company.get('primary_equity_model'):
            for col, value in {14: model_label, 17: '右侧为主候选模型研究值。分红、回购及最新财务尚未桥接为当前价值，不计算正式安全边际。'}.items():
                cell = ws.cell(row, col)
                if cell.value != value:
                    _cell(ws, row, col, value)
                    added.append(cell.coordinate)
    return {'04_估值跟踪': added} if added else {}


def _style_detail(ws):
    """Use one visual hierarchy while keeping values, units and warning fills."""
    header = 9 if ws.title == '05_仓位管理' else 3
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.freeze_panes = f'C{header + 1}'
    ws.sheet_view.topLeftCell = 'A1'
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 44
    ws.row_dimensions[header].height = 38
    for row in ws:
        for cell in row:
            if cell.value is None:
                continue
            font = copy(cell.font)
            font.name, font.sz = 'Microsoft YaHei', 11
            font.color = INK
            if cell.row == 1:
                font.color, font.sz, font.bold = 'FFFFFF', 18, True
                cell.fill = PatternFill('solid', fgColor=GREEN)
            elif cell.row == header:
                font.bold = True
                cell.fill = PatternFill('solid', fgColor='DCEBE5')
            elif cell.row == 2:
                cell.fill = PatternFill('solid', fgColor='EFF3F1')
            elif cell.hyperlink:
                font.color, font.underline = GREEN, 'single'
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.vertical, alignment.wrap_text = 'center', True
            cell.alignment = alignment
        if row[0].row > header and any(c.value is not None for c in row):
            ws.row_dimensions[row[0].row].height = (100 if ws.title == '09_公司研究' else
                32 if ws.title in ('05_仓位管理', '08_交易记录') else 56)
    _link(ws['A1'], ws['A1'].value, HOME)
    ws['A1'].font = Font(name='Microsoft YaHei', size=18, bold=True, color='FFFFFF', underline='single')
    widths = {'A': 12, 'B': 16}
    groups = []
    if ws.title == '21_决策验证':
        widths.update({'D': 20, 'O': 14, 'P': 14, 'Q': 14, 'S': 16, 'T': 48, 'U': 32, 'X': 45})
        groups = [('C', 'C'), ('E', 'N'), ('R', 'R'), ('V', 'W'), ('Y', 'AC')]
    elif ws.title == '04_估值跟踪':
        widths.update({'C': 14, 'K': 14, 'M': 24, 'N': 28, 'O': 27, 'Q': 42, 'R': 16, 'S': 16, 'T': 16})
        groups = [('D', 'J'), ('L', 'L'), ('P', 'P')]
    elif ws.title == '09_公司研究':
        widths.update({'C': 30, 'D': 30, 'G': 30, 'I': 32, 'K': 32, 'M': 26, 'N': 56, 'O': 48, 'P': 60, 'Q': 30})
        groups = [('E', 'F'), ('H', 'H'), ('J', 'J'), ('L', 'L')]
    elif ws.title == '18_指标证据':
        widths.update({'C': 30, 'D': 30, 'E': 18, 'F': 18, 'G': 14, 'H': 22, 'J': 44, 'P': 60})
        groups = [('I', 'I'), ('K', 'O')]
    elif ws.title == '05_仓位管理':
        widths.update({'C': 18, 'D': 18, 'E': 18, 'F': 18, 'G': 20, 'H': 18})
    elif ws.title == '08_交易记录':
        widths.update({'A': 22, 'B': 18, 'C': 14, 'D': 16, 'E': 12, 'F': 14, 'G': 14, 'H': 18, 'L': 42})
    for start, end in groups:
        ws.column_dimensions.group(start, end, hidden=True)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
        ws.column_dimensions[col].hidden = False
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _has_verified_research_card(ws):
    """Keep a separately WPS-verified research card intact during data refreshes."""
    if ws['A1'].value != '贵州茅台 | 单公司研究卡':
        return False
    labels = {str(ws.cell(row, 1).value) for row in range(1, ws.max_row + 1)}
    return {
        '主模型 / 价格要求',
        '资本配置 / 治理',
        '现金 / 低谷韧性',
        '当前交付边界',
    } <= labels


def apply_simple_overview(wb, *, root=None):
    root = Path(root) if root is not None else ROOT
    metadata = {row[0]: row[1] for row in wb[HOME].iter_rows(values_only=True)
                if row[0] in METADATA and len(row) > 1 and row[1] is not None}
    companies, refs = company_progress(wb, root)
    done = [row for row in companies if row['has_results']]
    pending = [row for row in companies if not row['has_results']]
    # Short-term research delivery must not refresh legacy mixed-date valuation cells.
    added = {}
    preserved_research_card = OVERVIEW in wb and _has_verified_research_card(wb[OVERVIEW])
    if not preserved_research_card:
        _results(wb, done)
    _pending(wb, pending)
    _home(wb, done, pending, metadata)
    _guide(wb, refs, metadata)
    visible = [name for name in PRIMARY if name in wb]
    for ws in wb:
        ws.sheet_state = 'visible' if ws.title in visible else 'hidden'
        ws.sheet_view.tabSelected = False
        if ws.title in visible and ws.title not in DERIVED:
            _style_detail(ws)
        ws.sheet_properties.tabColor = (GREEN if ws.title in [HOME, OVERVIEW] else
            BLUE if ws.title == PENDING else 'B2BCC2')
    for i, name in enumerate(visible):
        wb.move_sheet(wb[name], offset=i - wb.sheetnames.index(name))
    wb.active = 0
    repaired = repair_internal_links(wb)
    return {'version': VERSION, 'research_framework_version': 'single-company-brief-20260920',
            'visible_sheets': visible, 'company_count': len(companies),
            'research_result_count': len(done), 'pending_count': len(pending),
            'strategy_accepted_count': sum(c['strategy_accepted'] for c in companies),
            'groups': {c['code']: c['stage'] for c in companies},
            'source_references': refs, 'formal_trade_instructions': False,
            'added_case_cells': added, 'preserved_verified_research_card': preserved_research_card,
            'repaired_internal_links': repaired}


def _strategy_rules(home, start_row):
    """Keep all existing strategy explanations off the daily navigation page."""
    rules = [
        ('01 股票池与频率', '全市场身份核验 → 月度候选 → 6至10家深度研究 → 模拟准入 → 持仓。每月首个周末初筛；每日收盘跟踪已有公司，财报或重大事件触发重评。', '目标流程；候选不是推荐'),
        ('02 初筛通道', '旧版基线：PE≤25、PB≤3、市值≥50亿元。v2并集：行业相对盈利/资本回报/现金/杠杆质量通道，加分红与现金覆盖通道；不以低PB作为所有企业必备条件。', '旧版与v2分开；新通道待验收'),
        ('03 财务质量', '成熟经营企业候选实验：盈利/资本回报30%、现金25%、融资韧性25%、成长/每股配置20%，70/100待验证。不作所有路径硬筛；烟蒂、周期、成长和金融另定适用条件。', '旧版35/25/20/20保留；缺项不填中性分'),
        ('04 生意好坏', '检查量价、竞争壁垒、客户渠道、成本传导、再投资与治理；每项假设绑定支持证据、独立反证、观察指标与失效条件。商业质量14/20为实验门槛，不能由AI默认打分。', '首家公司研究中；非全池完成'),
        ('05 数据可信度', '财务：交易所、巨潮、公司IR原文；行情：AkShare及实际东方财富/腾讯/新浪上游。逐点保留URL、Hash、页码、报告期、可用时间与解析版本；同一原件的两个封装不算独立双源。', '逐字段门禁；旧报价不能当今日价'),
        ('06 企业如何估值', '行业与投资论点共同选模型：现金流/权益、正常化盈利或资产回收情景。茅台保留归母权益候选，金融业专用口径。核对现金、索偿及股份，互斥价值不相加。', '模型须按公司验收'),
        ('07 三情景与回报', '论证悲观/基准/乐观及增长、折现、再投资和终值敏感性。至少5年是现有长期持有实验，不套短期事件或资产处置；比较匹配期限的税费后回报，区间不是置信带。', '现有试算不等于批准价值'),
        ('08 股价贵不贵', '安全边际=(V中-P)/V中；上涨空间=(V中-P)/P。例如V中100、P70，分别为30%与42.86%。必须采用有效报价和已批准的正值估值，否则仅显示缺口。', '事实通过与模型通过分别判断'),
        ('09 首次买入', '现有候选实验上限=min(0.70×V中, V低)，同时过企业、数据、模型、回报、账户及成交门禁。V中100/V低65则上限65，仅为算式示例；30%不作所有路径通则。', '保留原参数；当前不批准买入'),
        ('10 分批与加仓', '目标金额按50%/30%/20%分批，不同交易日期执行。每次重新通过全部门禁，并有新财务/估值证据，或价格比上次实际成交低至少5%；下跌本身不足以批准加仓。', '实验初值；尚无可执行加仓股数'),
        ('11 组合上限', '虚拟100万元：单股≤8%、行业≤25%、关联风险组≤25%、股票≤70%、现金≥30%；不融资、不做空。真实账户与虚拟账户分开，不把模板金额视为你的本金。', '实验约束；模拟组合待验收'),
        ('12 拟买数量', '取目标缺口、含税费可用现金、组合剩余容量、流动性限额的最小值，再按合法手数取整。参考过去20个有效交易日成交额中位数×0.5%，测试0.1%/1%敏感性。', '不能用未来当日成交额'),
        ('13 持有与现金', '假设仍成立、回报及组合约束满足时持有，避免每日机械换仓。分红到账记现金，主实验不自动再投资；持仓未知不等于空仓，数据缺失不自动清仓。', '完整持仓状态机待验收'),
        ('14 风险优先级', '已证实的核心假设失效、不可接受的债务或治理风险 → 优先计划退出；其次处理组合上限与现金需要；最后处理估值减仓。未证实重大信息暂停新增风险并调查，不按传闻清仓。', '退出仍受可成交条件限制'),
        ('15 估值减仓', 'v2实验：首次P≥V中且P<V高，卖出当时可卖数量的1/3；必须有明确持仓。事件ID去重，不能每天重复卖1/3；阈值复位及版本迁移规则需冻结。', '区别旧版P>V的减仓研究提示'),
        ('16 清仓与再进入', 'v2实验：P≥V高退出剩余可卖股份；或已证实投资假设失败时按风险优先级退出。重新进入需失效原因已解决且重新通过全部买入门禁。', '完整退出与再入场未验收'),
        ('17 回撤处理', '不默认固定百分比止损。10%/20%回撤只触发风险调查，不直接产生卖单；用基本面、组合风险和可成交性作决定。', '实验告警；不保证控制最大亏损'),
        ('18 执行与失效', '收盘生成研究决策，下一可成交时点重核价格；跳空超过买入上限则不买。处理T+1、停牌、涨跌停、流动性及税费；拟单、模拟成交、真实交易独立记录。', '人工交易；不自动下单'),
        ('19 历史验证路径', '先按输入完整性登记首例连续窗口；随后保留2015—2025及三股扩展协议。逐时点资料、同口径全收益基准；新增路径含失败/不入选对照，不按收益挑样本。', '不以当前估值或当前股票池回填历史'),
        ('20 回测怎么看', '报告税费后收益、最大回撤、仓位、分红、交易账本及宽基/行业全收益对比。已有茅台持有诊断不等于价值策略；旧规则零交易照实保留，不能当作有效策略证明。', 'M3进行中；完整策略与基准未验收'),
        ('21 模拟到实盘', '当前模型、执行和账本通过可先做日常模拟；补真实历史后联合验收R1。加密备份/隔离恢复、至少30个真实会话及数据/模型/经济/运营/账户条件用于R2，不能拿研究零单摘要凑会话。', '日常模拟、R1、R2分开验收'),
        ('22 目前下一步', '先完成茅台适用模型→当前决策/执行/持久化账户→每日原表交付，再补历史窗口与R1；随后美的、神华及组合。框架展示不等于后端已实现，不并行开发全部路径。', '首例研究中；无实盘准入'),
    ]

    _band(home, start_row, '现有候选实验细则 | 非所有路径通则；不改变已登记参数', height=34)
    for r, (label, rule, state) in enumerate(rules, start_row + 1):
        for first, last, value in ((1, 2, label), (3, 6, rule), (7, 8, state)):
            _span(home, r, first, last, value, fill=AMBER if first == 7 else 'F3F6F8', bold=first == 1)
        home.row_dimensions[r].height = 106
