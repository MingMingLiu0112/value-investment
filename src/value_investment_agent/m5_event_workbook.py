"""Presentation-only simulated workbook for M5 event infrastructure."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .investment_decision import ACTION_NO_ORDER
from .m5_event_core import (
    EVENT_IDENTITY_CURRENT,
    EVENT_IDENTITY_LEGACY,
    EVENT_IDENTITY_SOURCE_ID_V1,
    EVENT_STATUS_SUPERSEDED,
    INGEST_CORRECTION_ACCEPTED,
    INGEST_DUPLICATE,
    INGEST_SUPERSEDES_ACCEPTED,
)
from .m5_event_run import M5EventRunReceipt, RUN_ATTENTION, RUN_HEALTHY
from .m5_event_run_state import M5EventRunState
from .m5_event_watermark import ScanWatermark


OVERVIEW_SHEET = "00_总览"
EVENT_SHEET = "01_事件账"
WATERMARK_SHEET = "02_水位与检查点"
INVALIDATION_SHEET = "03_依赖失效与重算"
OUTBOX_SHEET = "04_Outbox"
BOUNDARY_SHEET = "05_输入与边界"
OUTBOX_TRANSITION_SHEET = "05_Outbox迁移"
BOUNDARY_SHEET_WITH_TRANSITIONS = "06_输入与边界"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_HEALTH_LABELS = {
    RUN_HEALTHY: "健康",
    RUN_ATTENTION: "需关注",
}
_STATUS_LABELS = {
    "ACCEPTED": "已接受",
    "DUPLICATE": "重复忽略",
    "CORRECTION_ACCEPTED": "更正已接受",
    "SUPERSEDES_ACCEPTED": "替代已接受",
    "FUTURE_REJECTED": "未来事件拒绝",
    "CONFLICT_REJECTED": "冲突拒绝",
    "OBSERVED_TIME_REGRESSION_REJECTED": "观察时间回退拒绝",
    "SUPERSEDED": "已被替代",
}
_SEVERITY_LABELS = {
    "CRITICAL": "严重",
    "HIGH": "高",
    "MEDIUM": "中",
    "LOW": "低",
}
_CONFIDENCE_LABELS = {
    "HIGH": "高",
    "MEDIUM": "中",
    "LOW": "低",
}
_EVENT_LABELS = {
    "NEW_FINANCIAL_REPORT": "新财报",
    "MATERIAL_ANNOUNCEMENT": "重要公告",
    "DIVIDEND_CHANGE": "分红变化",
    "BUYBACK": "回购",
    "CAPITAL_ALLOCATION_CHANGE": "资本配置变化",
    "VALUATION_ZONE_CHANGED": "估值区间变化",
    "PRICE_ATTRACTIVENESS_CHANGED": "价格吸引力变化",
    "THESIS_WEAKENED": "论点弱化",
    "THESIS_BREAKER_TRIGGERED": "论点破坏触发",
    "MODEL_STALE": "模型过期",
    "POSITION_RISK_CHANGED": "仓位风险变化",
    "SOURCE_SCAN_FAILED": "数据源扫描失败",
    "SOURCE_SCAN_RECOVERED": "数据源扫描恢复",
}
_ALERT_LABELS = {
    "REVIEW_DUE": "待复核",
    "CRITICAL_BREAKER": "关键论点破坏",
    "MODEL_STALE": "模型过期",
    "THESIS_ALERT": "论点提醒",
    "DIVIDEND_ALERT": "股息提醒",
    "POSITION_RISK": "仓位风险",
    "SYSTEM_HEALTH": "系统健康",
}
_IDENTITY_LABELS = {
    EVENT_IDENTITY_LEGACY: "历史兼容（无来源ID）",
    EVENT_IDENTITY_SOURCE_ID_V1: "source-id-v1",
    EVENT_IDENTITY_CURRENT: "source-id-v2",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_row(ws, row: int, values: list[Any], fill: str = "FFFFFF") -> None:
    for column, value in enumerate(values, 1):
        _style(ws.cell(row, column, value), fill=fill)


def _name(symbol: str, names: Mapping[str, str]) -> str:
    return names.get(symbol, symbol)


def _overview(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
    *,
    current_state: M5EventRunState | None = None,
    include_outbox_transitions: bool = False,
) -> None:
    sheet = wb.create_sheet(OVERVIEW_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [22, 26, 66])
    _title(
        sheet,
        "M5 事件监控基础设施（模拟演示）",
        "显式模拟事件账、水位、有界重算和 Outbox；不连接生产调度或通知。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    correction_count = sum(
        item.status == INGEST_CORRECTION_ACCEPTED for item in receipt.ingest_results
    )
    duplicate_count = sum(
        item.status == INGEST_DUPLICATE for item in receipt.ingest_results
    )
    late_count = sum(item.is_late for item in receipt.ingest_results)
    alerts = (
        current_state.outbox.alerts()
        if current_state is not None
        else receipt.alerts
    )
    active_event_count = (
        len(current_state.event_ledger.active_events())
        if current_state is not None
        else len(receipt.active_events)
    )
    metrics = [
        ("命名空间", receipt.namespace, "公开工作簿只接受显式模拟输入。"),
        ("运行状态", _HEALTH_LABELS[receipt.health_status], "有需要人工复核的提醒时为需关注。"),
        ("事件输入数", len(receipt.ingest_results), "包含重复和更正输入。"),
        ("有效接受数", receipt.accepted_event_count, "重复输入不会增加新版本。"),
        ("当前有效事件数", active_event_count, "被更正或替代的旧事件不计入。"),
        ("更正数", correction_count, "更正显式绑定旧事件 ID。"),
        ("重复忽略数", duplicate_count, "相同来源与内容只保留一次。"),
        ("晚到事件数", late_count, "保留历史可用时间，不重写成当前时间。"),
        ("依赖失效记录数", len(receipt.invalidations), "只重算受影响的依赖节点。"),
        ("Outbox 提醒数", len(alerts), "不执行真实投递，仅记录状态。"),
        ("关键提醒数", sum(item.severity == "CRITICAL" for item in alerts), "关键提醒不参与普通去重。"),
        ("静默合法", "是" if receipt.silent_ok else "否", "正常日 0 提醒合法。"),
    ]
    if include_outbox_transitions:
        if current_state is None:
            raise ValueError("Outbox transitions require current state")
        metrics.extend(
            [
                (
                    "Outbox修订",
                    current_state.outbox_revision,
                    "提醒投递状态独立于事件批次 revision。",
                ),
                (
                    "迁移记录数",
                    len(current_state.outbox_transitions),
                    "只记录调用方确认的投递状态变化。",
                ),
            ]
        )
    metrics.append(("动作", ACTION_NO_ORDER, "所有复核仍由人工决定。"))
    for item in metrics:
        _write_row(sheet, row, list(item))
        row += 1


def _events(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
    *,
    include_identity_metadata: bool = False,
) -> None:
    sheet = wb.create_sheet(EVENT_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    base_columns = [
        "序号",
        "状态",
        "事件ID",
        "证券",
        "事件类型",
        "生效时间",
        "可用时间",
        "检测时间",
        "严重度",
        "置信度",
        "晚到",
        "人工复核",
        "说明",
    ]
    columns = (
        base_columns[:3] + ["身份版本", "来源ID"] + base_columns[3:]
        if include_identity_metadata
        else base_columns
    )
    widths = (
        [7, 18, 34, 22, 24, 16, 20, 18, 18, 18, 10, 9, 8, 10, 42]
        if include_identity_metadata
        else [7, 18, 34, 16, 18, 18, 18, 18, 10, 9, 8, 10, 42]
    )
    _widths(sheet, widths)
    _title(
        sheet,
        "事件账",
        "detected / effective / available 分开；重复、更正和晚到保留可审计状态。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    lifecycle_events = (
        {
            event.event_id: event
            for event in receipt.state.event_ledger.events()
        }
        if include_identity_metadata
        else {}
    )
    for index, result in enumerate(receipt.ingest_results, 1):
        event = result.event
        lifecycle_event = (
            lifecycle_events.get(event.event_id)
            if event is not None
            else None
        )
        values = [
            index,
            (
                _STATUS_LABELS["SUPERSEDED"]
                if include_identity_metadata
                and lifecycle_event
                and lifecycle_event.status == EVENT_STATUS_SUPERSEDED
                else _STATUS_LABELS.get(result.status, result.status)
            ),
            event.event_id if event else result.duplicate_event_id or "无",
            _name(event.symbol, names) if event else "无",
            _EVENT_LABELS.get(event.event_type, event.event_type) if event else "无",
            event.effective_at.isoformat() if event and event.effective_at else "未披露",
            event.available_at.isoformat() if event else "无",
            event.detected_at.isoformat() if event else "无",
            _SEVERITY_LABELS.get(event.severity, event.severity) if event else "无",
            _CONFIDENCE_LABELS.get(event.confidence, event.confidence) if event else "无",
            "是" if result.is_late else "否",
            "是" if event and event.requires_human_review else "否",
            result.message,
        ]
        if include_identity_metadata:
            values[3:3] = [
                _IDENTITY_LABELS.get(
                    event.event_identity_version,
                    event.event_identity_version,
                )
                if event
                else "无",
                event.source_id if event else "无",
            ]
        fill = AMBER if result.is_late else "FFFFFF"
        if event and event.status == EVENT_STATUS_SUPERSEDED:
            fill = GREY
        _write_row(sheet, row, values, fill=fill)
        row += 1


def _watermark_and_checkpoint(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    watermark: ScanWatermark,
) -> None:
    sheet = wb.create_sheet(WATERMARK_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [24, 28, 62])
    _title(
        sheet,
        "扫描水位与检查点",
        "水位单调前进，任务锁与检查点用于崩溃恢复和幂等补跑。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("扫描范围", watermark.scope, "ALL 表示本次演示的整批范围。"),
        ("来源", watermark.source, "演示源，不指向生产采集器。"),
        ("覆盖截至", watermark.coverage_through.isoformat(), "晚于该时点的输入不被当作历史。"),
        ("抓取时间", watermark.retrieved_at.isoformat(), "与覆盖边界分开记录。"),
        ("覆盖状态", watermark.coverage_status, "不完整时不得宣称正常日静默。"),
        ("来源健康", watermark.source_health, "源中断必须与无事件区分。"),
        ("检查点状态", receipt.checkpoint.status, "COMMITTED 表示本批已完成记录。"),
        ("检查点序号", receipt.checkpoint.last_sequence, "下一批可从此序号续接。"),
        ("水位ID", watermark.watermark_id, "检查点与水位共同形成恢复边界。"),
        ("动作", ACTION_NO_ORDER, "本工作簿不启动或修改生产调度。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def _invalidations(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(INVALIDATION_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "事件ID",
        "事件类型",
        "证券",
        "失效节点",
        "节点类型",
        "深度",
        "原因",
        "队列序号",
        "是否截断",
    ]
    _widths(sheet, [34, 20, 15, 22, 24, 8, 42, 9, 9])
    _title(
        sheet,
        "依赖失效与有界重算",
        "只失效受影响节点；价格事件不会把估值结果标记为需重算。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for invalidation in receipt.invalidations:
        for index, item in enumerate(invalidation.recalculation_queue(), 1):
            _write_row(
                sheet,
                row,
                [
                    invalidation.event_id,
                    _EVENT_LABELS.get(invalidation.event_type, invalidation.event_type),
                    _name(invalidation.symbol, names),
                    item.node_id,
                    item.kind,
                    item.depth,
                    item.reason,
                    index,
                    "是" if invalidation.truncated else "否",
                ],
            )
            row += 1


def _outbox(
    wb: Workbook,
    receipt: M5EventRunReceipt,
    names: Mapping[str, str],
    *,
    current_state: M5EventRunState | None = None,
    include_sent_at: bool = False,
) -> None:
    sheet = wb.create_sheet(OUTBOX_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    base_columns = [
        "提醒ID",
        "类型",
        "严重度",
        "证券",
        "事件ID",
        "人工复核",
        "状态",
        "尝试次数",
        "下次尝试",
        "投递时间",
        "错误",
    ]
    columns = (
        base_columns[:10] + ["发送时间"] + base_columns[10:]
        if include_sent_at
        else base_columns
    )
    widths = (
        [34, 18, 10, 15, 34, 10, 18, 9, 18, 18, 18, 28]
        if include_sent_at
        else [34, 18, 10, 15, 34, 10, 18, 9, 18, 18, 28]
    )
    _widths(sheet, widths)
    _title(
        sheet,
        "Outbox 提醒账",
        "仅记录投递状态；关键提醒不因普通去重被吞掉。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    alerts = (
        current_state.outbox.alerts()
        if current_state is not None
        else receipt.alerts
    )
    for alert in alerts:
        event = next(
            (
                result.event
                for result in receipt.ingest_results
                if result.event and result.event.event_id == alert.event_id
            ),
            None,
        )
        values = [
            alert.alert_id,
            _ALERT_LABELS.get(alert.alert_type, alert.alert_type),
            _SEVERITY_LABELS.get(alert.severity, alert.severity),
            _name(event.symbol, names) if event else "系统",
            alert.event_id or "无",
            "是" if alert.requires_human_review else "否",
            alert.status,
            alert.attempts,
            alert.next_attempt_at.isoformat() if alert.next_attempt_at else "无",
            alert.delivered_at.isoformat() if alert.delivered_at else "无",
            alert.last_error or "无",
        ]
        if include_sent_at:
            values[10:10] = [
                alert.sent_at.isoformat() if alert.sent_at else "无"
            ]
        _write_row(sheet, row, values)
        row += 1


def _outbox_transitions(
    wb: Workbook,
    state: M5EventRunState,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(OUTBOX_TRANSITION_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "修订",
        "迁移ID",
        "提醒ID",
        "来源事件ID",
        "来源ID",
        "身份版本",
        "证券",
        "事件ID",
        "前状态",
        "后状态",
        "发生时间",
        "错误",
        "动作",
    ]
    _widths(sheet, [8, 36, 34, 28, 22, 22, 15, 34, 18, 20, 20, 34, 12])
    _title(
        sheet,
        "Outbox 迁移日志",
        "追加记录提醒状态变化；本页只展示调用方确认的结果，不执行真实通知。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    alerts_by_id = {item.alert_id: item for item in state.outbox.alerts()}
    events_by_id = {
        item.event_id: item for item in state.event_ledger.events()
    }
    for transition in state.outbox_transitions:
        alert = alerts_by_id.get(transition.alert_id)
        event = events_by_id.get(alert.event_id) if alert and alert.event_id else None
        _write_row(
            sheet,
            row,
            [
                transition.outbox_revision,
                transition.transition_id,
                transition.alert_id,
                event.source_event_id if event else "无",
                event.source_id if event else "无",
                _IDENTITY_LABELS.get(
                    event.event_identity_version,
                    event.event_identity_version,
                )
                if event
                else "无",
                _name(event.symbol, names) if event else "系统",
                alert.event_id if alert and alert.event_id else "无",
                transition.from_status,
                transition.to_status,
                transition.occurred_at.isoformat(),
                transition.error or "无",
                transition.action,
            ],
            fill=AMBER if transition.error else "FFFFFF",
        )
        row += 1


def _boundaries(
    wb: Workbook,
    *,
    sheet_name: str = BOUNDARY_SHEET,
    include_outbox_transitions: bool = False,
) -> None:
    sheet = wb.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = False
    columns = ["边界", "状态", "说明"]
    _widths(sheet, [28, 14, 66])
    _title(
        sheet,
        "输入与运行边界",
        "本候选只证明离线合同；真实事件、生产调度与通知投递另需授权。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("事件命名空间", "SIMULATED", "真实事件和私人账户数据不进入公开工作簿。"),
        ("事件时间链", "已校验", "detected / effective / available 分别记录。"),
        ("重复与更正", "已校验", "相同事件忽略，更正显式引用旧事件 ID。"),
        ("晚到事件", "已校验", "保留历史可用时间，不用当前运行时间重标。"),
        ("扫描水位", "单调", "覆盖边界不允许回退。"),
        ("任务锁", "单锁/租约", "防止重复 worker；本候选不创建常驻任务。"),
        ("依赖重算", "有界", "超过上限时记录 deferred，不静默扩大范围。"),
        ("价格与估值", "隔离", "价格变化只影响桥接，不修改 Bear/Base/Bull。"),
    ]
    if include_outbox_transitions:
        rows.extend(
            [
                (
                    "事件身份",
                    "分版本",
                    "v1/legacy 保持历史 ID，新事件使用 source-id-v2。",
                ),
                (
                    "Outbox迁移",
                    "追加日志",
                    "迁移记录与事件批次 revision 分开，重复请求幂等。",
                ),
            ]
        )
    rows.extend(
        [
            ("通知投递", "未启用", "只记录 outbox 状态，不发送真实通知。"),
            ("生产调度", "未授权", "本批不修改服务器、PTA、数据库或计划任务。"),
            ("动作", ACTION_NO_ORDER, "所有决策仍由人工完成。"),
        ]
    )
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def build_m5_event_workbook(
    receipt: M5EventRunReceipt,
    watermark: ScanWatermark,
    *,
    security_names: Mapping[str, str] | None = None,
    current_state: M5EventRunState | None = None,
    include_outbox_transitions: bool = False,
    include_identity_metadata: bool = False,
) -> Workbook:
    if current_state is not None:
        if not isinstance(current_state, M5EventRunState):
            raise ValueError("current_state must be an M5EventRunState snapshot")
        current_payload = current_state.as_policy()
        receipt_payload = receipt.state.as_policy()
        for key in ("outbox", "outbox_revision", "outbox_transitions"):
            current_payload.pop(key)
            receipt_payload.pop(key)
        if current_payload != receipt_payload:
            raise ValueError(
                "current_state must share the receipt non-outbox payload"
            )
        if (
            not include_outbox_transitions
            and current_state.state_sha256() != receipt.state_sha256
        ):
            raise ValueError(
                "Changed outbox state requires transition history display"
            )
    if include_outbox_transitions:
        if not isinstance(current_state, M5EventRunState):
            raise ValueError(
                "Outbox transitions require an M5EventRunState snapshot"
            )
    names = dict(security_names or {})
    wb = Workbook()
    wb.remove(wb.active)
    _overview(
        wb,
        receipt,
        names,
        current_state=current_state,
        include_outbox_transitions=include_outbox_transitions,
    )
    _events(
        wb,
        receipt,
        names,
        include_identity_metadata=include_identity_metadata,
    )
    _watermark_and_checkpoint(wb, receipt, watermark)
    _invalidations(wb, receipt, names)
    _outbox(
        wb,
        receipt,
        names,
        current_state=current_state,
        include_sent_at=include_outbox_transitions,
    )
    if include_outbox_transitions:
        _outbox_transitions(wb, current_state, names)
    _boundaries(
        wb,
        sheet_name=(
            BOUNDARY_SHEET_WITH_TRANSITIONS
            if include_outbox_transitions
            else BOUNDARY_SHEET
        ),
        include_outbox_transitions=include_outbox_transitions,
    )
    return wb


def write_m5_event_workbook(
    receipt: M5EventRunReceipt,
    watermark: ScanWatermark,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
    current_state: M5EventRunState | None = None,
    include_outbox_transitions: bool = False,
    include_identity_metadata: bool = False,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Workbook output must remain inside the project root")
    if output.exists():
        raise ValueError("M5 event candidate output already exists")
    wb = build_m5_event_workbook(
        receipt,
        watermark,
        security_names=security_names,
        current_state=current_state,
        include_outbox_transitions=include_outbox_transitions,
        include_identity_metadata=include_identity_metadata,
    )
    wb.save(output)
    alerts = (
        current_state.outbox.alerts()
        if include_outbox_transitions and current_state is not None
        else receipt.alerts
    )
    active_event_count = (
        len(current_state.event_ledger.active_events())
        if include_outbox_transitions and current_state is not None
        else len(receipt.active_events)
    )
    result = {
        "workbook_path": str(output),
        "workbook_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "sheet_count": len(wb.sheetnames),
        "event_count": len(receipt.ingest_results),
        "active_event_count": active_event_count,
        "invalidation_count": len(receipt.invalidations),
        "alert_count": len(alerts),
        "action": receipt.action,
    }
    if include_outbox_transitions and current_state is not None:
        result.update(
            {
                "outbox_revision": current_state.outbox_revision,
                "outbox_transition_count": len(
                    current_state.outbox_transitions
                ),
                "state_sha256": current_state.state_sha256(),
            }
        )
    return result
