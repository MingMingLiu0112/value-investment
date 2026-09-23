"""Presentation-only simulated workbook for M4 position and dividend guidance."""
from __future__ import annotations

from decimal import Decimal
import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .dividend_income_projection import (
    BASIS_DECLARED,
    BASIS_FORWARD,
    BASIS_NORMALIZED,
    BASIS_PAID,
    PortfolioDividendIncomeProjection,
)
from .investment_decision import ACTION_NO_ORDER
from .position_guidance import (
    BUDGET_CONSTRAINED,
    BUDGET_EXHAUSTED,
    LINE_REVIEW_REQUIRED,
    LINE_WAIT,
    NAMESPACE_SIMULATED,
    STATUS_BUDGET_CONFLICT,
    STATUS_INCOMPLETE,
    STATUS_NO_ACTIONABLE_CAPACITY,
    STATUS_PARTIAL,
    STATUS_READY,
    STATUS_REVIEW_REQUIRED,
    TIER_MAX,
    TIER_NORMAL,
    TIER_STARTER,
    PositionGuidanceResult,
)


OVERVIEW_SHEET = "00_总览"
POSITION_SHEET = "01_仓位分层"
BUDGET_SHEET = "02_共同预算"
INCOME_SHEET = "03_股息收入"
BOUNDARY_SHEET = "04_输入与边界"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_GUIDANCE_LABELS = {
    STATUS_READY: "可复核",
    STATUS_PARTIAL: "部分候选等待",
    STATUS_BUDGET_CONFLICT: "共同预算冲突",
    STATUS_REVIEW_REQUIRED: "需人工复核",
    STATUS_NO_ACTIONABLE_CAPACITY: "无可用新增容量",
    STATUS_INCOMPLETE: "输入不完整",
}
_LINE_LABELS = {
    "ELIGIBLE": "符合上限条件",
    LINE_WAIT: "等待前置条件",
    LINE_REVIEW_REQUIRED: "需人工复核",
}
_TIER_LABELS = {
    "NONE": "无",
    TIER_STARTER: "Starter",
    TIER_NORMAL: "Normal",
    TIER_MAX: "Max",
}
_BUDGET_LABELS = {
    "AVAILABLE": "可用",
    BUDGET_CONSTRAINED: "共同预算约束",
    BUDGET_EXHAUSTED: "已耗尽",
    "NOT_APPLICABLE": "不适用",
}
_BASIS_LABELS = {
    BASIS_PAID: "已到账",
    BASIS_DECLARED: "已宣告未到账",
    BASIS_FORWARD: "Forward 估计",
    BASIS_NORMALIZED: "Normalized 情景",
}
_TAX_LABELS = {
    "CALCULATED": "已计算",
    "UNKNOWN": "未知",
    "UNKNOWN_PENDING_DISPOSAL": "待处置结算",
    "REVIEW_REQUIRED": "需复核",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_row(ws, row: int, values: list[Any], fill: str = "FFFFFF") -> None:
    for column, value in enumerate(values, 1):
        _style(ws.cell(row, column, value), fill=fill)


def _money(value: Decimal | None) -> str:
    if value is None:
        return "缺失"
    return f"{value:,.2f}"


def _pct(value: Decimal | None) -> str:
    if value is None:
        return "缺失"
    return f"{value * Decimal('100'):.2f}%"


def _join(values: tuple[str, ...]) -> str:
    return "；".join(value for value in values if value)


def _name(symbol: str, names: Mapping[str, str]) -> str:
    return names.get(symbol, symbol)


def _overview(
    wb: Workbook,
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
) -> None:
    sheet = wb.create_sheet(OVERVIEW_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [22, 22, 68])
    _title(
        sheet,
        "M4 仓位边界与股息收入（模拟演示）",
        "显式模拟组合仅演示预算、限制与收入口径；不生成订单或真实个人建议。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    metrics = [
        ("评估命名空间", "SIMULATED", "公开工作簿只允许模拟输入。"),
        ("仓位边界状态", _GUIDANCE_LABELS[guidance.status], "只输出上限和复核条件。"),
        ("股息收入状态", dividend.status, "四种收入基础分别展示。"),
        ("总资产（元）", _money(guidance.total_assets_cny()), "现金与持仓市值合计。"),
        ("最低保留现金（元）", _money(guidance.reserved_cash_cny()), "最低现金与应急/流动性需要取保守值。"),
        ("可用新增预算（元）", _money(guidance.free_budget_cny()), "扣除现有持仓与保留现金后的余额。"),
        ("年度股息目标（元）", _money(dividend.bundle.policy.dividend_income_goal_cny), "目标差额逐口径展示，不汇总成交易建议。"),
        ("动作", ACTION_NO_ORDER, "所有结论仍由人工决定。"),
    ]
    for name, value, explanation in metrics:
        _write_row(sheet, row, [name, value, explanation])
        row += 1


def _positions(
    wb: Workbook,
    guidance: PositionGuidanceResult,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(POSITION_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "证券",
        "复核意图",
        "状态",
        "分层",
        "上限",
        "当前权重",
        "剩余上限",
        "共同预算",
        "停止加仓条件",
        "减仓复核触发",
        "证据ID",
    ]
    _widths(sheet, [14, 14, 16, 12, 10, 10, 10, 15, 38, 32, 20])
    _title(
        sheet,
        "分层仓位边界",
        "Starter/Normal/Max 均为人工确认的上限，不生成交易量或分配结果。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for line in guidance.lines():
        fill = AMBER if line.status == LINE_REVIEW_REQUIRED else "FFFFFF"
        _write_row(
            sheet,
            row,
            [
                _name(line.symbol, names),
                line.review_intent,
                _LINE_LABELS[line.status],
                _TIER_LABELS[line.tier],
                _pct(line.ceiling_pct),
                _pct(line.current_weight_pct),
                _pct(line.remaining_ceiling_pct),
                _BUDGET_LABELS[line.budget_status],
                _join(line.stop_add_conditions),
                _join(line.reduce_review_triggers),
                _join(tuple(ref.get("id", "") for ref in line.evidence_refs)),
            ],
            fill=fill,
        )
        row += 1


def _budget(wb: Workbook, guidance: PositionGuidanceResult) -> None:
    sheet = wb.create_sheet(BUDGET_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [26, 20, 66])
    _title(
        sheet,
        "组合共同预算",
        "所有候选共享一个账户预算；只标记冲突，不代替人工分配。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("总资产（元）", _money(guidance.total_assets_cny()), "现金与已确认持仓合计。"),
        ("保留现金（元）", _money(guidance.reserved_cash_cny()), "应急与流动性需要未被当作投资预算。"),
        ("可用新增预算（元）", _money(guidance.free_budget_cny()), "多个候选共同争用这部分预算。"),
        ("预算状态", _GUIDANCE_LABELS[guidance.status], "候选上限之和超过可用预算时显式冲突。"),
        ("自动分配", "否", "系统不分配仓位，也不计算目标权重或订单。"),
        ("动作", ACTION_NO_ORDER, "人工决定是否复核或调整。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def _income(wb: Workbook, dividend: PortfolioDividendIncomeProjection) -> None:
    sheet = wb.create_sheet(INCOME_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["口径", "总收入（元）", "净收入（元）", "税费状态", "目标差额（元）", "解释"]
    _widths(sheet, [20, 18, 18, 18, 18, 50])
    _title(
        sheet,
        "股息收入四口径",
        "已到账、已宣告、Forward 与 Normalized 分别展示；特别分红不会自动年化。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    explanations = {
        BASIS_PAID: "只汇总已到账且权益确认的分配；税费仅在批次日期完整时计算。",
        BASIS_DECLARED: "已批准但尚未到账；净额须等处置结算后才能确定。",
        BASIS_FORWARD: "普通股息的显式未来估计，不是已到账事实。",
        BASIS_NORMALIZED: "压力情景输入，仅用于现金流复核，不冒充预测。",
    }
    for summary in dividend.summaries():
        _write_row(
            sheet,
            row,
            [
                _BASIS_LABELS[summary.basis],
                _money(summary.gross_income),
                _money(summary.net_income),
                _TAX_LABELS[summary.tax_status],
                _money(summary.goal_gap),
                explanations[summary.basis],
            ],
        )
        row += 1


def _boundary(
    wb: Workbook,
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
) -> None:
    sheet = wb.create_sheet(BOUNDARY_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["边界项", "值/说明", "证据ID"]
    _widths(sheet, [26, 64, 30])
    _title(
        sheet,
        "输入与边界",
        "所有输入均为模拟；真实 IPS、持仓和股息口径未提供时不进入公开工作簿。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    boundaries = [
        ("模拟组合标识", "SIMULATED_PUBLIC_DEMONSTRATION", "fixture:simulated-m4-guidance-income"),
        ("IPS 版本", dividend.bundle.policy.policy_version, "；".join(ref.get("id", "") for ref in dividend.bundle.policy.evidence_refs)),
        ("仓位分层政策版本", guidance.tier_policy.policy_version, "；".join(ref.get("id", "") for ref in guidance.tier_policy.evidence_refs)),
        ("Starter 上限", _pct(guidance.tier_policy.starter_cap_pct / Decimal("100")), "人工确认，不使用默认 20%。"),
        ("Normal 上限", _pct(guidance.tier_policy.normal_cap_pct / Decimal("100")), "人工确认的分层上限。"),
        ("Max 上限", _pct(guidance.tier_policy.max_cap_pct / Decimal("100")), "不得高于 IPS 单股上限。"),
        ("普通/特别分红", "分开显示", "特别分红不进入 Forward 或 Normalized。"),
        ("税费口径", "按已有批次可计算；未结算保持未知", "不虚构未来处置时点的净额。"),
        ("限制项", _join(dividend.bundle.policy.restrictions), "policy.restrictions"),
        ("缺失输入", _join((*guidance.missing_inputs(), *dividend.missing_inputs())), "空表示模拟结构完整"),
        ("动作", ACTION_NO_ORDER, "不生成订单"),
    ]
    for name, value, evidence in boundaries:
        _write_row(sheet, row, [name, value, evidence])
        row += 1


def build_guidance_income_workbook(
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
    *,
    security_names: Mapping[str, str] | None = None,
) -> Workbook:
    if guidance.assessment_namespace != NAMESPACE_SIMULATED:
        raise ValueError("Public guidance workbook requires simulated guidance")
    if dividend.assessment_namespace != NAMESPACE_SIMULATED:
        raise ValueError("Public guidance workbook requires simulated dividends")
    if guidance.action != ACTION_NO_ORDER or dividend.action != ACTION_NO_ORDER:
        raise ValueError("Guidance workbook requires action=no_order")
    if not guidance.can_guide() or not dividend.can_project():
        raise ValueError("Guidance workbook requires structurally complete simulated inputs")
    names = dict(security_names or {})
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, guidance, dividend)
    _positions(wb, guidance, names)
    _budget(wb, guidance)
    _income(wb, dividend)
    _boundary(wb, guidance, dividend)
    wb.active = 0
    return wb


def write_guidance_income_workbook(
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Guidance workbook output escapes its root")
    if output.exists():
        raise ValueError(f"Guidance workbook already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_guidance_income_workbook(
        guidance,
        dividend,
        security_names=security_names,
    )
    workbook.save(output)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    return {
        "workbook_path": str(output.relative_to(root)),
        "workbook_sha256": digest,
        "sheet_count": len(workbook.sheetnames),
        "candidate_count": len(guidance.candidates),
        "income_basis_count": len(dividend.summaries()),
        "guidance_status": guidance.status,
        "dividend_status": dividend.status,
        "assessment_namespace": NAMESPACE_SIMULATED,
        "action": ACTION_NO_ORDER,
    }
