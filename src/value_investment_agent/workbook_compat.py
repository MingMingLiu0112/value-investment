"""Read WPS non-stacked bar charts without changing the source workbook."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook


def load_wps_workbook(path, **kwargs):
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
    buffer = BytesIO()
    changed = False
    with ZipFile(path) as source, ZipFile(buffer, 'w') as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename.startswith('xl/charts/chart') and member.filename.endswith('.xml'):
                tree = ET.fromstring(content)
                repaired = False
                for kind in ('barChart', 'bar3DChart'):
                    for chart in tree.iter(ns + kind):
                        grouping = chart.find(ns + 'grouping')
                        if grouping is not None and grouping.get('val') == 'none':
                            grouping.set('val', 'standard')
                            repaired = True
                if repaired:
                    content = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
                    changed = True
            target.writestr(member, content)
    if changed:
        buffer.seek(0)
        return load_workbook(buffer, **kwargs)
    return load_workbook(path, **kwargs)
