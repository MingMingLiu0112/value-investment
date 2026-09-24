"""Presentation-only M2 Channel Verification human-review workbook."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .m2_channel_verification import (
    ACTION_NO_ORDER,
    STATUS_INSUFFICIENT,
    STATUS_REJECTED,
    STATUS_UNSUPPORTED,
    STATUS_VERIFIED,
    M2ChannelVerificationBatch,
)


OVERVIEW_SHEET = "00_验证总览"
RESOLUTION_SHEET = "01_二阶段结论"
EVIDENCE_SHEET = "02_证据追溯"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_STATUS_LABELS = {
    STATUS_VERIFIED: "进入深研队列",
    STATUS_REJECTED: "通道否决",
    STATUS_INSUFFICIENT: "证据不足",
    STATUS_UNSUPPORTED: "不支持",
}
_STATUS_FILLS = {
    STATUS_VERIFIED: "E4F3EE",
    STATUS_REJECTED: "F8E7E6",
    STATUS_INSUFFICIENT: "FFF1D6",
    STATUS_UNSUPPORTED: "EEEEEE",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _overview(wb: Workbook, batch: M2ChannelVerificationBatch) -> None:
    ws = wb.create_sheet(OVERVIEW_SHEET)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 100
    _title(
        ws,
        "M2 通道二阶段验证",
        "仅对预注册 LEAD 做可追溯二阶段 resolution；所有结果为研究队列结论，不代表便宜或值得买入。",
        2,
    )
    counts = batch.counts()
    rows = [
        ("接受状态", batch.acceptance_status),
        ("机器状态", batch.machine_status),
        ("LEAD 数", str(batch.resolution_count())),
        ("进入深研队列", str(counts[STATUS_VERIFIED])),
        ("通道否决", str(counts[STATUS_REJECTED])),
        ("证据不足", str(counts[STATUS_INSUFFICIENT])),
        ("不支持", str(counts[STATUS_UNSUPPORTED])),
        ("进入深研代码", "、".join(batch.verified_symbols())),
        ("进入深研通道", "、".join(batch.verified_channels())),
        ("规则", "低 PE/PB 或单一触发条件不得单独形成 VERIFIED；否决与证据不足均合法。"),
        ("动作", ACTION_NO_ORDER),
    ]
    row = 4
    for label, value in rows:
        _style(ws.cell(row, 1, label), fill=GREY, bold=True)
        _style(ws.cell(row, 2, str(value)))
        row += 1


def _resolution_sheet(wb: Workbook, batch: M2ChannelVerificationBatch) -> None:
    ws = wb.create_sheet(RESOLUTION_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    columns = [
        "证券代码",
        "公司",
        "通道",
        "二阶段结论",
        "原因",
        "证据数",
        "验证ID",
    ]
    _widths(ws, [12, 18, 24, 22, 72, 10, 44])
    _title(
        ws,
        "预注册 LEAD 二阶段结论",
        "全部 resolution 可追溯至固定 AC8 报告和原文 Hash；不允许为出现正向机会重新选股。",
        len(columns),
    )
    row = _header(ws, 4, columns)
    for item in batch.results:
        fill = _STATUS_FILLS[item.status]
        values = [
            item.symbol,
            item.name,
            item.channel,
            _STATUS_LABELS[item.status],
            item.reason,
            len(item.evidence),
            item.verification_id,
        ]
        for column, value in enumerate(values, 1):
            _style(ws.cell(row, column, value), fill=fill)
        row += 1


def _evidence_sheet(wb: Workbook, batch: M2ChannelVerificationBatch) -> None:
    ws = wb.create_sheet(EVIDENCE_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    columns = [
        "证券代码",
        "通道",
        "字段",
        "报告期",
        "值",
        "单位",
        "校验状态",
        "来源",
        "来源URL",
        "原件SHA-256",
    ]
    _widths(ws, [12, 24, 20, 15, 18, 13, 12, 32, 58, 64])
    _title(
        ws,
        "二阶段证据追溯",
        "进入深研队列必须保留来源、URL 和原件 Hash；证据不足不冒充拒绝或无机会。",
        len(columns),
    )
    row = _header(ws, 4, columns)
    for item in batch.results:
        for evidence in item.evidence:
            values = [
                item.symbol,
                item.channel,
                evidence.field_name,
                evidence.period_label,
                evidence.value,
                evidence.unit,
                evidence.validation_status,
                evidence.source_name,
                evidence.source_url,
                evidence.source_sha256 or "",
            ]
            for column, value in enumerate(values, 1):
                _style(ws.cell(row, column, value))
            if evidence.source_url.startswith(("http://", "https://")):
                url_cell = ws.cell(row, 9, evidence.source_url)
                url_cell.hyperlink = evidence.source_url
                url_cell.font = Font(
                    name="Microsoft YaHei",
                    size=11,
                    color=BLUE,
                    underline="single",
                )
            row += 1


def build_channel_verification_workbook(
    batch: M2ChannelVerificationBatch,
) -> Workbook:
    if batch.action != ACTION_NO_ORDER:
        raise ValueError("Channel verification workbook requires action=no_order")
    if batch.machine_status != "MACHINE_CHECKS_PASS":
        raise ValueError("Channel verification has not passed machine checks")
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, batch)
    _resolution_sheet(wb, batch)
    _evidence_sheet(wb, batch)
    wb.active = 0
    return wb


def write_channel_verification_workbook(
    batch: M2ChannelVerificationBatch,
    *,
    output: Path,
    root: Path,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Channel verification workbook escapes project root")
    if output.exists():
        raise ValueError(f"Channel verification workbook already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_channel_verification_workbook(batch)
    workbook.save(output)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    return {
        "workbook_path": str(output.relative_to(root)),
        "workbook_sha256": digest,
        "sheet_count": len(workbook.sheetnames),
        "action": ACTION_NO_ORDER,
    }
