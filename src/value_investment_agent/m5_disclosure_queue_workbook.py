"""Presentation-only workbook for the real CNINFO M5 review queue."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .event_scan import COVERAGE_COMPLETE, COVERAGE_INCOMPLETE
from .m5_disclosure_queue import (
    ACTION_NO_ORDER,
    PROVIDER_CNINFO,
    SOURCE_ARCHIVED,
    SOURCE_NAME_CNINFO,
    SOURCE_UNAVAILABLE,
    DisclosureReviewQueue,
)


OVERVIEW_SHEET = "00_总览"
PENDING_SHEET = "01_待复核公告"
COVERAGE_SHEET = "02_来源覆盖"
BOUNDARY_SHEET = "03_输入与边界"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
GREY = "EFF3F1"
AMBER = "FFF1D6"
RED = "C0392B"

_KIND_LABELS = {
    "financial_statement": "财务报告",
    "dividend": "利润分配",
    "buyback": "股份回购",
    "capital_structure": "资本结构",
    "asset_impairment": "资产减值",
    "accounting_policy": "会计政策",
    "guarantee": "对外担保",
    "operating_data": "经营数据",
    "governance": "治理/流程",
    "investor_relations": "投资者关系",
    "routine": "例行事项",
    "unknown": "待判断",
}
_REVIEW_LABELS = {
    "PENDING_HUMAN_REVIEW": "待人工复核",
    "REVIEWED_NO_MATERIAL_CANDIDATE": "标题规则不候选",
}
_COVERAGE_LABELS = {
    COVERAGE_COMPLETE: "完整",
    COVERAGE_INCOMPLETE: "不完整",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_row(ws, row: int, values: list[Any], fill: str = "FFFFFF") -> None:
    for column, value in enumerate(values, 1):
        _style(ws.cell(row, column, value), fill=fill)


def _name(symbol: str, names: Mapping[str, str]) -> str:
    return names.get(symbol, symbol)


def _pdf_hash(item) -> str:
    for ref in item.evidence_refs:
        if ref.get("source_status") == SOURCE_ARCHIVED and ref.get("sha256"):
            return str(ref["sha256"])
    return "待补原件"


def _overview(
    wb: Workbook,
    queue: DisclosureReviewQueue,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(OVERVIEW_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [24, 42, 64])
    _title(
        sheet,
        "M5 真实披露待复核队列",
        "真实 CNINFO 索引已归档；标题规则只形成候选，材料性结论仍由人工完成。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    pending = len(queue.pending_candidates)
    unavailable = queue.unavailable_source_count
    complete = sum(
        scan.coverage_status == COVERAGE_COMPLETE for scan in queue.scans
    )
    rows = [
        ("队列ID", queue.queue_id, "本轮真实索引的确定性运行标识。"),
        ("来源", SOURCE_NAME_CNINFO, "法定披露聚合来源，非自动投资结论。"),
        ("扫描区间", f"{queue.scan_from} 至 {queue.scan_to}", "仅纳入公告日落在该窗口内的原文。"),
        ("抓取时间", queue.retrieved_at.isoformat(), "所有公告均不得晚于该时间。"),
        ("公司数", len(queue.scans), "本轮使用小样本真实队列，不扩大全市场范围。"),
        ("待人工复核", pending, "标题规则候选，不是材料性判定。"),
        ("原件不可用", unavailable, "缺 PDF 时保持阻断，不降级为无事件。"),
        ("覆盖完整", f"{complete}/{len(queue.scans)}", "来源索引完整；失败公司单独标记。"),
        ("解析器", queue.parser_version, "规则版本固定，便于旧结果复算。"),
        ("动作", queue.action, "本队列始终不生成订单或交易意图。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def _pending(
    wb: Workbook,
    queue: DisclosureReviewQueue,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(PENDING_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "证券",
        "公司",
        "公告ID",
        "发布时间",
        "公告标题",
        "标题规则",
        "来源URL",
        "PDF SHA-256",
        "状态",
        "说明",
    ]
    _widths(sheet, [10, 16, 13, 20, 52, 13, 46, 28, 14, 42])
    _title(
        sheet,
        "待人工复核公告",
        "只列出标题规则候选；未知标题保留并进入人工复核，不静默丢弃。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for scan in queue.scans:
        for item in scan.announcements:
            if not item.materiality_candidate:
                continue
            unavailable = any(
                ref.get("source_status") == SOURCE_UNAVAILABLE
                for ref in item.evidence_refs
            )
            _write_row(
                sheet,
                row,
                [
                    scan.symbol,
                    _name(scan.symbol, names),
                    item.announcement_id,
                    item.published_at.isoformat(),
                    item.title,
                    _KIND_LABELS.get(item.rule_kind, item.rule_kind),
                    item.source_url,
                    _pdf_hash(item),
                    "原件缺失" if unavailable else _REVIEW_LABELS.get(
                        item.review_status,
                        item.review_status,
                    ),
                    item.notes,
                ],
                fill=AMBER if unavailable else "FFFFFF",
            )
            row += 1
    if row == 5:
        _write_row(sheet, row, ["无待复核公告"] + [""] * (len(columns) - 1), fill=GREY)


def _coverage(
    wb: Workbook,
    queue: DisclosureReviewQueue,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(COVERAGE_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "证券",
        "公司",
        "公告总数",
        "待复核候选",
        "非候选",
        "覆盖状态",
        "索引 SHA-256",
        "阻断原因",
    ]
    _widths(sheet, [10, 16, 10, 11, 9, 11, 66, 52])
    _title(
        sheet,
        "来源覆盖与证据索引",
        "每条扫描保存原始 CNINFO 索引与 Hash；源失败必须区别于无事件。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for scan in queue.scans:
        candidates = sum(item.materiality_candidate for item in scan.announcements)
        index_sha = next(
            (
                str(ref.get("sha256"))
                for ref in scan.evidence_refs
                if ref.get("sha256")
            ),
            "源失败",
        )
        _write_row(
            sheet,
            row,
            [
                scan.symbol,
                _name(scan.symbol, names),
                len(scan.announcements),
                candidates,
                len(scan.announcements) - candidates,
                _COVERAGE_LABELS.get(scan.coverage_status, scan.coverage_status),
                index_sha,
                "；".join(scan.blockers),
            ],
            fill=AMBER if scan.coverage_status == COVERAGE_INCOMPLETE else "FFFFFF",
        )
        row += 1


def _boundaries(wb: Workbook) -> None:
    sheet = wb.create_sheet(BOUNDARY_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["边界", "状态", "说明"]
    _widths(sheet, [26, 16, 70])
    _title(
        sheet,
        "输入与运行边界",
        "本候选只建立真实披露等待队列，不自动判定材料性、不接入通知或调度。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("材料性判定", "禁止自动", "标题规则只能产生候选；只有人工结论可进入 M5 事件桥。"),
        ("未知标题", "保留复核", "不能识别时不降级为非候选。"),
        ("缺失 PDF", "失败关闭", "候选原件缺失会写入阻断原因，不标记无事件。"),
        ("来源失败", "显式告警", "CNINFO 索引失败记录为来源健康异常。"),
        ("证据时间", "逐条校验", "未来公告、窗口外公告和重复 ID 直接拒绝。"),
        ("生产调度", "未启用", "不创建常驻服务或计划任务。"),
        ("通知投递", "未启用", "本队列不发送提醒。"),
        ("数据库", "未连接", "不连接生产 PostgreSQL。"),
        ("动作", ACTION_NO_ORDER, "所有后续决策仍由用户完成。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def build_m5_disclosure_queue_workbook(
    queue: DisclosureReviewQueue,
    *,
    security_names: Mapping[str, str] | None = None,
) -> Workbook:
    names = dict(security_names or {})
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, queue, names)
    _pending(wb, queue, names)
    _coverage(wb, queue, names)
    _boundaries(wb)
    return wb


def write_m5_disclosure_queue_workbook(
    queue: DisclosureReviewQueue,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Disclosure queue workbook must remain inside the project root")
    if output.exists():
        raise ValueError("Disclosure queue workbook output already exists")
    wb = build_m5_disclosure_queue_workbook(
        queue,
        security_names=security_names,
    )
    wb.save(output)
    return {
        "workbook_path": str(output),
        "workbook_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "sheet_count": len(wb.sheetnames),
        "company_count": len(queue.scans),
        "pending_candidate_count": len(queue.pending_candidates),
        "source_unavailable_count": queue.unavailable_source_count,
        "coverage_complete_count": sum(
            scan.coverage_status == COVERAGE_COMPLETE for scan in queue.scans
        ),
        "action": queue.action,
    }
