"""Excel renderer for the typed M7 product workbench read model.

This module lays out user-facing cards only. Every assessment, materiality
state, portfolio capacity and dividend state must already exist in the read
model; the renderer performs no investment calculation.
"""
from __future__ import annotations

import hashlib
import json
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..read_models.product_workbench import (
    ACTION_NO_ORDER,
    AssessmentView,
    CompanyCard,
    ProductWorkbenchReadModel,
    StatusView,
)


PRODUCT_WORKBENCH_CANDIDATE_SCHEMA_VERSION = "m7-product-workbench-candidate-v1"

SHEET_TODAY = "01_今日"
SHEET_OPPORTUNITIES = "02_机会"
SHEET_COMPANIES = "03_公司"
SHEET_PORTFOLIO = "04_我的组合"
SHEET_EVENTS = "05_事件"
SHEET_SYSTEM_AUDIT = "06_系统与审计"

USER_SHEETS = (
    SHEET_TODAY,
    SHEET_OPPORTUNITIES,
    SHEET_COMPANIES,
    SHEET_PORTFOLIO,
    SHEET_EVENTS,
)
SECONDARY_SHEETS = (SHEET_SYSTEM_AUDIT,)
WORKBOOK_SHEETS = (*USER_SHEETS, *SECONDARY_SHEETS)

AUDIT_EVIDENCE_START_ROW = 19

INK = "20312B"
WHITE = "FFFFFF"
GREEN = "176B55"
BLUE = "245B7A"
GREY = "EEF3F1"
AMBER = "FFF1D6"
RED_FILL = "F8E7E6"
NAVY = "1E3A4A"
MUTED = "5D6B66"

_THIN = Side(style="thin", color="C9D5D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style(
    cell: Cell,
    *,
    fill: str = WHITE,
    bold: bool = False,
    color: str = INK,
    size: int = 10,
    horizontal: str = "left",
    wrap: bool = True,
    border: bool = False,
) -> None:
    cell.font = Font(name="Microsoft YaHei", size=size, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(
        vertical="top",
        horizontal=horizontal,
        wrap_text=wrap,
    )
    if border:
        cell.border = _BORDER


def _title(ws: Worksheet, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, ws.title[3:] if "_" in ws.title else ws.title), fill=GREEN, bold=True, color=WHITE, size=16)
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 34


def _header(ws: Worksheet, row: int, labels: list[str]) -> int:
    for column, label in enumerate(labels, 1):
        _style(
            ws.cell(row, column, label),
            fill=BLUE,
            bold=True,
            color=WHITE,
            border=True,
        )
    return row + 1


def _section_title(ws: Worksheet, row: int, title: str, columns: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    _style(ws.cell(row, 1, title), fill=NAVY, bold=True, color=WHITE, size=11)
    return row + 1


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _glyph_width(character: str) -> int:
    """Return the column width a single character occupies in the rendered grid."""

    return 2 if unicodedata.east_asian_width(character) in ("W", "F") else 1


def _wrapped_lines(text: str, width_columns: int) -> int:
    """Count the visual lines ``text`` needs inside ``width_columns``."""

    limit = max(1, width_columns)
    lines = 0
    for paragraph in str(text).split("\n"):
        if not paragraph:
            lines += 1
            continue
        used = 1
        current = 0
        for character in paragraph:
            size = _glyph_width(character)
            if current + size > limit and current > 0:
                used += 1
                current = size
            else:
                current += size
        lines += used
    return max(lines, 1)


def _column_width(ws: Worksheet, column: int) -> int:
    dimension = ws.column_dimensions.get(get_column_letter(column))
    width = getattr(dimension, "width", None)
    return int(width) if width else 9


def _fit_rows(
    ws: Worksheet,
    row: int,
    cells: list[tuple[int, str, int]],
    *,
    minimum: int = 20,
) -> None:
    """Size a row from the wrapped height of the tallest cell in it.

    ``cells`` carries ``(first_column, text, column_span)`` so merged blocks and
    single cells are measured the way the renderer actually writes them. Counting
    only explicit newlines left every soft-wrapped Chinese sentence clipped.
    """

    lines = 1
    for column, text, span in cells:
        if text is None:
            continue
        width = sum(
            _column_width(ws, index) for index in range(column, column + span)
        )
        lines = max(lines, _wrapped_lines(text, width - 1))
    ws.row_dimensions[row].height = max(minimum, 15 * lines + 6)


def _hyperlink(cell: Cell, sheet: str, coordinate: str = "A1") -> None:
    cell.hyperlink = f"#'{sheet}'!{coordinate}"
    cell.font = Font(
        name="Microsoft YaHei",
        size=10,
        bold=True,
        color=BLUE,
        underline="single",
    )


def _status_text(status: StatusView) -> str:
    return status.user_label


def _assessment_text(assessment: AssessmentView) -> str:
    if assessment.available:
        return assessment.value_text or ""
    return (
        "暂不可评估\n"
        f"原因：{assessment.unavailable_reason}\n"
        f"需要：{assessment.needed_evidence}"
    )


def _evidence_cell(
    ws: Worksheet,
    row: int,
    column: int,
    refs: tuple[str, ...],
    audit_rows: dict[str, int],
) -> str:
    cell = ws.cell(row, column)
    if not refs:
        cell.value = "-"
        _style(cell, color=MUTED, border=True)
        return "-"
    text = f"查看证据 ({len(refs)})"
    cell.value = text
    _style(cell, border=True)
    _hyperlink(cell, SHEET_SYSTEM_AUDIT, f"A{audit_rows[refs[0]]}")
    return text


def _company_link(
    cell: Cell,
    symbol: str | None,
    company_rows: dict[str, int],
) -> None:
    if symbol is not None and symbol in company_rows:
        _hyperlink(cell, SHEET_COMPANIES, f"A{company_rows[symbol]}")
    else:
        _hyperlink(cell, SHEET_COMPANIES, "A1")


def _render_today(
    ws: Worksheet,
    model: ProductWorkbenchReadModel,
    company_rows: dict[str, int],
    audit_rows: dict[str, int],
) -> None:
    _widths(ws, [18, 28, 34, 34, 28, 16])
    _title(
        ws,
        "优先查看今天真正需要处理的事项；审计信息保留在次级页面。",
        6,
    )
    row = 4
    status_row = row
    _style(ws.cell(row, 1, "数据更新时间"), fill=GREY, bold=True, border=True)
    _style(ws.cell(row, 2, model.overview.data_updated_at.isoformat()), border=True)
    _style(ws.cell(row, 3, "系统健康"), fill=GREY, bold=True, border=True)
    _style(ws.cell(row, 4, _status_text(model.system_health.status)), border=True)
    _style(ws.cell(row, 5, "待处理事项"), fill=GREY, bold=True, border=True)
    _style(ws.cell(row, 6, str(model.overview.pending_count)), border=True)
    _fit_rows(
        ws,
        status_row,
        [
            (1, "数据更新时间", 1),
            (2, model.overview.data_updated_at.isoformat(), 1),
            (3, "系统健康", 1),
            (4, _status_text(model.system_health.status), 1),
            (5, "待处理事项", 1),
            (6, str(model.overview.pending_count), 1),
        ],
    )
    row += 2

    row = _section_title(ws, row, "我的组合", 6)
    if not model.portfolio.real_data_available:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        _style(ws.cell(row, 1, "尚未接入个人组合"), fill=AMBER, bold=True, size=12, border=True)
        link = ws.cell(row, 5, "查看接入说明")
        _style(link, fill=AMBER, border=True)
        _hyperlink(link, SHEET_SYSTEM_AUDIT, "A1")
        _style(ws.cell(row, 6, "-"), fill=AMBER, border=True, color=MUTED)
        row += 2
    else:
        header = _header(ws, row, ["项目", "当前情况"])
        for metric in model.portfolio.summary:
            _style(ws.cell(header, 1, metric.label), fill=GREY, bold=True, border=True)
            _style(ws.cell(header, 2, metric.value_text), border=True)
            header += 1
        row = header + 1

    row = _section_title(ws, row, "今日事项", 6)
    row = _header(
        ws,
        row,
        ["公司", "发生了什么", "为什么重要", "当前状态", "下一步", "证据"],
    )
    if not model.today_items:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _style(
            ws.cell(row, 1, "当前没有已分类的今日事项；这不表示没有风险。"),
            fill=GREY,
            border=True,
        )
        return

    for item in model.today_items:
        _style(ws.cell(row, 1, item.company), border=True)
        _company_link(ws.cell(row, 1), item.symbol, company_rows)
        values = (
            item.what_happened,
            item.why_it_matters,
            item.current_status,
            item.next_step,
        )
        for column, value in enumerate(values, 2):
            _style(ws.cell(row, column, value), border=True)
        evidence_text = _evidence_cell(ws, row, 6, item.evidence_refs, audit_rows)
        _fit_rows(
            ws,
            row,
            [
                (1, item.company, 1),
                (2, values[0], 1),
                (3, values[1], 1),
                (4, values[2], 1),
                (5, values[3], 1),
                (6, evidence_text, 1),
            ],
        )
        row += 1


def _render_opportunities(
    ws: Worksheet,
    model: ProductWorkbenchReadModel,
    company_rows: dict[str, int],
    audit_rows: dict[str, int],
) -> None:
    _widths(ws, [24, 34, 20, 20, 20, 20, 30, 28, 16])
    _title(ws, "只显示已经进入关注范围或研究队列的公司，不展示内部阶段码。", 9)
    row = _header(
        ws,
        4,
        [
            "公司",
            "为什么现在关注",
            "研究状态",
            "估值状态",
            "股息状态",
            "当前价格状态",
            "主要风险",
            "下一触发",
            "证据",
        ],
    )
    if not model.opportunities:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        _style(
            ws.cell(row, 1, "当前没有公司进入关注范围，研究队列仍可保持等待。"),
            fill=GREY,
            border=True,
        )
        return

    for card in model.opportunities:
        _style(ws.cell(row, 1, f"{card.company_name} / {card.symbol}"), border=True)
        _company_link(ws.cell(row, 1), card.symbol, company_rows)
        values = (
            card.why_now,
            _status_text(card.research_status),
            _status_text(card.valuation_status),
            _status_text(card.dividend_status),
            _status_text(card.price_status),
            card.main_risk,
            card.next_trigger,
        )
        for column, value in enumerate(values, 2):
            _style(ws.cell(row, column, value), border=True)
        evidence_text = _evidence_cell(ws, row, 9, card.evidence_refs, audit_rows)
        _fit_rows(
            ws,
            row,
            [
                (1, f"{card.company_name} / {card.symbol}", 1),
                (2, values[0], 1),
                (3, values[1], 1),
                (4, values[2], 1),
                (5, values[3], 1),
                (6, values[4], 1),
                (7, values[5], 1),
                (8, values[6], 1),
                (9, evidence_text, 1),
            ],
        )
        row += 1


def _render_portfolio(
    ws: Worksheet,
    model: ProductWorkbenchReadModel,
    audit_rows: dict[str, int],
) -> None:
    _widths(ws, [24, 34, 24, 42, 24, 16])
    _title(ws, "没有真实组合输入时只显示接入状态，不生成模拟资产数字。", 6)
    row = 4
    if not model.portfolio.real_data_available:
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=6)
        cell = ws.cell(row, 1, "尚未接入真实组合")
        _style(cell, fill=AMBER, bold=True, size=16, horizontal="center", border=True)
        row += 4
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _style(ws.cell(row, 1, model.portfolio.connection_hint), fill=GREY, border=True)
        row += 2
        link = ws.cell(row, 1, "查看接入说明")
        _style(link, border=True)
        _hyperlink(link, SHEET_SYSTEM_AUDIT, "A1")
        return

    row = _section_title(ws, row, "组合摘要", 6)
    row = _header(ws, row, ["指标", "当前情况", "指标", "当前情况", "指标", "当前情况"])
    metrics = model.portfolio.summary
    for index in range(0, len(metrics), 3):
        for offset, metric in enumerate(metrics[index : index + 3]):
            label_column = offset * 2 + 1
            _style(ws.cell(row, label_column, metric.label), fill=GREY, bold=True, border=True)
            _style(ws.cell(row, label_column + 1, metric.value_text), border=True)
        row += 1

    row += 1
    row = _section_title(ws, row, "股票级复核", 6)
    row = _header(
        ws,
        row,
        ["公司", "当前仓位", "允许容量", "当前风险", "继续复核", "证据"],
    )
    if not model.portfolio.positions:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _style(ws.cell(row, 1, "真实组合中暂无持仓。"), border=True)
        return
    for position in model.portfolio.positions:
        values = (
            f"{position.company_name} / {position.symbol}",
            position.current_position_text,
            position.allowed_capacity_text,
            position.current_risk_text,
        )
        for column, value in enumerate(values, 1):
            _style(ws.cell(row, column, value), border=True)
        review_text = (
            _status_text(position.continuation_review) + "\n" + position.reason
        )
        _style(ws.cell(row, 5, review_text), border=True)
        evidence_text = _evidence_cell(ws, row, 6, position.evidence_refs, audit_rows)
        _fit_rows(
            ws,
            row,
            [
                (1, values[0], 1),
                (2, values[1], 1),
                (3, values[2], 1),
                (4, values[3], 1),
                (5, review_text, 1),
                (6, evidence_text, 1),
            ],
        )
        row += 1


def _render_events(
    ws: Worksheet,
    model: ProductWorkbenchReadModel,
    audit_rows: dict[str, int],
) -> None:
    _widths(ws, [24, 54, 38, 28, 40, 16])
    _title(ws, "只显示已分类的重大、逻辑风险、股息、组合风险和系统数据风险。", 6)
    row = 4
    if not model.events:
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=6)
        _style(
            ws.cell(row, 1, "当前没有已分类的用户事件。"),
            fill=GREY,
            border=True,
        )
        return

    for event in model.events:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _style(
            ws.cell(row, 1, f"{_status_text(event.category)} | {event.company_name}"),
            fill=BLUE,
            bold=True,
            color=WHITE,
            size=12,
            border=True,
        )
        row += 1
        labels = (
            "发生什么",
            "影响哪里",
            "现在结论",
            "是否需要重新研究",
            "下一步",
            "证据",
        )
        values = (
            event.what_happened,
            event.impact_area,
            event.current_conclusion,
            _status_text(event.research_action),
            event.next_step,
        )
        for index, (label, value) in enumerate(zip(labels, values)):
            _style(ws.cell(row, index + 1, label), fill=GREY, bold=True, border=True)
            _style(ws.cell(row + 1, index + 1, value), border=True)
        _style(ws.cell(row, 6, labels[5]), fill=GREY, bold=True, border=True)
        evidence_text = _evidence_cell(ws, row + 1, 6, event.evidence_refs, audit_rows)
        _fit_rows(
            ws,
            row + 1,
            [
                (1, values[0], 1),
                (2, values[1], 1),
                (3, values[2], 1),
                (4, values[3], 1),
                (5, values[4], 1),
                (6, evidence_text, 1),
            ],
        )
        row += 3


def _render_companies(
    ws: Worksheet,
    model: ProductWorkbenchReadModel,
    audit_rows: dict[str, int],
) -> dict[str, int]:
    _widths(ws, [24, 32, 32, 32, 32, 16])
    _title(ws, "公司卡片只展示已封存的研究、估值和股息状态。", 6)
    company_rows: dict[str, int] = {}
    row = 4
    if not model.companies:
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=6)
        _style(
            ws.cell(row, 1, "当前没有公司进入详情页。"),
            fill=GREY,
            border=True,
        )
        return company_rows

    for company in model.companies:
        company_rows[company.symbol] = row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _style(
            ws.cell(row, 1, f"{company.company_name} / {company.symbol}"),
            fill=GREEN,
            bold=True,
            color=WHITE,
            size=15,
            border=True,
        )
        row += 1
        row = _header(
            ws,
            row,
            ["研究状态", "价格", "估值", "安全边际", "股息", "证据"],
        )
        assessment_values = (
            _status_text(company.research_status),
            _assessment_text(company.price),
            _assessment_text(company.valuation),
            _assessment_text(company.margin_of_safety),
            _assessment_text(company.dividend),
        )
        for column, value in enumerate(assessment_values, 1):
            _style(ws.cell(row, column, value), border=True)
        evidence_text = _evidence_cell(ws, row, 6, company.evidence_refs, audit_rows)
        _fit_rows(
            ws,
            row,
            [
                (1, assessment_values[0], 1),
                (2, assessment_values[1], 1),
                (3, assessment_values[2], 1),
                (4, assessment_values[3], 1),
                (5, assessment_values[4], 1),
                (6, evidence_text, 1),
            ],
        )
        row += 2

        row = _section_title(ws, row, "六块研究判断", 6)
        row = _header(ws, row, ["模块", "状态", "结论", "", "", "证据"])
        ws.merge_cells(start_row=row - 1, start_column=3, end_row=row - 1, end_column=5)
        for section in company.sections:
            _style(ws.cell(row, 1, section.title), fill=GREY, bold=True, border=True)
            _style(ws.cell(row, 2, _status_text(section.status)), border=True)
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
            _style(ws.cell(row, 3, section.summary), border=True)
            evidence_text = _evidence_cell(ws, row, 6, section.evidence_refs, audit_rows)
            _fit_rows(
                ws,
                row,
                [
                    (1, section.title, 1),
                    (2, _status_text(section.status), 1),
                    (3, section.summary, 3),
                    (6, evidence_text, 1),
                ],
            )
            row += 1
        row += 1

        row = _section_title(ws, row, "逻辑变化与下一触发", 6)
        details = (
            ("最新变化", company.latest_change),
            ("下一触发", company.next_trigger),
            ("原始投资逻辑", company.original_thesis),
            ("当前逻辑是否变化", _status_text(company.thesis_change)),
        )
        for label, value in details:
            _style(ws.cell(row, 1, label), fill=GREY, bold=True, border=True)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            _style(ws.cell(row, 2, value), border=True)
            _fit_rows(ws, row, [(1, label, 1), (2, value, 5)])
            row += 1
        row += 1

        row = _section_title(ws, row, "Bear / Base / Bull", 6)
        for index, scenario in enumerate(company.scenarios):
            start_column = index * 2 + 1
            ws.merge_cells(
                start_row=row,
                start_column=start_column,
                end_row=row,
                end_column=start_column + 1,
            )
            _style(
                ws.cell(row, start_column, scenario.title),
                fill=BLUE,
                bold=True,
                color=WHITE,
                horizontal="center",
                border=True,
            )
        row += 1
        scenario_texts = [
            _assessment_text(scenario.assessment)
            for scenario in company.scenarios
        ]
        for index, scenario in enumerate(company.scenarios):
            start_column = index * 2 + 1
            ws.merge_cells(
                start_row=row,
                start_column=start_column,
                end_row=row,
                end_column=start_column + 1,
            )
            _style(ws.cell(row, start_column, scenario_texts[index]), border=True)
        _fit_rows(
            ws,
            row,
            [
                (1, scenario_texts[0], 2),
                (3, scenario_texts[1], 2),
                (5, scenario_texts[2], 2),
            ],
        )
        row += 2
    return company_rows


def _audit_rows(model: ProductWorkbenchReadModel) -> dict[str, int]:
    return {
        record.evidence_id: AUDIT_EVIDENCE_START_ROW + index
        for index, record in enumerate(model.audit_evidence)
    }


def _render_audit(ws: Worksheet, model: ProductWorkbenchReadModel) -> None:
    _widths(ws, [18, 30, 22, 54, 70, 16])
    _title(ws, "次级页面：保留阶段码、证据路径和 SHA-256，供审计追溯。", 6)
    metadata = (
        ("生成时间", model.generated_at.isoformat()),
        ("数据截止", model.as_of.isoformat()),
        ("研究动作", model.action),
        ("最终用户签收", "未通过"),
        ("系统健康说明", model.system_health.message),
    )
    row = 4
    for label, value in metadata:
        _style(ws.cell(row, 1, label), fill=GREY, bold=True, border=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        _style(ws.cell(row, 2, value), border=True)
        row += 1

    row += 1
    row = _section_title(ws, row, "后台阶段状态", 6)
    row = _header(ws, row, ["模块", "用户说明", "技术状态", "说明", "", ""])
    ws.merge_cells(start_row=row - 1, start_column=4, end_row=row - 1, end_column=6)
    for stage in model.stage_summaries:
        _style(ws.cell(row, 1, stage.stage_label), fill=GREY, bold=True, border=True)
        _style(ws.cell(row, 2, _status_text(stage.status)), border=True)
        _style(ws.cell(row, 3, stage.status.code), border=True)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        _style(ws.cell(row, 4, stage.detail), border=True)
        _fit_rows(
            ws,
            row,
            [
                (1, stage.stage_label, 1),
                (2, _status_text(stage.status), 1),
                (3, stage.status.code, 1),
                (4, stage.detail, 3),
            ],
        )
        row += 1

    row += 1
    row = _section_title(ws, row, "证据与审计索引", 6)
    row = _header(
        ws,
        row,
        ["证据ID", "名称", "类型", "路径", "SHA-256", "可用时间"],
    )
    if not model.audit_evidence:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _style(ws.cell(row, 1, "当前候选没有证据记录。"), fill=GREY, border=True)
        return
    for record in model.audit_evidence:
        values = (
            record.evidence_id,
            record.title,
            record.artifact_type,
            record.path,
            record.sha256,
            record.available_at.isoformat() if record.available_at else "-",
        )
        for column, value in enumerate(values, 1):
            _style(ws.cell(row, column, value), border=True)
        _fit_rows(
            ws,
            row,
            [
                (1, values[0], 1),
                (2, values[1], 1),
                (3, values[2], 1),
                (4, values[3], 1),
                (5, values[4], 1),
                (6, values[5], 1),
            ],
            minimum=24,
        )
        row += 1

    if model.event_audit_decisions:
        row += 1
        row = _section_title(ws, row, "事件分类审计", 6)
        row = _header(ws, row, ["事件ID", "原始状态", "产品处置", "关联事件", "结论与时点", "证据ID"])
        for decision in model.event_audit_decisions:
            related = (decision.canonical_event_id or decision.correction_of_event_id
                       or decision.duplicate_event_id or "-")
            detail = "；".join(value for value in (
                decision.previous_conclusion, decision.corrected_conclusion,
                decision.published_at, decision.observed_at,
            ) if value) or "-"
            values = (
                decision.event_id, decision.state, decision.disposition,
                related, detail, ", ".join(decision.evidence_refs) or "-",
            )
            for column, value in enumerate(values, 1):
                _style(ws.cell(row, column, value), border=True)
            _fit_rows(ws, row, [(column, value, 1) for column, value in enumerate(values, 1)])
            row += 1


def build_product_workbench_workbook(
    model: ProductWorkbenchReadModel,
) -> Workbook:
    """Render the typed read model without adding investment logic."""

    if not isinstance(model, ProductWorkbenchReadModel):
        raise TypeError("model must be ProductWorkbenchReadModel")

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {title: workbook.create_sheet(title) for title in WORKBOOK_SHEETS}
    for sheet in sheets.values():
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B4"
    sheets[SHEET_TODAY].sheet_properties.tabColor = GREEN
    sheets[SHEET_SYSTEM_AUDIT].sheet_properties.tabColor = MUTED

    audit_rows = _audit_rows(model)
    company_rows = _render_companies(
        sheets[SHEET_COMPANIES],
        model,
        audit_rows,
    )
    _render_today(
        sheets[SHEET_TODAY],
        model,
        company_rows,
        audit_rows,
    )
    _render_opportunities(
        sheets[SHEET_OPPORTUNITIES],
        model,
        company_rows,
        audit_rows,
    )
    _render_portfolio(
        sheets[SHEET_PORTFOLIO],
        model,
        audit_rows,
    )
    _render_events(
        sheets[SHEET_EVENTS],
        model,
        audit_rows,
    )
    _render_audit(sheets[SHEET_SYSTEM_AUDIT], model)
    workbook.active = 0
    return workbook


def apply_product_workbench_to_existing_workbook(
    workbook: Workbook,
    model: ProductWorkbenchReadModel,
) -> tuple[str, ...]:
    """Insert the managed M7 surface without rebuilding retained workbook sheets.

    Only the six named product sheets are ever replaced.  Every other sheet,
    including manual, historical and evidence tabs, remains in the workbook and
    retains its relative order after the new front-door navigation.
    """

    if not isinstance(model, ProductWorkbenchReadModel):
        raise TypeError("model must be ProductWorkbenchReadModel")
    for title in WORKBOOK_SHEETS:
        if title in workbook.sheetnames:
            del workbook[title]
    sheets = {
        title: workbook.create_sheet(title, index=index)
        for index, title in enumerate(WORKBOOK_SHEETS)
    }
    for sheet in sheets.values():
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B4"
    sheets[SHEET_TODAY].sheet_properties.tabColor = GREEN
    sheets[SHEET_SYSTEM_AUDIT].sheet_properties.tabColor = MUTED

    audit_rows = _audit_rows(model)
    company_rows = _render_companies(sheets[SHEET_COMPANIES], model, audit_rows)
    _render_today(sheets[SHEET_TODAY], model, company_rows, audit_rows)
    _render_opportunities(sheets[SHEET_OPPORTUNITIES], model, company_rows, audit_rows)
    _render_portfolio(sheets[SHEET_PORTFOLIO], model, audit_rows)
    _render_events(sheets[SHEET_EVENTS], model, audit_rows)
    _render_audit(sheets[SHEET_SYSTEM_AUDIT], model)
    workbook.active = 0
    return tuple(workbook.sheetnames)


def write_product_workbench_candidate(
    model: ProductWorkbenchReadModel,
    *,
    output: Path,
    root: Path,
) -> dict[str, Any]:
    """Write a candidate workbook and its no-order manifest, never a pointer."""

    output = output.resolve()
    root = root.resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Product workbench candidate output must be .xlsx")
    if not output.is_relative_to(root):
        raise ValueError("Product workbench candidate output escapes project root")
    if output.exists():
        raise ValueError(f"Product workbench candidate already exists: {output}")

    workbook = build_product_workbench_workbook(model)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook_sha256 = hashlib.sha256(output.read_bytes()).hexdigest()

    manifest_path = output.with_name(
        output.stem + ".m7-product-workbench-candidate-manifest.json"
    )
    if manifest_path.exists():
        raise ValueError(
            f"Product workbench candidate manifest already exists: {manifest_path}"
        )
    manifest = {
        "schema_version": PRODUCT_WORKBENCH_CANDIDATE_SCHEMA_VERSION,
        "workbook_kind": "M7_PRODUCT_UX_CANDIDATE",
        "generated_at": model.generated_at.isoformat(),
        "as_of": model.as_of.isoformat(),
        "action": ACTION_NO_ORDER,
        "final_user_acceptance": "NOT_PASSED",
        "candidate_not_formal_workbook": True,
        "canonical_pointer_modified": False,
        "workbook_path": str(output.relative_to(root)),
        "workbook_sha256": workbook_sha256,
        "visible_user_sheets": list(USER_SHEETS),
        "secondary_sheets": list(SECONDARY_SHEETS),
        "counts": {
            "today_items": len(model.today_items),
            "opportunities": len(model.opportunities),
            "companies": len(model.companies),
            "portfolio_positions": len(model.portfolio.positions),
            "events": len(model.events),
            "evidence_records": len(model.audit_evidence),
        },
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    receipt = dict(manifest)
    receipt["manifest_path"] = str(manifest_path.relative_to(root))
    receipt["manifest_sha256"] = hashlib.sha256(
        manifest_path.read_bytes()
    ).hexdigest()
    return receipt


__all__ = [
    "PRODUCT_WORKBENCH_CANDIDATE_SCHEMA_VERSION",
    "SECONDARY_SHEETS",
    "SHEET_COMPANIES",
    "SHEET_EVENTS",
    "SHEET_OPPORTUNITIES",
    "SHEET_PORTFOLIO",
    "SHEET_SYSTEM_AUDIT",
    "SHEET_TODAY",
    "USER_SHEETS",
    "WORKBOOK_SHEETS",
    "build_product_workbench_workbook",
    "apply_product_workbench_to_existing_workbook",
    "write_product_workbench_candidate",
]
