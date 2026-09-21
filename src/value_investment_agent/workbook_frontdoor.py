"""Navigation over research results and pending companies; no signal inputs."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

VERSION = 'frontdoor-v5'
HOME = '00_首页Dashboard'
PRIMARY = [HOME, '00_公司总览', '00_待完成公司', '21_决策验证',
           '05_仓位管理', '08_交易记录', '09_公司研究', '04_估值跟踪',
           '18_指标证据', '00_使用说明']
GREEN, INK, GREY, AMBER = '18755D', '24312D', 'EFF3F1', 'FFF1D6'


def _rows(sheet):
    return {str(row[0]).zfill(6): (i, row) for i, row in enumerate(
        sheet.iter_rows(min_row=4, values_only=True), 4)
        if row[0] is not None and str(row[0]).zfill(6).isdigit()
        and len(str(row[0]).zfill(6)) == 6}


def _band(sheet, row, text, *, color=GREEN, size=12, height=28):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = sheet.cell(row, 1, text)
    cell.fill = PatternFill('solid', fgColor=color)
    cell.font = Font(name='Microsoft YaHei', size=size, bold=True, color='FFFFFF')
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[row].height = height


def _link(cell, text, sheet, row=1, column='A'):
    cell.value = text
    # A relationship target is an external file in WPS, even if it starts with #.
    cell.hyperlink = Hyperlink(ref=cell.coordinate,
        location=f"'{sheet.replace(chr(39), chr(39) * 2)}'!{column}{row}", display=str(text))
    cell.font = Font(name='Microsoft YaHei', size=11, color=GREEN, underline='single')


def repair_internal_links(wb):
    repaired = 0
    for ws in wb:
        for row in ws:
            for cell in row:
                link = cell.hyperlink
                if link and link.target and link.target.startswith('#'):
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=link.target[1:],
                        display=str(cell.value), tooltip=link.tooltip)
                    repaired += 1
    return repaired


def apply_frontdoor(wb, *, root=None):
    from .workbook_simple_overview import apply_simple_overview
    return apply_simple_overview(wb, root=root)
