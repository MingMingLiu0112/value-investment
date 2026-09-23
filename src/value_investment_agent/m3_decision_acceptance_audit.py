"""Recompute machine evidence for the non-personal M3 decision-card candidate.

This auditor never creates an Entry, Journal, Consistency result, private
portfolio snapshot or order.  It verifies the frozen M1 input, deterministic
negative-only cards, the standalone workbook and the WPS read-only receipt,
then leaves actual human comprehension as an explicit pending review.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

from .decision_read_model import (
    CARD_ENTRY_NOT_REQUIRED,
    CARD_PORTFOLIO_MISSING,
    CARD_REASON_RESEARCH_INCOMPLETE,
    DecisionCardCollection,
)
from .investment_decision import ACTION_NO_ORDER, STATUS_INSUFFICIENT_RESEARCH
from .m1_sample_preregistration import load_m1_sample_preregistration
from .m3_decision_application import (
    build_nonpersonal_decision_card_collection,
    load_integrated_runs,
)
from .research_artifacts import canonicalize_artifact_payload, sha256_text


SCHEMA_VERSION = "m3-decision-acceptance-audit-v1"
RULE_VERSION = "m3-decision-acceptance-v1"

DONE = "DONE"
PARTIAL = "PARTIAL"
PENDING_CI = "PENDING_CI"
PENDING_HUMAN_REVIEW = "PENDING_HUMAN_REVIEW"
FAILED = "FAILED"

EXPECTED_INPUT_SHA256 = (
    "b1123333f2b4caa6beae16102bdca613b0894ad7fb72aa17329e838cd32b0459"
)
EXPECTED_PREREGISTRATION_SHA256 = (
    "5b7df98f8781080f66e5e053b8d0015f0b7c7eee9e27e8a10f7b0fe7990ab5b1"
)
EXPECTED_CANDIDATE_SHA256 = (
    "589f19ef9e3d235401814e98450475d657c3e981b33637337ab5da9d33fb307d"
)
EXPECTED_MANIFEST_SHA256 = (
    "c2b69ea02a0b5bdb0734b41c416ee03147ffad0eea8c78febc37ebf6b0ea2cc6"
)
EXPECTED_WPS_RECEIPT_SHA256 = (
    "345ce9a2255f871eacbb80a670ad14b0bd38fe6eb33894c74c4adc2b102436f4"
)
EXPECTED_CANONICAL_WORKBOOK_SHA256 = (
    "64c8deff1a237076d2ba0b00afc8905d23bd9d117cb132dfc6757071b5659911"
)

_SHA256 = re.compile(r"^[0-9a-f]{64}$")
_FORBIDDEN_TEXT = (
    "买入",
    "加仓",
    "减仓",
    "目标仓位",
    "BUY",
    "ADD",
)
_EXPECTED_SHEETS = (
    "00_决策卡",
    "01_缺失与阻断",
    "02_来源哈希",
    "03_证据引用",
)
_EXPECTED_CARDS = (
    {
        "symbol": "000651",
        "name": "格力电器",
        "status": STATUS_INSUFFICIENT_RESEARCH,
        "reason_kind": CARD_REASON_RESEARCH_INCOMPLETE,
        "portfolio_status": CARD_PORTFOLIO_MISSING,
        "entry_status": CARD_ENTRY_NOT_REQUIRED,
    },
    {
        "symbol": "600741",
        "name": "华域汽车",
        "status": STATUS_INSUFFICIENT_RESEARCH,
        "reason_kind": CARD_REASON_RESEARCH_INCOMPLETE,
        "portfolio_status": CARD_PORTFOLIO_MISSING,
        "entry_status": CARD_ENTRY_NOT_REQUIRED,
    },
    {
        "symbol": "600887",
        "name": "伊利股份",
        "status": STATUS_INSUFFICIENT_RESEARCH,
        "reason_kind": CARD_REASON_RESEARCH_INCOMPLETE,
        "portfolio_status": CARD_PORTFOLIO_MISSING,
        "entry_status": CARD_ENTRY_NOT_REQUIRED,
    },
)
_M3_TEST_FILES = (
    "tests/test_investment_decision.py",
    "tests/test_decision_read_model.py",
    "tests/test_m3_decision_application.py",
    "tests/test_m3_decision_card_workbook.py",
    "tests/test_m3_decision_acceptance_audit.py",
)


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _require_sha256(value: str, field: str) -> str:
    normalized = str(value).lower()
    if not _SHA256.fullmatch(normalized):
        raise ValueError(f"{field} must be a SHA-256 hex digest")
    return normalized


def _inside(root: Path, value: Path, field: str) -> Path:
    root = root.resolve()
    target = value.resolve()
    if not target.is_relative_to(root):
        raise ValueError(f"{field} escapes the audit root")
    return target


def _load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Expected a JSON object: {path}")
    return payload


def _verify_pinned_bytes(path: Path, expected: str, label: str) -> Path:
    expected = _require_sha256(expected, label)
    if not path.is_file():
        raise ValueError(f"{label} is not a file: {path}")
    actual = _digest(path)
    if actual != expected:
        raise ValueError(f"{label} changed: expected {expected}, got {actual}")
    return path


@dataclass(frozen=True)
class M3DecisionAcceptanceSpec:
    """Immutable evidence locations and hashes for one M3 audit."""

    input_path: Path
    input_sha256: str
    preregistration_path: Path
    preregistration_sha256: str
    candidate_path: Path
    candidate_sha256: str
    manifest_path: Path
    manifest_sha256: str
    generated_at: datetime
    wps_candidate_path: Path | None = None
    wps_receipt_path: Path | None = None
    wps_receipt_sha256: str | None = None
    canonical_workbook_path: Path | None = None
    canonical_sha256: str | None = None

    def __post_init__(self) -> None:
        object.__setattr__(
            self,
            "input_sha256",
            _require_sha256(self.input_sha256, "input_sha256"),
        )
        object.__setattr__(
            self,
            "preregistration_sha256",
            _require_sha256(self.preregistration_sha256, "preregistration_sha256"),
        )
        object.__setattr__(
            self,
            "candidate_sha256",
            _require_sha256(self.candidate_sha256, "candidate_sha256"),
        )
        object.__setattr__(
            self,
            "manifest_sha256",
            _require_sha256(self.manifest_sha256, "manifest_sha256"),
        )
        if self.generated_at.utcoffset() is None:
            raise ValueError("generated_at must be timezone-aware")
        if self.wps_receipt_sha256 is not None:
            object.__setattr__(
                self,
                "wps_receipt_sha256",
                _require_sha256(self.wps_receipt_sha256, "wps_receipt_sha256"),
            )
        if self.canonical_sha256 is not None:
            object.__setattr__(
                self,
                "canonical_sha256",
                _require_sha256(self.canonical_sha256, "canonical_sha256"),
            )


def _run_pytest(root: Path) -> dict[str, Any]:
    test_paths = [str(root / item) for item in _M3_TEST_FILES]
    result = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            "-m",
            "pytest",
            "-q",
            "--basetemp=runtime/pytest-tmp-m3-acceptance",
            *test_paths,
        ],
        cwd=root,
        text=True,
        encoding="utf-8",
        capture_output=True,
    )
    passed_match = re.search(r"(\d+) passed", result.stdout or "")
    skipped_match = re.search(r"(\d+) skipped", result.stdout or "")
    return {
        "passed": result.returncode == 0,
        "passed_count": int(passed_match.group(1)) if passed_match else 0,
        "skipped_count": int(skipped_match.group(1)) if skipped_match else 0,
        "returncode": result.returncode,
        "output_tail": "\n".join((result.stdout or "").splitlines()[-8:]),
    }


def _criterion(
    status: str,
    checks: Sequence[Mapping[str, Any]],
    *,
    evidence: Mapping[str, Any] | None = None,
    blockers: Sequence[str] = (),
    human_review: Sequence[str] = (),
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "status": status,
        "checks": [dict(check) for check in checks],
        "blockers": list(blockers),
    }
    if evidence is not None:
        result["evidence"] = dict(evidence)
    if human_review:
        result["human_review"] = list(human_review)
    return result


def _check(label: str, passed: bool, detail: str = "") -> dict[str, Any]:
    return {"label": label, "passed": passed, "detail": detail}


def _criterion_status(checks: Sequence[Mapping[str, Any]]) -> str:
    return DONE if all(check["passed"] for check in checks) else PARTIAL


def _m3c1(
    root: Path,
    spec: M3DecisionAcceptanceSpec,
    payload: Sequence[Mapping[str, Any]],
    input_receipt: Mapping[str, Any],
) -> dict[str, Any]:
    checks = [
        _check(
            "frozen integrated-runs hash matches",
            input_receipt.get("sha256") == spec.input_sha256,
        ),
        _check(
            "frozen M1 preregistration hash matches",
            _digest(_inside(root, spec.preregistration_path, "preregistration_path"))
            == spec.preregistration_sha256,
        ),
        _check(
            "integrated runs are a non-empty JSON array",
            bool(payload) and all(isinstance(item, Mapping) for item in payload),
        ),
        _check(
            "all integrated runs remain no_order",
            all(item.get("action") == ACTION_NO_ORDER for item in payload),
        ),
    ]
    return _criterion(
        _criterion_status(checks),
        checks,
        evidence={
            "input_sha256": input_receipt.get("sha256"),
            "run_count": len(payload),
        },
    )


def _m3c2(collection: DecisionCardCollection, names: Mapping[str, str]) -> dict[str, Any]:
    cards = collection.by_symbol
    shape_checks = []
    for expected in _EXPECTED_CARDS:
        symbol = expected["symbol"]
        card = cards.get(symbol)
        shape_checks.append(_check(f"{symbol} card exists", card is not None))
        if card is None:
            continue
        shape_checks.append(
            _check(
                f"{symbol} status is {expected['status']}",
                card.status == expected["status"],
            )
        )
        shape_checks.append(
            _check(
                f"{symbol} reason kind is {expected['reason_kind']}",
                card.reason_kind == expected["reason_kind"],
            )
        )
        shape_checks.append(
            _check(f"{symbol} has no decision intent", card.decision_intent is None)
        )
        shape_checks.append(
            _check(
                f"{symbol} portfolio input is missing",
                card.portfolio_status == expected["portfolio_status"],
            )
        )
        shape_checks.append(
            _check(
                f"{symbol} entry is not required for this negative card",
                card.entry_status == expected["entry_status"],
            )
        )
        shape_checks.append(
            _check(
                f"{symbol} name matches preregistration",
                names.get(symbol) == expected["name"],
            )
        )
    checks = [
        _check("card count is exactly three", len(collection.cards) == 3),
        _check("input has no hidden failures", not collection.input_failures),
        _check(
            "every card requires human review",
            all(card.requires_human_review for card in collection.cards),
        ),
        _check(
            "every card remains action=no_order",
            all(card.action == ACTION_NO_ORDER for card in collection.cards),
        ),
        _check(
            "no positive decision card is emitted",
            not any(card.is_positive_review() for card in collection.cards),
        ),
        *shape_checks,
    ]
    return _criterion(
        _criterion_status(checks),
        checks,
        evidence={
            "symbols": [card.symbol for card in collection.cards],
            "positive_review_count": sum(
                card.is_positive_review() for card in collection.cards
            ),
        },
    )


def _m3c3(
    payload: Sequence[Mapping[str, Any]],
    collection: DecisionCardCollection,
    spec: M3DecisionAcceptanceSpec,
) -> dict[str, Any]:
    replay = build_nonpersonal_decision_card_collection(
        payload,
        generated_at=spec.generated_at,
        source_run_id=spec.input_path.resolve().parent.name,
    )
    raw_by_symbol = {
        str(item.get("symbol")): item for item in payload if isinstance(item, Mapping)
    }
    binding_checks = []
    for card in collection.cards:
        raw = raw_by_symbol.get(card.symbol)
        source_hash = next(
            (
                source.sha256
                for source in card.source_hashes
                if source.source_key == "m1_integrated_run"
            ),
            None,
        )
        expected = (
            sha256_text(canonicalize_artifact_payload(raw)) if raw is not None else None
        )
        binding_checks.append(
            _check(
                f"{card.symbol} binds its raw integrated run",
                source_hash is not None and source_hash == expected,
            )
        )
    checks = [
        _check(
            "replay is deterministic under the pinned clock",
            collection.to_json() == replay.to_json(),
        ),
        *binding_checks,
    ]
    return _criterion(
        _criterion_status(checks),
        checks,
        evidence={"rule_version": collection.cards[0].rule_version if collection.cards else None},
    )


def _m3c4(
    root: Path,
    spec: M3DecisionAcceptanceSpec,
    collection: DecisionCardCollection,
) -> dict[str, Any]:
    manifest = _load_json(_inside(root, spec.manifest_path, "manifest_path"))
    workbook_path = _inside(root, spec.candidate_path, "candidate_path")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    text = "\n".join(
        str(cell.value)
        for sheet in workbook
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    forbidden = [item for item in _FORBIDDEN_TEXT if item in text]
    checks = [
        _check(
            "candidate workbook hash is pinned",
            _digest(workbook_path) == spec.candidate_sha256,
        ),
        _check(
            "candidate manifest hash is pinned",
            _digest(_inside(root, spec.manifest_path, "manifest_path"))
            == spec.manifest_sha256,
        ),
        _check(
            "manifest points to the pinned workbook",
            manifest.get("workbook_sha256") == spec.candidate_sha256,
        ),
        _check("manifest contains three cards", manifest.get("card_count") == len(collection.cards)),
        _check("manifest reports no positive review", manifest.get("positive_review_count") == 0),
        _check("workbook has the expected four sheets", tuple(workbook.sheetnames) == _EXPECTED_SHEETS),
        _check("workbook has no forbidden order or position text", not forbidden),
    ]
    workbook.close()
    return _criterion(
        _criterion_status(checks),
        checks,
        evidence={
            "workbook_sha256": _digest(workbook_path),
            "manifest_sha256": _digest(_inside(root, spec.manifest_path, "manifest_path")),
            "forbidden_text_found": forbidden,
        },
    )


def _m3c5(
    root: Path,
    spec: M3DecisionAcceptanceSpec,
) -> dict[str, Any]:
    checks = []
    evidence: dict[str, Any] = {}
    if spec.wps_candidate_path is not None:
        wps = spec.wps_candidate_path.resolve()
        digest = _digest(wps)
        checks.append(
            _check(
                "WPS candidate is byte-identical to the repository candidate",
                digest == spec.candidate_sha256,
            )
        )
        evidence["wps_candidate_sha256"] = digest
    else:
        checks.append(_check("WPS candidate copy is provided", False))
    if spec.wps_receipt_path is not None and spec.wps_receipt_sha256 is not None:
        receipt_path = _inside(root, spec.wps_receipt_path, "wps_receipt_path")
        receipt = _load_json(receipt_path)
        checks.append(
            _check(
                "WPS receipt hash is pinned",
                _digest(receipt_path) == spec.wps_receipt_sha256,
            )
        )
        checks.append(_check("WPS receipt status is passed", receipt.get("status") == "passed"))
        checks.append(
            _check(
                "WPS receipt binds the candidate hash",
                str(receipt.get("sha256") or "").lower() == spec.candidate_sha256,
            )
        )
        evidence["wps_receipt_sha256"] = _digest(receipt_path)
    else:
        checks.append(_check("WPS read-only receipt is provided", False))
    if spec.canonical_workbook_path is not None and spec.canonical_sha256 is not None:
        canonical = spec.canonical_workbook_path.resolve()
        digest = _digest(canonical)
        checks.append(
            _check(
                "original production workbook is unchanged",
                digest == spec.canonical_sha256,
            )
        )
        evidence["canonical_sha256"] = digest
    else:
        checks.append(_check("original production workbook path is provided", False))
    return _criterion(_criterion_status(checks), checks, evidence=evidence)


def _m3c6(
    test_evidence: Mapping[str, Any] | None,
    ci_evidence: Mapping[str, Any] | None,
) -> dict[str, Any]:
    test_passed = bool(test_evidence and test_evidence.get("passed"))
    ci_passed = bool(ci_evidence and ci_evidence.get("status") == "success")
    checks = [
        _check(
            "offline M3 regression suite passed",
            test_passed,
            f"passed={test_evidence.get('passed_count') if test_evidence else None}",
        ),
        _check(
            "CI success is observed for the audited commit",
            ci_passed,
            str(ci_evidence.get("status") if ci_evidence else None),
        ),
    ]
    status = DONE if all(check["passed"] for check in checks) else PARTIAL
    if not test_passed or not ci_passed:
        status = PENDING_CI
    return _criterion(
        status,
        checks,
        evidence={
            "test_result": dict(test_evidence or {}),
            "ci": dict(ci_evidence or {}),
        },
    )


def _m3c7() -> dict[str, Any]:
    checks = [
        _check("no Entry, Journal or Consistency record was invented", True),
        _check("no personal portfolio input was used", True),
        _check("all cards remain no_order", True),
    ]
    return _criterion(
        PENDING_HUMAN_REVIEW,
        checks,
        human_review=(
            "在 WPS 中打开独立 M3 决策卡候选，确认三家公司均为研究不足、无个人化正向复核。",
            "确认能逐卡说明缺失的研究、组合输入或原始 Entry 为何阻断正向决策。",
            "抽查 02_来源哈希 与 03_证据引用 中至少一条链接和 Hash。",
            "确认没有买入、加仓、减仓、目标仓位或订单列。",
            "用户实际阅读至少三张卡并能复述理由与反证后，才可记录 Checkpoint B 的人工验收。",
        ),
    )


def audit(
    root: Path,
    spec: M3DecisionAcceptanceSpec,
    *,
    ci_evidence: Mapping[str, Any] | None = None,
    test_evidence: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    root = root.resolve()
    input_path = _verify_pinned_bytes(
        _inside(root, spec.input_path, "input_path"),
        spec.input_sha256,
        "frozen integrated-runs input",
    )
    _verify_pinned_bytes(
        _inside(root, spec.preregistration_path, "preregistration_path"),
        spec.preregistration_sha256,
        "M1 preregistration",
    )
    _verify_pinned_bytes(
        _inside(root, spec.candidate_path, "candidate_path"),
        spec.candidate_sha256,
        "M3 candidate workbook",
    )
    _verify_pinned_bytes(
        _inside(root, spec.manifest_path, "manifest_path"),
        spec.manifest_sha256,
        "M3 candidate manifest",
    )
    payload, input_receipt = load_integrated_runs(root, input_path)
    names = {
        entry.symbol: entry.name
        for entry in load_m1_sample_preregistration(
            _inside(root, spec.preregistration_path, "preregistration_path")
        ).companies
    }
    collection = build_nonpersonal_decision_card_collection(
        payload,
        generated_at=spec.generated_at,
        source_run_id=spec.input_path.resolve().parent.name,
    )
    criteria = {
        "m3c1_frozen_input_and_identity": _m3c1(root, spec, payload, input_receipt),
        "m3c2_nonpersonal_negative_card_semantics": _m3c2(collection, names),
        "m3c3_source_hash_and_deterministic_replay": _m3c3(
            payload,
            collection,
            spec,
        ),
        "m3c4_candidate_workbook_integrity": _m3c4(root, spec, collection),
        "m3c5_wps_and_canonical_workbook_boundary": _m3c5(root, spec),
        "m3c6_offline_regression_and_ci": _m3c6(test_evidence, ci_evidence),
        "m3c7_human_entry_journal_and_comprehension": _m3c7(),
    }
    partial = [
        key
        for key, item in criteria.items()
        if item["status"] in {PARTIAL, FAILED}
    ]
    pending_ci = [
        key for key, item in criteria.items() if item["status"] == PENDING_CI
    ]
    if partial:
        status = PARTIAL
    elif pending_ci:
        status = PENDING_CI
    else:
        status = PENDING_HUMAN_REVIEW
    return {
        "schema_version": SCHEMA_VERSION,
        "rule_version": RULE_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "action": ACTION_NO_ORDER,
        "status": status,
        "criteria": criteria,
        "summary": {
            "done": [
                key for key, item in criteria.items() if item["status"] == DONE
            ],
            "partial": partial,
            "pending_ci": pending_ci,
            "pending_human_review": [
                key
                for key, item in criteria.items()
                if item["status"] == PENDING_HUMAN_REVIEW
            ],
            "human_review_items": [
                item
                for criterion in criteria.values()
                for item in criterion.get("human_review") or []
            ],
            "next_action": "用户复核三张 M3 负向决策卡并记录 Checkpoint B 人工验收",
        },
        "blockers": [
            blocker
            for criterion in criteria.values()
            for blocker in criterion.get("blockers") or []
        ],
    }


def write_receipt(receipt: Mapping[str, Any], *, root: Path) -> dict[str, Any]:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    target = root / "runtime" / f"m3-decision-acceptance-audit-{timestamp}"
    target.mkdir(parents=True, exist_ok=False)
    evidence = target / "receipt.json"
    evidence.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    digest = _digest(evidence)
    pointer = {
        "path": str(target.relative_to(root)),
        "sha256": digest,
    }
    (root / "runtime/m3-decision-acceptance-audit-latest.json").write_text(
        json.dumps(pointer, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return {
        "receipt_path": str(evidence.relative_to(root)),
        "pointer_path": "runtime/m3-decision-acceptance-audit-latest.json",
        "receipt_sha256": digest,
    }
