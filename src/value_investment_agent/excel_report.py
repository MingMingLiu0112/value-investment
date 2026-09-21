"""Compact, repeatable Excel research views; never invent portfolio transactions."""
from __future__ import annotations

from collections import Counter, defaultdict
import gc
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from .workbook_compat import load_wps_workbook as load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .financial_institutions import provisional_financial_type
from .reminders import build_reminders
from .broker_regulatory_reference import reference, RULE_URL, RULE_SHA256
from .company_research_export import (load_historical_conditional_replay_status,
                                      load_cash_anchor_paper_contract_status,
                                      load_pe_mid_paper_contract_status,
                                      load_blocked_paper_ledger_status,
                                      load_simulation_closure_status,
                                      load_experimental_signal_coverage_status,
                                      load_historical_range_experiment_status,
                                      load_current_execution_contract_status,
                                      load_current_valuation_admission_status,
                                      load_current_conditional_observation,
                                      load_finance_cost_scope_sensitivity_status)

LABELS = {
    'borrowings_bonds_subtotal': '借款债券四项小计（非完整有息负债）',
    'provider_sales_net_margin': '接口销售净利率（利润口径待核）',
    'main_business_revenue_yoy': '主营业务收入增长率（非营业收入同比）',
    'provider_net_income_yoy': '接口净利润增长率（利润口径待核）',
    'provider_cashflow_profit_ratio': '接口现金利润比（单位口径待核）',
    'roe_simple': '普通净资产收益率（非加权评分口径）',
    'long_term_payables_noncurrent': '长期应付款总额（含专项，非纯融资负债）',
    'long_term_payables_excluding_special': '长期应付款（不含专项）',
    'special_payables_noncurrent': '专项应付款',
    'revenue': '营业收入', 'net_income': '归母净利润', 'roe': 'ROE（见来源口径）',
    'gross_margin': '毛利率', 'net_margin': '归母净利润/营业收入', 'revenue_yoy': '收入同比',
    'net_income_yoy': '利润同比', 'operating_cash_flow': '经营现金净流入',
    'free_cash_flow': '自由现金流', 'operating_cash_flow_to_net_income': '合并经营现金流/归母净利润',
    'cash': '货币资金', 'interest_bearing_debt': '有息负债', 'debt_ratio': '资产负债率',
    'lease_liabilities_noncurrent': '非流动租赁负债',
    'eps_ttm': '每股收益TTM', 'eps_annual': '年度每股收益', 'eps_reported': '当期每股收益',
    'bvps': '每股净资产', 'dps_ttm': '每股分红TTM', 'payout_ratio': '分红支付率',
    'npl_ratio': '不良贷款率', 'provision_coverage': '拨备覆盖率',
    'cet1_ratio': '核心一级资本充足率', 'net_interest_margin': '净息差',
    'risk_coverage': '风险覆盖率', 'capital_leverage': '资本杠杆率',
    'liquidity_coverage': '流动性覆盖率', 'net_stable_funding': '净稳定资金率',
    'core_solvency': '核心偿付能力充足率', 'comprehensive_solvency': '综合偿付能力充足率',
    'segment_roe': '分部ROE', 'segment_equity': '分部净资产',
}
MODELS = {'bank': '银行', 'broker': '券商', 'insurer': '保险', 'financial_group': '金融控股/多元金融', 'general_enterprise': '非金融企业'}
SPECIAL_FIELDS = {
    'bank': ('npl_ratio', 'provision_coverage', 'cet1_ratio', 'net_interest_margin'),
    'broker': ('risk_coverage', 'capital_leverage', 'liquidity_coverage', 'net_stable_funding'),
    'insurer': ('core_solvency', 'comprehensive_solvency'),
    'financial_group': ('segment_roe', 'segment_equity'),
}
GENERAL = ('roe', 'gross_margin', 'net_margin', 'operating_cash_flow_to_net_income',
           'cash', 'interest_bearing_debt', 'debt_ratio', 'revenue_yoy', 'net_income_yoy')
MONEY = {'CNY': Decimal('0.00000001'), 'CNY 100M': Decimal(1)}
GREEN, AMBER, RED = 'E7F3EB', 'FFF3CF', 'FBE3E3'


def number(value):
    try:
        result = Decimal(str(value))
        return float(result) if result.is_finite() else None
    except (InvalidOperation, TypeError, ValueError):
        return None


def symbol(value):
    return str(value or '').zfill(6)


def accepted(point):
    from .financial_quality import _accepted
    return _accepted(point)


def point_status(point):
    metadata = point.get('metadata') or {}
    if metadata.get('evidence_quarantine'):
        return '证据已隔离，不参与决策'
    if (point.get('field_name') == 'interest_bearing_debt' and
            metadata.get('derivation_formula') ==
            'short_term_borrowings + current_portion_long_term_debt + long_term_borrowings + bonds_payable'):
        return '历史四项债务代理，完整口径未验证'
    if (point.get('field_name') == 'interest_bearing_debt'
            and 'complete_debt_verified' in metadata
            and metadata['complete_debt_verified'] is not True):
        return '完整债务口径未验证，不参与决策'
    if accepted(point):
        return '已交叉验证'
    if point.get('validation_status') in {'conflict', 'rejected', 'failed'}:
        return '来源冲突'
    if (point.get('metadata') or {}).get('official_extraction'):
        return '官方原文提取，待交叉验证'
    return '补充来源，待交叉验证'


def choose_points(points):
    selected = {}
    for p in points:
        key = (symbol(p['symbol']), p['field_name'])
        rank = (str(p.get('period_label', '')), accepted(p), str(p.get('created_at', '')))
        if key not in selected or rank > selected[key][0]:
            selected[key] = (rank, p)
    return {key: item[1] for key, item in selected.items()}


def _safe(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)[:32767] if isinstance(value, str) else value


def _cell(ws, row, column, value):
    cell = ws.cell(row, column, _safe(value))
    if isinstance(value, str):
        cell.data_type = 's'  # Source text must not become an Excel formula.
    return cell


def _sheet(wb, name, headers, description, rows):
    index = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
    if name in wb:
        del wb[name]
    ws = wb.create_sheet(name, index)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = 'C4'
    for r, text in [(1, name.split('_', 1)[-1]), (2, description)]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        cell = _cell(ws, r, 1, text)
        cell.font = Font(name='Microsoft YaHei', size=18 if r == 1 else 10,
                         bold=r == 1, color='183D34' if r == 1 else '555555')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[r].height = 32 if r == 1 else 38
    for j, text in enumerate(headers, 1):
        cell = _cell(ws, 3, j, text)
        cell.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill('solid', fgColor='245344')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        width = 13 if j == 1 else 17
        if any(word in text for word in ('缺口', '状态', '原因', '行动', '说明', '证据')):
            width = 32
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[3].height = 34
    for i, row in enumerate(rows, 4):
        for j, value in enumerate(row, 1):
            cell = _cell(ws, i, j, value)
            cell.font = Font(name='Microsoft YaHei', size=10, color='222222')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if i % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F2F6F4')
            if isinstance(value, (float, int)):
                cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
        ws.row_dimensions[i].height = 34
    ws.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:3'
    return ws


def read_holdings(worksheet):
    """Read explicit share counts; duplicate rows cannot establish a position."""
    holdings = {}
    seen = set()
    for row in worksheet.iter_rows(min_row=10, values_only=True):
        if not row or row[0] is None:
            continue
        code = symbol(row[0])
        if not re.fullmatch(r'\d{6}', code):
            continue
        if code in seen:
            holdings[code] = None
            continue
        seen.add(code)
        value = row[4] if len(row) > 4 else None
        quantity = None
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            parsed = Decimal(str(value))
            if parsed.is_finite() and parsed >= 0 and parsed == parsed.to_integral_value():
                quantity = parsed
        holdings[code] = quantity
    return holdings


def build_report(template: Path, payload: dict, output: Path, debt_research: dict | None = None,
                 company_research: dict | None = None) -> dict:
    from .candidate_tracking import tracking_candidate_view
    payload = dict(payload, market_candidates=tracking_candidate_view(payload))
    debt_research = debt_research or {'companies': {}, 'evidence': {}}
    candidates = payload.get('market_candidates')
    if not candidates or len({symbol(c['symbol']) for c in candidates}) != len(candidates):
        raise ValueError('Missing or duplicate screened company list; workbook unchanged')
    if any(not re.fullmatch(r'\d{6}', symbol(c['symbol'])) or not c.get('source_id') for c in candidates):
        raise ValueError('Invalid company code or missing market evidence ID')
    expected = payload.get('market_candidate_count')
    if expected is not None and len(candidates) != int(expected):
        raise ValueError('Export candidate count mismatch')
    wb = load_workbook(template)
    for source, archive in [('02_质量评分','历史_初始评分'),('04_估值跟踪','历史_估值假设'),('10_年报跟踪','历史_报告备注')]:
        if source in wb and archive not in wb:
            saved = wb.copy_worksheet(wb[source])
            saved.title = archive
            saved.sheet_state = 'hidden'
    points = choose_points(payload.get('points', []))
    annual_points = choose_points(payload.get('annual_points', []))
    quality = {symbol(q['symbol']): q for q in payload.get('financial_quality', [])}
    valuations = {symbol(v['symbol']): v for v in payload.get('valuations', [])}
    disclosures = {symbol(d['symbol']): d for d in payload.get('disclosures', [])}
    company = {symbol(c['symbol']): c for c in candidates}
    generated = str(payload['generated_at'])
    now = datetime.now(timezone.utc)
    holdings = read_holdings(wb['05_仓位管理'])
    alerts=build_reminders(payload,holdings,now)
    alerts_by_symbol={symbol(a['symbol']):a for a in alerts}
    export_stale = (now - datetime.fromisoformat(generated)).total_seconds() > 36 * 3600
    market_date = min(str(c.get('screen_date', ''))[:10] for c in candidates)
    try:
        quote_stale = (now.date() - datetime.fromisoformat(market_date).date()).days > 3
    except ValueError:
        quote_stale = True
    if 'candidate_tracking' in payload:
        try:
            quote_times = [datetime.fromisoformat(str(c.get('quote_as_of')).replace('Z', '+00:00')) for c in candidates]
            quote_stale = any(t.tzinfo is None or not -300 <= (now-t).total_seconds() <= 30*3600 for t in quote_times)
        except (ValueError, TypeError):
            quote_stale = True
    stale = export_stale or quote_stale
    snapshots = {s: [p for (code, _), p in points.items() if code == s] for s in company}
    all_special = [f for fields in SPECIAL_FIELDS.values() for f in fields]

    def get(s, f, period=None):
        p = points.get((s, f))
        return p if p and (period is None or str(p['period_label']) == period) else None

    def value(s, f, period=None, money=False):
        p = get(s, f, period)
        if not p:
            return None
        if p.get('validation_status') in {'failed', 'conflict'}:
            return None
        if (p.get('metadata') or {}).get('evidence_quarantine'):
            return None
        if (f == 'operating_cash_flow_to_net_income'
                and p.get('source_name') == 'AkShare / Sina financial indicators'):
            # Legacy provider ratios lack a proven scale and consolidated profit scope.
            return None
        v = number(p.get('value'))
        if money:
            return v * float(MONEY[p['unit']]) if v is not None and p.get('unit') in MONEY else None
        return v

    def model(c):
        return provisional_financial_type(c.get('name'), c.get('sector')) or 'general_enterprise'

    def period_for(s):
        periods = [str(p['period_label']) for p in snapshots[s]
                   if p['field_name'] in (*GENERAL, *all_special, 'revenue', 'net_income', 'bvps')]
        return max(periods, default=str(disclosures.get(s, {}).get('report_period', '')))

    def gap(s, fields, period):
        missing = [LABELS.get(f, f) for f in fields if not get(s, f, period)]
        unverified = [LABELS.get(f, f) for f in fields if get(s, f, period) and not accepted(get(s, f, period))]
        return '；'.join(filter(None, ['缺值：' + '、'.join(missing) if missing else '',
                                       '待比对：' + '、'.join(unverified) if unverified else ''])) or '本组指标已交叉验证'

    def annotate(ws, row, col, s, f, period=None):
        p = get(s, f, period)
        cell = ws.cell(row, col)
        if not p or cell.value is None:
            return
        meta = p.get('metadata') or {}
        cell.fill = PatternFill('solid', fgColor=GREEN if accepted(p) else AMBER)
        cell.comment = Comment('\n'.join([LABELS.get(f, f), point_status(p),
            f"报告期：{p.get('period_label')}；原单位：{p.get('unit')}",
            f"source_id：{p.get('source_id')}", f"页码：{meta.get('page_number', meta.get('page', ''))}",
            str(p.get('source_url', '')), f"SHA-256：{p.get('sha256', '')}",
            str(meta.get('candidate_excerpt', ''))]), '数据证据')

    # Preserve manual status/notes across refreshes. Historical positions are not candidates.
    old_watch = {}
    if '01_观察名单' in wb:
        for row in wb['01_观察名单'].iter_rows(min_row=4, values_only=True):
            if row[0]:
                old_watch[symbol(row[0])] = row
    rows, qrows, frows, vrows, srows = [], [], [], [], []
    for c in candidates:
        s = symbol(c['symbol']); q = quality.get(s, {}); v = valuations.get(s, {})
        m = model(c); period = period_for(s); fields = SPECIAL_FIELDS.get(m, GENERAL)
        verified = sum(accepted(get(s, f, period)) for f in fields)
        available = sum(get(s, f, period) is not None for f in fields)
        missing = gap(s, fields, period)
        qperiod = max((str(p['period_label']) for (code, f), p in annual_points.items()
                      if code == s and f in fields), default=period) if m == 'general_enterprise' else period
        qpoints = annual_points if m == 'general_enterprise' and any(code == s for code, _ in annual_points) else points
        qcurrent = {f: qpoints.get((s, f)) for f in fields}
        qcurrent = {f: p for f, p in qcurrent.items() if p and str(p['period_label']) == qperiod}
        qverified = sum(accepted(p) for p in qcurrent.values())
        qmissing = '；'.join(filter(None, [
            '缺值：' + '、'.join(LABELS.get(f, f) for f in fields if f not in qcurrent) if len(qcurrent) < len(fields) else '',
            '待比对：' + '、'.join(LABELS.get(f, f) for f, p in qcurrent.items() if not accepted(p)) if qverified < len(qcurrent) else '',
        ])) or '本组指标已交叉验证'
        score = number(q.get('total_score')) if (q.get('quality_status') == '已验证' and m == 'general_enterprise'
            and (q.get('calculation_details') or {}).get('period') == qperiod and qverified == len(fields)) else None
        signal = v.get('build_signal', '待数据')
        if stale or verified != len(fields) or score is None or alerts_by_symbol[s]['signal_blocked']:
            signal = '仅研究，不生成建仓信号'
        old = old_watch.get(s, ())
        manual_status = old[4] if len(old) > 4 else '初筛候选'
        rows.append([s, c['name'], c.get('sector'), MODELS[m], manual_status,
            number(c.get('current_price')), number(v.get('fair_value')),
            number(v.get('safety_margin')), v.get('valuation_status', '尚无合理价'), score,
            signal, number(v.get('target_weight')) if signal == '建仓候选' else None,
            old[12] if len(old) > 12 else None, missing, old[14] if len(old) > 14 else None,
            f'同报告期已验证 {verified}/{len(fields)}；已取得 {available}/{len(fields)}',
            c.get('board'), c.get('screen_date'), period, number(c.get('pe')), number(c.get('pb'))])
        qrows.append([s,c['name'],c.get('sector'),MODELS[m],qperiod,
            f'{len(qcurrent)}/{len(fields)}',f'{qverified}/{len(fields)}',score,
            number(q.get('profitability_score')) if score is not None else None,
            number(q.get('cash_flow_score')) if score is not None else None,
            number(q.get('balance_sheet_score')) if score is not None else None,
            number(q.get('growth_score')) if score is not None else None, qmissing] + [
                (number(qcurrent[f].get('value')) * float(MONEY[qcurrent[f]['unit']])
                 if f in {'cash', 'interest_bearing_debt'} and qcurrent[f].get('unit') in MONEY
                 else number(qcurrent[f].get('value')) if f not in {'cash', 'interest_bearing_debt'} else None)
                if f in qcurrent and number(qcurrent[f].get('value')) is not None
                and qcurrent[f].get('validation_status') not in {'failed','conflict'} else None
                for f in GENERAL])
        ff = ['revenue','revenue_yoy','net_income','net_income_yoy','roe','roic','gross_margin','net_margin',
              'operating_cash_flow','free_cash_flow','operating_cash_flow_to_net_income','debt_ratio',
              'cash','interest_bearing_debt','eps_reported','dps_ttm','payout_ratio']
        not_applicable = {'free_cash_flow','operating_cash_flow_to_net_income','debt_ratio','gross_margin','roic'} if m in SPECIAL_FIELDS else set()
        frows.append([s,c['name'],period] + [
            '不适用' if f in not_applicable else value(s,f,period,money=f in {'revenue','net_income','operating_cash_flow','free_cash_flow','cash','interest_bearing_debt'})
            for f in ff] + [MODELS[m],missing])
        vrows.append([s,c['name'],number(c.get('current_price')),value(s,'eps_ttm'),value(s,'bvps',period),
            None if m in SPECIAL_FIELDS else value(s,'fcf_per_share'),value(s,'dps_ttm'),
            number(c.get('pe')),number(c.get('pb')),value(s,'dividend_yield'),
            number(v.get('fair_value')),number(v.get('safety_margin')),signal,MODELS[m],
            c.get('quote_as_of') if 'candidate_tracking' in payload else c.get('screen_date'),period,
            '没有经验证的估值假设时，不用低PE/PB直接推算合理价'])
        if m in SPECIAL_FIELDS:
            for f in fields:
                p = get(s,f,period)
                regulatory = reference(f, period, value(s,f,period)) if m == 'broker' else None
                srows.append([s,c['name'],MODELS[m],period,LABELS[f],value(s,f,period),'CNY' if f=='segment_equity' else '%',
                    point_status(p) if p else '未取得，自动补证',
                    (p.get('metadata') or {}).get('statement_scope') if p else None,
                    (p.get('metadata') or {}).get('page_number') if p else None,
                    p.get('source_url') if p else disclosures.get(s,{}).get('source_url'),
                    p.get('sha256') if p else None,
                    float(regulatory['minimum']) if regulatory else None,
                    float(regulatory['warning']) if regulatory else None,
                    regulatory['position'] if regulatory else None])

    ws = _sheet(wb,'01_观察名单', ['代码','公司','行业','研究模型','个人关注状态','现价(元)','合理价(元)',
        '安全边际','估值状态','质量分/100','研究信号','建议权重','当前权重','自动补证缺口','个人复核日期',
        '证据状态','板块','初筛日期','财务报告期','PE(行情口径)','PB'],
        '初筛候选不等于买入清单。先看研究信号和证据状态；黄色数值尚未交叉验证。',rows)
    for col in ('H','L','M'):
        for cell in ws[col][3:]: cell.number_format = '0.0%'
    # Daily scanning stays compact; valuation and personal weight details live
    # on their dedicated pages, with the original column contract retained.
    for col in ('D','E','G','H','I','L','M','O','T','U'):
        ws.column_dimensions[col].hidden=True
    if all(r[9] is None for r in rows): ws.column_dimensions['J'].hidden=True
    for i,r in enumerate(rows,4):
        candidate = candidates[i-4]
        if 'candidate_tracking' in payload:
            ws.cell(i,6).comment = Comment(
                f"行情采集时间：{candidate.get('quote_as_of') or '缺失'}\n"
                f"行情source_id：{candidate.get('quote_source_id') or '缺失'}\n"
                f"PE口径：{candidate.get('pe_basis', 'unspecified')}", '行情溯源')
        full=r[13]
        ws.cell(i,14).comment=Comment(full,'完整补证缺口')
        ws.cell(i,14).value=full[:38]+'…' if len(full)>38 else full
        ws.row_dimensions[i].height=46
    qs = _sheet(wb,'02_质量评分',['代码','公司','行业','适用模型','报告期','已取得/必需','已验证/必需','质量分/100',
        '盈利/35','现金/25','稳健/20','成长/20','具体缺口',
        *[LABELS.get(f,f) + ('(亿元)' if f in {'cash','interest_bearing_debt'} else '(%)') for f in GENERAL]],
        '非金融企业按年度报告评分，中期数据见财务指标页；金融机构按最新专用指标跟踪。分数仅用于同类研究排序，缺证不出总分。',qrows)
    qs.column_dimensions['M'].width=66
    qs.column_dimensions['M'].hidden=True
    if all(r[7] is None for r in qrows):
        for col in ('H','I','J','K','L'): qs.column_dimensions[col].hidden=True
    for i, c in enumerate(candidates,4):
        s=symbol(c['symbol']); qperiod=str(qs.cell(i,5).value)
        for j,f in enumerate(GENERAL,14):
            p=annual_points.get((s,f)) or points.get((s,f))
            if p and str(p['period_label']) == qperiod and qs.cell(i,j).value is not None:
                qs.cell(i,j).fill=PatternFill('solid',fgColor=GREEN if accepted(p) else AMBER)
                qs.cell(i,j).comment=Comment('\n'.join([point_status(p),f'报告期：{qperiod}',
                    f"source_id：{p.get('source_id')}",str(p.get('source_url','')),f"SHA-256：{p.get('sha256','')}"]), '数据证据')
    for i in range(4,qs.max_row+1): qs.row_dimensions[i].height=80
    fs = _sheet(wb,'03_财务指标',['代码','公司','报告期','收入(亿元)','收入同比(%)','归母利润(亿元)',
        '利润同比(%)','ROE(%)','ROIC(%)','毛利率(%)','归母净利润/营业收入(%)','经营现金流(亿元)',
        '自由现金流(亿元)','合并经营现金流/归母净利润(%)','负债率(%)','货币资金(亿元)','有息负债(亿元)',
        '当期EPS(元)','DPS TTM(元)','分红支付率(%)','适用模型','证据缺口'],
        '各行仅展示同一报告期。空白=未取得或口径不符，不是0；金融机构不套用工业企业现金流/负债率判断。',frows)
    fs.row_dimensions[3].height=52
    fs.column_dimensions['K'].width=28
    fs.column_dimensions['N'].width=34
    for i,c in enumerate(candidates,4):
        s=symbol(c['symbol'])
        for j,f in enumerate(ff,4): annotate(fs,i,j,s,f,period_for(s))
    for col in ('I','M','N','Q','R','S','T','V'): fs.column_dimensions[col].hidden=True
    vs = _sheet(wb,'04_估值跟踪',['代码','公司','现价(元)','EPS TTM(元)','每股净资产(元)',
        '自由现金流/股(元)','DPS TTM(元)','PE(行情口径)','PB','股息率(%)','合理价(元)','安全边际',
        '研究信号','适用模型','行情时点','净资产报告期','估值说明'],
        'TTM=最近十二个月。半年报EPS不当作TTM；PE/PB只是价格比率，不代表便宜或安全。',vrows)
    for i,c in enumerate(candidates,4):
        s=symbol(c['symbol'])
        for col,f in [(4,'eps_ttm'),(5,'bvps'),(6,'fcf_per_share'),(7,'dps_ttm'),(10,'dividend_yield')]: annotate(vs,i,col,s,f)
        vs.cell(i,12).number_format='0.0%'
    special = _sheet(wb,'19_金融专用指标',['代码','公司','类型','报告期','专用指标','披露值','单位','验证状态',
        '报表主体口径','原文页码','官方URL','原文SHA-256',
        '一般监管下限(%)','一般预警线(%)','数值对照（非合规结论）'],
        '银行看信贷质量与资本；券商看风险和流动性；保险看偿付能力。监管比例不是投资收益评分。',srows)
    special.column_dimensions['E'].width=28
    special.column_dimensions['I'].width=32
    special.column_dimensions['K'].hidden=True
    special.column_dimensions['L'].hidden=True
    special.column_dimensions['M'].width=22
    special.column_dimensions['N'].width=22
    special.column_dimensions['O'].width=34
    for i,row in enumerate(srows,4):
        if row[12] is not None:
            special.cell(i,15).comment=Comment(
                '仅为一般标准的数值对照，不代表数据已验证、公司合规或值得买入。'
                '\n并表监管及个别要求可能不同；适用规则仍需持续检查。'
                '\n依据：2024年第13号公告，附件第20-22页；2025-01-01起施行。'
                f'\n{RULE_URL}\n附件SHA-256：{RULE_SHA256}', '监管参考')
        if row[5] is not None:
            special.cell(i,6).fill=PatternFill('solid',fgColor=GREEN if row[7]=='已交叉验证' else AMBER)
            special.cell(i,6).comment=Comment(f'原文：{row[10]}\nSHA-256：{row[11]}\n页码：{row[9]}','官方证据')
            special.cell(i,9).value=row[8] if row[8] and row[8] != 'issuer regulatory disclosure; entity scope requires corroboration' else '发行人监管披露；主体口径待比对'
        special.row_dimensions[i].height=44

    research = wb['09_公司研究']
    existing = {symbol(row[0]): list(row[:13]) for row in research.iter_rows(min_row=4,values_only=True) if row[0]}
    research_rows=[]
    for c in candidates:
        s=symbol(c['symbol'])
        if s in existing:
            research_rows.append(existing.pop(s))
        else:
            research_rows.append([s,c['name'],None,None,None,None,None,None,
                '、'.join(LABELS.get(f,f) for f in SPECIAL_FIELDS.get(model(c),GENERAL)),
                None,None,None,'已入研究队列；尚无公司定性结论'])
    research_rows.extend(existing.values())
    for row in research_rows:
        row.extend([None] * (13 - len(row)))
        row.extend((company_research or {}).get(symbol(row[0]), [None] * 4))
    _sheet(wb,'09_公司研究',['代码','公司','商业模式','核心护城河','跟踪理由','长期驱动','主要风险',
        '竞争对手','关键KPI','投资逻辑','证伪条件','研究日期','研究状态',
        '自动研究摘要','自动研究反证','自动研究适用范围','研究原文'],
        '保留原有个人研究；新公司不自动编造护城河和投资逻辑。未入本期初筛的历史研究保留在末尾。',research_rows)
    research = wb['09_公司研究']
    for column in ('N', 'O', 'P'):
        research.column_dimensions[column].width = 48
    for i, row in enumerate(research_rows, 4):
        if row[13]:
            research.row_dimensions[i].height = 100
            for col in range(14, 18):
                research.cell(i, col).alignment = Alignment(wrap_text=True, vertical='top')
            research.cell(i, 17).hyperlink = row[16]
    reports=[]
    for c in candidates:
        s=symbol(c['symbol']); d=disclosures.get(s,{})
        reports.append([s,c['name'],d.get('report_period'),d.get('title','尚无归档报告'),d.get('published_at'),
            '已有官方原件；指标逐项验证' if d else '自动补充官方原件',d.get('source_url'),d.get('sha256')])
    _sheet(wb,'10_年报跟踪',['代码','公司','最新报告期','报告名称','发布时间','状态','官方URL','SHA-256'],
        '包括年报、半年报和季报；原件归档不代表全部财务数字已验证。',reports)

    evidence=[]
    evidence_companies = {symbol(r[0]): r[1] for r in research_rows if r[0]}
    evidence_companies.update({s:c['name'] for s,c in company.items()})
    evidence_points = {(s, f, str(p['period_label'])): p
                       for source in (annual_points, points) for (s, f), p in source.items()}
    for (s,f,_),p in sorted(evidence_points.items()):
        if s not in evidence_companies: continue
        meta=p.get('metadata') or {}
        evidence.append([s,evidence_companies[s],LABELS.get(f,f),f,p.get('period_label'),number(p.get('value')),
            p.get('unit'),point_status(p),str(p.get('source_id','')),p.get('source_url'),p.get('sha256'),
            meta.get('page_number',meta.get('page')),p.get('parser_version'),p.get('fetched_at'),
            str(meta.get('input_source_ids', '')),meta.get('candidate_excerpt', '')])
    research_first_row = {}
    for s, records in sorted(debt_research['evidence'].items()):
        if s not in evidence_companies or not records:
            continue
        research_first_row[s] = len(evidence) + 4
        for row in records:
            evidence.append([s,evidence_companies[s],row['label'],row['field_name'],row['period'],
                number(row['value']),row['unit'],row['status'],row['source_id'],row['source_url'],
                row['sha256'],row['pages'],row['parser'],row['fetched_at'],'',row['excerpt']])
    _sheet(wb,'18_指标证据',['代码','公司','指标','字段ID','报告期','原始数值','原始单位','验证状态','source_id',
        '来源URL','SHA-256','页码','解析版本','抓取时间','推导输入source_id','原文与单位依据'],
        '数值追溯：财务页的单元格批注 → 此页source_id → 数据库原始文档 → 官方URL与Hash。',evidence)
    # Append-only audits: search all populated rows, including beyond the former 1000-row cap.
    audit=wb['11_数据源审计']
    populated=[row for row in audit.iter_rows(min_row=4) if row[0].value]
    keys={tuple(str(c.value or '') for c in row[:4]) for row in populated}
    next_row=max([row[0].row for row in populated],default=3)+1
    for p in payload.get('points',[]) + payload.get('annual_points',[]):
        key=(str(p.get('source_id','')),symbol(p['symbol']),p['field_name'],str(p['period_label']))
        if key in keys: continue
        vals=[*key,number(p.get('value')),p.get('unit'),p.get('source_name'),p.get('source_url'),None,None,
              p.get('fetched_at'),p.get('published_at'),p.get('parser_version'),p.get('sha256'),None,
              point_status(p),p.get('human_reviewed'),None]
        for j,val in enumerate(vals,1): _cell(audit,next_row,j,val)
        keys.add(key); next_row+=1
    monthly=wb['06_月度跟踪']
    existing_months={(str(row[0].value),symbol(row[1].value)) for row in monthly.iter_rows(min_row=4) if row[0].value}
    last=max([row[0].row for row in monthly.iter_rows(min_row=4) if row[0].value],default=3)+1
    for snap in payload.get('monthly_snapshots',[]):
        key=(str(snap['snapshot_month']),symbol(snap['symbol']))
        if key in existing_months: continue
        vals=[*key,snap.get('name'),number(snap.get('current_price')),None,number(snap.get('safety_margin')),
            snap.get('valuation_status'),snap.get('build_signal'),None,number(snap.get('revenue_yoy')),
            number(snap.get('net_income_yoy')),number(snap.get('roe')),None,None,snap.get('data_status'),snap.get('snapshot_at')]
        for j,val in enumerate(vals,1): _cell(monthly,last,j,val)
        existing_months.add(key); last+=1
    _cell(monthly,2,1,'仅追加已发生的月末快照；不把今天的数据伪装成历史。')
    verified_cells=sum(accepted(p) for (s,_),p in points.items() if s in company)
    special_values=sum(row[5] is not None for row in srows)
    ar=[]
    for alert in alerts:
        ar.append([alert['symbol'],alert['name'],alert.get('board'),alert.get('sector'),alert['category'],
            alert['priority'],alert['action'],'；'.join(alert['reasons']) or '数据门禁通过，仍需投资决策',
            number(alert.get('current_price')),number(alert.get('safety_margin')),alert['quality_status'],
            str(alert['as_of']),alert['price_source_id'],alert['fair_value_source_id'],
            '历史回测未完成','查看验证过程'])
    reminder_sheet=_sheet(wb,'12_提醒',['代码','公司','板块','行业','提醒类别','研究优先级','建议行动','依据或阻断原因',
        '现价(元)','安全边际','质量状态','行情抓取时间','行情source_id','合理价source_id',
        '策略验证状态','决策验证'],
        '买卖均为研究提示，不执行交易。重点观察不是买入批准；未知持仓不生成卖出指令。',ar)
    for col in ('G','H'): reminder_sheet.column_dimensions[col].width=54
    for col in ('M','N'): reminder_sheet.column_dimensions[col].hidden=True
    reminder_sheet.column_dimensions['O'].width=26
    reminder_sheet.column_dimensions['P'].width=22
    for i in range(4,reminder_sheet.max_row+1):
        reminder_sheet.row_dimensions[i].height=72
        reminder_sheet.cell(i,10).number_format='0.0%'
        reminder_sheet.cell(i,16).hyperlink=f"#'21_决策验证'!A{i}"
        reminder_sheet.cell(i,16).style='Hyperlink'
    wb.move_sheet(reminder_sheet,offset=1-wb.sheetnames.index('12_提醒'))
    # Keep the Moutai experiment visible in the existing decision page, while
    # preserving its strict exclusion from any executable decision.
    research_root = Path(__file__).resolve().parents[2]
    moutai_replay_status = (load_cash_anchor_paper_contract_status(research_root)
                            + load_pe_mid_paper_contract_status(research_root)
                            + load_current_valuation_admission_status(research_root)
                            + load_historical_conditional_replay_status(research_root)
                            + load_blocked_paper_ledger_status(research_root)
                            + load_simulation_closure_status(research_root)
                            + load_experimental_signal_coverage_status(research_root)
                            + load_historical_range_experiment_status(research_root)
                            + load_current_execution_contract_status(research_root)
                            + load_current_conditional_observation(research_root)
                            + load_finance_cost_scope_sensitivity_status(research_root))
    decision_rows=[]
    for alert in alerts:
        checks=alert['decision_checks']
        decision_rows.append([alert['symbol'],alert['name'],alert.get('board'),alert['category'],
            sum(c['passed'] for c in checks),len(checks),
            *['通过' if c['passed'] else '未通过' for c in checks],
            number(alert['decision_price']),number(alert['decision_fair_value']),
            number(alert.get('safety_margin')),alert['upstream_build_signal'],
            number(alert.get('known_holding')),
            '；'.join(alert['reasons']) or '证据门禁通过；按安全边际及持仓判断',
            alert['action'],alert['price_source_id'],alert['fair_value_source_id'],
            moutai_replay_status if re.sub(r'\D', '', str(alert['symbol'])).zfill(6)[-6:] == '600519' else '历史回测未完成',None,
            debt_research['companies'].get(alert['symbol'], '尚未纳入负债附注研究批次'),
            alert['quote_session']['reason'], alert['quote_session']['expected_session'],
            '；'.join(f'{provider}: {stamp}' for provider,stamp in
                     alert['quote_session']['provider_times'].items()) or None])
    if company_research and '600519' in company_research and not any(re.sub(r'\D', '', str(alert['symbol'])).zfill(6)[-6:] == '600519' for alert in alerts):
        # Moutai is the fixed single-stock validation case, even when it is
        # outside today's market-candidate reminder list.
        decision_rows.append([
            '600519', '贵州茅台', '主板', '研究验证（非买卖提醒）', 0, 8,
            *['未通过'] * 8,
            None, None, None, '研究实验', None,
            '实验历史估值不得作为正式合理价、买卖信号或策略收益。',
            '查看09_公司研究的证据与阻断项；不生成订单。',
            None, None, moutai_replay_status, None,
            '单股票研究案例，非负债附注研究', None, None, None,
        ])
    decision=_sheet(wb,'21_决策验证',['代码','公司','板块','最终提醒','通过项','总门禁',
        '每日行情一致性','行情验证与时效','合理价证据','财务质量评分','最新同报告期指标',
        '行业模型适用性','估值计算与证据关联','现价与合理价数值','决策现价(元)',
        '合理价(元)','安全边际','上游建仓信号','已知持仓数量','阻断原因','研究行动',
        '行情source_id','合理价source_id','策略验证状态','回测运行ID','负债附注研究（非交易批准）',
        '收盘交易日核验','最近已完成交易日','逐源实际报价时间'],
        '数据门禁通过不代表策略有效。历史回测尚未完成。当前研究规则：安全边际≥30%且上游为建仓候选→买入研究；现价高于合理价且已知持仓>0→减仓研究。空白持仓表示未知；不自动交易。',decision_rows)
    decision.column_dimensions['X'].width=26
    decision.column_dimensions['Y'].width=24
    decision.column_dimensions['Z'].width=50
    decision.column_dimensions['AA'].width=52
    decision.column_dimensions['AB'].width=23
    decision.column_dimensions['AC'].width=48
    for col in ('T','U'): decision.column_dimensions[col].width=58
    for col in ('V','W'): decision.column_dimensions[col].hidden=True
    for i,alert in enumerate(alerts,4):
        decision.row_dimensions[i].height=84
        decision.cell(i,17).number_format='0.0%'
        for j,check in enumerate(alert['decision_checks'],7):
            cell=decision.cell(i,j)
            cell.fill=PatternFill('solid',fgColor='E2F0D9' if check['passed'] else 'FCE4D6')
            cell.comment=Comment('；'.join(check['reasons']) or '本项门禁通过；不代表批准交易','Value Agent')
        decision.cell(i,1).hyperlink="#'18_指标证据'!A1"
        if alert['symbol'] in research_first_row:
            decision.cell(i,26).hyperlink=f"#'18_指标证据'!A{research_first_row[alert['symbol']]}"
            decision.cell(i,26).font=Font(name='Microsoft YaHei',size=10,color='0563C1',underline='single')
    wb.move_sheet(decision,offset=2-wb.sheetnames.index('21_决策验证'))
    audit_data=payload.get('market_audit') or {}
    coverage=audit_data.get('coverage') or {}
    coverage_rows=[['全量接收',coverage.get('received_rows'),'接口返回条数，不直接等于官方公司总数'],
        ['去重证券代码',coverage.get('unique_symbols'),'当前快照，不是历史证券库累加数'],
        ['交易所清单逐代码对账','已完成' if coverage.get('official_universe_reconciled') else '尚未完成','未逐代码对账前不声称100%覆盖A股']]
    official = payload.get('official_coverage') or {}
    if official.get('expected_count'):
        coverage_rows.extend([
            ['官方研究证券范围',official['expected_count'],official.get('scope')],
            ['逐代码匹配',official['matched_count'],f"官方清单日 {official.get('official_date')}；行情快照日 {official.get('market_date')}"],
            ['官方有、行情缺失',len(official.get('missing_in_market',[])),'缺失代码不能被静默排除'],
            ['行情有、官方未匹配',len(official.get('not_in_official',[])),'需核对证券状态与代码变更'],
            ['板块冲突',len(official.get('board_conflicts',[])),'逐代码比较官方与行情的板块归类'],
            ['官方清单证据ID',official.get('source_id'),'可通过数据库追溯本次官方清单'],
            ['官方清单Hash',official.get('sha256'),'全部交易所原始响应均已归档'],
        ])
        for source in official.get('sources',[]):
            coverage_rows.append([source['label'],source['sha256'],source['url']])
        for key,label in [('missing_in_market','行情缺失'),('not_in_official','官方未匹配'),('board_conflicts','板块冲突')]:
            for row in official.get(key,[]):
                coverage_rows.append([label,row['symbol'],str(row)])
    for board,count in (coverage.get('board_counts') or {}).items():
        coverage_rows.append([board,count,f"本期初筛候选 {(coverage.get('candidate_board_counts') or {}).get(board,0)} 家"])
    outcome_labels={'candidate':'初筛候选','invalid_identity':'代码/名称无效','outside_a_share_boards':'历史未识别板块',
        'unmapped_board':'板块未识别，待官方清单映射',
        'risk_warning_st':'ST风险警示','price_conflict':'双源价格冲突','price_cross_check_missing':'缺少第二行情来源',
        'missing_or_nonpositive_inputs':'指标缺失或非正值','pe_above_25':'PE高于25','market_cap_below_5bn':'市值小于50亿元','pb_above_3':'PB高于3'}
    for reason,count in (coverage.get('outcome_counts') or {}).items():
        coverage_rows.append([outcome_labels.get(reason,reason),count,'依初筛规则的首个未通过条件计数，互不重复'])
    coverage_rows.extend([['原始快照Hash',audit_data.get('sha256'),'完整响应及逐股票排除原因已归档'],
                          ['快照时间',str(audit_data.get('fetched_at','')),'UTC时间含时区'],
                          ['来源',audit_data.get('source_url'),'公共行情需要和交易所证券清单进一步对账']])
    cv=_sheet(wb,'20_市场覆盖',['检查项','数量/状态','口径说明'],'遍历范围、板块分类及排除原因；所有计数来自当前快照。',coverage_rows)
    for col,width in [('A',34),('B',42),('C',82)]: cv.column_dimensions[col].width=width
    blocks=[['初筛研究公司',len(candidates),'已进入观察、质量、财务、估值、公司研究和报告页'],
        ['初筛日期',market_date,'入池依据日期，不代表最新行情日期'],
        ['行情快照时间',
         min((str(c.get('quote_as_of') or '') for c in candidates), default='') or '跟踪缺失'
         if 'candidate_tracking' in payload else market_date,
         '每日跟踪显示本批最早采集时间；不是实时成交时间'],
        ['服务导出时间',generated,'导出时间不等于行情时间'],
        ['数据时效','已过期，信号阻断' if stale else '请结合各字段报告期',
         '超过36小时未导出或每日行情超过30小时则阻断' if 'candidate_tracking' in payload
         else '超过36小时未导出或行情超过3日则阻断'],
        ['已交叉验证指标',verified_cells,'黄色仅供研究；绿色为逐项自动交叉验证'],
        ['金融专用指标',f'{special_values}/{len(srows)}','展示原文数值、页码、主体与待验证状态'],
        ['有效建仓候选',sum(r[10]=='建仓候选' for r in rows),'质量、估值、时效未通过时不输出建仓信号'],
        ['第一步','01_观察名单','筛选行业、板块、研究信号，定位证据缺口'],
        ['当日提醒','12_提醒','买入研究、减仓研究、重点观察和数据过期分开显示'],
        ['市场覆盖','20_市场覆盖','核对本次遍历、板块及排除原因；不把候选数当成市场总数'],
        ['第二步','03_财务指标 / 19_金融专用指标','查看盈利、现金与风险指标；不同行业不同口径'],
        ['第三步','04_估值跟踪','合理价为空表示估值证据不足，不是0元'],
        ['第四步','09_公司研究 / 10_年报跟踪','先理解生意与风险，再考虑交易；不会自动下单'],
        ['个人记录','05_仓位管理 / 06_月度跟踪 / 08_交易记录','原有内容保留；初始资产等样例不等于券商真实资产'],
        ['完整证据','18_指标证据','保留source_id、报告期、单位、URL、页码和Hash'],
    ]
    dash=_sheet(wb,'00_首页Dashboard',['项目','当前结果','怎么看'],
        '价值投资研究台账 | 先看证据，再谈价格；初筛、财务验证和买入判断分开。',blocks)
    for col,width in [('A',26),('B',42),('C',82)]: dash.column_dimensions[col].width=width
    for i in range(4,dash.max_row+1): dash.row_dimensions[i].height=40
    for i,row in enumerate(blocks,4):
        dest=str(row[1]).split(' / ')[0]
        if dest in wb: dash.cell(i,2).hyperlink=f"#'{dest}'!A1"
    for name in ['13_全市场初筛','14_财报候选','15_公司财务覆盖','16_板块行业汇总',
                 '17_财务质量评分','11_数据源审计','12_系统设置']:
        if name in wb: wb[name].sheet_state='hidden'
    for name in ['13_全市场初筛','14_财报候选','15_公司财务覆盖','16_板块行业汇总','17_财务质量评分']:
        if name in wb: _cell(wb[name],2,1,'历史技术视图（停止维护）；当前结果请看01/02/03/18/19页。')
    for ws in wb:
        ws.sheet_view.tabSelected=False
    wb.active=wb.sheetnames.index('00_首页Dashboard')
    from .workbook_frontdoor import apply_frontdoor
    apply_frontdoor(wb)
    output.parent.mkdir(parents=True,exist_ok=True)
    # The canonical workbook contains retained history and is rebuilt alongside a
    # full market payload. Drop transient indexes before openpyxl serializes it.
    del snapshots, points, annual_points, quality, valuations, disclosures, company
    gc.collect()
    wb.save(output)
    return {'candidates':len(candidates),'verified_points':verified_cells,
            'specialized_values':special_values,'specialized_required':len(srows),'stale':stale}
