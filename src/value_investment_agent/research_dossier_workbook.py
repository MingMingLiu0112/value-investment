"""Standalone W3 candidate workbook for local inspection.

This workbook is a disposable presentation artifact derived only from
ResearchDossierCollection.  It is not the WPS production workbook and contains
no valuation formula, price signal, order, or position column.
"""
from __future__ import annotations

from datetime import timezone
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .research_read_model import (
    ACTION_NO_ORDER,
    ResearchDossier,
    ResearchDossierCollection,
)


OVERVIEW_SHEET = "00_研究档案总览"
GAP_SHEET = "00_缺口与阻断"
EVIDENCE_SHEET = "00_证据目录"
INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_DISPLAY_STATUS = {
    "COMPLETE": "已完成",
    "PARTIAL": "部分完成",
    "UNKNOWN": "未知",
    "MISSING": "缺失",
    "NOT_APPLICABLE": "不适用",
    "BLOCKED": "受阻",
    "NOT_STARTED": "未开始",
    "READABLE": "可阅读",
}


def _style_cell(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int = 8) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style_cell(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style_cell(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 28


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style_cell(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _overview(wb: Workbook, collection: ResearchDossierCollection) -> None:
    ws = wb.create_sheet(OVERVIEW_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D4"
    widths = [10, 14, 24, 30, 13, 12, 15, 16, 60]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    _title(
        ws,
        "M1 固定样本研究档案候选",
        f"生成时间 {collection.generated_at.isoformat()} | action={collection.action} | 仅研究，不交易",
        len(widths),
    )
    row = _header(
        ws,
        4,
        [
            "证券代码",
            "公司",
            "经济画像",
            "注册模型",
            "档案状态",
            "财务时点",
            "证据数",
            "阻塞项",
            "首个缺口/阻塞",
        ],
    )
    for dossier in collection.dossiers:
        first = dossier.research_gaps.blockers[0] if dossier.research_gaps.blockers else "无已登记阻塞"
        _style_cell(ws.cell(row, 1, dossier.symbol), bold=True)
        _style_cell(ws.cell(row, 2, dossier.name))
        _style_cell(ws.cell(row, 3, dossier.profile_id))
        _style_cell(ws.cell(row, 4, dossier.primary_model or "未注册"))
        _style_cell(ws.cell(row, 5, _DISPLAY_STATUS.get(dossier.readiness, dossier.readiness)))
        _style_cell(
            ws.cell(row, 6, dossier.financial_period.isoformat() if dossier.financial_period else "未登记")
        )
        _style_cell(ws.cell(row, 7, len(dossier.evidence_refs)))
        _style_cell(ws.cell(row, 8, len(dossier.visible_blockers)))
        _style_cell(ws.cell(row, 9, first))
        code = ws.cell(row, 1)
        code.hyperlink = f"#'{dossier.symbol}'!A1"
        code.font = Font(name="Microsoft YaHei", size=11, bold=True, color=GREEN, underline="single")
        ws.row_dimensions[row].height = 30
        row += 1


def _company_sheet(wb: Workbook, dossier: ResearchDossier) -> None:
    ws = wb.create_sheet(dossier.symbol)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 100
    _title(
        ws,
        f"{dossier.name} {dossier.symbol}",
        (
            f"as_of={dossier.as_of.isoformat()} | profile={dossier.profile_id} | "
            f"readiness={dossier.readiness} | action={dossier.action}"
        ),
        3,
    )
    row = 4
    rows = [
        ("研究版本", dossier.research_version, dossier.run_id),
        ("模型", dossier.primary_model or "未注册", dossier.research_status),
        (
            "财务时点",
            dossier.financial_period.isoformat() if dossier.financial_period else "未登记",
            dossier.valuation_status,
        ),
    ]
    for label, value, note in rows:
        _style_cell(ws.cell(row, 1, label), fill=GREY, bold=True)
        _style_cell(ws.cell(row, 2, value))
        _style_cell(ws.cell(row, 3, note))
        row += 1
    row += 1
    sections = (
        dossier.business_quality,
        dossier.financial_quality,
        dossier.capital_allocation,
        dossier.thesis,
        dossier.counter_evidence,
        dossier.thesis_breakers,
        dossier.next_events,
        dossier.research_gaps,
    )
    for item in sections:
        _style_cell(ws.cell(row, 1, item.title), fill=BLUE, bold=True, color="FFFFFF")
        _style_cell(ws.cell(row, 2, _DISPLAY_STATUS.get(item.status, item.status)), bold=True)
        _style_cell(ws.cell(row, 3, item.explanation), fill=GREY)
        ws.row_dimensions[row].height = 28
        row += 1
        for blocker in item.blockers:
            _style_cell(ws.cell(row, 2, "阻塞"), fill=AMBER, bold=True, color=RED)
            _style_cell(ws.cell(row, 3, blocker))
            row += 1
        for finding in item.findings:
            _style_cell(ws.cell(row, 2, finding.status))
            _style_cell(
                ws.cell(row, 3, f"[{finding.kind}] {finding.text}")
            )
            row += 1
        for dimension in item.dimensions:
            _style_cell(ws.cell(row, 1, dimension.key), fill=GREY)
            _style_cell(ws.cell(row, 2, _DISPLAY_STATUS.get(dimension.status, dimension.status)))
            _style_cell(ws.cell(row, 3, dimension.observation))
            row += 1
        row += 1
    _style_cell(ws.cell(row, 1, "证据引用"), fill=BLUE, bold=True, color="FFFFFF")
    row += 1
    for ref in dossier.evidence_refs:
        _style_cell(ws.cell(row, 1, ref.get("id") or ""))
        _style_cell(ws.cell(row, 2, ref.get("sha256", "")[:12]))
        _style_cell(ws.cell(row, 3, ref.get("path") or ""))
        row += 1


def _gaps(wb: Workbook, collection: ResearchDossierCollection) -> None:
    ws = wb.create_sheet(GAP_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 100
    _title(ws, "研究缺口与阻断", "缺失、未知和阻断必须可见；不得用公司数量掩盖真实状态。", 3)
    row = _header(ws, 4, ["证券代码", "公司", "缺口/阻断"])
    for dossier in collection.dossiers:
        blockers = dossier.visible_blockers
        if not blockers:
            blockers = ["当前候选档案未登记阻塞，但仍不等于研究已完成。"]
        for blocker in blockers:
            _style_cell(ws.cell(row, 1, dossier.symbol))
            _style_cell(ws.cell(row, 2, dossier.name))
            _style_cell(ws.cell(row, 3, blocker))
            row += 1


def _evidence(wb: Workbook, collection: ResearchDossierCollection) -> None:
    ws = wb.create_sheet(EVIDENCE_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 68
    _title(ws, "证据目录", "只列当前候选档案已引用、带 Hash 的本地证据。", 3)
    row = _header(ws, 4, ["证据 ID", "路径", "SHA-256"])
    seen: set[str] = set()
    for dossier in collection.dossiers:
        for ref in dossier.evidence_refs:
            ref_id = str(ref.get("id") or "")
            if ref_id in seen:
                continue
            seen.add(ref_id)
            _style_cell(ws.cell(row, 1, ref_id))
            _style_cell(ws.cell(row, 2, str(ref.get("path") or "")))
            _style_cell(ws.cell(row, 3, str(ref.get("sha256") or "")))
            row += 1


def build_dossier_workbook(collection: ResearchDossierCollection) -> Workbook:
    if collection.action != ACTION_NO_ORDER:
        raise ValueError("Dossier workbook cannot consume a non-no_order collection")
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, collection)
    _gaps(wb, collection)
    _evidence(wb, collection)
    for dossier in collection.dossiers:
        _company_sheet(wb, dossier)
    wb.active = 0
    return wb


def write_dossier_workbook(
    collection: ResearchDossierCollection,
    *,
    root: Path,
) -> dict[str, Any]:
    timestamp = collection.generated_at.astimezone(timezone.utc).strftime(
        "%Y%m%dT%H%M%SZ"
    )
    target_dir = root / "runtime" / f"m1-research-dossier-candidate-{timestamp}"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / "m1-research-dossier-candidate.xlsx"
    workbook = build_dossier_workbook(collection)
    workbook.save(target)
    manifest_path = target_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["workbook_sha256"] = hashlib.sha256(target.read_bytes()).hexdigest()
    manifest["workbook_path"] = str(target.relative_to(root))
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {
        "workbook_path": str(target.relative_to(root)),
        "workbook_sha256": manifest["workbook_sha256"],
    }
