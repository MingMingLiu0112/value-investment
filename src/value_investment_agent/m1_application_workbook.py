"""Build a presentation-only workbook from the shared M1 Application output.

This module deliberately contains no valuation arithmetic.  It reads the
already-computed Application evidence bundle, the bounded reverse-valuation
results and the archived dossier read model, then formats those states for a
human reviewer.  The workbook is a disposable candidate and is not the WPS
production workbook.
"""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .research_read_model import ACTION_NO_ORDER, ResearchDossierCollection


APPLICATION_SCHEMA = "m1-research-application-run-v1"
APPLICATION_POINTER = "runtime/m1-research-application-latest.json"
ROOT = Path(__file__).resolve().parents[2]

OVERVIEW_SHEET = "00_M1应用总览"
VALUATION_SHEET = "01_估值与价格"
DIVIDEND_SHEET = "02_股利评估"
REVERSE_SHEET = "03_反向估值"
BLOCKERS_SHEET = "04_阻断与证据"
SAMPLE_SHEET = "05_研究样本"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_DISPLAY = {
    "COMPLETED_WITH_BLOCKERS": "完成，有阻断",
    "COMPLETED": "完成",
    "UNSUPPORTED": "模型不支持",
    "PENDING_EXTERNAL_DATA": "等待外部数据",
    "READY": "可用",
    "INVALID": "不可用",
    "STALE_MODEL": "模型过期",
    "PARTIAL": "部分完成",
    "LOW": "低",
    "HIGH": "高",
    "MEDIUM": "中",
    "UNKNOWN": "未知",
    "READABLE": "可阅读",
    "BLOCKED": "受阻",
    "NOT_STARTED": "未开始",
    "NOT_READY": "未就绪",
    "paid": "已支付",
    "proposed": "已提议",
    "approved": "已批准",
    "ordinary": "普通股利",
    "trailing_paid": "滚动已支付",
    "declared": "已宣告",
    "normalized_scenario": "正常化情景",
    "current": "当前",
    "normalized": "正常化",
    "above_registered_envelope": "高于登记包络",
    "below_registered_envelope": "低于登记包络",
    "conditional_solution": "包络内有解",
    "not_ready": "未就绪",
    "conditional_research_only": "条件研究，非交易",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 28


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _set_widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _display(value: Any) -> str:
    text = str(value if value is not None else "")
    return _DISPLAY.get(text, text)


def _decimal(value: Any) -> str:
    if value is None:
        return ""
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return str(value)


def _optional_text(mapping: Mapping[str, Any], key: str) -> str:
    value = mapping.get(key)
    if value is None:
        return ""
    text = str(value)
    return text if text else ""


def _percent(value: Any) -> str:
    if value is None:
        return ""
    try:
        return f"{float(value) * 100:.2f}%"
    except (TypeError, ValueError):
        return str(value)


def _capital_allocation_context(context: Any) -> str:
    if not isinstance(context, Mapping):
        return ""
    ordered = (
        ("coverage_context", "生命周期事实"),
        ("fact_policy_forecast_separation", "事实/政策/预测"),
        ("policy_2025_2027", "回报政策"),
        ("long_term_policy_observation", "长期政策"),
        ("special_dividend_observation", "特别股利"),
        ("h1_2026_decision", "2026 H1 决策"),
    )
    return "\n".join(
        f"{label}: {value}"
        for key, label in ordered
        if (value := context.get(key))
    )


def _company_name(
    symbol: str,
    dossiers: dict[str, Any],
) -> str:
    dossier = dossiers.get(symbol)
    return str(dossier.name) if dossier is not None else symbol


def _company_profile(symbol: str, dossiers: dict[str, Any]) -> str:
    dossier = dossiers.get(symbol)
    return str(dossier.profile_id) if dossier is not None else ""


def _model_text(item: Mapping[str, Any]) -> str:
    valuation = item.get("valuation") or {}
    if not isinstance(valuation, Mapping):
        return ""
    return str(valuation.get("model_type") or valuation.get("model_id") or "")


def _distribution_sustainability(distribution: Any) -> str:
    if not isinstance(distribution, Mapping):
        return ""
    sustainability = distribution.get("sustainability")
    if not isinstance(sustainability, Mapping):
        return ""
    return _display(sustainability.get("status"))


def _reverse_summary(item: Mapping[str, Any]) -> str:
    values = item.get("reverse_valuations") or []
    if not isinstance(values, list):
        return ""
    parts = []
    for value in values:
        if not isinstance(value, Mapping):
            continue
        parts.append(
            f"{value.get('driver')}:{_display(value.get('status'))}"
        )
    return " | ".join(parts)


def _overview(wb: Workbook, payload: Mapping[str, Any], dossiers: dict[str, Any]) -> None:
    ws = wb.create_sheet(OVERVIEW_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D4"
    widths = [10, 15, 20, 24, 14, 20, 18, 16, 16, 17, 17, 14, 34]
    _set_widths(ws, widths)
    _title(
        ws,
        "M1 三公司研究 Application 候选",
        f"生成时间 {payload['generated_at']} | action={payload['action']} | 仅研究，不交易",
        len(widths),
    )
    row = _header(
        ws,
        4,
        [
            "证券代码",
            "公司",
            "经济画像",
            "模型",
            "运行状态",
            "估值状态",
            "Bear/Base/Bull",
            "当前价",
            "价格桥",
            "当前数据",
            "股利可持续性",
            "反向估值",
            "首个阻断",
        ],
    )
    for item in payload.get("results") or []:
        symbol = str(item.get("symbol") or "")
        valuation = item.get("valuation") or {}
        bridge = item.get("price_bridge") or {}
        current = item.get("current_data_status") or {}
        scenarios = ""
        if isinstance(valuation, Mapping):
            scenarios = " / ".join(
                _decimal(valuation.get(key))
                for key in ("bear_value", "base_value", "bull_value")
                if valuation.get(key) is not None
            )
        blockers = item.get("blockers") or []
        first_blocker = blockers[0] if blockers else "无已登记阻断"
        _style(ws.cell(row, 1, symbol), bold=True)
        _style(ws.cell(row, 2, _company_name(symbol, dossiers)))
        _style(ws.cell(row, 3, _company_profile(symbol, dossiers)))
        _style(ws.cell(row, 4, _model_text(item)))
        _style(ws.cell(row, 5, _display(item.get("run_status"))))
        _style(ws.cell(row, 6, _display(valuation.get("status")) if isinstance(valuation, Mapping) else ""))
        _style(ws.cell(row, 7, scenarios))
        _style(ws.cell(row, 8, _decimal(bridge.get("current_price")) if isinstance(bridge, Mapping) else ""))
        _style(ws.cell(row, 9, _display(bridge.get("bridge_status")) if isinstance(bridge, Mapping) else ""))
        _style(ws.cell(row, 10, _display(current.get("status")) if isinstance(current, Mapping) else ""))
        _style(ws.cell(row, 11, _distribution_sustainability(item.get("distribution"))))
        _style(ws.cell(row, 12, _reverse_summary(item)))
        _style(ws.cell(row, 13, first_blocker), fill=AMBER if blockers else "FFFFFF", color=RED if blockers else INK)
        ws.row_dimensions[row].height = 42
        row += 1


def _valuation_sheet(wb: Workbook, payload: Mapping[str, Any], dossiers: dict[str, Any]) -> None:
    ws = wb.create_sheet(VALUATION_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    _set_widths(ws, [10, 14, 20, 14, 14, 14, 14, 16, 18, 18, 14, 24])
    _title(
        ws,
        "估值与价格桥",
        "数值均为研究算例，不是目标价、买卖点或仓位建议；人工估值门仍为 false。",
        12,
    )
    row = _header(
        ws,
        4,
        [
            "证券代码",
            "公司",
            "模型",
            "Bear",
            "Base",
            "Bull",
            "置信度",
            "当前价",
            "相对 Bear",
            "相对 Base",
            "桥状态",
            "估值说明",
        ],
    )
    for item in payload.get("results") or []:
        symbol = str(item.get("symbol") or "")
        valuation = item.get("valuation") or {}
        bridge = item.get("price_bridge") or {}
        if not isinstance(valuation, Mapping) or not isinstance(bridge, Mapping):
            continue
        margin_bear = bridge.get("margin_to_bear")
        margin_base = bridge.get("margin_to_base")
        margin_bear_text = (
            f"{float(margin_bear):.1%}" if margin_bear is not None else ""
        )
        margin_base_text = (
            f"{float(margin_base):.1%}" if margin_base is not None else ""
        )
        assumptions = valuation.get("assumptions") or {}
        scope = str(assumptions.get("scope") or "") if isinstance(assumptions, Mapping) else ""
        _style(ws.cell(row, 1, symbol), bold=True)
        _style(ws.cell(row, 2, _company_name(symbol, dossiers)))
        _style(ws.cell(row, 3, valuation.get("model_type")))
        _style(ws.cell(row, 4, _decimal(valuation.get("bear_value"))))
        _style(ws.cell(row, 5, _decimal(valuation.get("base_value"))))
        _style(ws.cell(row, 6, _decimal(valuation.get("bull_value"))))
        _style(ws.cell(row, 7, _display(valuation.get("confidence"))))
        _style(ws.cell(row, 8, _decimal(bridge.get("current_price"))))
        _style(ws.cell(row, 9, margin_bear_text))
        _style(ws.cell(row, 10, margin_base_text))
        _style(ws.cell(row, 11, _display(bridge.get("bridge_status"))))
        _style(ws.cell(row, 12, scope))
        row += 1


def _dividend_sheet(wb: Workbook, payload: Mapping[str, Any], dossiers: dict[str, Any]) -> None:
    ws = wb.create_sheet(DIVIDEND_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    _set_widths(ws, [10, 14, 18, 18, 18, 18, 62, 56, 18, 18, 18, 48])
    _title(
        ws,
        "股利可持续性评估",
        "已支付/已提议、普通/特别、事实/政策/预测、当前/正常化均分开登记；LOW 是保留真实缺口后的低置信度结论。",
        12,
    )
    row = _header(
        ws,
        4,
        [
            "证券代码",
            "公司",
            "历史状态",
            "能力状态",
            "可持续性",
            "置信度",
            "政策/时点事实",
            "主要缺口",
            "",
            "",
            "",
            "",
        ],
    )
    lifecycle_rows = []
    snapshot_rows = []
    for item in payload.get("results") or []:
        distribution = item.get("distribution")
        if not isinstance(distribution, Mapping):
            continue
        symbol = str(item.get("symbol") or "")
        history = distribution.get("history") or {}
        capacity = distribution.get("capacity") or {}
        sustainability = distribution.get("sustainability") or {}
        context = capacity.get("capital_allocation_context") or {}
        coverage_text = _optional_text(sustainability, "coverage_context") if isinstance(sustainability, Mapping) else ""
        context_text = _capital_allocation_context(context)
        fact_text = "\n".join(
            part for part in (coverage_text, context_text) if part
        )
        blockers = distribution.get("blockers") or []
        _style(ws.cell(row, 1, symbol), bold=True)
        _style(ws.cell(row, 2, _company_name(symbol, dossiers)))
        _style(ws.cell(row, 3, _display(history.get("status")) if isinstance(history, Mapping) else ""))
        _style(ws.cell(row, 4, _display(capacity.get("status")) if isinstance(capacity, Mapping) else ""))
        _style(ws.cell(row, 5, _display(sustainability.get("status")) if isinstance(sustainability, Mapping) else ""))
        _style(ws.cell(row, 6, _display(sustainability.get("confidence")) if isinstance(sustainability, Mapping) else ""))
        _style(ws.cell(row, 7, fact_text))
        _style(ws.cell(row, 8, "；".join(str(value) for value in blockers)), fill=AMBER if blockers else "FFFFFF")
        ws.row_dimensions[row].height = 96
        if isinstance(history, Mapping):
            lifecycle_rows.extend(
                (symbol, _company_name(symbol, dossiers), record)
                for record in history.get("records") or []
                if isinstance(record, Mapping)
            )
        snapshot_rows.extend(
            (symbol, _company_name(symbol, dossiers), snapshot)
            for snapshot in distribution.get("yield_snapshots") or []
            if isinstance(snapshot, Mapping)
        )
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    _style(ws.cell(row, 1, "法定股利生命周期台账"), fill=GREY, bold=True)
    row = _header(
        ws,
        row + 1,
        [
            "证券代码",
            "公司",
            "报告期",
            "类型",
            "状态",
            "每股股利",
            "公告日",
            "批准日",
            "除息日",
            "支付日",
            "信息可用日",
            "阻断",
        ],
    )
    for symbol, company, record in lifecycle_rows:
        blockers = record.get("blockers") or []
        _style(ws.cell(row, 1, symbol), bold=True)
        _style(ws.cell(row, 2, company))
        _style(ws.cell(row, 3, record.get("fiscal_period")))
        _style(ws.cell(row, 4, _display(record.get("dividend_type"))))
        _style(ws.cell(row, 5, _display(record.get("status"))))
        _style(ws.cell(row, 6, _decimal(record.get("dividend_per_share"))))
        _style(ws.cell(row, 7, record.get("announcement_date")))
        _style(ws.cell(row, 8, record.get("approval_date")))
        _style(ws.cell(row, 9, record.get("ex_date")))
        _style(ws.cell(row, 10, record.get("payment_date")))
        _style(ws.cell(row, 11, record.get("known_at")))
        _style(
            ws.cell(row, 12, "；".join(str(value) for value in blockers)),
            fill=AMBER if blockers else "FFFFFF",
        )
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    _style(
        ws.cell(row, 1, "当前价股息收益率快照：当前与正常化分别登记，正常化未评估保持未就绪"),
        fill=GREY,
        bold=True,
    )
    row = _header(
        ws,
        row + 1,
        [
            "证券代码",
            "公司",
            "计算基准",
            "收益率类型",
            "基准期间",
            "每股股利",
            "股利已知日",
            "行情日",
            "当前价",
            "股息率",
            "状态",
            "阻断",
        ],
    )
    for symbol, company, snapshot in snapshot_rows:
        status = snapshot.get("status")
        blockers = snapshot.get("blockers") or []
        not_ready = status not in (None, "READY")
        _style(ws.cell(row, 1, symbol), bold=True)
        _style(ws.cell(row, 2, company))
        _style(ws.cell(row, 3, _display(snapshot.get("basis_type"))))
        _style(ws.cell(row, 4, _display(snapshot.get("yield_type"))))
        _style(ws.cell(row, 5, snapshot.get("dividend_basis_period")))
        _style(ws.cell(row, 6, _decimal(snapshot.get("dividend_per_share"))))
        _style(ws.cell(row, 7, snapshot.get("dividend_known_at")))
        _style(ws.cell(row, 8, snapshot.get("quote_date")))
        _style(ws.cell(row, 9, _decimal(snapshot.get("current_price"))))
        _style(ws.cell(row, 10, _percent(snapshot.get("dividend_yield"))))
        _style(ws.cell(row, 11, _display(status)), fill=AMBER if not_ready else "FFFFFF")
        _style(
            ws.cell(row, 12, "；".join(str(value) for value in blockers)),
            fill=AMBER if blockers else "FFFFFF",
        )
        row += 1


def _reverse_sheet(wb: Workbook, payload: Mapping[str, Any], dossiers: dict[str, Any]) -> None:
    ws = wb.create_sheet(REVERSE_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    _set_widths(ws, [10, 14, 22, 24, 14, 18, 18, 16, 18, 30])
    _title(
        ws,
        "反向估值",
        "只反转预登记参数包络；无解时报告包络外，不扩大范围制造解。",
        10,
    )
    row = _header(
        ws,
        4,
        [
            "证券代码",
            "公司",
            "模型",
            "反转驱动",
            "当前价",
            "下界",
            "上界",
            "状态",
            "包络内解",
            "阻断",
        ],
    )
    for item in payload.get("results") or []:
        symbol = str(item.get("symbol") or "")
        for reverse in item.get("reverse_valuations") or []:
            if not isinstance(reverse, Mapping):
                continue
            _style(ws.cell(row, 1, symbol), bold=True)
            _style(ws.cell(row, 2, _company_name(symbol, dossiers)))
            _style(ws.cell(row, 3, reverse.get("model_type")))
            _style(ws.cell(row, 4, reverse.get("driver")))
            _style(ws.cell(row, 5, _decimal(reverse.get("target_price"))))
            _style(ws.cell(row, 6, _decimal(reverse.get("lower_bound"))))
            _style(ws.cell(row, 7, _decimal(reverse.get("upper_bound"))))
            _style(ws.cell(row, 8, _display(reverse.get("status"))))
            _style(ws.cell(row, 9, _decimal(reverse.get("solution"))))
            _style(ws.cell(row, 10, "；".join(str(value) for value in (reverse.get("blockers") or []))))
            row += 1


def _blockers_sheet(wb: Workbook, payload: Mapping[str, Any], dossiers: dict[str, Any]) -> None:
    ws = wb.create_sheet(BLOCKERS_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    _set_widths(ws, [10, 14, 22, 100])
    _title(ws, "阻断与证据", "所有阻断保留原文字，不用汇总数量代替证据链。", 4)
    row = _header(ws, 4, ["证券代码", "公司", "来源域", "阻断"])
    for item in payload.get("results") or []:
        symbol = str(item.get("symbol") or "")
        entries: list[tuple[str, Any]] = [
            ("Application", item.get("blockers") or []),
            ("估值", (item.get("valuation") or {}).get("blockers") or []),
            ("价格桥", (item.get("price_bridge") or {}).get("blockers") or []),
            ("当前数据", (item.get("current_data_status") or {}).get("blockers") or []),
            ("股利", (item.get("distribution") or {}).get("blockers") or []),
        ]
        for reverse in item.get("reverse_valuations") or []:
            if isinstance(reverse, Mapping):
                entries.append(("反向估值", reverse.get("blockers") or []))
        for domain, blockers in entries:
            for blocker in blockers:
                _style(ws.cell(row, 1, symbol))
                _style(ws.cell(row, 2, _company_name(symbol, dossiers)))
                _style(ws.cell(row, 3, domain))
                _style(ws.cell(row, 4, str(blocker)), fill=AMBER)
                row += 1


def _sample_sheet(wb: Workbook, collection: ResearchDossierCollection) -> None:
    ws = wb.create_sheet(SAMPLE_SHEET)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    _set_widths(ws, [10, 14, 22, 24, 16, 60])
    _title(
        ws,
        "固定样本研究状态",
        "本次 Application 只计算三家公司；其余样本仍为研究档案状态。",
        6,
    )
    row = _header(ws, 4, ["证券代码", "公司", "画像", "注册模型", "档案状态", "首个研究阻断"])
    for dossier in collection.dossiers:
        first = dossier.research_gaps.blockers[0] if dossier.research_gaps.blockers else ""
        _style(ws.cell(row, 1, dossier.symbol), bold=True)
        _style(ws.cell(row, 2, dossier.name))
        _style(ws.cell(row, 3, dossier.profile_id))
        _style(ws.cell(row, 4, dossier.primary_model or "未注册"))
        _style(ws.cell(row, 5, _display(dossier.readiness)))
        _style(ws.cell(row, 6, first))
        row += 1


def build_application_workbook(
    application_payload: Mapping[str, Any],
    collection: ResearchDossierCollection,
) -> Workbook:
    if application_payload.get("schema_version") != APPLICATION_SCHEMA:
        raise ValueError("Application bundle has an unsupported schema")
    if application_payload.get("action") != ACTION_NO_ORDER:
        raise ValueError("Application workbook cannot consume a non-no_order bundle")
    if collection.action != ACTION_NO_ORDER:
        raise ValueError("Dossier collection action must remain no_order")
    dossiers = collection.by_symbol
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, application_payload, dossiers)
    _valuation_sheet(wb, application_payload, dossiers)
    _dividend_sheet(wb, application_payload, dossiers)
    _reverse_sheet(wb, application_payload, dossiers)
    _blockers_sheet(wb, application_payload, dossiers)
    _sample_sheet(wb, collection)
    wb.active = 0
    return wb


def _load_pointer_payload(
    root: Path,
    pointer_path: Path,
    *,
    expected_schema: str,
) -> dict[str, Any]:
    pointer = json.loads(pointer_path.read_text(encoding="utf-8"))
    target = (root / pointer["path"] / "evidence.json").resolve()
    if not target.is_relative_to(root.resolve()):
        raise ValueError(f"Runtime pointer escapes project root: {pointer_path}")
    expected = str(pointer.get("sha256") or "").lower()
    if not expected or hashlib.sha256(target.read_bytes()).hexdigest() != expected:
        raise ValueError(f"Runtime evidence hash changed: {pointer_path}")
    payload = json.loads(target.read_text(encoding="utf-8"))
    if payload.get("schema_version") != expected_schema:
        raise ValueError(f"Runtime payload has unsupported schema: {pointer_path}")
    return payload


def load_application_payload(root: Path = ROOT) -> dict[str, Any]:
    return _load_pointer_payload(
        root,
        root / APPLICATION_POINTER,
        expected_schema=APPLICATION_SCHEMA,
    )


def write_application_workbook(
    *,
    root: Path,
    application_payload: Mapping[str, Any],
    collection: ResearchDossierCollection,
    output: Path | None = None,
) -> dict[str, Any]:
    workbook = build_application_workbook(application_payload, collection)
    if output is None:
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        output = (
            root
            / "runtime"
            / f"m1-research-application-workbook-{timestamp}"
            / "m1-research-application-candidate.xlsx"
        )
    output = output.resolve()
    if not output.is_relative_to(root.resolve()):
        raise ValueError("Application workbook output escapes project root")
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise ValueError(f"Application workbook candidate already exists: {output}")
    workbook.save(output)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    return {
        "workbook_path": str(output.relative_to(root)),
        "workbook_sha256": digest,
        "action": ACTION_NO_ORDER,
    }
