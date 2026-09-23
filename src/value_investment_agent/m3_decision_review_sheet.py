"""Single-sheet presentation bridge for M3 decision cards.

This module produces a replacement for the derived frontend sheet
``00_决策复核``.  It consumes the same non-personal DecisionCardCollection as
the standalone M3 candidate and keeps every row fail-closed, human-review
required and ``action=no_order``.  It contains no formula, order, position or
execution column.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .decision_read_model import (
    ACTION_NO_ORDER,
    CARD_ENTRY_MISSING,
    CARD_ENTRY_NOT_REQUIRED,
    CARD_ENTRY_RECONSTRUCTED,
    CARD_JOURNAL_ABSENT,
    CARD_JOURNAL_AVAILABLE,
    CARD_PORTFOLIO_CONFIRMED,
    CARD_PORTFOLIO_MISSING,
    CARD_PORTFOLIO_PROVIDED_UNCONFIRMED,
    CARD_REASON_ENTRY_MISSING,
    CARD_REASON_HUMAN_REVIEW_REQUIRED,
    CARD_REASON_OTHER,
    CARD_REASON_PERSONAL_INPUT_MISSING,
    CARD_REASON_PRICE_UNAVAILABLE,
    CARD_REASON_RESEARCH_INCOMPLETE,
    DecisionCardCollection,
)


TARGET_SHEET = "00_决策复核"
SCHEMA_VERSION = "m3-decision-review-sheet-v1"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_STATUS_LABELS = {
    "INSUFFICIENT_RESEARCH": "研究证据不足",
    "WAIT_FOR_PRICE": "等待有效价格",
    "RESEARCH_CANDIDATE": "研究候选",
    "WATCH": "等待人工复核",
    "MANUAL_BUY_REVIEW": "待人工复核",
    "MANUAL_ADD_REVIEW": "待人工复核",
    "HOLD": "待人工复核",
    "MANUAL_REDUCE_REVIEW": "待人工复核",
    "MANUAL_EXIT_REVIEW": "待人工复核",
}
_REASON_LABELS = {
    CARD_REASON_RESEARCH_INCOMPLETE: "研究证据不足",
    CARD_REASON_PRICE_UNAVAILABLE: "价格证据不可用",
    CARD_REASON_PERSONAL_INPUT_MISSING: "个人组合输入缺失",
    CARD_REASON_ENTRY_MISSING: "原始 Entry 缺失",
    CARD_REASON_HUMAN_REVIEW_REQUIRED: "需要人工复核",
    CARD_REASON_OTHER: "其他",
}
_PORTFOLIO_LABELS = {
    CARD_PORTFOLIO_MISSING: "缺失",
    CARD_PORTFOLIO_PROVIDED_UNCONFIRMED: "已提供，未确认",
    CARD_PORTFOLIO_CONFIRMED: "已确认",
}
_ENTRY_LABELS = {
    CARD_ENTRY_NOT_REQUIRED: "本卡不需要",
    CARD_ENTRY_MISSING: "缺失",
    CARD_ENTRY_RECONSTRUCTED: "已标注事后重建",
    "AVAILABLE": "已提供",
}
_JOURNAL_LABELS = {
    CARD_JOURNAL_ABSENT: "无",
    CARD_JOURNAL_AVAILABLE: "有",
}
_MISSING_EXPLANATIONS = {
    "research_or_valuation_inputs": "研究与估值前置条件未完成，不能进入正向决策复核。",
    "verified_price_or_quote": "尚无通过模型有效期与双源校验的当前价格或报价链。",
    "portfolio_input": "用户尚未提供个人组合输入，不能形成个人化正向复核。",
    "portfolio_capacity_confirmation": "最小组合容量尚未由用户人工确认。",
    "original_entry": "缺少原始 Entry Thesis，不能形成持有、追加或退出复核。",
    "human_decision": "最终动作仍须用户明确做出人工决定。",
}


def _style(cell, *, fill="FFFFFF", bold=False, color=INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _section(ws, row, title, columns) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    _style(ws.cell(row, 1, title), fill=BLUE, bold=True, color="FFFFFF")
    ws.row_dimensions[row].height = 24
    return row + 1


def _header(ws, row, columns) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=GREY, bold=True)
    return row + 1


def _widths(ws, widths) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _name(symbol, names) -> str:
    return str(names.get(symbol) or symbol)


def _cards(ws, row, collection, names) -> int:
    columns = [
        "证券代码",
        "公司",
        "决策时点",
        "系统状态",
        "原因类别",
        "决策意图",
        "组合输入",
        "原始 Entry",
        "决策日志",
        "阻断",
    ]
    row = _header(ws, row, columns)
    for card in collection.cards:
        values = [
            card.symbol,
            _name(card.symbol, names),
            card.decision_as_of.isoformat(),
            _STATUS_LABELS.get(card.status, card.status),
            _REASON_LABELS.get(card.reason_kind, card.reason_kind),
            card.decision_intent or "未提交",
            _PORTFOLIO_LABELS.get(card.portfolio_status, card.portfolio_status),
            _ENTRY_LABELS.get(card.entry_status, card.entry_status),
            _JOURNAL_LABELS.get(card.journal_status, card.journal_status),
            "；".join(card.blockers) or "无",
        ]
        for column, value in enumerate(values, 1):
            _style(ws.cell(row, column, value), fill=AMBER if column in {4, 5, 10} else "FFFFFF")
        row += 1
    return row + 1


def _missing(ws, row, collection, names) -> int:
    columns = ["证券代码", "公司", "缺失项", "说明", "相关阻断"]
    row = _header(ws, row, columns)
    for card in collection.cards:
        if not card.missing_inputs:
            _style(ws.cell(row, 1, card.symbol))
            _style(ws.cell(row, 2, _name(card.symbol, names)))
            _style(ws.cell(row, 3, "无自动可判定缺失；仍须人工复核"), fill=GREY)
            row += 1
            continue
        for missing in card.missing_inputs:
            values = [
                card.symbol,
                _name(card.symbol, names),
                missing,
                _MISSING_EXPLANATIONS.get(missing, "缺少完成该决策复核所需输入。"),
                "；".join(card.blockers) if missing in {
                    "research_or_valuation_inputs",
                    "verified_price_or_quote",
                } else "",
            ]
            for column, value in enumerate(values, 1):
                _style(ws.cell(row, column, value), fill=AMBER if value else "FFFFFF")
            row += 1
    return row + 1


def _sources(ws, row, collection) -> int:
    columns = ["证券代码", "来源键", "制品类型", "制品 ID", "SHA-256", "可用时点"]
    row = _header(ws, row, columns)
    for card in collection.cards:
        for source in card.source_hashes:
            values = [
                card.symbol,
                source.source_key,
                source.artifact_type,
                source.artifact_id or "",
                source.sha256,
                source.available_at.isoformat() if source.available_at else "",
            ]
            for column, value in enumerate(values, 1):
                _style(ws.cell(row, column, value), fill=GREY)
            row += 1
    return row + 1


def _evidence(ws, row, collection) -> int:
    columns = ["证券代码", "证据 ID", "类型", "路径", "来源 URL", "SHA-256"]
    row = _header(ws, row, columns)
    for card in collection.cards:
        for ref in card.evidence_refs:
            values = [
                card.symbol,
                str(ref.get("id") or ""),
                str(ref.get("kind") or ref.get("provider") or ""),
                str(ref.get("path") or ref.get("location") or ""),
                str(ref.get("source_url") or ref.get("url") or ""),
                str(ref.get("sha256") or ""),
            ]
            for column, value in enumerate(values, 1):
                _style(ws.cell(row, column, value))
            row += 1
    return row + 1


def build_m3_decision_review_sheet(
    collection: DecisionCardCollection,
    *,
    security_names: Mapping[str, str] | None = None,
) -> Workbook:
    if collection.action != ACTION_NO_ORDER:
        raise ValueError("M3 decision review sheet requires action=no_order")
    if any(card.action != ACTION_NO_ORDER for card in collection.cards):
        raise ValueError("M3 decision cards must remain no_order")
    names = dict(security_names or {})
    wb = Workbook()
    ws = wb.active
    ws.title = TARGET_SHEET
    ws.sheet_view.showGridLines = False
    _widths(ws, [11, 16, 14, 18, 22, 18, 18, 20, 13, 46])

    _style(ws.cell(1, 1, "决策复核（M3 非个人化负向卡）"), fill=GREEN, bold=True, color="FFFFFF")
    _style(
        ws.cell(
            2,
            1,
            f"生成 {collection.generated_at.isoformat()} | 研究输入 "
            f"{max((card.decision_as_of.isoformat() for card in collection.cards), default="未知")}，"
            "非今日交易 | 自动决策未接通 | action=no_order",
        ),
        fill=GREY,
        bold=True,
    )
    ws.merge_cells("A1:J1")
    ws.merge_cells("A2:J2")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 32
    row = 4
    row = _section(ws, row, "当前三张决策卡", 10)
    row = _cards(ws, row, collection, names)
    row = _section(ws, row, "缺失与阻断明细", 5)
    row = _missing(ws, row, collection, names)
    row = _section(ws, row, "来源 Hash 绑定", 6)
    row = _sources(ws, row, collection)
    row = _section(ws, row, "证据引用", 6)
    row = _evidence(ws, row, collection)
    row = _section(ws, row, "人工复核", 5)
    review_rows = [
        "系统只登记研究证据不足、个人组合输入缺失和原始 Entry 缺失；不会自动形成任何正向决策。",
        "三张卡全部要求人工复核；持有、退出或继续等待的最终结论都必须由用户明确记录。",
        "来源 Hash 用于追溯 M1 集成运行、决策复核和证据包，Hash 一致只证明字节一致，不证明研究结论正确。",
        "Checkpoint B：未完成。",
    ]
    for text in review_rows:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        _style(ws.cell(row, 1, text), fill=GREY)
        row += 1
    ws.freeze_panes = "A4"
    wb.calculation.fullCalcOnLoad = True
    return wb


def write_m3_decision_review_addon(
    collection: DecisionCardCollection,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("M3 decision review addon escapes its root")
    if output.exists():
        raise ValueError(f"M3 decision review addon already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_m3_decision_review_sheet(collection, security_names=security_names)
    workbook.save(output)
    return {
        "addon_path": str(output.relative_to(root)),
        "target_sheet": TARGET_SHEET,
        "schema_version": SCHEMA_VERSION,
        "card_count": len(collection.cards),
        "positive_review_count": sum(card.is_positive_review() for card in collection.cards),
        "action": ACTION_NO_ORDER,
    }
