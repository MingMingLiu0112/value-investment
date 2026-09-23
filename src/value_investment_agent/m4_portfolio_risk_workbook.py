"""Presentation-only simulated workbook for M4 portfolio risk assessment."""
from __future__ import annotations

from decimal import Decimal
import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .investment_decision import ACTION_NO_ORDER
from .portfolio_risk import (
    ASSESSMENT_NAMESPACE_SIMULATED,
    FINDING_CYCLICAL,
    FINDING_INDUSTRY,
    FINDING_LIQUIDITY,
    FINDING_MINIMUM_CASH,
    FINDING_RESERVED_CASH,
    FINDING_SINGLE_SECURITY,
    LIQUIDITY_LIQUID,
    LIQUIDITY_RESTRICTED,
    LIQUIDITY_UNKNOWN,
    STATUS_INCOMPLETE,
    STATUS_PASS,
    STATUS_REVIEW_REQUIRED,
    STATUS_VIOLATION,
    PortfolioRiskAssessment,
)


OVERVIEW_SHEET = "00_组合风险"
HOLDINGS_SHEET = "01_持仓与集中度"
FINDINGS_SHEET = "02_风险发现"
BOUNDARY_SHEET = "03_输入与边界"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_STATUS_LABELS = {
    STATUS_PASS: "通过",
    STATUS_VIOLATION: "触发上限",
    STATUS_REVIEW_REQUIRED: "需人工复核",
    STATUS_INCOMPLETE: "输入不完整",
}
_LIQUIDITY_LABELS = {
    LIQUIDITY_LIQUID: "正常",
    LIQUIDITY_RESTRICTED: "受限",
    LIQUIDITY_UNKNOWN: "未知",
}
_FINDING_LABELS = {
    FINDING_SINGLE_SECURITY: "单股集中度",
    FINDING_INDUSTRY: "行业集中度",
    FINDING_CYCLICAL: "周期暴露",
    FINDING_MINIMUM_CASH: "最低现金",
    FINDING_RESERVED_CASH: "应急/流动性现金",
    FINDING_LIQUIDITY: "流动性复核",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_row(ws, row: int, values: list[Any], fill: str = "FFFFFF") -> None:
    for column, value in enumerate(values, 1):
        _style(ws.cell(row, column, value), fill=fill)


def _label(mapping: Mapping[str, str], value: str) -> str:
    return mapping.get(value, value)


def _name(symbol: str, names: Mapping[str, str]) -> str:
    return names.get(symbol, symbol)


def _money(value: Decimal | None) -> str:
    if value is None:
        return "缺失"
    return f"{value:,.2f}"


def _pct(value: Decimal | None) -> str:
    if value is None:
        return "缺失"
    return f"{value * Decimal('100'):.2f}%"


def _join(values: tuple[str, ...]) -> str:
    return "；".join(value for value in values if value)


def _overview(ws: Workbook, assessment: PortfolioRiskAssessment) -> None:
    sheet = ws.create_sheet(OVERVIEW_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [22, 22, 68])
    _title(
        sheet,
        "M4 组合风险与集中度（模拟演示）",
        "模拟组合仅演示计算结构，不读取真实账户，不生成订单。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    metrics = [
        ("评估命名空间", "SIMULATED", "公开仓库只允许模拟演示。"),
        ("风险状态", _label(_STATUS_LABELS, assessment.status), "基于人工确认 IPS 的上限计算。"),
        ("总资产（元）", _money(assessment.total_assets_cny()), "现金与持仓市值合计。"),
        ("现金（元）", _money(assessment.cash_cny()), "不把缺失现金填成 0。"),
        ("最低保留现金（元）", _money(assessment.reserved_cash_cny()), "最低现金与应急/流动性需要取保守值。"),
        ("单股上限", _pct(assessment.policy.max_single_security_pct), "由 IPS 人工确认，不是系统默认 20%。"),
        ("行业上限", _pct(assessment.policy.max_single_industry_pct), "同一行业组合市值权重。"),
        ("周期暴露上限", _pct(assessment.policy.max_cyclical_exposure_pct), "周期性证券组合市值权重。"),
        ("持仓数量", str(len(assessment.snapshot.holdings)), "每只证券均需人工确认数量和市值。"),
        ("动作", ACTION_NO_ORDER, "本页不产生买入、卖出、仓位或订单。"),
    ]
    for name, value, explanation in metrics:
        _write_row(sheet, row, [name, value, explanation])
        row += 1


def _holdings(
    ws: Workbook,
    assessment: PortfolioRiskAssessment,
    names: Mapping[str, str],
) -> None:
    sheet = ws.create_sheet(HOLDINGS_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "证券代码",
        "公司",
        "行业",
        "市值（元）",
        "当前权重",
        "周期",
        "流动性",
        "共同因子",
        "数量状态",
        "动作",
    ]
    _widths(sheet, [12, 18, 16, 18, 14, 10, 12, 24, 16, 12])
    _title(
        sheet,
        "持仓与集中度",
        "当前权重是组合事实；超限只提示复核。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    weights = assessment.security_weights()
    for holding in assessment.snapshot.holdings:
        attributes = assessment.security_attributes[holding.symbol]
        _write_row(
            sheet,
            row,
            [
                holding.symbol,
                _name(holding.symbol, names),
                attributes.industry,
                _money(holding.market_value_cny),
                _pct(weights[holding.symbol]),
                "是" if attributes.cyclical else "否",
                _label(_LIQUIDITY_LABELS, attributes.liquidity_profile),
                _join(attributes.common_factors),
                holding.quantity_source,
                ACTION_NO_ORDER,
            ],
        )
        row += 1


def _findings(ws: Workbook, assessment: PortfolioRiskAssessment) -> None:
    sheet = ws.create_sheet(FINDINGS_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = [
        "风险类型",
        "对象",
        "实测",
        "上限",
        "说明",
        "证据ID",
        "动作",
    ]
    _widths(sheet, [18, 16, 14, 14, 48, 30, 12])
    _title(
        sheet,
        "风险发现",
        "每个发现保留可复核的实测值与上限；流动性风险没有上限时仅要求人工复核。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for finding in assessment.findings():
        measured = (
            _pct(finding.measured)
            if finding.kind not in {FINDING_MINIMUM_CASH, FINDING_RESERVED_CASH}
            else _money(finding.measured)
        )
        limit = (
            _pct(finding.limit)
            if finding.kind not in {FINDING_MINIMUM_CASH, FINDING_RESERVED_CASH}
            else _money(finding.limit)
        )
        _write_row(
            sheet,
            row,
            [
                _label(_FINDING_LABELS, finding.kind),
                finding.subject,
                measured,
                limit,
                finding.message,
                "；".join(ref.get("id", "") for ref in finding.evidence_refs),
                ACTION_NO_ORDER,
            ],
            fill=AMBER if finding.kind == FINDING_LIQUIDITY else "FFFFFF",
        )
        row += 1


def _boundary(ws: Workbook, assessment: PortfolioRiskAssessment) -> None:
    sheet = ws.create_sheet(BOUNDARY_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["边界项", "值/说明", "证据ID"]
    _widths(sheet, [26, 64, 30])
    _title(
        sheet,
        "输入与边界",
        "所有输入均为模拟；真实 IPS/持仓未提供时不会进入本公开工作簿。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    boundaries = [
        ("模拟组合标识", "SIMULATED_PUBLIC_DEMONSTRATION", "fixture:simulated-portfolio-v1"),
        ("IPS 版本", assessment.policy.policy_version, "；".join(ref.get("id", "") for ref in assessment.policy.evidence_refs)),
        ("组合快照版本", assessment.snapshot.snapshot_version, "；".join(ref.get("id", "") for ref in assessment.snapshot.evidence_refs)),
        ("账户范围", assessment.snapshot.account_scope, "不读取真实账户"),
        ("最低现金", _money(assessment.policy.minimum_cash_cny), "policy.minimum_cash_cny"),
        ("应急现金", _money(assessment.policy.emergency_cash_cny), "policy.emergency_cash_cny"),
        ("流动性需要", _money(assessment.policy.liquidity_needs_cny), "policy.liquidity_needs_cny"),
        ("单股上限", _pct(assessment.policy.max_single_security_pct), "policy.max_single_security_pct"),
        ("行业上限", _pct(assessment.policy.max_single_industry_pct), "policy.max_single_industry_pct"),
        ("周期上限", _pct(assessment.policy.max_cyclical_exposure_pct), "policy.max_cyclical_exposure_pct"),
        ("投资期限（年）", str(assessment.policy.time_horizon_years), "policy.time_horizon_years"),
        ("风险承受口径", assessment.policy.risk_tolerance, "policy.risk_tolerance"),
        ("限制项", _join(assessment.policy.restrictions), "policy.restrictions"),
        ("缺失输入", _join(assessment.missing_inputs()), "空表示模拟结构完整"),
        ("动作", ACTION_NO_ORDER, "不生成订单"),
    ]
    for name, value, evidence in boundaries:
        _write_row(sheet, row, [name, value, evidence])
        row += 1


def build_portfolio_risk_workbook(
    assessment: PortfolioRiskAssessment,
    *,
    security_names: Mapping[str, str] | None = None,
) -> Workbook:
    if assessment.assessment_namespace != ASSESSMENT_NAMESPACE_SIMULATED:
        raise ValueError("Public risk workbook requires a simulated assessment")
    if assessment.action != ACTION_NO_ORDER:
        raise ValueError("Risk workbook requires action=no_order")
    if not assessment.can_assess():
        raise ValueError("Risk workbook requires structurally complete simulated inputs")
    names = dict(security_names or {})
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, assessment)
    _holdings(wb, assessment, names)
    _findings(wb, assessment)
    _boundary(wb, assessment)
    wb.active = 0
    return wb


def write_portfolio_risk_workbook(
    assessment: PortfolioRiskAssessment,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Risk workbook output escapes its root")
    if output.exists():
        raise ValueError(f"Risk workbook already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_portfolio_risk_workbook(
        assessment,
        security_names=security_names,
    )
    workbook.save(output)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    return {
        "workbook_path": str(output.relative_to(root)),
        "workbook_sha256": digest,
        "sheet_count": len(workbook.sheetnames),
        "holding_count": len(assessment.snapshot.holdings),
        "finding_count": len(assessment.findings()),
        "assessment_namespace": assessment.assessment_namespace,
        "status": assessment.status,
        "action": ACTION_NO_ORDER,
    }
