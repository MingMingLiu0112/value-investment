"""Presentation-only simulated workbook for the M4/M5 joint checkpoint."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .dividend_income_projection import PortfolioDividendIncomeProjection
from .investment_decision import ACTION_NO_ORDER
from .m4_m5_integration import (
    M4M5IntegrationResult,
    STATUS_NEGATIVE,
    STATUS_PARTIAL,
    STATUS_PAUSED,
    STATUS_READY,
)
from .portfolio_risk import PortfolioRiskAssessment
from .position_guidance import PositionGuidanceResult


OVERVIEW_SHEET = "00_总览"
M4_SHEET = "01_M4组合基线"
EVENT_SHEET = "02_M5事件批"
INVALIDATION_SHEET = "03_依赖失效映射"
ARTIFACT_SHEET = "04_联合产品状态"
BOUNDARY_SHEET = "05_输入与边界"

INK = "24312D"
GREEN = "18755D"
BLUE = "245D83"
AMBER = "FFF1D6"
GREY = "EFF3F1"
RED = "C0392B"

_STATUS_LABELS = {
    STATUS_READY: "READY",
    STATUS_PARTIAL: "PARTIAL",
    STATUS_PAUSED: "PAUSED",
    STATUS_NEGATIVE: "NEGATIVE",
}
_STATUS_FILL = {
    STATUS_READY: "E7F4EC",
    STATUS_PARTIAL: AMBER,
    STATUS_PAUSED: "FCE4E4",
    STATUS_NEGATIVE: RED,
}
_EVENT_LABELS = {
    "NEW_FINANCIAL_REPORT": "新财报",
    "MATERIAL_ANNOUNCEMENT": "重要公告",
    "DIVIDEND_CHANGE": "分红变化",
    "CAPITAL_ALLOCATION_CHANGE": "资本配置变化",
    "VALUATION_ZONE_CHANGED": "估值区间变化",
    "PRICE_ATTRACTIVENESS_CHANGED": "价格吸引力变化",
    "THESIS_WEAKENED": "论点弱化",
    "THESIS_BREAKER_TRIGGERED": "论点破坏触发",
    "MODEL_STALE": "模型过期",
    "POSITION_RISK_CHANGED": "仓位风险变化",
    "SOURCE_SCAN_FAILED": "数据源扫描失败",
    "SOURCE_SCAN_RECOVERED": "数据源扫描恢复",
}
_SEVERITY_LABELS = {
    "CRITICAL": "严重",
    "HIGH": "高",
    "MEDIUM": "中",
    "LOW": "低",
}
_KIND_LABELS = {
    "portfolio_risk": "组合风险",
    "position_guidance": "仓位边界",
    "distribution_history": "分配历史",
    "dividend_sustainability": "股息可持续性",
    "financial_facts": "财务事实",
    "valuation_inputs": "估值输入",
    "model_validity": "模型有效性",
    "research_thesis": "研究论点",
    "decision_review": "决策复核",
    "entry_consistency": "Entry一致性",
    "current_research_status": "当前研究状态",
}


def _style(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = INK) -> None:
    cell.font = Font(name="Microsoft YaHei", size=11, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _title(ws, title: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    _style(ws.cell(1, 1, title), fill=GREEN, bold=True, color="FFFFFF")
    _style(ws.cell(2, 1, subtitle), fill=GREY, bold=True)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 34


def _header(ws, row: int, columns: list[str]) -> int:
    for column, value in enumerate(columns, 1):
        _style(ws.cell(row, column, value), fill=BLUE, bold=True, color="FFFFFF")
    return row + 1


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_row(ws, row: int, values: list[Any], fill: str = "FFFFFF") -> None:
    for column, value in enumerate(values, 1):
        _style(ws.cell(row, column, value), fill=fill)


def _join(values: tuple[str, ...] | list[str]) -> str:
    return "；".join(str(value) for value in values if value)


def _name(symbol: str, names: Mapping[str, str]) -> str:
    if symbol in {"SYSTEM", "*"}:
        return "组合整体"
    return names.get(symbol, symbol)


def _overview(
    wb: Workbook,
    result: M4M5IntegrationResult,
    risk: PortfolioRiskAssessment,
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
) -> None:
    sheet = wb.create_sheet(OVERVIEW_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    columns = ["项目", "数值", "解释"]
    _widths(sheet, [22, 24, 68])
    _title(
        sheet,
        "M4/M5 联合检查点（模拟演示）",
        "把 M5 事件失效精确投影到 M4 组合基线；不重算仓位，也不生成订单。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    joint_status = "READY"
    if any(item.status == STATUS_NEGATIVE for item in result.artifacts):
        joint_status = STATUS_NEGATIVE
    elif any(item.status == STATUS_PAUSED for item in result.artifacts):
        joint_status = STATUS_PAUSED
    elif any(item.status == STATUS_PARTIAL for item in result.artifacts):
        joint_status = STATUS_PARTIAL
    metrics = [
        ("命名空间", "SIMULATED", "公开工作簿只接受显式模拟输入。"),
        ("联合状态", _STATUS_LABELS[joint_status], "负向事件优先于暂停，暂停优先于未完成基线。"),
        ("M4组合风险基线", risk.status, "只读展示 M4 已计算状态。"),
        ("M4仓位边界基线", guidance.status, "只输出复核条件与上限。"),
        ("M4股息收入基线", dividend.status, "已到账、已宣告、Forward、Normalized 分开。"),
        ("已接受事件", len(result.receipt.active_events), "重复、更正或未来事件按 M5 规则处理。"),
        ("受影响产品", result.affected_artifact_count(), "PAUSED 或 NEGATIVE 的联合产品数。"),
        ("M5运行状态", result.receipt.health_status, "有需人工复核提醒时不为静默健康。"),
        ("动作", ACTION_NO_ORDER, "本候选不产生交易或私人投资建议。"),
    ]
    for item in metrics:
        _write_row(sheet, row, list(item))
        row += 1


def _m4_baseline(
    wb: Workbook,
    risk: PortfolioRiskAssessment,
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
) -> None:
    sheet = wb.create_sheet(M4_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["产品", "基线状态", "缺失或阻断", "说明"]
    _widths(sheet, [20, 20, 34, 64])
    _title(
        sheet,
        "M4 组合基线",
        "这些是事件批运行前的已计算状态；本页不重新计算或修正。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("组合风险", risk.status, _join(risk.missing_inputs()), "集中度、现金储备和流动性边界。"),
        ("仓位边界", guidance.status, _join(guidance.missing_inputs()), "分层上限与共同预算。"),
        ("股息收入", dividend.status, _join(dividend.blockers()), "四类收入口径与目标缺口。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def _events(
    wb: Workbook,
    result: M4M5IntegrationResult,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(EVENT_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["来源ID", "公司", "事件", "严重度", "状态", "原因", "人工复核"]
    _widths(sheet, [24, 18, 18, 12, 16, 52, 12])
    _title(
        sheet,
        "M5 事件批",
        "事件账已由 M5 run-once 协调器接受；本页不发送通知或修改生产状态。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for ingest in result.receipt.ingest_results:
        event = ingest.event
        if event is None:
            continue
        _write_row(
            sheet,
            row,
            [
                event.source_event_id,
                _name(event.symbol, names),
                _EVENT_LABELS.get(event.event_type, event.event_type),
                _SEVERITY_LABELS.get(event.severity, event.severity),
                ingest.status,
                event.reason,
                "是" if event.requires_human_review else "否",
            ],
        )
        row += 1


def _invalidations(
    wb: Workbook,
    result: M4M5IntegrationResult,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(INVALIDATION_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["事件", "公司", "失效节点", "产品", "深度", "原因"]
    _widths(sheet, [22, 18, 26, 24, 10, 42])
    _title(
        sheet,
        "M5 依赖失效映射",
        "只有事件政策命中的节点进入此表；未命中的产品保持原状态。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for invalidation in result.receipt.invalidations:
        for item in invalidation.affected_nodes:
            _write_row(
                sheet,
                row,
                [
                    _EVENT_LABELS.get(invalidation.event_type, invalidation.event_type),
                    _name(invalidation.symbol, names),
                    item.node_id,
                    _KIND_LABELS.get(item.kind, item.kind),
                    item.depth,
                    item.reason,
                ],
            )
            row += 1


def _artifacts(
    wb: Workbook,
    result: M4M5IntegrationResult,
    names: Mapping[str, str],
) -> None:
    sheet = wb.create_sheet(ARTIFACT_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["产品", "类型", "范围", "基线", "失效节点", "触发事件", "联合状态"]
    _widths(sheet, [24, 22, 18, 16, 30, 32, 16])
    _title(
        sheet,
        "M4/M5 联合产品状态",
        "READY 表示无失效且基线可用；PAUSED 等待有界重算；NEGATIVE 为论点破坏。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    for artifact in result.artifacts:
        _write_row(
            sheet,
            row,
            [
                artifact.artifact_id,
                _KIND_LABELS.get(artifact.kind, artifact.kind),
                _name(artifact.symbol, names),
                _STATUS_LABELS[artifact.baseline_status],
                _join(artifact.invalidated_node_ids),
                _join(artifact.triggering_event_ids),
                _STATUS_LABELS[artifact.status],
            ],
            fill=_STATUS_FILL[artifact.status],
        )
        row += 1


def _boundaries(wb: Workbook) -> None:
    sheet = wb.create_sheet(BOUNDARY_SHEET)
    sheet.sheet_view.showGridLines = False
    columns = ["边界", "状态", "说明"]
    _widths(sheet, [22, 18, 70])
    _title(
        sheet,
        "输入与边界",
        "联合候选仅用于 Checkpoint C 的离线工程演示；真实组合仍需用户确认。",
        len(columns),
    )
    row = _header(sheet, 4, columns)
    rows = [
        ("命名空间", "SIMULATED", "不得导入或混用真实 IPS、持仓或账户信息。"),
        ("仓位语义", "仅上限/复核", "不输出仓位数值、加仓数量或交易指令。"),
        ("事件依赖", "有界", "超过重算上限记录 deferred，不静默扩大范围。"),
        ("通知投递", "未启用", "只维护 outbox 状态，不发送真实通知。"),
        ("生产调度", "未授权", "不修改服务器、PTA、数据库或计划任务。"),
        ("动作", ACTION_NO_ORDER, "所有买卖持有判断仍由人工完成。"),
    ]
    for item in rows:
        _write_row(sheet, row, list(item))
        row += 1


def build_m4_m5_integration_workbook(
    result: M4M5IntegrationResult,
    risk: PortfolioRiskAssessment,
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
    *,
    security_names: Mapping[str, str] | None = None,
) -> Workbook:
    names = dict(security_names or {})
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, result, risk, guidance, dividend)
    _m4_baseline(wb, risk, guidance, dividend)
    _events(wb, result, names)
    _invalidations(wb, result, names)
    _artifacts(wb, result, names)
    _boundaries(wb)
    return wb


def write_m4_m5_integration_workbook(
    result: M4M5IntegrationResult,
    risk: PortfolioRiskAssessment,
    guidance: PositionGuidanceResult,
    dividend: PortfolioDividendIncomeProjection,
    *,
    output: Path,
    root: Path,
    security_names: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    output = output.resolve()
    root = root.resolve()
    if not output.is_relative_to(root):
        raise ValueError("Workbook output must remain inside the project root")
    if output.exists():
        raise ValueError("M4/M5 integration candidate output already exists")
    workbook = build_m4_m5_integration_workbook(
        result,
        risk,
        guidance,
        dividend,
        security_names=security_names,
    )
    workbook.save(output)
    return {
        "workbook_path": str(output),
        "workbook_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "sheet_count": len(workbook.sheetnames),
        "artifact_count": len(result.artifacts),
        "affected_artifact_count": result.affected_artifact_count(),
        "event_count": len(result.receipt.active_events),
        "invalidation_count": len(result.receipt.invalidations),
        "namespace": result.namespace,
        "action": result.action,
    }
