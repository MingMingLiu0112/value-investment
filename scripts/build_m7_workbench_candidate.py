#!/usr/bin/env python3
"""Build a protected M7 read-only workbench display candidate.

This packages the already verified M3 history-overlay workbook with renamed
M4/M5 read-model candidates and one navigation overview.  It is display
preparation only: it never publishes canonical, imports a real IPS/portfolio,
changes any source sheet, or emits an order.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import runpy
import shutil
import sys
import tempfile
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


STAGE_FRONTEND = runpy.run_path(
    str(ROOT / "scripts" / "stage_frontend_package.py"),
    run_name="stage_frontend_package",
)


MANIFEST_SCHEMA = "m7-workbench-candidate-v1"
DEFAULT_BASE = (
    ROOT / "A股价值投资_Agent前端智能跟踪模板_M3历史链叠加候选_20260924.xlsx"
)
DEFAULT_OUTPUT = (
    ROOT / "A股价值投资_Agent前端智能跟踪模板_M7统一工作台候选_20260924.xlsx"
)
DEFAULT_CANONICAL = ROOT / "A股价值投资_Agent前端智能跟踪模板.xlsx"
DEFAULT_GENERATED_AT = datetime(2026, 9, 24, 0, 0, tzinfo=timezone.utc)

BASE_SHA256 = (
    "67e720f2326443bb3d36003db707a86169483bcd2f2be10a97dbda6d3bfacd4d"
)
CANONICAL_SHA256 = (
    "64c8deff1a237076d2ba0b00afc8905d23bd9d117cb132dfc6757071b5659911"
)

# Addons are listed in the desired final workbook order.  The cumulative graft
# helper inserts each new batch at the front, so they are processed in reverse.
ADDON_SPECS = (
    {
        "id": "m4_portfolio_risk",
        "path": ROOT / "A股价值投资_M4组合风险候选_20260924.xlsx",
        "sha256": "0594db6981a78063883271cd2fa45e117487fe6a9657a97e14665dc6bf8b07a7",
        "prefix": "M4风险_",
        "group": "M4组合风险",
        "overview_sheet": "M4风险_00_组合风险",
    },
    {
        "id": "m4_guidance_income",
        "path": ROOT / "A股价值投资_M4仓位与股息候选_20260924.xlsx",
        "sha256": "764f8d201dfc798012a6e27f9080d927d2b6f7b0ada6bb53bf947c6a5ff2e45e",
        "prefix": "M4仓位_",
        "group": "M4仓位与股息",
        "overview_sheet": "M4仓位_00_总览",
    },
    {
        "id": "m5_event_infrastructure",
        "path": ROOT / "A股价值投资_M5事件监控候选_20260924.xlsx",
        "sha256": "2b86953f793df46e199c40c614f3291e19b249cc0e53e2a670b0000506403dae",
        "prefix": "M5事件_",
        "group": "M5事件监控",
        "overview_sheet": "M5事件_00_总览",
    },
    {
        "id": "m5_materiality_bridge",
        "path": ROOT / "A股价值投资_M5材料性接入候选_20260924.xlsx",
        "sha256": "e976e330ae517f06ddd341220ce71fb9b6c0753ff4c7f421ef39baed5e9ce1df",
        "prefix": "M5材料性_",
        "group": "M5材料性接入",
        "overview_sheet": "M5材料性_00_总览",
    },
    {
        "id": "m5_disclosure_queue",
        "path": ROOT / "A股价值投资_M5真实披露待复核队列_20260924.xlsx",
        "sha256": "58b16bf00dd7ea57ee9cdcd6d7d7d00d80c0fc9669cd047f5171f500a9b16ec5",
        "prefix": "M5披露队列_",
        "group": "M5真实披露待复核队列",
        "overview_sheet": "M5披露队列_00_总览",
    },
    {
        "id": "m5_disclosure_review",
        "path": ROOT / "A股价值投资_M5真实披露人工复核回填_20260924.xlsx",
        "sha256": "1ae75325fe02c93011201c3a44af73d739a49680e24af35a8f33fcd362a6c420",
        "prefix": "M5披露复核_",
        "group": "M5真实披露人工复核回填",
        "overview_sheet": "M5披露复核_00_总览",
    },
)

FORBIDDEN_OVERVIEW_TEXT = (
    "买入",
    "加仓",
    "减仓",
    "目标仓位",
    "BUY",
    "ADD",
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _verify_pinned(path: Path, expected: str, label: str) -> None:
    actual = digest(path)
    if actual != expected:
        raise ValueError(f"{label} changed: expected {expected}, got {actual}")


def _inside(root: Path, path: Path, field: str) -> Path:
    target = path.resolve()
    if not target.is_relative_to(root.resolve()):
        raise ValueError(f"{field} must stay under project root: {target}")
    return target


def build_overview(
    output: Path,
    *,
    generated_at: datetime,
    destinations: tuple[tuple[str, str], ...],
) -> dict:
    """Create the one-page navigation sheet with neutral status text."""
    if output.exists():
        raise ValueError(f"M7 overview already exists: {output}")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "00_M7总览"
    sheet.freeze_panes = "A5"
    sheet.sheet_view.zoomScale = 90
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 68
    sheet.column_dimensions["C"].width = 22

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9E2F3")
    title_font = Font(bold=True, color="FFFFFF", size=14)
    section_font = Font(bold=True, color="1F4E78")
    link_font = Font(color="0563C1", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")

    sheet["A1"] = "M7 统一工作台候选（只读展示准备）"
    sheet["A1"].fill = title_fill
    sheet["A1"].font = title_font
    sheet["A2"] = (
        "只叠加已完成保护的 M3/M4/M5 只读候选；simulated/no_order；"
        "不替代 Checkpoint A-D，不构成真实组合或投资建议。"
    )
    sheet["A2"].alignment = wrap

    status_rows = (
        ("候选类型", "受保护展示候选", "read_model"),
        ("M2 机器门", "AC1-AC7、AC11 DONE；AC8-AC10、AC12 等待用户复核", "PENDING_HUMAN_REVIEW"),
        ("M3 决策复核", "只读候选已形成；Checkpoint B 等待用户复核", "PARTIAL"),
        ("M4 组合能力", "显式模拟；真实 IPS/持仓授权未提供", "PARTIAL"),
        ("M5 事件能力", "模拟/真实披露队列已形成；生产调度未授权", "PARTIAL"),
        ("M6 运营准入", "未开始；不把展示叠加当作运营证据", "NOT_STARTED"),
        ("M7 最终交付", "仅展示准备；Checkpoint D 不提前签收", "PARTIAL"),
        ("动作边界", "no_order；人工确认；不自动执行", "no_order"),
    )
    row = 4
    for column, value in enumerate(("项目", "当前事实", "状态"), start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.fill = section_fill
        cell.font = section_font
    for item in status_rows:
        row += 1
        for column, value in enumerate(item, start=1):
            sheet.cell(row=row, column=column, value=value).alignment = wrap

    row += 2
    sheet.cell(row=row, column=1, value="入口导航").fill = section_fill
    sheet.cell(row=row, column=1).font = section_font
    row += 1
    for column, value in enumerate(("入口", "查看内容", "工作表"), start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.fill = section_fill
        cell.font = section_font
    for label, target in destinations:
        row += 1
        link = sheet.cell(row=row, column=1, value=label)
        link.hyperlink = f"'{target}'!A1"
        link.font = link_font
        sheet.cell(row=row, column=2, value=f"打开 {target}").alignment = wrap
        sheet.cell(row=row, column=3, value=target).alignment = wrap

    row += 2
    sheet.cell(row=row, column=1, value="边界说明").fill = section_fill
    sheet.cell(row=row, column=1).font = section_font
    for offset, value in enumerate(
        (
            "所有上游候选均保持只读；本页只负责统一入口，不重算研究、估值、仓位或事件结论。",
            "真实 M7 交付仍需 M2/M3 用户复核、真实 IPS/组合授权、M5/M6 生产观察与 Checkpoint D。",
            "生成时间固定为 {generated_at}；源文件均以 SHA-256 绑定。".format(
                generated_at=generated_at.isoformat()
            ),
        ),
        start=1,
    ):
        cell = sheet.cell(row=row + offset, column=1, value=value)
        cell.alignment = wrap

    text = " ".join(
        str(cell.value)
        for sheet_row in sheet.iter_rows()
        for cell in sheet_row
        if cell.value is not None
    )
    if any(item in text for item in FORBIDDEN_OVERVIEW_TEXT):
        raise ValueError("M7 overview contains forbidden decision presentation text")

    workbook.save(output)
    return {
        "overview_sha256": digest(output),
        "overview_sheets": ["00_M7总览"],
        "overview_destinations": [item[1] for item in destinations],
    }


def build_candidate(
    *,
    base: Path,
    canonical: Path | None,
    output: Path,
    expected_base_sha256: str,
    expected_canonical_sha256: str | None,
    addon_specs: tuple[dict, ...],
    generated_at: datetime,
    project_root: Path | None = None,
) -> dict:
    """Overlay all candidates, preserving every source worksheet byte."""
    project_root = (project_root or ROOT).resolve()
    base = _inside(project_root, base, "base")
    output = _inside(project_root, output, "output")
    if canonical is not None:
        canonical = _inside(project_root, canonical, "canonical")
        if expected_canonical_sha256 is None:
            raise ValueError("canonical hash is required when canonical path is set")
        _verify_pinned(canonical, expected_canonical_sha256, "canonical")
    if output.exists():
        raise ValueError(f"M7 workbench candidate already exists: {output}")
    _verify_pinned(base, expected_base_sha256, "M7 base")

    addons = tuple(addon_specs)
    for spec in addons:
        for field in ("id", "path", "sha256", "prefix", "group", "overview_sheet"):
            if field not in spec:
                raise ValueError(f"M7 addon {spec.get('id')!r} is missing {field}")
        _inside(project_root, Path(spec["path"]), f"{spec['id']} path")
        _verify_pinned(Path(spec["path"]), spec["sha256"], spec["id"])

    runtime_dir = project_root / "runtime"
    runtime_dir.mkdir(parents=True, exist_ok=True)
    build_dir = Path(tempfile.mkdtemp(prefix="m7-workbench-build-", dir=runtime_dir))
    current = base
    current_sha256 = expected_base_sha256
    layer_receipts = []
    destinations = [
        (spec["group"], spec["overview_sheet"])
        for spec in addons
    ] + [
        ("M3 论点连续性历史链", "00_历史链"),
        ("M3 决策复核", "00_决策复核"),
        ("原投资工作台", "00_投资工作台"),
    ]
    try:
        for index, spec in enumerate(reversed(addons)):
            source = Path(spec["path"])
            renamed = build_dir / f"renamed-{spec['id']}.xlsx"
            rename_receipt = STAGE_FRONTEND["rename_workbook_sheets"](
                source,
                renamed,
                spec["sha256"],
                prefix=spec["prefix"],
            )
            next_candidate = build_dir / f"stage-{index}.xlsx"
            graft_receipt = STAGE_FRONTEND["graft"](
                current,
                renamed,
                next_candidate,
                current_sha256,
            )
            layer_receipts.append(
                {
                    "id": spec["id"],
                    "group": spec["group"],
                    "source": str(source.relative_to(project_root)),
                    "source_sha256": spec["sha256"],
                    "renamed_candidate_sha256": rename_receipt["candidate_sha256"],
                    "renamed_sheets": rename_receipt["renamed_sheets"],
                    "new_sheets": graft_receipt["new_sheets"],
                    "original_sheets_preserved": graft_receipt["original_sheets_preserved"],
                }
            )
            current = next_candidate
            current_sha256 = graft_receipt["candidate_sha256"]

        overview = build_dir / "m7-overview.xlsx"
        overview_receipt = build_overview(
            overview,
            generated_at=generated_at,
            destinations=tuple(destinations),
        )
        final_graft = STAGE_FRONTEND["graft"](
            current,
            overview,
            output,
            current_sha256,
        )

        with ZipFile(output) as package:
            _, sheets = STAGE_FRONTEND["sheets"](package)
            final_sheet_order = [sheet.get("name") for sheet, _path in sheets]
        candidate_sha256 = digest(output)
        manifest = {
            "schema_version": MANIFEST_SCHEMA,
            "generated_at": generated_at.isoformat(),
            "action": "no_order",
            "presentation_namespace": "read_model_candidate",
            "status": "candidate_verified_not_published",
            "human_review_required": True,
            "base": str(base.relative_to(project_root)),
            "base_sha256": expected_base_sha256,
            "canonical": str(canonical.relative_to(project_root)) if canonical else None,
            "canonical_sha256": expected_canonical_sha256,
            "candidate": str(output.relative_to(project_root)),
            "candidate_sha256": candidate_sha256,
            "sheet_count": len(final_sheet_order),
            "final_sheet_order": final_sheet_order,
            "layers": list(reversed(layer_receipts)),
            "overview": overview_receipt,
            "final_original_sheets_preserved": final_graft["original_sheets_preserved"],
            "final_original_parts_unchanged": final_graft["original_parts_unchanged"],
            "checkpoints": {
                "m2": "PENDING_HUMAN_REVIEW",
                "m3": "PARTIAL",
                "m4": "PARTIAL",
                "m5": "PARTIAL",
                "m6": "NOT_STARTED",
                "m7": "PARTIAL",
                "final": "PENDING_HUMAN_REVIEW",
            },
        }
        manifest_path = output.with_name(output.stem + ".candidate.manifest.json")
        if manifest_path.exists():
            raise ValueError(f"M7 workbench manifest already exists: {manifest_path}")
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        manifest["manifest_path"] = str(manifest_path.relative_to(project_root))
        manifest["manifest_sha256"] = digest(manifest_path)
        return manifest
    finally:
        shutil.rmtree(build_dir, ignore_errors=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE)
    parser.add_argument("--canonical", type=Path, default=DEFAULT_CANONICAL)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--expected-base-sha256", default=BASE_SHA256)
    parser.add_argument("--expected-canonical-sha256", default=CANONICAL_SHA256)
    parser.add_argument(
        "--generated-at",
        type=lambda value: datetime.fromisoformat(value).astimezone(timezone.utc),
        default=DEFAULT_GENERATED_AT,
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    result = build_candidate(
        base=args.base,
        canonical=args.canonical,
        output=args.output,
        expected_base_sha256=args.expected_base_sha256,
        expected_canonical_sha256=args.expected_canonical_sha256,
        addon_specs=ADDON_SPECS,
        generated_at=args.generated_at,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
