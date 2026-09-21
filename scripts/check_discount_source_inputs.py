"""Extract original author inputs; candidate evidence, not company cost of capital."""
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
import hashlib
import json
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]


def main():
    source=ROOT/'runtime/valuation-research/20260908T075333929647Z/fcffsimpleginzu.xlsx'
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    if digest!='d6ffb67d965dc22463e4d013636befab3da431e7fc608936a82612fa73198e72':
        raise ValueError('Author workbook changed')
    book=load_workbook(source,read_only=True,data_only=False)
    country=book['Country equity risk premiums']
    industry=book['Industry Averages (Global)']
    if country['A42'].value!='China' or country['D42'].value!='=$B$1+E42':
        raise ValueError('Country identity or ERP formula changed')
    if country['C1'].value!='Updated January 1, 2026':
        raise ValueError('Country vintage changed')
    if industry['A10'].value!='Beverage (Alcoholic)' or industry['G1'].value!='Unlevered Beta':
        raise ValueError('Industry identity or beta type changed')
    base=Decimal(str(country['B1'].value))
    premium=Decimal(str(country['E42'].value))
    values={}
    for sheet,coords in [('Country equity risk premiums',['B1','C1','A42','B42','C42','D42','E42','F42']),
        ('Industry Averages (Global)',['A10','B10','G10','H10','L10','M10','N10','T10','U10','V10'])]:
        values[sheet]={cell:book[sheet][cell].value for cell in coords}
    formulas={cell.coordinate:cell.value for row in book['Cost of capital worksheet'].iter_rows(min_row=15,max_row=62)
              for cell in row[:5] if cell.value is not None}
    book.close()
    out=ROOT/'runtime/valuation-research'/('discount-candidates-'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ'))
    out.mkdir(parents=True,exist_ok=False)
    result={'source_url':'https://pages.stern.nyu.edu/~adamodar/pc/fcffsimpleginzu.xlsx',
        'source_sha256':digest,'cells':values,'cost_of_capital_formulas':formulas,
        'china_total_erp_recomputed':str(base+premium),
        'country_table_vintage':'2026-01-01','industry_table_vintage':'not separately confirmed',
        'company_wacc_approved':False,'company_beta_approved':False,
        'remaining_inputs':['Dated nominal CNY risk-free reference and credit-risk convention',
            'Issuer geographic operating exposure and country-premium treatment',
            'Comparable beta composition, leverage and cash adjustment',
            'Matched industrial debt/equity weights and borrowing cost',
            'Independent methodological corroboration and sensitivity registration'],
        'limitations':['China total ERP already contains country premium; do not add it again',
            'Country default spread is not automatically a CNY sovereign-yield deduction',
            'Global alcoholic industry average is not a baijiu-specific or issuer beta',
            'Global table WACC and default workbook 9% are not issuer parameter approval',
            'Workbook numbers cannot be backfilled into 2015 historical estimates']}
    raw=json.dumps(result,ensure_ascii=False,indent=2).encode('utf-8')
    (out/'evidence.json').write_bytes(raw)
    (out/'manifest.json').write_text(json.dumps({'evidence_sha256':hashlib.sha256(raw).hexdigest(),
        'source_sha256':digest,'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest()},indent=2),encoding='utf-8')
    print(json.dumps({'output':str(out),'china_total_erp':str(base+premium),
        'global_unlevered_beta':values['Industry Averages (Global)']['G10'],'company_wacc_approved':False}))


if __name__=='__main__':
    main()
