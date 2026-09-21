#!/usr/bin/env python3
"""Add the latest 600519 paper-execution status to an existing workbook copy."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from value_investment_agent.company_research_export import load_current_execution_contract_status


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def row_for_symbol(sheet, symbol: str) -> int:
    rows = [cell.row for cell in sheet["A"] if str(cell.value).zfill(6) == symbol]
    if len(rows) != 1:
        raise ValueError(f"Expected exactly one {symbol} row in {sheet.title}")
    return rows[0]


def prepend(cell, status: str) -> bool:
    old = "" if cell.value is None else str(cell.value)
    if status in old:
        return False
    value = status + old
    if len(value) > 32767:
        raise ValueError("Execution status would exceed Excel cell text limit")
    cell.value = value
    return True


def build(source: Path, output: Path) -> dict:
    source, output = source.resolve(), output.resolve()
    if source == output or source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("Source must be an existing, distinct workbook")
    status = load_current_execution_contract_status(ROOT)
    shutil.copy2(source, output)
    workbook = load_workbook(output)
    research = workbook["09_公司研究"]
    research_row = row_for_symbol(research, "600519")
    decision = workbook["21_决策验证"]
    decision_row = row_for_symbol(decision, "600519")
    headers = [cell.value for cell in decision[3]]
    try:
        decision_column = headers.index("策略验证状态") + 1
    except ValueError as error:
        raise ValueError("Decision validation status column is missing") from error
    changed = []
    if prepend(research.cell(research_row, 16), status):
        changed.append(f"09_公司研究!P{research_row}")
    if prepend(decision.cell(decision_row, decision_column), status):
        changed.append(f"21_决策验证!{decision.cell(decision_row, decision_column).coordinate}")
    workbook.save(output)
    return {"source_sha256": digest(source), "candidate_sha256": digest(output),
            "status": status, "changed_cells": changed}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    args = parser.parse_args()
    result = build(args.source, args.output)
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
