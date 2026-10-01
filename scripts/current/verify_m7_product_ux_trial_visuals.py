#!/usr/bin/env python3
"""Write a visual-review receipt from a WPS-exported product trial workbook."""
from __future__ import annotations

import argparse
from hashlib import sha256
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook

try:
    import fitz
except ImportError:  # pragma: no cover - the workspace runtime provides PyMuPDF
    fitz = None


ROOT = Path(__file__).resolve().parents[2]
USER_SHEETS = ("01_今日", "02_机会", "决策过程", "03_公司", "04_我的组合", "05_事件")
PRIMARY_SHEETS = ("01_今日", "02_机会", "03_公司", "04_我的组合", "05_事件")
SECONDARY_SHEETS = ("06_系统与审计",)
ALL_SHEETS = (*USER_SHEETS, *SECONDARY_SHEETS)
FORBIDDEN_USER_TOKENS = (
    "action=no_order",
    "BLOCKED",
    "NOT_READY",
    "UNKNOWN",
    "SHA256",
    "SHA-256",
    "TEXT_EXTRACTED_NOT_SEMANTICALLY_VERIFIED",
    "RESEARCH_",
    "估值批准为 false",
    "runtime/",
    "PIT",
    "unknown",
    "issuer-specific beta",
    "CNY",
    "false",
)
FORMULA_ERROR = re.compile(r"^#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A)$", re.IGNORECASE)


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def resolve_inside(root: Path, path: Path) -> Path:
    resolved = path.resolve() if path.is_absolute() else (root / path).resolve()
    if not resolved.is_relative_to(root.resolve()):
        raise ValueError(f"path escapes repository root: {path}")
    return resolved


def visible_text(sheet) -> str:
    return "\n".join(
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


def audit_workbook(workbook_path: Path) -> dict[str, object]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    try:
        if tuple(workbook.sheetnames[: len(ALL_SHEETS)]) != ALL_SHEETS:
            raise ValueError("product workbook sheet contract changed")
        retained_sheets = workbook.sheetnames[len(ALL_SHEETS):]
        if any(workbook[name].sheet_state != "hidden" for name in retained_sheets):
            raise ValueError("retained legacy sheets must remain hidden")
        page_checks = {}
        user_tokens = {}
        formula_errors = []
        evidence_links = 0
        for sheet_name in ALL_SHEETS:
            sheet = workbook[sheet_name]
            if sheet.freeze_panes != "B4":
                raise ValueError(f"{sheet_name} does not freeze the first data column")
            if sheet.sheet_view.showGridLines:
                raise ValueError(f"{sheet_name} still shows grid lines")
            if sheet.max_column > 6:
                raise ValueError(f"{sheet_name} requires horizontal scrolling beyond six columns")
            text = visible_text(sheet)
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and FORMULA_ERROR.fullmatch(cell.value):
                        formula_errors.append(f"{sheet_name}!{cell.coordinate}={cell.value}")
                    if cell.hyperlink:
                        evidence_links += 1
            page_checks[sheet_name] = {
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "freeze_panes": sheet.freeze_panes,
                "grid_lines": sheet.sheet_view.showGridLines,
                "formula_errors": [],
            }
            if sheet_name in USER_SHEETS:
                hits = [token for token in FORBIDDEN_USER_TOKENS if token in text]
                user_tokens[sheet_name] = hits
        if formula_errors:
            raise ValueError("formula errors: " + ", ".join(formula_errors))
        token_hits = {
            sheet: hits for sheet, hits in user_tokens.items() if hits
        }
        if token_hits:
            raise ValueError(f"user pages expose internal tokens: {token_hits}")
        return {
            "page_contract": True,
            "five_product_pages_present": all(
                sheet in workbook.sheetnames for sheet in PRIMARY_SHEETS
            ),
            "sheet_checks": page_checks,
            "user_token_hits": token_hits,
            "evidence_hyperlink_count": evidence_links,
        }
    finally:
        workbook.close()


def audit_pdfs(pdf_directory: Path) -> dict[str, object]:
    if fitz is None:
        raise RuntimeError("PyMuPDF is required for PDF visual review")
    checks: dict[str, object] = {}
    total_pages = 0
    for sheet_name in ALL_SHEETS:
        pdf = pdf_directory / f"{sheet_name}.pdf"
        if not pdf.is_file():
            raise ValueError(f"missing WPS PDF for {sheet_name}")
        document = fitz.open(pdf)
        try:
            page_text = [document[index].get_text() for index in range(document.page_count)]
            if any(not text.strip() for text in page_text):
                raise ValueError(f"blank PDF page in {sheet_name}")
            if sheet_name in USER_SHEETS:
                joined = "\n".join(page_text)
                hits = [token for token in FORBIDDEN_USER_TOKENS if token in joined]
                if hits:
                    raise ValueError(f"user PDF {sheet_name} exposes internal tokens: {hits}")
            checks[sheet_name] = {
                "page_count": document.page_count,
                "pdf_sha256": file_sha256(pdf),
                "extracted_text_characters": sum(len(text) for text in page_text),
            }
            total_pages += document.page_count
        finally:
            document.close()
    return {"pdf_checks": checks, "total_pages": total_pages}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--wps-receipt", type=Path, required=True)
    parser.add_argument("--pdf-directory", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--expected-workbook-sha256", required=True)
    args = parser.parse_args()

    root = args.root.resolve()
    workbook_path = resolve_inside(root, args.workbook)
    manifest_path = resolve_inside(root, args.manifest)
    wps_receipt_path = resolve_inside(root, args.wps_receipt)
    pdf_directory = resolve_inside(root, args.pdf_directory)
    output_path = resolve_inside(root, args.output)
    for path in (workbook_path, manifest_path, wps_receipt_path):
        if not path.is_file():
            raise ValueError(f"missing review input: {path}")

    workbook_sha = file_sha256(workbook_path)
    if workbook_sha != args.expected_workbook_sha256:
        raise ValueError("workbook SHA-256 does not match the expected final candidate")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    wps = json.loads(wps_receipt_path.read_text(encoding="utf-8"))
    if manifest.get("workbook_sha256") != workbook_sha:
        raise ValueError("manifest does not bind the final workbook")
    if wps.get("sha256") != workbook_sha or wps.get("status") != "passed" or wps.get("read_only") is not True:
        raise ValueError("WPS receipt is not a passed read-only review of the final workbook")
    if output_path.exists():
        raise FileExistsError(output_path)

    workbook_audit = audit_workbook(workbook_path)
    pdf_audit = audit_pdfs(pdf_directory)
    receipt = {
        "schema_version": "m7-product-ux-visual-review-v1",
        "status": "PASS",
        "publication_allowed": True,
        "visual_review": "PASS",
        "review_method": "WPS_OFFICE_NATIVE_PDF_EXPORT_PLUS_PYMUPDF_TEXT_AND_LAYOUT_AUDIT",
        "workbook": str(workbook_path.relative_to(root)),
        "workbook_sha256": workbook_sha,
        "manifest": str(manifest_path.relative_to(root)),
        "manifest_sha256": file_sha256(manifest_path),
        "wps_receipt": str(wps_receipt_path.relative_to(root)),
        "wps_receipt_sha256": file_sha256(wps_receipt_path),
        "wps_application_path": wps.get("application_path"),
        "user_facing_sheets": list(USER_SHEETS),
        "primary_product_sheets": list(PRIMARY_SHEETS),
        "secondary_sheets": list(SECONDARY_SHEETS),
        "checks": {
            "five_product_pages_present": workbook_audit["five_product_pages_present"],
            "six_user_pages_have_frozen_company_column": True,
            "no_horizontal_scroll_beyond_six_columns": True,
            "no_formula_errors": True,
            "no_internal_tokens_on_user_pages": True,
            "no_blank_wps_pdf_pages": True,
            "evidence_hyperlinks_present": workbook_audit["evidence_hyperlink_count"] > 0,
        },
        "workbook_audit": workbook_audit,
        "pdf_audit": pdf_audit,
        "final_user_acceptance": "NOT_PASSED",
        "initial_assisted_use": "NOT_REACHED",
        "action": "no_order",
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": receipt["status"], "output": str(output_path), "workbook_sha256": workbook_sha}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(str(error), file=sys.stderr)
        raise SystemExit(1)
