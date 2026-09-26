#!/usr/bin/env python3
"""Publish the M7 product surface into the one configured canonical workbook.

This deliberately never creates a second user-facing workbook.  It copies the
configured workbook to a staging file beside it, proves retained sheets did not
change at cell level, then atomically replaces the original.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from scripts.build_m7_daily_workbench import build_packet  # noqa: E402
from value_investment_agent.application.product.product_workbench_candidate import (  # noqa: E402
    build_product_workbench_candidate_payload,
)
from value_investment_agent.presentation.excel.product_workbench import (  # noqa: E402
    WORKBOOK_SHEETS,
    apply_product_workbench_to_existing_workbook,
)
from value_investment_agent.presentation.read_models.product_workbench import (  # noqa: E402
    product_workbench_from_payload,
)


CANONICAL_NAME = "A股价值投资_Agent前端智能跟踪模板.xlsx"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _workbook_path() -> Path:
    value = os.environ.get("WORKBOOK_PATH")
    if not value:
        for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
            if line.strip().startswith("WORKBOOK_PATH"):
                value = line.split("=", 1)[1].strip()
                break
    if not value:
        raise ValueError("CANONICAL_WORKBOOK_NOT_RESOLVED")
    path = Path(value).expanduser().resolve()
    if not path.is_file() or path.suffix.lower() != ".xlsx" or path.name != CANONICAL_NAME:
        raise ValueError("CANONICAL_WORKBOOK_NOT_RESOLVED")
    if (path.parent / f"~${path.name}").exists():
        raise RuntimeError("CANONICAL_PUBLICATION_BLOCKED_BY_OPEN_WORKBOOK")
    return path


def _cell_fingerprint(sheet: Any) -> dict[str, Any]:
    cells = []
    formulas = hyperlinks = 0
    # ``iter_rows`` materializes every coordinate within WPS's styled range.
    # The internal cell map contains only real cells, so it protects the same
    # user content without turning a 13 MB workbook into a multi-GB scan.
    for cell in sheet._cells.values():
        if cell.value is None and not cell.hyperlink:
            continue
        is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
        formulas += int(is_formula)
        hyperlinks += int(bool(cell.hyperlink))
        cells.append((cell.coordinate, cell.data_type, cell.value, cell.hyperlink.target if cell.hyperlink else None))
    encoded = json.dumps(cells, ensure_ascii=False, default=str, separators=(",", ":")).encode("utf-8")
    return {
        "state": sheet.sheet_state,
        "formula_count": formulas,
        "hyperlink_count": hyperlinks,
        "cell_sha256": hashlib.sha256(encoded).hexdigest(),
    }


def _snapshot(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, keep_links=True)
    try:
        return {
            "sheet_order": workbook.sheetnames,
            "defined_names": sorted(name.name for name in workbook.defined_names.values()),
            "sheets": {sheet.title: _cell_fingerprint(sheet) for sheet in workbook.worksheets},
        }
    finally:
        workbook.close()


def _assert_retained(before: dict[str, Any], after: dict[str, Any]) -> None:
    retained = [name for name in before["sheet_order"] if name not in WORKBOOK_SHEETS]
    actual_retained = [name for name in after["sheet_order"] if name not in WORKBOOK_SHEETS]
    if retained != actual_retained:
        raise ValueError("retained sheet order changed")
    if before["defined_names"] != after["defined_names"]:
        raise ValueError("defined names changed")
    changed = [name for name in retained if before["sheets"][name] != after["sheets"].get(name)]
    if changed:
        raise ValueError("protected sheets changed: " + ", ".join(changed))
    if tuple(after["sheet_order"][: len(WORKBOOK_SHEETS)]) != WORKBOOK_SHEETS:
        raise ValueError("product navigation is not the first six sheets")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--publish", action="store_true", help="Required explicit confirmation to replace the canonical workbook.")
    args = parser.parse_args()
    if not args.publish:
        raise ValueError("refusing publish without --publish")

    canonical = _workbook_path()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup_dir = ROOT / "runtime" / "workbook-backups"
    receipt_dir = ROOT / "runtime" / "publication-receipts"
    backup_dir.mkdir(parents=True, exist_ok=True)
    receipt_dir.mkdir(parents=True, exist_ok=True)
    before_sha = _sha256(canonical)
    before = _snapshot(canonical)
    backup = backup_dir / f"canonical-before-m7-product-ux-{stamp}.xlsx"
    shutil.copy2(canonical, backup)
    if _sha256(backup) != before_sha:
        raise ValueError("backup hash mismatch")

    staging = canonical.with_name(f".{canonical.stem}.m7-staging-{uuid4().hex}{canonical.suffix}")
    shutil.copy2(canonical, staging)
    try:
        packet = build_packet(datetime.now(timezone.utc))
        model = product_workbench_from_payload(build_product_workbench_candidate_payload(packet, root=ROOT))
        workbook = load_workbook(staging, data_only=False, keep_links=True)
        try:
            apply_product_workbench_to_existing_workbook(workbook, model)
            workbook.save(staging)
        finally:
            workbook.close()
        after_staging = _snapshot(staging)
        _assert_retained(before, after_staging)
        staging_sha = _sha256(staging)
        os.replace(staging, canonical)
        after_sha = _sha256(canonical)
        after = _snapshot(canonical)
        _assert_retained(before, after)
    except Exception:
        if staging.exists():
            staging.unlink()
        raise

    receipt = {
        "schema_version": "canonical-product-publication-v1",
        "status": "PUBLISHED_PENDING_WPS_VISUAL_REVIEW",
        "workbook_source": "WORKBOOK_PATH",
        "canonical_filename": canonical.name,
        "canonical_file_before_sha256": before_sha,
        "backup": str(backup.relative_to(ROOT)),
        "backup_sha256": _sha256(backup),
        "staging_sha256": staging_sha,
        "canonical_file_after_sha256": after_sha,
        "workbook_path_unchanged": True,
        "product_sheets_present": True,
        "user_managed_sheets_preserved": True,
        "frozen_sheets_unchanged": True,
        "sheet_contract": {name: ("PRODUCT_MANAGED" if name in WORKBOOK_SHEETS else "PRESERVED") for name in after["sheet_order"]},
        "action": "no_order",
    }
    receipt_path = receipt_dir / f"canonical-m7-product-publication-{stamp}.json"
    receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({**receipt, "receipt": str(receipt_path.relative_to(ROOT))}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
