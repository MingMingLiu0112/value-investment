#!/usr/bin/env python3
"""Audit the M7 product workbook for clipped rows and emit a review receipt."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys
import unicodedata


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


LINE_POINTS = 15
LINE_PADDING = 6
USER_SHEETS = ("01_今日", "02_机会", "03_公司", "04_我的组合", "05_事件")
SECONDARY_SHEETS = ("06_系统与审计",)


def _glyph_width(character: str) -> int:
    return 2 if unicodedata.east_asian_width(character) in ("W", "F") else 1


def _required_lines(text: str, width_columns: int) -> int:
    limit = max(1, width_columns)
    lines = 0
    for paragraph in str(text).split("\n"):
        if not paragraph:
            lines += 1
            continue
        used = 1
        current = 0
        for character in paragraph:
            size = _glyph_width(character)
            if current + size > limit and current > 0:
                used += 1
                current = size
            else:
                current += size
        lines += used
    return max(lines, 1)


def _column_width(sheet, column: int) -> int:
    dimension = sheet.column_dimensions.get(get_column_letter(column))
    width = getattr(dimension, "width", None)
    return int(width) if width else 9


def _merge_spans(sheet) -> dict[tuple[int, int], int]:
    spans: dict[tuple[int, int], int] = {}
    for merged in sheet.merged_cells.ranges:
        spans[(merged.min_row, merged.min_col)] = merged.max_col - merged.min_col + 1
    return spans


def audit_sheet(sheet) -> dict[str, object]:
    spans = _merge_spans(sheet)
    covered = {
        (row, column)
        for merged in sheet.merged_cells.ranges
        for row in range(merged.min_row, merged.max_row + 1)
        for column in range(merged.min_col, merged.max_col + 1)
    }
    short_rows: list[dict[str, object]] = []
    inspected = 0
    for row in sheet.iter_rows():
        if not row:
            continue
        row_index = row[0].row
        required = 1
        driving = ""
        for cell in row:
            if cell.value is None:
                continue
            if (cell.row, cell.column) in covered and (
                cell.row,
                cell.column,
            ) not in spans:
                continue
            if not getattr(cell.alignment, "wrap_text", False):
                continue
            span = spans.get((cell.row, cell.column), 1)
            width = sum(
                _column_width(sheet, index)
                for index in range(cell.column, cell.column + span)
            )
            needed = _required_lines(str(cell.value), width - 1)
            if needed > required:
                required = needed
                driving = cell.coordinate
        if required <= 1:
            continue
        inspected += 1
        height = sheet.row_dimensions[row_index].height
        available = ((height - LINE_PADDING) / LINE_POINTS) if height else None
        if available is None or available + 0.05 < required:
            short_rows.append(
                {
                    "row": row_index,
                    "driving_cell": driving,
                    "required_lines": required,
                    "row_height": height,
                }
            )
    return {
        "rows_needing_multiple_lines": inspected,
        "clipped_rows": short_rows,
        "verdict": "PASS" if not short_rows else "FAIL",
    }


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    parser.add_argument("--rendered-pdf", type=Path, default=None)
    parser.add_argument("--rendered-pages", type=int, default=None)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    workbook = load_workbook(workbook_path)
    pages: dict[str, object] = {}
    for name in workbook.sheetnames:
        pages[name] = audit_sheet(workbook[name])

    user_failures = [
        name for name in USER_SHEETS if pages.get(name, {}).get("verdict") != "PASS"
    ]
    freeze_ok = all(
        workbook[name].freeze_panes == "B4" for name in workbook.sheetnames
    )
    rendered: dict[str, object] | None = None
    if args.rendered_pdf is not None:
        pdf_path = args.rendered_pdf.resolve()
        rendered = {
            "path": str(pdf_path.relative_to(ROOT))
            if pdf_path.is_relative_to(ROOT)
            else str(pdf_path),
            "sha256": _sha256(pdf_path),
            "pages": args.rendered_pages,
        }

    passed = not user_failures and freeze_ok
    receipt = {
        "schema_version": "m7-product-ux-readability-audit-v1",
        "status": "passed" if passed else "failed",
        "workbook": str(workbook_path.relative_to(ROOT)),
        "workbook_sha256": _sha256(workbook_path),
        "user_sheets": list(USER_SHEETS),
        "secondary_sheets": list(SECONDARY_SHEETS),
        "frozen_first_column": freeze_ok,
        "pages": pages,
        "failing_user_sheets": user_failures,
        "rendered_evidence": rendered,
        "action": "no_order",
        "final_user_acceptance": "NOT_PASSED",
    }
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(receipt, ensure_ascii=False, indent=2))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
