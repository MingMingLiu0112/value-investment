#!/usr/bin/env python3
"""Read-only acceptance check for the published 600519 Excel case summary."""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    if "09_公司研究" not in workbook.sheetnames or "21_决策验证" not in workbook.sheetnames:
        raise ValueError("Required research and decision sheets are missing")
    rows = [row for row in workbook["09_公司研究"].iter_rows(min_row=4, values_only=True)
            if str(row[0]).zfill(6) == "600519"]
    if len(rows) != 1:
        raise ValueError("Expected exactly one Moutai research row")
    scope = str(rows[0][15])
    # This is an acceptance gate for decision boundaries, not a frozen copy check of
    # superseded P1 explanatory wording. A current P1 explanation may change while
    # the no-value/no-order boundary remains mandatory.
    required = ("单股票流程案例已生成", "禁止交易", "逐日模拟决策已重放2674个交易日", "0笔订单", "不能证明策略收益", "旧PE18/PB4等权参考价", "历史策略准入审计：2674日中已批准的点时历史价值为0日", "2025回购计划后", "历史条件估值输入包：已按当时可用日期冻结12个年报版本，原公告URL、归档路径及Hash均已逐份核验", "逐日历史输入时间线：2674个交易日均仅绑定当时已披露的12个年度输入版本", "历史条件估值重放：2674个交易日均只使用当时已披露年报；2603日得到固定有限情景的实验DCF范围", "逐日阻断式纸面账本：2674日均保存决策证据及下一可成交日", "模拟执行闭环验收：真实历史输入覆盖2674个交易日和15条现金分派", "日线纸面执行契约：2026-09-18会话对应2026-09-21下一会话", "仅执行机械条件已登记，不构成估值、订单或实盘准入", "上交所交易日历和四段官方停复牌查询已核验", "旧全拒绝执行契约", "保守日线模拟契约", "仅验证T+1、费用、现金、持仓与幂等性，不是历史策略收益、合理价或实盘建议", "条件估值反证预检（区间覆盖审计）：2603日有实验范围", "历史区间逐笔研究实验：下端三档均0笔成交", "全部成交均在2015-2019开发期", "2020-2022验证期和2023-2025封存测试期均为0笔", "现金分派锚定研究模拟：364日因尚无已实施年度现金观察而阻断，2310日按点时保守端点完成模型评估", "\u4e2d\u4f4d\u5386\u53f2PE\u7814\u7a76\u5b9e\u9a8cv2\uff1a\u4fdd\u7559v1\uff0c\u53e6\u5c06\u76f8\u540c\u89c4\u5219\u5ef6\u81f32015-2025", "\u56fa\u5b9a30%\u5b89\u5168\u8fb9\u9645\u4ec51\u6b21\u7814\u7a76\u5165\u573a\u30010\u6b21\u9000\u51fa", "\u5e74\u62a5\u4e8b\u540e\u786e\u8ba4\u4e0d\u5f97\u5012\u704c", "正式合理价准入审计", "候选主模型为合并归母权益残余收益/分红能力模型", "当前合理价为空，不得生成买卖指令")
    missing = [value for value in required if value not in scope]
    if missing:
        raise ValueError(f"Moutai single-stock case status is not published: missing={missing!r}")
    headers = [cell.value for cell in next(workbook["21_决策验证"].iter_rows(min_row=3, max_row=3))]
    try:
        status_index = headers.index("策略验证状态")
    except ValueError as error:
        raise ValueError("Decision validation status column is missing") from error
    decision_rows = [row for row in workbook["21_决策验证"].iter_rows(min_row=4, values_only=True)
                     if str(row[0]).zfill(6) == "600519"]
    if len(decision_rows) != 1 or "2603日得到固定有限情景的实验DCF范围" not in str(decision_rows[0][status_index]) or "逐日阻断式纸面账本" not in str(decision_rows[0][status_index]) or "模拟执行闭环验收" not in str(decision_rows[0][status_index]) or "条件估值反证预检" not in str(decision_rows[0][status_index]) or "历史区间逐笔研究实验" not in str(decision_rows[0][status_index]) or "现金分派锚定研究模拟" not in str(decision_rows[0][status_index]) or "中位历史PE研究实验" not in str(decision_rows[0][status_index]) or "正式合理价准入审计" not in str(decision_rows[0][status_index]):
        observed = None if not decision_rows else decision_rows[0][status_index]
        raise ValueError(f"Moutai experimental replay status is not visible in decision validation: rows={len(decision_rows)} observed={observed!r}")
    print(f"WORKBOOK_OK sheets={len(workbook.sheetnames)} symbol={rows[0][0]} company={rows[0][1]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
