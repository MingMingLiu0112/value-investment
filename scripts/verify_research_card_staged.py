#!/usr/bin/env python3
"""Verify a WPS-staged research-card update before canonical publication."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


DERIVED = {"00_首页Dashboard", "00_公司总览", "00_待完成公司", "00_使用说明"}


def sheet_digest(sheet) -> str:
    digest = hashlib.sha256()
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                digest.update(f"{cell.coordinate}={cell.value!r}\n".encode("utf-8"))
    return digest.hexdigest()


def sheet_cells(sheet) -> dict[str, object]:
    return {
        cell.coordinate: cell.value
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("staged", type=Path)
    parser.add_argument("receipt", type=Path)
    args = parser.parse_args()
    source = load_workbook(args.source, read_only=True, data_only=False)
    staged = load_workbook(args.staged, read_only=True, data_only=False)
    if source.sheetnames != staged.sheetnames:
        raise ValueError("Sheet order changed while staging")
    changed = []
    examples = {}
    for name in source.sheetnames:
        if name in DERIVED:
            continue
        before = sheet_digest(source[name])
        after = sheet_digest(staged[name])
        if before != after:
            changed.append(name)
            before_cells = sheet_cells(source[name])
            after_cells = sheet_cells(staged[name])
            coordinates = sorted(set(before_cells) | set(after_cells))
            examples[name] = [
                {"cell": coordinate, "before": before_cells.get(coordinate), "after": after_cells.get(coordinate)}
                for coordinate in coordinates
                if before_cells.get(coordinate) != after_cells.get(coordinate)
            ][:5]
    if changed:
        raise ValueError(f"Non-derived sheet values/formulas changed: {changed}; examples={examples}")
    card = staged["00_公司总览"]
    rows = {str(card.cell(row, 1).value): row for row in range(1, card.max_row + 1)}
    label = "主模型 / 价格要求"
    capital_label = "资本配置 / 治理"
    resilience_label = "现金 / 低谷韧性"
    if label not in rows or capital_label not in rows or resilience_label not in rows:
        raise ValueError("Updated research-card rows are missing")
    body = str(card.cell(rows[label], 3).value)
    required = ("归母权益剩余收益/分配能力模型", "2026-09-18", "不是唯一市场预测", "合理价或卖出信号")
    if any(token not in body for token in required):
        raise ValueError("Updated research-card conclusion is incomplete")
    capital_body = str(card.cell(rows[capital_label], 3).value)
    capital_required = ("3,927,585", "92.06亿元", "不证明实际价格公允")
    if any(token not in capital_body for token in capital_required):
        raise ValueError("Capital-allocation and governance conclusion is incomplete")
    resilience_body = str(card.cell(rows[resilience_label], 3).value)
    resilience_required = ("469.54亿元", "254.26亿元", "不称无债或低谷安全")
    if any(token not in resilience_body for token in resilience_required):
        raise ValueError("Resilience conclusion is incomplete")
    delivery_label = "当前交付边界"
    delivery_body = "R0研究卡与P1研究日条件估值准入已通过；正式合理价、模拟准入和实盘指令仍未通过。"
    if delivery_label not in rows or card.cell(rows[delivery_label], 3).value != delivery_body:
        raise ValueError("R0 delivery boundary is incomplete")
    receipt = {
        "status": "passed",
        "source": str(args.source),
        "staged": str(args.staged),
        "non_derived_sheets_preserved": True,
        "updated_labels": [label, capital_label, resilience_label, delivery_label],
        "required_conclusion_tokens": list(required),
        "required_capital_tokens": list(capital_required),
        "required_resilience_tokens": list(resilience_required),
    }
    args.receipt.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
