"""Independently reopen the published canonical file after atomic replacement."""
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def main():
    receipt = json.loads((ROOT / 'runtime/excel-sync-status.json').read_text(encoding='utf-8-sig'))
    if receipt['status'] != 'published':
        raise ValueError('No canonical publication receipt')
    payload_raw = (ROOT / 'runtime/server-export-payload.json').read_bytes()
    if hashlib.sha256(payload_raw).hexdigest() != receipt['payloadSha256']:
        raise ValueError('Published source no longer matches cached payload')
    payload = json.loads(payload_raw)
    expected = {row['symbol'] for row in payload['market_candidates']}
    path = Path(receipt['workbook'])
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb['21_决策验证']
        headers = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
        assert headers[26:29] == ('收盘交易日核验', '最近已完成交易日', '逐源实际报价时间')
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        assert {row[0] for row in rows} == expected
        assert len(rows) == len(expected)
        statuses = Counter(row[26] for row in rows)
        for row in rows:
            if row[26] == '缺少带实际日期的逐股行情及交易日历证据':
                assert row[7] == '未通过' and row[27] is None and row[28] is None
                assert row[3] not in ('买入研究候选', '减仓研究候选')
            assert row[23] == '历史回测未完成'
        report = {'verified_at': datetime.now(timezone.utc).isoformat(),
                  'canonical_workbook': str(path), 'workbook_sha256': before,
                  'published_at': receipt['at'], 'payload_sha256': receipt['payloadSha256'],
                  'payload_generated_at': payload['generated_at'],
                  'company_count': len(rows), 'sheet_count': len(wb.sheetnames),
                  'quote_session_reasons': dict(statuses),
                  'historical_backtest_incomplete': len(rows),
                  'post_publication_readback': True}
    finally:
        wb.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != before:
        raise ValueError('Canonical workbook changed during readback')
    output = ROOT / 'runtime/quote-session-workbook-readback-20260909.json'
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(output), **report}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
