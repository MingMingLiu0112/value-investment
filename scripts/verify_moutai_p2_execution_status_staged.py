#!/usr/bin/env python3
"""Verify the P2 Excel candidate modifies only the two status fields."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from stage_moutai_p2_execution_status import UPDATES


def values(sheet) -> dict[str, object]:
    return {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("staged", type=Path)
    parser.add_argument("receipt", type=Path)
    args = parser.parse_args()
    source = load_workbook(args.source, read_only=True, data_only=False)
    staged = load_workbook(args.staged, read_only=True, data_only=False)
    if source.sheetnames != staged.sheetnames:
        raise ValueError("Sheet order changed")
    expected = {(sheet, address) for sheet, cells in UPDATES.items() for address in cells}
    actual = set()
    for name in source.sheetnames:
        before, after = values(source[name]), values(staged[name])
        for address in set(before) | set(after):
            if before.get(address) != after.get(address):
                actual.add((name, address))
    if actual != expected:
        raise ValueError(f"Unexpected changes: {sorted(actual)}")
    required = ("【P2机制】", "【P2日期一致零单】", "2026-09-16", "1258.00元",
                "安全边际-163.18%", "现金100万元、0股、0成交", "未伪造开盘成交",
                "【未通过】正式合理价、R1历史策略验证、交易指令和实盘准入均未通过")
    for sheet, address in expected:
        value = str(staged[sheet][address].value or "")
        if any(token not in value for token in required):
            raise ValueError(f"P2 status is incomplete at {sheet}!{address}")
    receipt = {"status": "passed", "source_sha256": digest(args.source), "staged_sha256": digest(args.staged),
               "changed_cells": sorted(f"{sheet}!{address}" for sheet, address in actual)}
    args.receipt.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
