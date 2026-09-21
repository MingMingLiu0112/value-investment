#!/usr/bin/env python3
"""Verify that a P1 research-status candidate changes only two status cells."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from stage_moutai_p1_research_status import UPDATES


def values(sheet) -> dict[str, object]:
    return {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("staged", type=Path)
    parser.add_argument("receipt", type=Path)
    args = parser.parse_args()
    source = load_workbook(args.source, read_only=True, data_only=False)
    staged = load_workbook(args.staged, read_only=True, data_only=False)
    if source.sheetnames != staged.sheetnames:
        raise ValueError("Sheet order changed")
    expected = {(sheet, address) for sheet, cells in UPDATES.items() for address in cells}
    actual = set()
    for name in source.sheetnames:
        before, after = values(source[name]), values(staged[name])
        for address in set(before) | set(after):
            if before.get(address) != after.get(address):
                actual.add((name, address))
    if actual != expected:
        raise ValueError(f"Unexpected changes: {sorted(actual)}")
    for sheet, address in expected:
        value = str(staged[sheet][address].value or "")
        for token in ("【当前操作】研究观察", "【当前P1】截至2026-09-20", "五项门禁已通过", "【未通过】正式合理价、模拟准入、纸面订单和实盘指令均未通过", "【下一步】P2"):
            if token not in value:
                raise ValueError(f"Missing {token!r} at {sheet}!{address}")
    receipt = {"status": "passed", "source_sha256": digest(args.source), "staged_sha256": digest(args.staged),
               "changed_cells": sorted(f"{sheet}!{address}" for sheet, address in actual)}
    args.receipt.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
