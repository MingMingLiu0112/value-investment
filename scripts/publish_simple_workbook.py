from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import os
import sys
from openpyxl import load_workbook

ROOT = Path('D:/GPTProject/value-investment')
sys.path[:0] = [str(ROOT/'src'), str(ROOT/'scripts')]
from value_investment_agent.workbook_frontdoor import apply_frontdoor, HOME
from value_investment_agent.workbook_simple_overview import OVERVIEW
from preview_workbook_frontdoor import digest_sheet

original = Path('C:/Users/we/WPSDrive/197617831/WPS云盘/价投跟踪/A股价值投资_Agent前端智能跟踪模板.xlsx')
out = ROOT/'runtime/workbook-backups'/('simple-'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ'))
out.mkdir(parents=True)
raw = original.read_bytes()
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
before_hash = hashlib.sha256(raw).hexdigest()
backup = out/'before.xlsx'
backup.write_bytes(raw)
wb = load_workbook(backup)
before = {s.title: digest_sheet(s) for s in wb if s.title not in {HOME, OVERVIEW}}
result = apply_frontdoor(wb)
staged = out/'ready.xlsx'
wb.save(staged)
wb.close()
print('Staged; verifying all source cells and overview mappings', flush=True)
checked = load_workbook(staged)
after = {s.title: digest_sheet(s) for s in checked if s.title not in {HOME, OVERVIEW}}
assert before == after, 'Source cells changed'
assert checked.active.title == HOME
assert [s.title for s in checked if s.sheet_state == 'visible'] == result['visible_sheets']
overview = checked[OVERVIEW]
assert overview.max_row-3 == result['candidate_count']
assert len({overview.cell(r,1).value for r in range(4,overview.max_row+1)}) == result['candidate_count']
for row in overview.iter_rows(min_row=4):
    assert row[7].value == '研究观察，暂不交易'
    if row[4].value is not None and row[5].value != '日期已核验':
        assert row[5].value == '历史报价，日期未核验'
    for cell in row:
        if cell.hyperlink:
            assert cell.hyperlink.target.split("'")[1] in checked.sheetnames
checked.close()
assert sha(original) == before_hash, 'Original changed during preparation'
pending = original.with_suffix('.publishing.xlsx')
with pending.open('xb') as f:
    f.write(staged.read_bytes())
assert sha(pending) == sha(staged)
assert sha(original) == before_hash
os.replace(pending, original)
assert sha(original) == sha(staged)
result.update({'status':'published', 'backup':str(backup), 'source_cells_preserved':True,
               'original_sha256':before_hash, 'published_sha256':sha(original),
               'native_wps_render_checked':False, 'market_data_refreshed':False})
(out/'publication.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False),flush=True)
