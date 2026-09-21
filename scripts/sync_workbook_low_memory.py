#!/usr/bin/env python3
"""Publish a compact, evidence-linked server payload to the WPS workbook."""
from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_FILING_CANDIDATES = 2500


def scalar(value, limit: int = 1000):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).replace("\x00", "").replace("\r", " ").replace("\n", " ")[:limit]


def replace_rows(workbook, name: str, headers: list[str], rows: list[list[object]]):
    sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([scalar(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(28, max(11, len(header) + 2))
    return sheet


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--payload", required=True, type=Path)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    payload = json.loads(args.payload.read_text(encoding="utf-8-sig"))
    candidates = payload.get("market_candidates", [])
    points_by_symbol = defaultdict(dict)
    for point in payload.get("points", []):
        old = points_by_symbol[point["symbol"]].get(point["field_name"])
        if old is None or (point.get("validation_status") == "verified" and old.get("validation_status") != "verified"):
            points_by_symbol[point["symbol"]][point["field_name"]] = point
    quality = {item["symbol"]: item for item in payload.get("financial_quality", [])}
    valuation = {item["symbol"]: item for item in payload.get("valuations", [])}

    workbook = load_workbook(args.workbook)
    dashboard = workbook["00_首页Dashboard"]
    dashboard["I4"], dashboard["J4"] = "全市场初筛候选", len(candidates)
    dashboard["I5"], dashboard["J5"] = "当前估值记录", len(valuation)
    dashboard["I6"], dashboard["J6"] = "已加载数据点", len(payload.get("points", []))
    dashboard["I7"], dashboard["J7"] = "载荷生成时间", scalar(payload.get("generated_at"), 100)

    replace_rows(workbook, "13_全市场初筛",
        ["股票代码", "公司名称", "板块", "行业", "当前价", "PE", "PB", "市值(亿元)", "初筛分", "状态", "行情来源", "筛选日期"],
        [[row.get("symbol"), row.get("name"), row.get("board"), row.get("sector"), row.get("current_price"),
          row.get("pe"), row.get("pb"), row.get("market_cap"), row.get("initial_score"), row.get("status"),
          row.get("market_source"), row.get("screen_date")] for row in candidates])

    coverage_rows = []
    quality_rows = []
    for candidate in candidates:
        symbol = candidate["symbol"]
        fields = points_by_symbol[symbol]
        verified = sum(item.get("validation_status") == "verified" for item in fields.values())
        pending = len(fields) - verified
        q = quality.get(symbol, {})
        v = valuation.get(symbol, {})
        coverage_rows.append([symbol, candidate.get("name"), candidate.get("board"), candidate.get("sector"),
                              verified, pending, q.get("quality_status"), q.get("total_score"),
                              q.get("coverage_ratio"), v.get("valuation_status"), v.get("build_signal"),
                              v.get("current_price"), v.get("fair_value"), v.get("safety_margin")])
        quality_rows.append([symbol, candidate.get("name"), candidate.get("board"), candidate.get("sector"),
                             q.get("model_type"), q.get("quality_status"), q.get("total_score"), q.get("coverage_ratio"),
                             q.get("profitability_score"), q.get("cash_flow_score"), q.get("balance_sheet_score"),
                             q.get("growth_score"), q.get("reasons")])
    replace_rows(workbook, "15_公司财务覆盖",
        ["股票代码", "公司名称", "板块", "行业", "已验证字段", "待验证字段", "质量状态", "质量分", "质量覆盖率", "估值状态", "研究信号", "当前价", "合理价", "安全边际"], coverage_rows)
    replace_rows(workbook, "17_财务质量评分",
        ["股票代码", "公司名称", "板块", "行业", "质量模型", "质量状态", "质量分", "覆盖率", "盈利分", "现金流分", "资产负债表分", "成长分", "阻断原因"], quality_rows)

    recommendations = []
    for candidate in candidates:
        symbol = candidate["symbol"]
        q = quality.get(symbol, {})
        v = valuation.get(symbol, {})
        coverage = float(q.get("coverage_ratio") or 0)
        margin = v.get("safety_margin")
        margin_value = float(margin) if margin is not None else None
        quality_ok = q.get("quality_status") in {"通过", "qualified", "verified"}
        valuation_ok = v.get("fair_value") is not None and margin_value is not None
        paper_entry = quality_ok and coverage >= 0.8 and valuation_ok and margin_value >= 0.25
        blockers = []
        if coverage < 0.8:
            blockers.append(f"财报验证覆盖率 {coverage:.0%}，低于 80%")
        if not quality_ok:
            blockers.append(f"质量状态：{q.get('quality_status') or '待补证'}")
        if not valuation_ok:
            blockers.append("尚无经验证合理价/安全边际")
        if paper_entry:
            suggestion = "纸面建仓条件具备（非实盘）"
            reason = "数据、质量与估值门禁已同时通过；仍需人工确认交易执行条件。"
        elif float(candidate.get("initial_score") or 0) >= 55:
            suggestion = "重点研究，暂不交易"
            reason = "；".join(blockers) if blockers else "等待下一次数据更新"
        else:
            suggestion = "继续观察，暂不交易"
            reason = "初筛分不足；" + "；".join(blockers)
        priority = float(candidate.get("initial_score") or 0) + coverage * 20 + max(0, margin_value or 0) * 10
        recommendations.append((priority, [symbol, candidate.get("name"), candidate.get("board"), candidate.get("sector"),
            candidate.get("initial_score"), sum(item.get("validation_status") == "verified" for item in points_by_symbol[symbol].values()),
            coverage, v.get("current_price"), v.get("fair_value"), margin, suggestion, reason, payload.get("generated_at")]))
    recommendations.sort(key=lambda item: item[0], reverse=True)
    recommendation_rows = [[index, *row] for index, (_, row) in enumerate(recommendations[:50], 1)]
    replace_rows(workbook, "18_研究建议",
        ["排名", "股票代码", "公司名称", "板块", "行业", "初筛分", "已验证字段", "财务覆盖率", "当前价", "合理价", "安全边际", "当前建议", "门禁与原因", "数据时间"], recommendation_rows)

    sectors = defaultdict(lambda: [0, 0, 0.0, 0])
    for candidate in candidates:
        key = (candidate.get("board"), candidate.get("sector"))
        bucket = sectors[key]
        bucket[0] += 1
        bucket[1] += int(float(candidate.get("initial_score") or 0) >= 55)
        bucket[2] += float(candidate.get("initial_score") or 0)
        bucket[3] += int(float(quality.get(candidate["symbol"], {}).get("coverage_ratio") or 0) > 0)
    replace_rows(workbook, "16_板块行业汇总",
        ["板块", "行业", "初筛候选", "重点观察", "平均初筛分", "有财务覆盖公司"],
        [[board, sector, total, focus, round(score / total, 2), covered]
         for (board, sector), (total, focus, score, covered) in sorted(sectors.items())])

    filings = payload.get("filing_candidates", [])[:MAX_FILING_CANDIDATES]
    replace_rows(workbook, "14_财报候选",
        ["候选ID", "股票代码", "公司名称", "报告期", "报告类型", "字段", "候选值", "单位", "页码", "验证状态", "官方URL", "文件SHA-256", "原文摘录"],
        [[item.get("candidate_id"), item.get("symbol"), item.get("name"), item.get("report_period"), item.get("report_kind"),
          item.get("field_name"), item.get("value"), item.get("unit"), item.get("page_number"), item.get("status"),
          item.get("source_url"), item.get("sha256"), scalar(item.get("excerpt"), 1000)] for item in filings])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(json.dumps({"output": str(args.output), "market_candidates": len(candidates),
                      "filing_candidates_published": len(filings)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
