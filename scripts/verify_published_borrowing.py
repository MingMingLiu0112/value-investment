"""Read back the canonical workbook and verify the newly promoted PDF evidence."""
import json
import sys
from decimal import Decimal
from openpyxl import load_workbook

book = load_workbook(sys.argv[1], read_only=True, data_only=False)
evidence = book['18_指标证据']
rows = [r for r in evidence.iter_rows(min_row=4, values_only=True)
        if str(r[0])=='600900' and r[3]=='long_term_borrowings' and str(r[4])=='2025-12-31']
assert len(rows)==1, 'Missing or duplicate annual borrowing evidence'
row = rows[0]
assert Decimal(str(row[5])) == Decimal('172310624303.31')
assert row[11]==81
assert row[10]=='aa0cdb92e3aad81b1e6454594a64cacbcfeb9013b10dd31f47a93a94e7a70f91'
assert row[9]=='https://static.cninfo.com.cn/finalpage/2026-04-30/1225262036.PDF'
assert row[8]
dashboard = {r[0]:r[1] for r in book['00_首页Dashboard'].iter_rows(min_row=4,values_only=True)}
assert dashboard['初筛日期'] and dashboard['行情快照时间']
print(json.dumps({'workbook':sys.argv[1],'symbol':'600900','field':'long_term_borrowings',
    'value':str(row[5]), 'page':row[11], 'source_id':row[8], 'status':row[7],
    'screen_date':dashboard['初筛日期'],'quote_as_of':dashboard['行情快照时间'],
    'canonical_evidence_verified':True},ensure_ascii=False))
book.close()
