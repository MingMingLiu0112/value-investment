#!/usr/bin/env python3
"""Inspect a pinned ChinaBond curve workbook without admitting a rate input."""
from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "runtime/strategy-validation/moutai-historical-market-rate-20260920T053000Z"
SOURCE = DIRECTORY / "chinabond-2014-12-31.bin"
EXPECTED_SHA256 = "1d99673eb8733bc5ea724268f378cfe2da1f42f3997e5cae79d0316c45fa9b2f"
MATERIALIZED = DIRECTORY / "chinabond-2014-12-31.xlsx"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build() -> dict[str, object]:
    if digest(SOURCE) != EXPECTED_SHA256:
        raise ValueError("Pinned ChinaBond source changed")
    shutil.copyfile(SOURCE, MATERIALIZED)
    workbook = load_workbook(MATERIALIZED, read_only=True, data_only=True)
    sheets: dict[str, list[list[object]]] = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True):
            rows.append(list(row[: min(sheet.max_column, 12)]))
        sheets[sheet.title] = rows
    populated = [value for rows in sheets.values() for row in rows for value in row if value is not None]
    status = (
        "rejected_metadata_only_no_curve_points"
        if populated == ["日期"]
        else "inspected_not_admitted_pending_curve_definition_and_tenor_validation"
    )
    return {
        "source_id": "chinabond:historical-curve-download",
        "source_url": "https://yield.chinabond.com.cn/cbweb-mn/yc/downYearBzqx?ycDefId=2c9081a24af9dca9014b013b2a7f0005&workTime=2014-12-31&zblx=0&ycArea=0&nameType=1&locale=zh_CN",
        "requested_date": "2014-12-31",
        "raw_file": str(SOURCE.relative_to(ROOT)),
        "raw_file_sha256": EXPECTED_SHA256,
        "workbook_sheets": sheets,
        "rate_input_status": status,
        "valuation_approved": False,
        "trade_approved": False,
    }


def main() -> int:
    result = build()
    output = DIRECTORY / "inspection.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "sheets": list(result["workbook_sheets"]), "status": result["rate_input_status"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
