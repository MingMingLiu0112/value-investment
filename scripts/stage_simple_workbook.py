#!/usr/bin/env python3
"""Stage a verified research-card workbook without changing the canonical file."""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path("D:/GPTProject/value-investment")
WORKBOOK = Path("C:/Users/we/WPSDrive/197617831/WPS云盘/价投跟踪/A股价值投资_Agent前端智能跟踪模板.xlsx")
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "scripts")]

from preview_workbook_frontdoor import digest_sheet
from value_investment_agent.workbook_frontdoor import HOME, apply_frontdoor
from value_investment_agent.workbook_simple_overview import DERIVED, OVERVIEW


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    output = ROOT / "runtime" / "workbook-backups" / (
        "research-card-v3-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    )
    output.mkdir(parents=True)
    before = output / "before.xlsx"
    before.write_bytes(WORKBOOK.read_bytes())
    original_hash = sha256(before)

    workbook = load_workbook(before)
    preserved_before = {sheet.title: digest_sheet(sheet) for sheet in workbook if sheet.title not in DERIVED}
    result = apply_frontdoor(workbook, root=ROOT)
    staged = output / "ready.xlsx"
    workbook.save(staged)
    workbook.close()

    check = load_workbook(staged, data_only=False)
    preserved_after = {sheet.title: digest_sheet(sheet) for sheet in check if sheet.title not in DERIVED}
    if preserved_before != preserved_after:
        raise ValueError("Non-derived workbook content changed while staging")
    card = check[OVERVIEW]
    labels = {str(card.cell(row, 1).value) for row in range(1, card.max_row + 1)}
    required = {"主模型及假设边界", "市场对价格的要求", "价格如何理解"}
    if not required.issubset(labels):
        raise ValueError("Updated research-card conclusions are missing")
    if check.active.title != HOME:
        raise ValueError("Homepage must remain active")
    check.close()
    if sha256(WORKBOOK) != original_hash:
        raise ValueError("Canonical workbook changed during staging")

    receipt = {
        "status": "staged",
        "canonical_workbook": str(WORKBOOK),
        "canonical_sha256": original_hash,
        "staged_workbook": str(staged),
        "staged_sha256": sha256(staged),
        "source_sheets_preserved": True,
        "updated_research_card_labels": sorted(required),
        "publication_result": result,
    }
    (output / "staging.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
