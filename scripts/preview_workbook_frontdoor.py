"""Stage a navigation-only update and verify every retained source cell."""
import argparse
import gc
import hashlib
import json
from pathlib import Path
import sys
from datetime import datetime, timezone
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from value_investment_agent.workbook_frontdoor import apply_frontdoor, HOME, PRIMARY
from value_investment_agent.workbook_simple_overview import DERIVED, OVERVIEW, PENDING, GUIDE
from value_investment_agent.workbook_compat import load_wps_workbook


def digest_sheet(sheet, exclude=()):
    digest = hashlib.sha256()
    count = 0
    for row in sheet:
        for cell in row:
            if cell.coordinate in exclude:
                continue
            if cell.value == '返回首页' and cell.row == 1:
                continue
            if cell.value is None and not cell.hyperlink and not cell.comment:
                continue
            value = format(cell.value, '.14g') if isinstance(cell.value, float) else cell.value
            link = cell.hyperlink
            destination = (link.location or (link.target[1:] if link.target.startswith('#') else link.target)) if link else None
            item = [cell.coordinate, value, cell.data_type,
                    destination if cell.row != 1 else None,
                    cell.comment.text if cell.comment else None]
            digest.update(json.dumps(item, ensure_ascii=False, default=str).encode('utf-8'))
            count += 1
    return {'sha256': digest.hexdigest(), 'nonempty_cells': count}


def render_sheet(sheet, path, end_row, end_col, start_row=1):
    """Render a layout preview and reject text that does not fit its cells."""
    from PIL import Image, ImageDraw, ImageFont
    widths = [round(sheet.column_dimensions[chr(64 + c)].width * 7 + 5) for c in range(1, end_col + 1)]
    heights = [0 if sheet.row_dimensions[r].hidden else round((sheet.row_dimensions[r].height or 15) * 4 / 3)
               for r in range(1, end_row + 1)]
    offset_y = sum(heights[:start_row - 1])
    image = Image.new('RGB', (sum(widths) + 32, sum(heights) - offset_y + 32), 'white')
    draw = ImageDraw.Draw(image)
    merges = {(m.min_row, m.min_col): m for m in sheet.merged_cells.ranges}
    for row in range(start_row, end_row + 1):
        if heights[row - 1] == 0:
            continue
        for col in range(1, end_col + 1):
            cell = sheet.cell(row, col)
            if cell.value is None:
                continue
            merged = merges.get((row, col))
            end = merged.max_col if merged else col
            x, y = 16 + sum(widths[:col - 1]), 16 + sum(heights[:row - 1]) - offset_y
            width, height = sum(widths[col - 1:end]), heights[row - 1]
            fill = cell.fill.fgColor.rgb
            if cell.fill.patternType == 'solid' and isinstance(fill, str):
                draw.rectangle((x, y, x + width, y + height), fill='#' + fill[-6:])
            font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', round((cell.font.sz or 11) * 4 / 3))
            color = cell.font.color.rgb if cell.font.color and cell.font.color.type == 'rgb' else '263330'
            lines, line = [], ''
            display = '（公式，WPS重算）' if cell.data_type == 'f' else str(cell.value)
            for char in display:
                if char == '\n':
                    lines.append(line)
                    line = ''
                elif line and draw.textlength(line + char, font=font) > width - 12:
                    lines.append(line)
                    line = char
                else:
                    line += char
            lines.append(line)
            line_height = round((cell.font.sz or 11) * 4 / 3) + 4
            if len(lines) * line_height > height:
                raise ValueError(f'Text overflow at {sheet.title}!{cell.coordinate}: {cell.value}')
            for n, line in enumerate(lines):
                draw.text((x + 6, y + (height - len(lines) * line_height) / 2 + n * line_height),
                          line, font=font, fill='#' + color[-6:])
    image.save(path)


def verify_navigation(book, result):
    done = {str(r[0]) for r in book[OVERVIEW].iter_rows(min_row=4, values_only=True)
            if str(r[0]).isdigit() and len(str(r[0])) == 6}
    pending = {str(r[0]) for r in book[PENDING].iter_rows(min_row=4, values_only=True)
               if str(r[0]).isdigit() and len(str(r[0])) == 6}
    assert not done & pending, 'A company appears in both groups'
    assert done | pending == set(result['groups']), 'A company was lost from navigation'
    assert len(done) == result['research_result_count']
    assert len(pending) == result['pending_count']
    assert book.active.title == HOME
    assert [s.title for s in book if s.sheet_state == 'visible'] == [s for s in PRIMARY if s in book]
    internal = 0
    for sheet in book:
        for row in sheet:
            for cell in row:
                link = cell.hyperlink
                if not link:
                    continue
                assert not (link.target or '').startswith('#'), 'Internal link serialized as an external file'
                if link.location:
                    internal += 1
                    assert link.target is None and link.id is None, 'Internal location must not use an external relationship'
                    target_name, address = link.location.rsplit('!', 1)
                    target_name = target_name.strip("'").replace("''", "'")
                    assert book[target_name].sheet_state == 'visible', link.location
                    assert book[target_name][address].value is not None, link.location
    assert internal > 0
    return internal


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', type=Path, default=Path(
        'C:/Users/we/WPSDrive/197617831/WPS云盘/价投跟踪/A股价值投资_Agent前端智能跟踪模板.xlsx'))
    args = parser.parse_args()
    original = args.workbook.resolve()
    out = ROOT / 'runtime/workbook-backups' / ('frontdoor-' + datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ'))
    out.mkdir(parents=True, exist_ok=False)
    raw = original.read_bytes()
    before_hash = hashlib.sha256(raw).hexdigest()
    backup = out / 'before.xlsx'
    backup.write_bytes(raw)
    print(json.dumps({'stage': 'snapshot', 'output': str(out)}), flush=True)
    w = load_wps_workbook(backup)
    before = {s.title: digest_sheet(s) for s in w if s.title not in DERIVED}
    result = apply_frontdoor(w)
    staged = out / 'ready.xlsx'
    w.save(staged)
    w.close()
    del w
    gc.collect()
    print(json.dumps({'stage': 'candidate_saved', 'file': str(staged)}), flush=True)
    checked = load_workbook(staged)
    after = {s.title: digest_sheet(s, result.get('added_case_cells', {}).get(s.title, ()))
             for s in checked if s.title not in DERIVED}
    if before != after:
        raise ValueError('Source-cell preservation failed: ' + str([k for k in before if before[k] != after.get(k)]))
    result['verified_internal_links'] = verify_navigation(checked, result)
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    with ZipFile(staged) as package:
        for filename in package.namelist():
            if filename.startswith('xl/worksheets/_rels/'):
                assert not any(e.get('Target', '').startswith('#') for e in ET.fromstring(package.read(filename))), filename
    for name, filename, end_row, end_col in [
            (HOME, 'homepage-preview.png', 23, 8),
            (OVERVIEW, 'results-preview.png', checked[OVERVIEW].max_row, 8),
            (PENDING, 'pending-preview.png', 8, 9)]:
        render_sheet(checked[name], out / filename, end_row, end_col)
    for start in range(1, checked[GUIDE].max_row + 1, 20):
        render_sheet(checked[GUIDE], out / f'guide-preview-{start:03d}.png',
                     min(start + 19, checked[GUIDE].max_row), 8, start)
    result.update({'original': str(original), 'original_sha256': before_hash, 'backup': str(backup),
                   'staged': str(staged), 'staged_sha256': hashlib.sha256(staged.read_bytes()).hexdigest(),
                   'preserved_sheets': before, 'checked_at': datetime.now(timezone.utc).isoformat(),
                   'status': 'verified_ready_for_atomic_publication',
                   'render_note': 'Pillow layout previews; not a WPS application screenshot.'})
    checked.close()
    (out / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: result[k] for k in ('staged', 'staged_sha256', 'research_result_count',
        'pending_count', 'company_count', 'status')}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
