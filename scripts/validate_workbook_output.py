"""Verify a generated research workbook before publishing it to WPS."""
import json
from pathlib import Path
import sys
import math
import hashlib
from collections import Counter

from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from value_investment_agent.debt_research_export import load_debt_research
from value_investment_agent.company_research_export import load_company_research
from value_investment_agent.candidate_tracking import tracking_candidate_view
from value_investment_agent.excel_report import read_holdings
from value_investment_agent.reminders import build_reminders

template, payload_path, report_path = map(Path,sys.argv[1:])
payload=json.loads(payload_path.read_text(encoding='utf-8-sig'))
from value_investment_agent.workbook_compat import load_wps_workbook
original=load_wps_workbook(template)
report=load_workbook(report_path)
expected={str(c['symbol']).zfill(6) for c in payload['market_candidates']}
def same_value(a,b):
    if a == b: return True
    # WPS empty strings serialize as blank cells in openpyxl; both are empty.
    if a in (None, '') and b in (None, ''): return True
    # XLSX writes finite decimal precision; allow only float serialization noise.
    return isinstance(a,(int,float)) and isinstance(b,(int,float)) and math.isclose(a,b,rel_tol=1e-15,abs_tol=0)
counts={}
for name in ['01_观察名单','02_质量评分','03_财务指标','04_估值跟踪','09_公司研究','10_年报跟踪']:
    actual={str(row[0]).zfill(6) for row in report[name].iter_rows(min_row=4,values_only=True) if row[0]}
    assert expected <= actual, (name, sorted(expected-actual))
    counts[name]=len(actual)
from value_investment_agent.workbook_frontdoor import primary_sheet_names
assert [s.title for s in report if s.sheet_state=='visible']==primary_sheet_names(report), 'Frontdoor structure missing or regressed'
assert report.active.title=='00_首页Dashboard'
assert report['00_公司总览']['A1'].value=='贵州茅台 | 单公司研究卡'
assert report['00_待完成公司']['D3'].value=='工作进度'
assert report['00_首页Dashboard']['A1'].value=='价值投资 | 先研究一家公司'
assert report['00_首页Dashboard']['G17'].hyperlink.location=="'21_决策验证'!A1"
guide_text = '\n'.join(str(cell.value) for row in report['00_使用说明'].iter_rows() for cell in row if cell.value)
assert all(required in guide_text for required in (
    '投资路径 | 可以重叠；行业口径另行匹配',
    '财务五问 | 所有适用公司必答，不用单一总分替代',
    '日常模拟、R1、R2分开验收',
)), 'Research guide structure missing or regressed'
assert report['12_提醒']['O3'].value=='策略验证状态'
for i in range(4,report['12_提醒'].max_row+1):
    assert report['12_提醒'].cell(i,15).value==report['21_决策验证'].cell(i,24).value=='历史回测未完成'
    assert report['12_提醒'].cell(i,1).value==report['21_决策验证'].cell(i,1).value
    assert report['12_提醒'].cell(i,16).hyperlink.location==f"'21_决策验证'!A{i}"
assert {r[0] for r in report['12_提醒'].iter_rows(min_row=4,values_only=True) if r[0]}==expected
assert report['21_决策验证']['AA3'].value == '收盘交易日核验'
session_alerts = {row['symbol']: row for row in build_reminders(
    dict(payload, market_candidates=tracking_candidate_view(payload)),
    read_holdings(original['05_仓位管理']))}
for row in report['21_决策验证'].iter_rows(min_row=4, values_only=True):
    if row[0] not in session_alerts:
        assert row[0] == '600519' and row[3] == '研究验证（非买卖提醒）'
        assert '2603日得到固定有限情景的实验DCF范围' in str(row[23])
        continue
    session = session_alerts[row[0]]['quote_session']
    assert row[26] == session['reason'], 'Closing-session reason mismatch'
    assert row[27] == session['expected_session'], 'Expected exchange session mismatch'
    observed = '；'.join(f'{provider}: {stamp}' for provider,stamp in session['provider_times'].items()) or None
    assert row[28] == observed, 'Provider timestamp evidence mismatch'
    if not session['passed']:
        assert row[7] == '未通过', 'Missing quote date must block the freshness gate'
        assert row[3] not in ('买入研究候选', '减仓研究候选'), 'Unproven close cannot release a trade hint'
coverage=(payload.get('market_audit') or {}).get('coverage') or {}
if coverage:
    assert coverage['unique_symbols']==coverage['received_rows']
    assert sum(coverage['outcome_counts'].values())==coverage['received_rows']
    assert coverage['outcome_counts'].get('candidate',0)==len(expected)
official=payload.get('official_coverage') or {}
if official.get('reconciled'):
    assert official['expected_count']==official['matched_count']==coverage['unique_symbols']
    assert not any(official.get(k) for k in ('missing_in_market','not_in_official','board_conflicts'))
    assert official['source_id'] and official['sha256'] and official['sources']
    assert report['20_市场覆盖']['B6'].value=='已完成'
for name in ['05_仓位管理','07_月度复盘','08_交易记录']:
    for row in original[name]:
        for cell in row:
            assert same_value(cell.value, report[name][cell.coordinate].value), (name,cell.coordinate)
for name in ['06_月度跟踪','11_数据源审计']:
    for row in original[name].iter_rows(min_row=3):
        for cell in row:
            if cell.value is not None:
                assert same_value(cell.value, report[name][cell.coordinate].value), (name,cell.coordinate)
notes={str(row[0]).zfill(6):row for row in report['09_公司研究'].iter_rows(min_row=4,values_only=True) if row[0]}
for row in original['09_公司研究'].iter_rows(min_row=4,values_only=True):
    if row[0]: assert tuple(row[:13])==tuple(notes[str(row[0]).zfill(6)][:13])
company_research = load_company_research(Path(__file__).resolve().parents[1])
for code, values in company_research.items():
    assert list(notes[code][13:17]) == values, ('Company research display mismatch', code)
for row in report['09_公司研究'].iter_rows(min_row=4):
    if str(row[0].value).zfill(6) in company_research:
        assert row[16].hyperlink.target == company_research[str(row[0].value).zfill(6)][3]
assert not any(c.data_type=='f' for row in report['02_质量评分'] for c in row)
special=[row for row in report['19_金融专用指标'].iter_rows(min_row=4,values_only=True) if row[5] is not None]
assert len(special)>0
assert all(row[10] and row[11] and row[9] for row in special)
research = load_debt_research(Path(__file__).resolve().parents[1], payload,
                             hashlib.sha256(payload_path.read_bytes()).hexdigest())
research_rows = [(i, list(row)) for i, row in enumerate(report['18_指标证据'].iter_rows(values_only=True), 1)
                 if len(row) > 3 and str(row[3]).startswith('research_current_maturity_')]
expected_research = []
for code, rows in sorted(research['evidence'].items()):
    if code not in notes:
        continue
    for row in rows:
        expected_research.append([code, notes[code][1], row['label'], row['field_name'], row['period'],
            float(row['value']) if row['value'] is not None else None, row['unit'], row['status'],
            row['source_id'] or None, row['source_url'], row['sha256'], row['pages'], row['parser'],
            row['fetched_at'], None, row['excerpt']])
assert len(expected_research) == len(research_rows), 'Research evidence count mismatch'
for expected_row, (_, actual_row) in zip(expected_research, research_rows):
    assert all(same_value(a, b) for a, b in zip(expected_row, actual_row[:16])), 'Research evidence mismatch'
for row in report['21_决策验证'].iter_rows(min_row=4):
    code = row[0].value
    if code not in session_alerts:
        assert code == '600519' and row[3].value == '研究验证（非买卖提醒）'
        continue
    assert row[25].value == research['companies'].get(code, '尚未纳入负债附注研究批次')
    first = next((i for i, values in research_rows if values[0] == code), None)
    if first is not None:
        assert row[25].hyperlink.location == f"'18_指标证据'!A{first}"
result={'main_company_counts':counts,'specialized_values':len(special),
        'annual_input_cells':sum(isinstance(c.value,(int,float)) for row in report['02_质量评分'].iter_rows(min_row=4,min_col=14,max_col=22) for c in row),
        'specialized_companies':len({r[0] for r in special}),
        'visible_tabs':[s.title for s in report if s.sheet_state=='visible'],
        'manual_records_preserved':True,'history_preserved':True,
        'payload_generated_at':payload['generated_at']}
result['research_evidence_rows'] = len(research_rows)
result['research_affects_trade_rules'] = False
result['official_universe_reconciled']=official.get('reconciled',False)
result['quote_session_status_counts'] = dict(Counter(
    row['quote_session']['status'] for row in session_alerts.values()))
print(json.dumps(result,ensure_ascii=False,indent=2))
