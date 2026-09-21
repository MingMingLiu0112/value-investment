#!/usr/bin/env python3
"""Verify the narrow 600519 action-status update before WPS publication."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


RESEARCH = "09_公司研究"
DECISIONS = "21_决策验证"
SYMBOL = "600519"
STATUS_HEADER = "策略验证状态"
ACTION_PREFIX = "【当前操作】研究观察：不新增买入、不减仓、不卖出。"
REASON = "【为什么】正式合理价、模拟准入和实盘指令均未通过；条件试算不是买卖点。"


def symbol_row(sheet) -> int:
    for row in range(4, sheet.max_row + 1):
        if str(sheet.cell(row, 1).value).zfill(6) == SYMBOL:
            return row
    raise ValueError(f"{sheet.title} lacks symbol {SYMBOL}")


def status_column(sheet) -> int:
    headers = [cell.value for cell in sheet[3]]
    try:
        return headers.index(STATUS_HEADER) + 1
    except ValueError as error:
        raise ValueError(f"{sheet.title} lacks header {STATUS_HEADER!r}") from error


def populated_cells(sheet) -> dict[str, object]:
    return {
        cell.coordinate: cell.value
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify(source_path: Path, staged_path: Path) -> dict[str, object]:
    source = load_workbook(source_path, read_only=True, data_only=False)
    staged = load_workbook(staged_path, read_only=True, data_only=False)
    if source.sheetnames != staged.sheetnames:
        raise ValueError("Sheet order changed while staging")
    research_cell = source[RESEARCH].cell(symbol_row(source[RESEARCH]), 16).coordinate
    decision_cell = source[DECISIONS].cell(
        symbol_row(source[DECISIONS]), status_column(source[DECISIONS])
    ).coordinate
    allowed = {(RESEARCH, research_cell), (DECISIONS, decision_cell)}
    changes: list[dict[str, object]] = []
    for name in source.sheetnames:
        before, after = populated_cells(source[name]), populated_cells(staged[name])
        for coordinate in sorted(set(before) | set(after)):
            if before.get(coordinate) == after.get(coordinate):
                continue
            if (name, coordinate) not in allowed:
                raise ValueError(f"Unexpected workbook change at {name}!{coordinate}")
            changes.append({"sheet": name, "cell": coordinate})
    if {(item["sheet"], item["cell"]) for item in changes} != allowed:
        raise ValueError(f"Expected exactly two status changes, observed={changes!r}")
    for name, coordinate in allowed:
        value = str(staged[name][coordinate].value or "")
        if not value.startswith(ACTION_PREFIX) or REASON not in value:
            raise ValueError(f"Action status missing at {name}!{coordinate}")

    receipt = {
        "status": "passed",
        "source": str(source_path),
        "staged": str(staged_path),
        "source_sha256": sha256(source_path),
        "staged_sha256": sha256(staged_path),
        "changed_cells": changes,
        "note": "Only the 600519 research and decision status cells changed; formulas elsewhere are preserved.",
    }
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("staged", type=Path)
    parser.add_argument("receipt", type=Path)
    args = parser.parse_args()
    receipt = verify(args.source, args.staged)
    args.receipt.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
