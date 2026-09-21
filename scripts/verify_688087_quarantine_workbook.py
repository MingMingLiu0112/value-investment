"""Verify known truncated facts are not presented as trusted in the original workbook."""
import json
import sys
from decimal import Decimal
from openpyxl import load_workbook

book = load_workbook(sys.argv[1], read_only=True, data_only=False)
bad = {'revenue':Decimal('3546216993'), 'net_income':Decimal('285734507.5'),
       'operating_cash_flow':Decimal('568565377.2')}
corrected = {'revenue':Decimal('3546216993.66'), 'net_income':Decimal('285734507.57'),
             'operating_cash_flow':Decimal('568565377.28'), 'net_margin':Decimal('8.0574'),
             'operating_cash_flow_to_net_income':Decimal('198.9838')}
fields = corrected if '--corrected' in sys.argv else bad
rows = [r for r in book['18_指标证据'].iter_rows(min_row=4, values_only=True)
        if str(r[0])=='688087' and r[3] in fields and str(r[4])=='2025-12-31']
assert {r[3] for r in rows}==set(fields), 'Annual evidence rows missing'
for r in rows:
    if '--corrected' in sys.argv:
        assert Decimal(str(r[5]))==corrected[r[3]] and r[7]=='已交叉验证'
        if r[3]=='net_income':
            assert r[11]==125 and r[10]=='d5ea80d0a678586b178bb95bdedcabd00459fd23c98ea5601d772ec55adfaa9a'
    elif r[5] is not None and Decimal(str(r[5]))==bad[r[3]]:
        assert r[7] != '已交叉验证', 'Truncated fact still displayed as verified'
alerts = [r for r in book['12_提醒'].iter_rows(min_row=4,values_only=True) if str(r[0])=='688087']
assert len(alerts)==1 and alerts[0][4] not in ('买入研究候选','减仓研究候选')
assert alerts[0][7] and alerts[0][7]!='数据门禁通过，仍需投资决策'
print(json.dumps({'symbol':'688087','annual_evidence':[
    {'field':r[3],'value':r[5],'status':r[7]} for r in rows],
    'alert':alerts[0][4],'reasons':alerts[0][7],
    'known_truncations_not_trusted':True},ensure_ascii=False))
book.close()
