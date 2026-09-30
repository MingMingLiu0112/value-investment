from __future__ import annotations

from pathlib import Path
import importlib.util
import json
import subprocess
import sys

from openpyxl import load_workbook
import pytest

from product_workbench_candidate_fixture import (
    materialize_synthetic_legacy_packet,
    synthetic_legacy_packet,
)


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "current" / "build_simulated_product_user_trial.py"


def _synthetic_cli(tmp_path, monkeypatch):
    spec = importlib.util.spec_from_file_location("synthetic_trial_cli", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(module, "ROOT", tmp_path)
    monkeypatch.setattr(module, "build_packet", lambda *_: pytest.fail("production packet read"))
    source = tmp_path / "runtime" / "synthetic-input"
    packet = synthetic_legacy_packet()
    packet["simulation_only"] = True
    materialize_synthetic_legacy_packet(source, packet)
    packet_path = source / "packet.json"
    packet_path.write_text(json.dumps(packet), encoding="utf-8")
    output = tmp_path / "runtime" / "preview.xlsx"
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "--simulated-user-trial", "--synthetic-packet", str(packet_path),
        "--evidence-root", str(source), "--output", str(output),
    ])
    return module, packet, packet_path, output


def test_synthetic_cli_runs_without_production_artifacts(tmp_path, monkeypatch):
    module, _, _, output = _synthetic_cli(tmp_path, monkeypatch)
    assert module.main() == 0
    workbook = load_workbook(output)
    try:
        text = "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value)
        assert "模拟产品体验" in text
        home_text = "\n".join(str(cell.value) for row in workbook["01_今日"] for cell in row if cell.value)
        assert "模拟产品体验" in home_text
        assert "模拟产品体验" in workbook["决策过程"]["A2"].value
        assert "加仓纪律" in text
        assert "no_order" in text
        assert len(workbook.sheetnames) == 7
        decision_text = "\n".join(
            str(cell.value) for row in workbook["决策过程"] for cell in row if cell.value
        )
        assert "available_at" in decision_text
        assert "不默认采用 FCFF" in decision_text
        assert "未提供时不输出个性化仓位" in decision_text
        assert "NOT_READY" in decision_text
    finally:
        workbook.close()
    manifest = json.loads(output.with_name(output.stem + ".m7-product-workbench-candidate-manifest.json").read_text(encoding="utf-8"))
    assert manifest["canonical_pointer_modified"] is False
    assert manifest["counts"]["portfolio_positions"] == 0


def test_synthetic_cli_rejects_undeclared_simulation(tmp_path, monkeypatch):
    module, packet, packet_path, output = _synthetic_cli(tmp_path, monkeypatch)
    packet.pop("simulation_only")
    packet_path.write_text(json.dumps(packet), encoding="utf-8")
    with pytest.raises(ValueError, match="simulation_only"):
        module.main()
    assert not output.exists()


def test_synthetic_cli_rejects_tampered_evidence(tmp_path, monkeypatch):
    module, packet, packet_path, output = _synthetic_cli(tmp_path, monkeypatch)
    evidence = packet_path.parent / packet["audit"]["artifacts"][0]["path"]
    evidence.write_bytes(b"changed")
    with pytest.raises(ValueError, match="hash|SHA|changed|mismatch"):
        module.main()
    assert not output.exists()


def test_synthetic_cli_rejects_external_evidence_root(tmp_path, monkeypatch):
    module, _, _, output = _synthetic_cli(tmp_path, monkeypatch)
    arguments = list(sys.argv)
    arguments[arguments.index("--evidence-root") + 1] = str(tmp_path / "external")
    monkeypatch.setattr(sys, "argv", arguments)
    with pytest.raises(ValueError, match="under runtime"):
        module.main()
    assert not output.exists()


def test_synthetic_cli_rejects_order_action(tmp_path, monkeypatch):
    module, packet, packet_path, output = _synthetic_cli(tmp_path, monkeypatch)
    packet["action"] = "order"
    packet_path.write_text(json.dumps(packet), encoding="utf-8")
    with pytest.raises(ValueError, match="no_order"):
        module.main()
    assert not output.exists()


def test_simulated_product_preview_requires_explicit_acknowledgement(tmp_path: Path) -> None:
    completed = subprocess.run(
        [sys.executable, str(SCRIPT), "--output", str(tmp_path / "preview.xlsx")],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )

    assert completed.returncode != 0
    assert "--simulated-user-trial" in (completed.stderr + completed.stdout)
    assert not (tmp_path / "preview.xlsx").exists()


def test_simulated_product_preview_refuses_output_outside_runtime() -> None:
    output = ROOT.parent / "simulated-product-preview-outside-runtime.xlsx"
    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--simulated-user-trial",
            "--output",
            str(output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )

    assert completed.returncode != 0
    assert "must remain under runtime" in (completed.stderr + completed.stdout)
    assert not output.exists()
