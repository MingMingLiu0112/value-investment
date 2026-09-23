from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.build_m5_materiality_bridge_candidate import build_models
from value_investment_agent.m5_materiality_workbook import (
    build_m5_materiality_workbook,
    write_m5_materiality_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "m5_materiality_bridge_demo.json"


def test_simulated_materiality_workbook_has_expected_sheets_and_boundary():
    batch, receipt, watermark = build_models(FIXTURE)
    workbook = build_m5_materiality_workbook(
        batch,
        receipt,
        watermark,
        security_names={"600887": "模拟公司"},
    )

    assert workbook.sheetnames == [
        "00_总览",
        "01_材料性映射",
        "02_事件账",
        "03_依赖失效与重算",
        "04_Outbox",
        "05_输入与边界",
    ]
    overview = workbook["00_总览"]
    assert overview["B5"].value == "SIMULATED"
    assert overview["B9"].value == "需关注"
    assert overview["B10"].value == 6
    assert overview["B11"].value == 3
    assert overview["B12"].value == 3
    assert overview["B13"].value == 3
    assert overview["B14"].value == 3
    assert overview["B15"].value == "no_order"
    assert workbook["01_材料性映射"].max_row == 10
    assert workbook["02_事件账"].max_row == 7
    assert workbook["03_依赖失效与重算"].max_row == 17
    assert workbook["04_Outbox"].max_row == 7
    assert workbook["05_输入与边界"].max_row == 12

    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "自动交易" not in text
    assert "下单" not in text
    assert "目标仓位" not in text


def test_materiality_workbook_writer_pins_hash_and_rejects_existing_output(tmp_path):
    batch, receipt, watermark = build_models(FIXTURE)
    output = tmp_path / "materiality-candidate.xlsx"
    result = write_m5_materiality_workbook(
        batch,
        receipt,
        watermark,
        output=output,
        root=tmp_path,
        security_names={"600887": "模拟公司"},
    )
    loaded = load_workbook(output)

    assert result["sheet_count"] == 6
    assert result["materiality_decision_count"] == 6
    assert result["silent_decision_count"] == 3
    assert result["event_count"] == 3
    assert result["invalidation_count"] == 3
    assert result["alert_count"] == 3
    assert result["action"] == "no_order"
    assert loaded.sheetnames[0] == "00_总览"
    assert result["workbook_sha256"] == result["workbook_sha256"].lower()

    with pytest.raises(ValueError, match="already exists"):
        write_m5_materiality_workbook(
            batch,
            receipt,
            watermark,
            output=tmp_path / "materiality-candidate.xlsx",
            root=tmp_path,
            security_names={"600887": "模拟公司"},
        )
