from decimal import Decimal

from openpyxl import Workbook
import pytest

from value_investment_agent.excel_report import read_holdings


def sheet(values):
    ws = Workbook().active
    for index, (code, quantity) in enumerate(values, start=10):
        ws.cell(index, 1, code)
        ws.cell(index, 5, quantity)
    return ws


def test_zero_is_known_and_odd_lot_is_allowed():
    assert read_holdings(sheet([('600519', 0), ('000333', 51)])) == {
        '600519': Decimal(0), '000333': Decimal(51)}


@pytest.mark.parametrize('values', [(100, 0), (0, 100), (100, 100), (100, None, 100)])
def test_duplicate_rows_never_establish_positive_holding(values):
    result = read_holdings(sheet([('600519', value) for value in values]))
    assert result['600519'] is None


@pytest.mark.parametrize('value', [True, False, -1, 1.5, float('inf'), float('nan'), '=SUM(E11:E12)', '100', None])
def test_unknown_or_invalid_quantity_is_not_a_holding(value):
    assert read_holdings(sheet([('600519', value)]))['600519'] is None


def test_non_security_rows_are_ignored():
    assert read_holdings(sheet([('TOTAL', 100), (None, 100)])) == {}
