from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path

import pytest

from test_bound_valuation_step import result
from value_investment_agent.application.product.existing_valuation import load_existing_valuation
from value_investment_agent.domain.research.research_run_contract import valuation_result_sha256


def fixture(tmp_path):
    original = tmp_path / 'original.txt'
    original.write_bytes(b'original')
    valuation = replace(result(), evidence_refs=[{'id': 'fixture', 'sha256': hashlib.sha256(original.read_bytes()).hexdigest()}])
    artifact = tmp_path / 'result.json'
    artifact.write_text(json.dumps(dict(action='no_order', valuation_result=json.loads(valuation.to_json()),
                                       valuation_result_sha256=valuation_result_sha256(valuation))), encoding='utf-8')
    return dict(root=tmp_path, artifact_path=artifact,
                expected_sha256=hashlib.sha256(artifact.read_bytes()).hexdigest(),
                evidence_paths={'fixture': original}, observed_at=datetime(2026, 9, 30, tzinfo=timezone.utc))


def test_existing_read_is_observed_only_without_gate_admission(tmp_path):
    kwargs = fixture(tmp_path)
    valuation, records = load_existing_valuation(**kwargs)
    assert valuation.status == 'conditional_research_only'
    assert records[0]['available_at'] == kwargs['observed_at'].date().isoformat()
    assert records[0]['action'] == 'no_order'


@pytest.mark.parametrize('fault', [None, 'input_hash', 'source_hash', 'scenario', 'symbol', 'basis', 'model_version',
    'primary_valid', 'primary_value', 'primary_unit', 'primary_period', 'primary_source',
    'primary_admission', 'primary_future', 'primary_timezone', 'primary_missing', 'primary_excerpt'])
def test_reconstructed_arithmetic_replay_is_bound_and_not_pit_admission(tmp_path, fault):
    from value_investment_agent.application.product.existing_valuation import replay_residual_income_input
    from value_investment_agent.valuation_models.residual_income import MODEL_VERSION
    valuation = replace(result(), model_type='residual_income_or_equity_value', model_version=MODEL_VERSION,
                        bear_value=Decimal('10'), base_value=Decimal('10'), bull_value=Decimal('10'))
    if fault == 'model_version': valuation = replace(valuation, model_version='retained-older-model')
    source = tmp_path / 'review.json'
    source.write_text('{}', encoding='utf-8')
    packet = dict(schema_version='residual-income-arithmetic-replay-input-v1', action='no_order',
        symbol=valuation.symbol, model_version=MODEL_VERSION, input_basis='reconstructed_from_pinned_review',
        valuation_result_sha256=valuation_result_sha256(valuation), valuation_date=valuation.valuation_date.isoformat(),
        source_bindings=[dict(path=source.name, sha256=hashlib.sha256(source.read_bytes()).hexdigest())],
        start_book_equity='100', ordinary_shares='10', scenarios={key:dict(cost_of_equity='0.1',
            forecast_roe=['0.1'], terminal_roe='0.1', terminal_growth='0', retention='0') for key in ('bear','base','bull')})
    if fault == 'source_hash': packet['source_bindings'][0]['sha256'] = 'a' * 64
    if fault == 'scenario': packet['scenarios']['bear']['forecast_roe'] = ['0.2']
    if fault == 'symbol': packet['symbol'] = '999999'
    if fault == 'basis': packet['valuation_date'] = '2000-01-01'
    if fault and fault.startswith('primary_'):
        observed = '2026-09-30T00:00:00+00:00'
        facts = [dict(symbol=valuation.symbol, fact_name=key, value=packet[key], unit=unit,
            period=valuation.valuation_date.isoformat(), source_id='primary-fixture',
            source_url='https://example.test/report', source_path=source.name,
            source_file_hash=hashlib.sha256(source.read_bytes()).hexdigest(),
            available_at=observed, verification_status='NUMERIC_SOURCE_REVIEW_ONLY',
            availability_basis='current_observation_only', current_admission=False,
            historical_availability_verified=False, evidence_excerpt='synthetic excerpt',
            parser_version='fixture', physical_page=1)
            for key, unit in [('start_book_equity', 'CNY'), ('ordinary_shares', 'shares')]]
        review = dict(schema_version='primary-equity-basis-numeric-review-v1', symbol=valuation.symbol,
            action='no_order', current_admission=False, financial_gate_admission=False,
            strict_pit='NOT_PROVEN', observed_at=observed, facts=facts)
        if fault == 'primary_value': facts[0]['value'] = '101'
        if fault == 'primary_unit': facts[0]['unit'] = 'shares'
        if fault == 'primary_period': facts[0]['period'] = '2000-01-01'
        if fault == 'primary_source': facts[0]['source_file_hash'] = 'a' * 64
        if fault == 'primary_admission': review['current_admission'] = True
        if fault == 'primary_future': review['observed_at'] = '2100-01-01T00:00:00+00:00'
        if fault == 'primary_timezone': review['observed_at'] = '2026-09-30T00:00:00'
        if fault == 'primary_missing': review['facts'] = facts[:1]
        if fault == 'primary_excerpt': facts[0]['evidence_excerpt'] = ''
        reviewed = tmp_path / 'facts.json'
        reviewed.write_text(json.dumps(review), encoding='utf-8')
        packet['source_bindings'].append(dict(path=reviewed.name, sha256=hashlib.sha256(reviewed.read_bytes()).hexdigest()))
        packet['primary_numeric_review'] = dict(path=reviewed.name,
            scope='two_input_values_verified_against_primary_report', current_admission=False)
    path = tmp_path / 'input.json'
    path.write_text(json.dumps(packet), encoding='utf-8')
    digest = 'a' * 64 if fault == 'input_hash' else hashlib.sha256(path.read_bytes()).hexdigest()
    if fault and fault != 'primary_valid':
        with pytest.raises(ValueError): replay_residual_income_input(root=tmp_path, path=path, expected_sha256=digest, valuation=valuation)
    else:
        replay = replay_residual_income_input(root=tmp_path, path=path, expected_sha256=digest, valuation=valuation)
        assert replay['status'] == 'ARITHMETIC_MATCH'
        assert len(replay['scenarios']) == 3
        assert replay['original_run_input_descriptor_verified'] is False
        assert replay['current_admission'] is False
        if fault == 'primary_valid':
            assert replay['primary_numeric_review']['status'] == 'INPUT_VALUES_BOUND_TO_REVIEWED_FACTS'
            assert replay['primary_numeric_review']['financial_gate_admission'] is False
            assert replay['primary_numeric_review']['source_excerpt_semantics_verified'] is False
        else:
            assert replay['primary_numeric_review'] is None


@pytest.mark.parametrize('field,value', [
    ('available_at', '2026-10-01T00:00:00+08:00'),
    ('computed_at', '2026-10-01T00:00:00+08:00'),
    ('available_at', '2026-09-29T00:00:00'),
    ('research_as_of', '2026-10-01'),
])
def test_existing_read_rejects_future_or_timezone_free_followup(tmp_path, field, value):
    from value_investment_agent.application.product.existing_valuation import read_existing_research_result
    kwargs = fixture(tmp_path)
    payload = json.loads(kwargs['artifact_path'].read_bytes())
    payload['run'] = {field: value}
    kwargs['artifact_path'].write_text(json.dumps(payload), encoding='utf-8')
    manifest = tmp_path / 'manifest.json'
    manifest.write_text(json.dumps(dict(status='PASS', action='no_order',
        artifact_path=kwargs['artifact_path'].name,
        artifact_sha256=hashlib.sha256(kwargs['artifact_path'].read_bytes()).hexdigest(),
        verified_at=kwargs['observed_at'].isoformat(),
        evidence=[dict(evidence_id='fixture', path='original.txt')])), encoding='utf-8')
    with pytest.raises(ValueError):
        read_existing_research_result(root=tmp_path, symbol='600887', manifest_path=manifest,
            manifest_sha256=hashlib.sha256(manifest.read_bytes()).hexdigest(), output_path=tmp_path / 'output.json')
    assert not (tmp_path / 'output.json').exists()


@pytest.mark.parametrize('fault', [None, 'hash', 'symbol', 'future'])
def test_pinned_original_research_gate_is_historical_not_current_admission(tmp_path, fault):
    from value_investment_agent.application.product.existing_valuation import read_existing_research_result
    from value_investment_agent.domain.research.research_gate import GATE_VALUATION

    kwargs = fixture(tmp_path)
    root = Path(__file__).resolve().parents[1]
    raw = json.loads((root / 'config/m1-valuation-packages-v1/600887-quality-compounder.json').read_text(encoding='utf-8'))
    if fault == 'symbol': raw['research_case']['symbol'] = '000333'
    if fault == 'future': raw['research_case']['as_of'] = '2026-10-01'
    package = tmp_path / 'package.json'
    package.write_text(json.dumps(raw), encoding='utf-8')
    payload = json.loads(kwargs['artifact_path'].read_bytes())
    payload['run'] = {'source_package': {'path': package.name,
        'sha256': 'a' * 64 if fault == 'hash' else hashlib.sha256(package.read_bytes()).hexdigest()}}
    kwargs['artifact_path'].write_text(json.dumps(payload), encoding='utf-8')
    manifest = tmp_path / 'manifest.json'
    manifest.write_text(json.dumps(dict(status='PASS', action='no_order',
        artifact_path=kwargs['artifact_path'].name,
        artifact_sha256=hashlib.sha256(kwargs['artifact_path'].read_bytes()).hexdigest(),
        verified_at=kwargs['observed_at'].isoformat(),
        evidence=[dict(evidence_id='fixture', path='original.txt')])), encoding='utf-8')
    call = dict(root=tmp_path, symbol='600887', manifest_path=manifest,
                manifest_sha256=hashlib.sha256(manifest.read_bytes()).hexdigest())
    if fault:
        with pytest.raises(ValueError): read_existing_research_result(**call)
    else:
        output = read_existing_research_result(**call)
        gate = output['historical_research_gate']
        assert gate['results'][GATE_VALUATION] is False
        assert gate['current_admission'] is False
        assert gate['research_as_of'] == '2026-09-22'
        assert output['suggested_state'] == 'NOT_READY'
        assert output['action'] == 'no_order'
        negative = output['negative_decision_review']
        assert negative['scope'] == 'CURRENT_FAIL_CLOSED_READ_OF_EXISTING_RESEARCH'
        assert negative['pre_decision_eligibility']['status'] == 'NOT_ELIGIBLE'
        assert 'human_research_approval_missing' in negative['pre_decision_eligibility']['blockers']
        assert 'event_materiality_review_missing' in negative['pre_decision_eligibility']['blockers']
        assert negative['investment_decision_review']['status'] == 'INSUFFICIENT_RESEARCH'
        assert negative['portfolio_input_status'] == 'BLOCKED_PRIVATE_INPUT'
        assert negative['position_guidance'] is None
        assert negative['action'] == 'no_order'
        from test_product_workbench_read_model import _payload as base_payload
        from value_investment_agent.presentation.read_models.product_workbench import product_workbench_from_payload
        from value_investment_agent.presentation.read_models.existing_research_report import project_existing_research_workbench
        from value_investment_agent.presentation.excel.product_workbench import build_product_workbench_workbook
        base = base_payload()
        base['as_of'] = negative['investment_decision_review']['decision_as_of']
        base['generated_at'] = negative['investment_decision_review']['created_at']
        base['companies'][0]['symbol'] = '600887'
        base['opportunities'][0]['symbol'] = '600887'
        base['companies'][0]['price'].update(
            available=False, value_text=None, status='UNAVAILABLE',
            unavailable_reason='No admitted current quote',
            needed_evidence='Current verified PriceBridge',
        )
        wrapped = dict(schema_version='product-existing-research-workbench-v1',
                       action='no_order', scope='EXISTING_RESEARCH_ONLY_NOT_CURRENT_ADVICE',
                       symbol='600887', suggested_state='NOT_READY', position_guidance=None,
                       research=output)
        projected = project_existing_research_workbench(product_workbench_from_payload(base), wrapped)
        workbook = build_product_workbench_workbook(projected)
        try:
            opportunity_text = '\n'.join(str(cell.value) for row in workbook['02_机会'] for cell in row if cell.value)
            company_text = '\n'.join(str(cell.value) for row in workbook['03_公司'] for cell in row if cell.value)
            assert '研究不足' in opportunity_text
            assert '暂不可评估' in opportunity_text
            assert '人工买入复核条件' in company_text
            assert '尚未接入真实组合' in company_text
        finally:
            workbook.close()
        dependencies = output['dependency_view']
        assert dependencies['original_research_as_of'] == gate['research_as_of']
        assert dependencies['source_package_sha256'] == gate['source_package_sha256']
        assert dependencies['original_case_is_current_approval'] is False


@pytest.mark.parametrize('fault', ['artifact', 'original', 'missing', 'escape', 'date', 'timezone'])
def test_existing_read_fails_closed(tmp_path, fault):
    kwargs = fixture(tmp_path)
    if fault == 'artifact': kwargs['expected_sha256'] = 'a' * 64
    elif fault == 'original': kwargs['evidence_paths']['fixture'].write_bytes(b'changed')
    elif fault == 'missing': kwargs['evidence_paths'] = {}
    elif fault == 'escape': kwargs['evidence_paths'] = {'fixture': tmp_path.parent / 'outside.txt'}
    elif fault == 'date': kwargs['observed_at'] = datetime(2020, 1, 1, tzinfo=timezone.utc)
    else: kwargs['observed_at'] = datetime(2026, 9, 30)
    with pytest.raises(ValueError): load_existing_valuation(**kwargs)


def test_existing_application_result_renders_scenarios_without_passing_other_gates(tmp_path):
    from test_product_workbench_read_model import _payload
    from value_investment_agent.presentation.read_models.product_workbench import EvidenceRecord, product_workbench_from_payload
    from value_investment_agent.presentation.read_models.valuation_step import workbench_with_bound_valuation
    from value_investment_agent.presentation.excel.product_workbench import build_product_workbench_workbook
    from openpyxl import load_workbook

    kwargs = fixture(tmp_path)
    valuation, records = load_existing_valuation(**kwargs)
    model = product_workbench_from_payload(_payload())
    from dataclasses import asdict
    from value_investment_agent.presentation.read_models.existing_research_report import public_workbench_payload_from_snapshot
    snapshot = json.loads(json.dumps(asdict(model), default=lambda value: value.isoformat()))
    restored = product_workbench_from_payload(public_workbench_payload_from_snapshot(snapshot))
    assert restored.companies == model.companies
    next(stage for stage in snapshot['stage_summaries'] if stage['stage_key'] == 'm6')['status']['code'] = 'PARTIAL'
    snapshot['companies'][0]['price']['status']['code'] = 'NOT_READY'
    snapshot['companies'][0]['price']['available'] = False
    snapshot['companies'][0]['price']['value_text'] = None
    snapshot['companies'][0]['price']['unavailable_reason'] = 'Missing quote'
    snapshot['companies'][0]['price']['needed_evidence'] = 'Current quote'
    migrated = product_workbench_from_payload(public_workbench_payload_from_snapshot(snapshot))
    assert next(stage for stage in migrated.stage_summaries if stage.stage_key == 'm6').status.code == 'OPERATIONAL_NOT_STARTED'
    assert not migrated.companies[0].price.available
    assert snapshot['companies'][0]['price']['status']['code'] == 'NOT_READY'
    snapshot['portfolio']['real_data_available'] = True
    with pytest.raises(ValueError, match='nonpersonalized'):
        public_workbench_payload_from_snapshot(snapshot)
    missing_price = replace(model.companies[0].price, available=False, value_text=None,
                            unavailable_reason='No admitted current quote or quote date',
                            needed_evidence='Verified current PriceBridge')
    model = replace(model, companies=(replace(model.companies[0], symbol=valuation.symbol, price=missing_price),),
                    opportunities=(replace(model.opportunities[0], symbol=valuation.symbol),),
                    audit_evidence=(*model.audit_evidence, *(EvidenceRecord(
                        item['evidence_id'], item['title'], item['artifact_type'], item['path'], item['sha256'],
                        datetime.fromisoformat(item['available_at']).date(), source_url=item['source_url']) for item in records)),
                    as_of=kwargs['observed_at'].date())
    projected = workbench_with_bound_valuation(model, valuation,
                                               expected_sha256=valuation_result_sha256(valuation), assessment_id='fixture-run')
    assert projected.companies[0].decision_process[3].status == 'CONDITIONAL'
    assert all(step.status == 'BLOCKED' for step in projected.companies[0].decision_process if step.key != 'valuation')
    workbook = build_product_workbench_workbook(projected)
    output = tmp_path / 'preview.xlsx'
    workbook.save(output)
    with_workbook = load_workbook(output)
    try:
        values = [str(cell.value) for row in with_workbook['03_公司'] for cell in row if cell.value is not None]
        assert '人民币 8.00 / 股' in values
        assert '人民币 11.00 / 股' in values
        assert '人民币 13.00 / 股' in values
        scenario_row = next(cell.row for row in with_workbook['03_公司'] for cell in row
                            if cell.value == 'Bear / Base / Bull')
        review_rows = [cell.row for row in with_workbook['03_公司'] for cell in row if cell.value == '决策复核']
        assert all(scenario_row < row for row in review_rows)
        opportunity = with_workbook['02_机会']
        labels = {
            str(cell.value): cell.row
            for row in opportunity.iter_rows(min_col=1, max_col=1)
            for cell in row
            if cell.value is not None
        }

        def opportunity_value(label: str) -> str:
            return opportunity.cell(row=labels[label], column=2).value

        valuation_text = opportunity_value('估值区间 / 日期 / 置信度')
        assert '8.00/11.00/13.00' in valuation_text
        assert '2026-06-30' in valuation_text
        assert '置信度' in valuation_text
        assert '暂不可评估' in opportunity_value('当前价格 / 日期')
        assert opportunity_value('当前建议与原因').startswith('暂未通过')
        assert not any(
            'BLOCKED' in str(cell.value)
            for row in opportunity.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert len(with_workbook.sheetnames) == 7
        for sheet in with_workbook:
            assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
            assert sheet.page_setup.fitToWidth == 1
            assert sheet.page_setup.orientation == 'landscape'
    finally:
        with_workbook.close()


def test_existing_workbench_keeps_pending_gates_and_renders_verified_sources(tmp_path, monkeypatch):
    from value_investment_agent.application.product.workbench import build_current_workbench_for_symbol
    from value_investment_agent.presentation.read_models.existing_research_report import render_existing_research_report

    kwargs = fixture(tmp_path)
    manifest = tmp_path / 'manifest.json'
    manifest.write_text(json.dumps(dict(status='PASS', action='no_order',
        artifact_path=kwargs['artifact_path'].name, artifact_sha256=kwargs['expected_sha256'],
        verified_at=kwargs['observed_at'].isoformat(), evidence=[dict(evidence_id='fixture', path='original.txt')])), encoding='utf-8')
    digest = hashlib.sha256(manifest.read_bytes()).hexdigest()
    call = dict(root=tmp_path, symbol='600887', existing_manifest_path=manifest,
                existing_manifest_sha256=digest, output_path=tmp_path / 'workbench.json')
    output = build_current_workbench_for_symbol(**call)
    payload = output['result']
    assert payload['research']['research_rerun'] is False
    assert payload['research']['price_bridge']['current_price'] is None
    assert payload['suggested_state'] == 'NOT_READY'
    assert payload['canonical_workbook_written'] is False
    text = render_existing_research_report(payload)
    assert 'NOT current investment advice' in text
    assert kwargs['expected_sha256'] == hashlib.sha256(kwargs['artifact_path'].read_bytes()).hexdigest()
    assert payload['research']['source_records'][0]['sha256'] in text
    assert payload['research']['valuation']['base_value'] in text
    from test_product_workbench_read_model import _payload
    from value_investment_agent.presentation.read_models.product_workbench import product_workbench_from_payload
    from value_investment_agent.presentation.read_models.existing_research_report import project_existing_research_workbench
    model = product_workbench_from_payload(_payload())
    missing = replace(model.companies[0].price, available=False, value_text=None,
                      unavailable_reason='No current quote', needed_evidence='Current admitted quote')
    model = replace(model, as_of=kwargs['observed_at'].date(),
                    companies=(replace(model.companies[0], symbol='600887', price=missing),),
                    opportunities=(replace(model.opportunities[0], symbol='600887'),))
    projected = project_existing_research_workbench(model, payload)
    assert projected.companies[0].decision_process[3].status == 'CONDITIONAL'
    assert not projected.companies[0].price.available
    assert all(step.status == 'BLOCKED' for step in projected.companies[0].decision_process if step.key != 'valuation')
    assert model.companies[0].scenarios != projected.companies[0].scenarios
    from value_investment_agent.application.product.workbench import load_existing_workbench_for_presentation
    loaded = load_existing_workbench_for_presentation(
        root=tmp_path, path=call['output_path'], expected_sha256=output['receipt']['output_sha256'])
    assert loaded == payload
    import scripts.current.build_product_workbench_candidate as preview_cli
    from argparse import Namespace
    base = _payload()
    base['as_of'] = kwargs['observed_at'].date().isoformat()
    base['companies'][0]['symbol'] = '600887'
    base['opportunities'][0]['symbol'] = '600887'
    base['companies'][0]['price']['available'] = False
    base['companies'][0]['price']['value_text'] = None
    base['companies'][0]['price']['unavailable_reason'] = 'No admitted quote'
    base['companies'][0]['price']['needed_evidence'] = 'Verified PriceBridge'
    base_path = tmp_path / 'base.json'
    base_path.write_text(json.dumps(base), encoding='utf-8')
    preview = tmp_path / 'runtime' / 'historical-preview.xlsx'
    monkeypatch.setattr(preview_cli, 'ROOT', tmp_path)
    monkeypatch.setattr(preview_cli, 'build_packet', lambda _: pytest.fail('legacy packet must not be used'))
    monkeypatch.setattr(preview_cli, 'parse_args', lambda: Namespace(
        output=preview, historical_preview=True, generated_at=kwargs['observed_at'],
        base_payload=base_path, base_payload_sha256=hashlib.sha256(base_path.read_bytes()).hexdigest(),
        base_read_model_snapshot=False,
        existing_workbench=call['output_path'], existing_workbench_sha256=output['receipt']['output_sha256']))
    assert preview_cli.main() == 0
    assert preview.is_file()
    from test_canonical_workbook_preservation import _workbook
    import scripts.current.publish_product_workbench_to_canonical as publisher
    canonical = tmp_path / 'user-canonical.xlsx'
    _workbook(canonical)
    canonical_hash = publisher._sha256(canonical)
    monkeypatch.setattr(publisher, '_workbook_path', lambda: canonical)
    integrated = tmp_path / 'runtime/integrated/canonical-integration-historical-preview.xlsx'
    monkeypatch.setattr(preview_cli, 'parse_args', lambda: Namespace(
        historical_preview=True, integrate_canonical=True, output=integrated, generated_at=datetime.now(timezone.utc),
        base_payload=base_path, base_payload_sha256=hashlib.sha256(base_path.read_bytes()).hexdigest(),
        base_read_model_snapshot=False,
        existing_workbench=call['output_path'], existing_workbench_sha256=output['receipt']['output_sha256']))
    assert preview_cli.main() == 0
    assert publisher._sha256(canonical) == canonical_hash
    bindings = json.loads(integrated.with_name(integrated.stem + '.source-bindings.json').read_text())
    assert bindings['integrated_canonical'] is True
    assert bindings['canonical_written'] is False
    proof = json.loads((integrated.parent / 'canonical-preservation.json').read_text())
    assert proof['preserved_sheet_count'] == 1
    bad = json.loads(json.dumps(payload))
    bad['research']['valuation']['base_value'] = '12'
    with pytest.raises(ValueError, match='valuation hash mismatch'):
        project_existing_research_workbench(model, bad)
    with pytest.raises((ValueError, FileExistsError)):
        build_current_workbench_for_symbol(**call)
    with pytest.raises(ValueError):
        build_current_workbench_for_symbol(**(call | {'package_path': tmp_path / 'package.json'}))
    payload['research']['suggested_state'] = 'ADD'
    with pytest.raises(ValueError, match='cannot admit'):
        render_existing_research_report(payload)
    kwargs['evidence_paths']['fixture'].write_bytes(b'changed')
    with pytest.raises(ValueError, match='original hash mismatch'):
        load_existing_workbench_for_presentation(
            root=tmp_path, path=call['output_path'], expected_sha256=output['receipt']['output_sha256'])


def test_existing_command_service_reads_only_and_refuses_overwrite_or_wrong_symbol(tmp_path):
    from value_investment_agent.application.product.existing_valuation import read_existing_research_result
    kwargs = fixture(tmp_path)
    manifest = tmp_path / 'manifest.json'
    manifest.write_text(json.dumps(dict(status='PASS', action='no_order',
        artifact_path=kwargs['artifact_path'].name, artifact_sha256=kwargs['expected_sha256'],
        verified_at=kwargs['observed_at'].isoformat(), evidence=[dict(evidence_id='fixture', path='original.txt')])), encoding='utf-8')
    digest = hashlib.sha256(manifest.read_bytes()).hexdigest()
    output = tmp_path / 'read-result.json'
    result_payload = read_existing_research_result(root=tmp_path, symbol='600887', manifest_path=manifest,
                                                   manifest_sha256=digest, output_path=output)
    assert result_payload['research_rerun'] is False
    assert result_payload['canonical_workbook_written'] is False
    assert result_payload['suggested_state'] == 'NOT_READY'
    bridge = result_payload['price_bridge']
    assert bridge['bridge_status'] == 'PENDING_EXTERNAL_DATA'
    assert bridge['model_validity_status'] == 'UNKNOWN'
    assert bridge['quote_status'] == 'PENDING_EXTERNAL_DATA'
    assert bridge['current_price'] is None and bridge['quote_date'] is None
    assert bridge['margin_to_base'] is None and bridge['margin_to_bear'] is None
    assert bridge['valuation_date'] == result_payload['valuation']['valuation_date']
    assert bridge['valuation_base_value'] == result_payload['valuation']['base_value']
    assert bridge['model_version'] == result_payload['valuation']['model_version']
    assert result_payload['model_validity'] == 'NOT_ESTABLISHED'
    assert result_payload['model_validity_result']['status'] == 'UNKNOWN'
    assert result_payload['model_validity_result']['last_material_event_check'] is None
    assert result_payload['dependency_view']['current_price_admission'] is False
    assert result_payload['position_guidance'] is None
    assert result_payload['strict_pit'] == 'NOT_PROVEN'
    assert json.loads(output.read_bytes())['price_bridge'] == bridge
    assert hashlib.sha256(kwargs['artifact_path'].read_bytes()).hexdigest() == kwargs['expected_sha256']
    from test_product_workbench_read_model import _payload
    from value_investment_agent.presentation.read_models.product_workbench import product_workbench_from_payload
    from value_investment_agent.presentation.read_models.valuation_step import company_with_pending_price_bridge
    from value_investment_agent.price_bridge import price_bridge_from_payload
    from value_investment_agent.presentation.excel.product_workbench import build_product_workbench_workbook
    base = product_workbench_from_payload(_payload())
    valuation, _ = load_existing_valuation(**kwargs)
    original = replace(base.companies[0], symbol=valuation.symbol)
    original = replace(original,
        price=replace(original.price, available=False, value_text=None, unavailable_reason='No admitted quote',
                      needed_evidence='Verified quote required'),
        margin_of_safety=replace(original.margin_of_safety, available=False, value_text=None,
                                unavailable_reason='No admitted margins', needed_evidence='Verified bridge required'))
    typed_bridge = price_bridge_from_payload(valuation, bridge)
    projected = company_with_pending_price_bridge(original, valuation, typed_bridge)
    assert projected.price.available is False
    assert projected.price.status.code == 'PENDING_EXTERNAL_DATA'
    assert projected.margin_of_safety.value_text is None
    assert projected.decision_process[-1] == original.decision_process[-1]
    assert projected.original_thesis == original.original_thesis
    assert projected.scenarios == original.scenarios
    with pytest.raises(ValueError, match='available price assessment'):
        company_with_pending_price_bridge(replace(projected, price=base.companies[0].price), valuation, typed_bridge)
    with pytest.raises(ValueError, match='blocked price and decision gates'):
        passing = replace(projected, decision_process=tuple(
            replace(step, status='PASS', evidence_refs=('fixture',), assessment_id='not-approved')
            if step.key == 'decision_gate' else step for step in projected.decision_process))
        company_with_pending_price_bridge(passing, valuation, typed_bridge)
    with pytest.raises(ValueError, match='binding mismatch'):
        company_with_pending_price_bridge(projected, replace(valuation, base_value=Decimal('12')),
                                          typed_bridge)
    workbook = build_product_workbench_workbook(replace(base, companies=(projected,)))
    try:
        assert len(workbook.sheetnames) == 7
        values = [str(cell.value) for row in workbook['决策过程'] for cell in row if cell.value is not None]
        assert any('模型有效性 暂不确定' in value for value in values)
        assert all('模型有效性 UNKNOWN' not in value for value in values)
    finally:
        workbook.close()
    with pytest.raises(FileExistsError):
        read_existing_research_result(root=tmp_path, symbol='600887', manifest_path=manifest, manifest_sha256=digest, output_path=output)
    with pytest.raises(ValueError, match='symbol mismatch'):
        read_existing_research_result(root=tmp_path, symbol='000333', manifest_path=manifest, manifest_sha256=digest)
    with pytest.raises(ValueError, match='manifest hash mismatch'):
        read_existing_research_result(root=tmp_path, symbol='600887', manifest_path=manifest, manifest_sha256='a' * 64)


@pytest.mark.parametrize('extra, expected_error', [
    ([], 'existing mode requires manifest and hash'),
    (['--existing-manifest-sha256', 'a' * 64, '--package', 'unused.json'],
     'existing mode excludes package/schedule/market/review inputs'),
])
def test_cli_existing_mode_requires_explicit_binding_and_excludes_rerun_inputs(extra, expected_error):
    from pathlib import Path
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    completed = subprocess.run([sys.executable, '-X', 'utf8', str(root / 'scripts/current/run_company_research.py'),
                                '--symbol', '600887', '--existing-manifest', 'unused.json', *extra],
                               cwd=root, capture_output=True, text=True, encoding='utf-8')
    assert completed.returncode == 2
    assert expected_error in completed.stderr
@pytest.mark.parametrize("symbol", ["000333", "600887", "601088"])
def test_current_workbench_explains_stopped_cases_without_rerunning_research(tmp_path, monkeypatch, symbol):
    from value_investment_agent.application.product.workbench import build_current_workbench_for_symbol
    from value_investment_agent.application.product import company_research
    from value_investment_agent.presentation.read_models.existing_research_report import render_current_research_readiness

    repository = Path(__file__).resolve().parents[1]
    config = tmp_path / "config"
    config.mkdir()
    source = repository / "config/research-evidence-stop-ledger-v1.json"
    (config / source.name).write_bytes(source.read_bytes())
    def forbidden(*args, **kwargs):
        raise AssertionError("stopped research must not rebuild a descriptor")
    monkeypatch.setattr(company_research, "build_descriptor", forbidden)
    result = build_current_workbench_for_symbol(root=tmp_path, symbol=symbol,
                                                output_path=tmp_path / "runtime/result.json")
    payload = result["result"]
    assert payload["research_status"] == "BLOCKED_BY_RESEARCH_SCHEDULER"
    assert payload["valuation"] is None
    assert payload["price_bridge"] is None
    assert payload["current_status"] is None
    assert payload["evidence_stops"]
    assert all(stop["symbol"] == symbol for stop in payload["evidence_stops"])
    report = render_current_research_readiness(payload)
    assert symbol in report
    assert "重新开展条件" in report
    assert "停止台账 SHA-256" in report
    assert "action=no_order" in report
    from value_investment_agent.application.product.workbench import load_stopped_workbench_for_presentation
    from value_investment_agent.presentation.read_models.research_readiness import project_research_readiness
    from value_investment_agent.presentation.read_models.product_workbench import product_workbench_from_payload
    from test_product_workbench_read_model import _payload
    source_path = tmp_path / "runtime/result.json"
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    packet = load_stopped_workbench_for_presentation(root=tmp_path, path=source_path, expected_sha256=source_hash)
    model = product_workbench_from_payload(_payload())
    card = replace(model.companies[0], symbol=symbol)
    model = replace(model, as_of=datetime.now(timezone.utc).date(), companies=(card,))
    projected = project_research_readiness(model, packet)
    assert projected.companies[0].valuation == card.valuation
    for before, after in zip(card.decision_process, projected.companies[0].decision_process):
        assert after.status == before.status
        assert after.assessment_id == before.assessment_id
        if before.key != 'research_gate':
            assert after == before
        else:
            assert '原始重开条件' in after.next_action
            assert after.evidence_refs
            for stop in payload['evidence_stops']:
                assert stop['research_question_id'] in after.reason
                assert stop['reopen_condition'] in after.next_action
    from value_investment_agent.presentation.read_models.conditional_expectations import render_company_review_cards
    guide = render_company_review_cards(projected)
    assert '下一动作：' in guide
    assert '买入、加仓与退出的一致性边界' in guide
    assert source_hash in guide
    assert 'action=no_order' in guide
    assert projected.portfolio == model.portfolio
    assert projected.today_items == model.today_items
    assert len(projected.companies[0].decision_review) > len(card.decision_review)
    assert len(projected.audit_evidence) == len(model.audit_evidence) + 2
    (config / source.name).write_bytes(source.read_bytes() + b"\n")
    with pytest.raises(ValueError, match="ledger hash mismatch"):
        load_stopped_workbench_for_presentation(root=tmp_path, path=source_path, expected_sha256=source_hash)
    payload["valuation"] = {"status": "READY"}
    with pytest.raises(ValueError, match="cannot expose newly admitted"):
        render_current_research_readiness(payload)
