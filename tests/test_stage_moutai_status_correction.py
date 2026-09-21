import importlib.util
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "stage_moutai_status_correction", ROOT / "scripts" / "stage_moutai_status_correction.py")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def build_fixture(path: Path) -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "00_首页Dashboard"
    overview["A1"] = "Unchanged dashboard"
    research = workbook.create_sheet("09_公司研究")
    research.append(["placeholder"] * 16)
    research.append(["placeholder"] * 16)
    research.append(["代码", "公司"] + ["字段"] * 13 + ["研究状态"])
    research.append(["600519", "贵州茅台"] + [""] * 13 + ["原状态"])
    research.append(["000651", "格力电器"] + [""] * 13 + ["其他公司状态"])
    decisions = workbook.create_sheet("21_决策验证")
    decisions.append(["placeholder"] * 4)
    decisions.append(["placeholder"] * 4)
    decisions.append(["代码", "公司", "公式", "策略验证状态"])
    decisions.append(["600519", "贵州茅台", "=1+1", "原验证状态"])
    decisions.append(["000651", "格力电器", "=2+2", "其他公司验证状态"])
    workbook.save(path)


def status_cell_values(path: Path) -> tuple[str, str]:
    workbook = load_workbook(path, data_only=False)
    return (
        workbook["09_公司研究"]["P4"].value,
        workbook["21_决策验证"]["D4"].value,
    )


def test_stage_changes_only_moutai_status_and_is_idempotent(tmp_path):
    source, candidate, rerun = (tmp_path / name for name in ("source.xlsx", "candidate.xlsx", "rerun.xlsx"))
    build_fixture(source)
    before = load_workbook(source, data_only=False)
    original_status_alignment = copy(before["09_公司研究"]["P4"].alignment)

    MODULE.stage(source, candidate)
    staged = load_workbook(candidate, data_only=False)
    research, decisions = staged["09_公司研究"], staged["21_决策验证"]
    assert MODULE.STATUS in research["P4"].value
    assert MODULE.STATUS in decisions["D4"].value
    assert decisions["D4"].value.startswith("【当前操作】研究观察：不新增买入、不减仓、不卖出。")
    assert "【为什么】正式合理价、模拟准入和实盘指令均未通过" in decisions["D4"].value
    assert research["P5"].value == "其他公司状态"
    assert decisions["D5"].value == "其他公司验证状态"
    assert decisions["C4"].value == "=1+1"
    assert staged["00_首页Dashboard"]["A1"].value == "Unchanged dashboard"
    assert research["P4"].alignment.wrap_text is True
    assert research["P4"].alignment.vertical == "top"
    assert original_status_alignment != research["P4"].alignment

    MODULE.stage(candidate, rerun)
    first, second = status_cell_values(candidate), status_cell_values(rerun)
    assert first == second
    assert first[0].count(MODULE.STATUS) == 1
    assert first[1].count(MODULE.STATUS) == 1
