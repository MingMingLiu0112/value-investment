from __future__ import annotations

from pathlib import Path
import json

import pytest
from openpyxl import load_workbook

from scripts.build_m4_guidance_income_candidate import build_models, load_input
from value_investment_agent.m4_guidance_income_workbook import (
    build_guidance_income_workbook,
    write_guidance_income_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "m4_guidance_income_demo.json"


def test_simulated_guidance_income_workbook_has_expected_sheets_and_boundary(tmp_path):
    payload, _ = load_input(FIXTURE)
    guidance, dividend = build_models(payload)
    workbook = build_guidance_income_workbook(
        guidance,
        dividend,
        security_names=payload["security_names"],
    )

    assert workbook.sheetnames == [
        "00_总览",
        "01_仓位分层",
        "02_共同预算",
        "03_股息收入",
        "04_输入与边界",
    ]
    overview = workbook["00_总览"]
    assert overview["B5"].value == "SIMULATED"
    assert overview["B6"].value == "共同预算冲突"
    assert overview["B7"].value == "READY"
    assert overview["B12"].value == "no_order"
    assert workbook["01_仓位分层"].max_row == 8
    assert workbook["03_股息收入"].max_row == 8

    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "目标仓位" not in text
    assert "下单" not in text
    assert "position_size" not in text


def test_writer_pins_hash_and_rejects_existing_output(tmp_path):
    payload, _ = load_input(FIXTURE)
    guidance, dividend = build_models(payload)
    output = tmp_path / "candidate.xlsx"
    result = write_guidance_income_workbook(
        guidance,
        dividend,
        output=output,
        root=tmp_path,
        security_names=payload["security_names"],
    )
    loaded = load_workbook(output)

    assert result["sheet_count"] == 5
    assert result["candidate_count"] == 4
    assert result["income_basis_count"] == 4
    assert result["action"] == "no_order"
    assert loaded.sheetnames[0] == "00_总览"
    assert result["workbook_sha256"] == result["workbook_sha256"].lower()

    with pytest.raises(ValueError, match="already exists"):
        write_guidance_income_workbook(
            guidance,
            dividend,
            output=tmp_path / "candidate.xlsx",
            root=tmp_path,
            security_names=payload["security_names"],
        )
    assert "target_weight" not in json.dumps(result, ensure_ascii=False)
