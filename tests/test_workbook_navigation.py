import hashlib
import json

from openpyxl import Workbook, load_workbook
import pytest
from decimal import Decimal

from value_investment_agent.workbook_simple_overview import company_progress, _links, _valuation_cases, _valuation_view, ROOT


def test_research_card_quote_and_formulas_do_not_create_fair_value():
    from value_investment_agent.workbook_simple_overview import _company_brief, _research_quote
    company = {'daily_paper': {'observed_at': '2026-09-16T16:00:00+08:00', 'observed_close_cny': '1258'},
               'observation': {'quote_session_verified': True, 'as_of': '2026-09-14T16:00:00+08:00',
                               'scenarios': [{'observed_price_cny': '1277.96'}]},
               'current_equity_research': {'profit_anchor_cny': '80000000000', 'facts': {
                   'parent_equity_cny': '250000000000', 'issued_shares': '1250000000',
                   'period_end': '2026-06-30', 'source_url': 'https://example.invalid/fixture',
                   'raw_file_hash': 'fixture'}}}
    assert _research_quote(company) == ('2026-09-16', Decimal('1258'))
    w = Workbook()
    _company_brief(w.active, 7, company)
    assert w.active['C20'].value == 250000000000
    assert w.active['C21'].value == 80000000000
    assert w.active['C22'].value == 1250000000
    assert w.active['C23'].value == 1258
    assert w.active['C24'].value == '=C21/C22'
    assert w.active['C25'].value == '=IF(AND(ISNUMBER(C23),C24>0),C23/C24,"")'
    assert w.active['C26'].value == '=C20/C22'
    assert w.active['C24'].data_type == 'f'
    text = '\n'.join(str(c.value) for row in w.active for c in row if c.value)
    assert '不提供建议买入价' in text and '不是公司恶化证据' in text
    assert '1277.96' not in text
    assert _research_quote({}) is None


def companies():
    w = Workbook()
    w.active.title = '01_观察名单'
    w.create_sheet('09_公司研究')
    w['01_观察名单']['A4'], w['01_观察名单']['B4'] = '601088', '中国神华'
    w['09_公司研究']['A4'], w['09_公司研究']['B4'] = '600519', '贵州茅台'
    w['09_公司研究']['A5'], w['09_公司研究']['B5'] = '000333', '美的集团'
    return w


def test_existing_research_row_refreshes_quote_and_unchanged_value_provenance(tmp_path):
    w = Workbook()
    w.active.title = '04_估值跟踪'
    company = {'code': '600519', 'name': '贵州茅台',
               'case': {'data_snapshot': {'period_end': '2025-12-31'}},
               'primary_equity_model': {'results': [
                   {'scenario': name, 'conditional_value_per_2025_issued_share_cny': value}
                   for name, value in [('bear', '670'), ('base', '1053.144906431247413452885390'), ('bull', '1810')]]},
               'primary_equity_model_sha256': 'old-model',
               'observation': {'as_of': '2026-09-11T16:14:50+08:00', 'quote_session_verified': True,
                   'scenarios': [{'scenario': name, 'observed_price_cny': '1275.16', 'conditional_value_per_share_cny': value}
                                 for name, value in [('bear', '670'), ('base', '1053.144906431247413452885390'), ('bull', '1810')]]}}
    _valuation_cases(w, [company])
    w.active['U4'] = 'User research note'
    path = tmp_path / 'saved.xlsx'
    w.save(path)
    w.close()
    w = load_workbook(path)
    company['observation']['as_of'] = '2026-09-14T16:14:50+08:00'
    for row in company['observation']['scenarios']:
        row['observed_price_cny'] = '1277.96'
    company['primary_equity_model_sha256'] = 'corrected-model'
    _valuation_cases(w, [company])
    assert w.active['C4'].value == 1277.96
    assert w.active['O4'].value == '2026-09-14T16:14:50+08:00'
    assert w.active['R4'].value == 670
    assert 'corrected-model' in w.active['R4'].comment.text
    assert '2026-09-14' in w.active['R4'].comment.text
    assert w.active['K4'].value is None and w.active['L4'].value is None
    assert w.active['U4'].value == 'User research note'
    assert _valuation_cases(w, [company]) == {}


def pinned(root, name, filename, data):
    folder = root / 'runtime' / name.removesuffix('-latest.json')
    folder.mkdir(parents=True)
    target = folder / filename
    target.write_text(json.dumps(data), encoding='utf-8')
    pointer = root / 'runtime' / name
    pointer.write_text(json.dumps({'path': str(folder.relative_to(root)),
        'sha256': hashlib.sha256(target.read_bytes()).hexdigest()}), encoding='utf-8')
    return target


def test_fixed_case_outside_candidates_is_listed_once_and_does_not_imply_strategy_acceptance(tmp_path):
    pinned(tmp_path, 'strategy-validation/moutai-simulation-closure-latest.json', 'summary.json', {
        'symbol': '600519', 'execution_mechanics_verified': True,
        'real_history': {'second_run_new_journal_rows': 0},
        'cash_anchor_research_history': {'research_model_decision_sessions': 10, 'second_run_new_journal_rows': 0},
        'simulation_eligible': False, 'trade_approved': False})
    pinned(tmp_path, 'company-research/600519-end-to-end-case-latest.json', 'evidence.json', {'symbol': '600519'})
    rows, _ = company_progress(companies(), tmp_path)
    assert {c['code'] for c in rows} == {'600519', '601088', '000333'}
    assert [c['code'] for c in rows if c['has_results']] == ['600519']
    assert not any(c['strategy_accepted'] for c in rows)


def test_research_labels_without_execution_evidence_do_not_claim_completion(tmp_path):
    w = companies()
    w['09_公司研究']['P4'] = '已跑通 全部通过'
    rows, _ = company_progress(w, tmp_path)
    assert len(rows) == 3
    assert not any(c['has_results'] for c in rows)


def test_current_date_research_does_not_silently_replace_primary_model_or_admit_strategy():
    rows, refs = company_progress(companies(), ROOT)
    company = next(row for row in rows if row['code'] == '600519')
    current = company['current_equity_research']
    primary, label, _ = _valuation_view(company)
    assert '研究日已披露股数' in label
    assert primary['base'] != float(current['results'][1]['conditional_value_per_current_disclosed_share_cny'])
    assert current['facts']['period_end'] == '2026-06-30'
    assert company['strategy_accepted'] is False
    assert current['simulation_eligible'] is False
    assert any('residual-income-current-' in ref['path'] for ref in refs['600519'])
    card = company['research_card']
    assert card['research_card_version'] == 'moutai-research-card-evidence-v6'
    assert card['facts']['liquor_revenue_growth_pct'] == '-1.08'
    assert card['facts']['liquor_sales_volume_growth_pct'] == '2.13'
    assert card['trade_approved'] is False
    assert any('research-card-evidence-' in ref['path'] for ref in refs['600519'])


def test_p05_status_rows_expose_assumptions_materiality_and_gap_types():
    rows, refs = company_progress(companies(), ROOT)
    by_code = {row['code']: row for row in rows}

    assert by_code['600519']['assumption_status'] == 'READY'
    assert by_code['601088']['assumption_status'] == 'PARTIAL'
    assert by_code['000333']['materiality']['materiality'] == 'LOW'
    assert by_code['000333']['materiality']['treatment'] == 'MODEL_AS_RANGE'
    assert 'MODEL_NOT_APPLICABLE' in by_code['000333']['gap_types']
    assert 'ASSUMPTION_MISSING' in by_code['601088']['gap_types']
    assert 'ASSUMPTION_LOW_CONFIDENCE' in by_code['600519']['gap_types']
    assert 'FACT_MISSING' in by_code['600519']['gap_types']
    assert 'UNCLASSIFIED' not in by_code['600519']['gap_types']
    assert any('valuation-assumptions/' in ref['path'].replace('\\', '/') for ref in refs['600519'])


def test_changed_run_evidence_rejects_navigation_promotion(tmp_path):
    target = pinned(tmp_path, 'strategy-validation/moutai-simulation-closure-latest.json', 'summary.json', {})
    target.write_text('{"execution_mechanics_verified": true}', encoding='utf-8')
    with pytest.raises(ValueError, match='evidence changed'):
        company_progress(companies(), tmp_path)


def test_research_link_uses_personal_record_when_automatic_summary_is_empty():
    w = Workbook()
    indexes = {'09_公司研究': {'000333': (19, ('000333', '美的集团'))},
               '04_估值跟踪': {}, '21_决策验证': {}}
    _links(w.active, 4, '000333', indexes)
    assert w.active['F4'].hyperlink.location == "'09_公司研究'!A19"
    assert w.active['F4'].hyperlink.target is None


def test_internal_links_serialize_as_locations_and_public_urls_remain_external(tmp_path):
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    from openpyxl import load_workbook
    from value_investment_agent.workbook_frontdoor import _link, repair_internal_links
    w = Workbook()
    w.active.title = '公司研究'
    w.active['B4'] = '证据'
    _link(w.active['A1'], '查看公司', '公司研究', 4, 'B')
    w.active['A2'].hyperlink = "#'公司研究'!B4"
    w.active['A3'].hyperlink = 'https://static.cninfo.com.cn/example.PDF'
    assert repair_internal_links(w) == 1
    path = tmp_path / 'wps-navigation.xlsx'
    w.save(path)
    with ZipFile(path) as z:
        sheet = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        links = sheet.findall('.//{*}hyperlink')
        assert [x.get('location') for x in links[:2]] == ["'公司研究'!B4"] * 2
        assert all(not any(k.endswith('}id') for k in x.attrib) for x in links[:2])
        relationships = ET.fromstring(z.read('xl/worksheets/_rels/sheet1.xml.rels'))
        assert len(relationships) == 1
        assert relationships[0].get('Target') == 'https://static.cninfo.com.cn/example.PDF'
    saved = load_workbook(path)
    assert saved.active['A1'].hyperlink.target is None
    assert saved.active['A2'].hyperlink.location == "'公司研究'!B4"


def test_retained_case_valuation_does_not_become_approved_price_or_order():
    from value_investment_agent.workbook_simple_overview import _valuation_cases
    w = Workbook()
    w.active.title = '04_估值跟踪'
    w.active['A4'], w.active['B4'], w.active['K4'] = '601088', '中国神华', 20
    case = {'code': '600519', 'name': '贵州茅台', 'case': {'data_snapshot': {'period_end': '2026-06-30'}},
        'observation': {'as_of': '2026-09-11', 'quote_session_verified': True,
            'scenarios': [{'scenario': s, 'conditional_value_per_share_cny': v, 'observed_price_cny': '1275.16'}
                          for s, v in [('bear', '735'), ('base', '956'), ('bull', '1086')]]}}
    _valuation_cases(w, [case])
    assert w.active['K4'].value == 20
    assert w.active['A5'].value == '600519'
    assert w.active['K5'].value is None and w.active['L5'].value is None
    assert w.active['M5'].value == '研究观察；不生成订单'
    assert w.active['R5'].value == 735
    assert _valuation_cases(w, [case]) == {}


def test_research_framework_is_repeatable_and_does_not_touch_source_cells(tmp_path):
    from value_investment_agent.workbook_simple_overview import _guide, GUIDE, _research_contract
    w = companies()
    w['09_公司研究']['C4'] = 'Personal thesis'
    w['09_公司研究']['D4'] = '=1+2'
    _guide(w, {}, {})
    first = tuple(w[GUIDE].values)
    _guide(w, {}, {})
    assert tuple(w[GUIDE].values) == first
    assert w['09_公司研究']['C4'].value == 'Personal thesis'
    assert w['09_公司研究']['D4'].value == '=1+2'
    body = '\n'.join(str(c.value) for row in w[GUIDE] for c in row if c.value)
    for label in ('资产折价 / 烟蒂', '成熟优质复利', '成长价值', '现金回报',
                  '周期正常化', '困境反转', '特殊情形', '财务五问', '市场隐含预期'):
        assert label in body
    assert '暂不支持套利建议' in body
    assert '30%不作所有路径通则' in body
    assert '日常模拟、R1、R2分开验收' in body
    case = w.create_sheet('Case preview')
    end = _research_contract(case, 1, {'code': '600519'})
    assert end == 13
    contract = '\n'.join(str(c.value) for row in case for c in row if c.value)
    assert '错价尚未证明' in contract and '仅作约束，不是目标价' in contract
    assert '不改变当前no_order' in contract
    assert not any(c.data_type == 'f' for row in case for c in row)
    other = w.create_sheet('Unclassified')
    assert _research_contract(other, 1, {'code': '000333'}) == 1
    assert other['A1'].value is None
    path = tmp_path / 'framework.xlsx'
    w.save(path)
    saved = load_workbook(path)
    assert tuple(saved[GUIDE].values) == first
    assert saved['09_公司研究']['D4'].value == '=1+2'
    saved.close()
