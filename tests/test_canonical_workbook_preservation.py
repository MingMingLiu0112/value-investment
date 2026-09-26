from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
import pytest

from scripts.current.publish_product_workbench_to_canonical import _assert_retained, _snapshot
from value_investment_agent.presentation.excel.product_workbench import WORKBOOK_SHEETS


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人工持仓"
    sheet["A1"] = "用户输入"
    sheet["B2"] = "=1+1"
    sheet["C3"].hyperlink = "https://example.test/evidence"
    sheet["D4"].comment = Comment("人工备注", "user")
    sheet.merge_cells("E1:F1")
    sheet.freeze_panes = "B3"
    sheet.row_dimensions[3].height = 24
    sheet.column_dimensions["A"].width = 22
    sheet.protection.sheet = True
    validation = DataValidation(type="list", formula1='"A,B"')
    sheet.add_data_validation(validation)
    validation.add(sheet["G1"])
    sheet.conditional_formatting.add("H1", CellIsRule(operator="greaterThan", formula=["0"]))
    workbook.create_named_range("manual_input", sheet, "A1")
    workbook.save(path)


def test_canonical_snapshot_covers_advanced_preservation_contract(tmp_path: Path):
    path = tmp_path / "canonical.xlsx"
    _workbook(path)

    before = _snapshot(path)
    after = _snapshot(path)
    before["sheet_order"] = list(WORKBOOK_SHEETS) + before["sheet_order"]
    after["sheet_order"] = list(WORKBOOK_SHEETS) + after["sheet_order"]

    _assert_retained(before, after)
    protected = before["sheets"]["人工持仓"]
    assert protected["merged_ranges"]
    assert protected["row_dimensions"]
    assert protected["column_dimensions"]
    assert protected["freeze_panes"] == "B3"
    assert protected["protection"]
    assert protected["comments"]
    assert protected["data_validations"]["count"] == 1
    assert protected["conditional_formatting"]["count"] == 1


def test_canonical_snapshot_fails_when_protected_structure_changes(tmp_path: Path):
    path = tmp_path / "canonical.xlsx"
    _workbook(path)
    before = _snapshot(path)

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook["人工持仓"].freeze_panes = "C3"
    workbook.save(path)
    after = _snapshot(path)
    before["sheet_order"] = list(WORKBOOK_SHEETS) + before["sheet_order"]
    after["sheet_order"] = list(WORKBOOK_SHEETS) + after["sheet_order"]

    with pytest.raises(ValueError, match="protected sheets changed"):
        _assert_retained(before, after)
