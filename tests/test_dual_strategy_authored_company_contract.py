"""Read-only regression against actual v3/v5 authored company pages.

Hotspot observations do not expand the registered company research universe.
These tests never invoke publication or alter source workbooks.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
import pytest


ROOT = Path(__file__).resolve().parents[1]
SHEET = '03_\u516c\u53f8'
HOTSPOT_LABEL = '\u70ed\u70b9\u8bc1\u5238\u89c2\u5bdf'
SYMBOL_PATTERN = re.compile(r'\b\d{6}\b')
CONTRACT_ERROR = 'authored company pages do not match the research model'


def _company_symbols(rows):
    # Same A-column contract as build_protected_research_preview. B/C are
    # observation detail, not registered research-company identity columns.
    return {symbol for row in rows if row
            for symbol in SYMBOL_PATTERN.findall(str(row[0] or ''))}


def _validate_company_contract(rows, expected):
    if _company_symbols(rows) != expected:
        raise ValueError(CONTRACT_ERROR)


def _actual(version):
    folder = ROOT / 'runtime' / 'daily-trade-assistant' / (
        f'20261012-productization-canonical-publish-v{version}')
    path = folder / 'trade-assistant-preview.xlsx'
    if not path.is_file():
        pytest.skip(f'actual v{version} authored workbook is not available')
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    envelope = json.loads((folder / 'product-read-model.json').read_text(encoding='utf-8'))
    expected = {company['symbol'] for company in envelope['snapshot']['companies']}
    binding = envelope['product_extensions']['hotspot_report']
    report_path = (ROOT / binding['path']).resolve()
    assert report_path.is_relative_to(ROOT)
    report_bytes = report_path.read_bytes()
    assert hashlib.sha256(report_bytes).hexdigest() == binding['sha256']
    report = json.loads(report_bytes)
    assert report['action'] == 'no_order'
    hotspot_symbols = {row['symbol'] for row in report['observation_list']}
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = tuple(workbook[SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before_hash
    return rows, expected, hotspot_symbols


def test_actual_v3_rejects_hotspot_symbols_in_registered_company_column():
    rows, expected, hotspot_symbols = _actual(3)
    assert len(expected) == 6
    assert len(expected | hotspot_symbols) == 18
    assert expected <= _company_symbols(rows)
    assert _company_symbols(rows) - expected == hotspot_symbols - expected
    with pytest.raises(ValueError, match=re.escape(CONTRACT_ERROR)):
        _validate_company_contract(rows, expected)


def test_actual_v5_keeps_six_registered_companies_without_expanding_universe():
    rows, expected, _ = _actual(5)
    assert len(expected) == 6
    _validate_company_contract(rows, expected)


def test_actual_v5_preserves_all_source_bound_hotspots_in_column_b():
    rows, expected, hotspot_symbols = _actual(5)
    assert len(expected | hotspot_symbols) == 18
    observation_rows = [row for row in rows if row and row[0] == HOTSPOT_LABEL]
    assert observation_rows
    b_symbols = {symbol for row in observation_rows if len(row) > 1
                 for symbol in SYMBOL_PATTERN.findall(str(row[1] or ''))}
    assert hotspot_symbols <= b_symbols
    for symbol in hotspot_symbols:
        matching = [row for row in observation_rows if len(row) > 1
                    and symbol in SYMBOL_PATTERN.findall(str(row[1] or ''))]
        for period in ('\u65e5\u7ebf', '60 \u5206\u949f', '15 \u5206\u949f'):
            assert any(period in str(row[1]) for row in matching), (symbol, period)


@pytest.mark.parametrize('version', [3, 5])
def test_actual_contract_rejects_unknown_or_missing_registered_symbol_in_memory(version):
    rows, expected, _ = _actual(version)
    if version == 3:
        # Remove the already-known v3 defect in memory to isolate the new tamper.
        rows = tuple(row for row in rows if not row
                     or set(SYMBOL_PATTERN.findall(str(row[0] or ''))) <= expected)
    _validate_company_contract(rows, expected)
    assert '999999' not in expected
    with pytest.raises(ValueError, match=re.escape(CONTRACT_ERROR)):
        _validate_company_contract((*rows, ('999999', 'isolated tamper', None)), expected)
    omitted = sorted(expected)[0]
    missing_rows = tuple(row for row in rows if not row
                         or omitted not in SYMBOL_PATTERN.findall(str(row[0] or '')))
    with pytest.raises(ValueError, match=re.escape(CONTRACT_ERROR)):
        _validate_company_contract(missing_rows, expected)
