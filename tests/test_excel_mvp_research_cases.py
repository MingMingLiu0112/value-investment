import hashlib
import json
from pathlib import Path

from openpyxl import Workbook

from value_investment_agent.current_research_status import (
    current_research_status_from_payloads,
)
from value_investment_agent.workbook_simple_overview import _mvp_research_card


ROOT = Path(__file__).resolve().parents[1]


def test_stage_a_cases_are_pinned_and_not_promoted_to_valuation():
    pointer = json.loads((ROOT / "runtime/excel-mvp-research-cases-latest.json").read_text(encoding="utf-8"))
    evidence = ROOT / pointer["path"] / "evidence.json"
    assert hashlib.sha256(evidence.read_bytes()).hexdigest() == pointer["sha256"]
    payload = json.loads(evidence.read_text(encoding="utf-8"))
    assert payload["formal_trade_instructions"] is False
    records = {record["case"]["symbol"]: record for record in payload["records"]}
    assert set(records) == {"600519", "000333", "601088"}
    for record in records.values():
        case, gate = record["case"], record["gate"]
        assert case["financial_summary"]
        assert case["thesis"] and case["return_driver"] and case["mispricing_hypothesis"]
        assert case["positives"] and case["counter_evidence"]
        assert case["thesis_breakers"] and case["next_events"]
        assert case["evidence_refs"]
        assert case["valuation_status"] == "not_ready"
        assert gate["conclusion"] in {"数据不足", "研究未完成", "估值未就绪"}
        assert gate["conclusion"] != "估值具备研究吸引力"
        assert gate["blockers"]

    midea = records["000333"]
    assert len(midea["case"]["positives"]) >= 3
    assert len(midea["case"]["counter_evidence"]) >= 3
    assert len(midea["case"]["thesis_breakers"]) >= 3
    assert midea["gate"]["results"]["G2_商业论点门"] is True
    assert midea["gate"]["results"]["G0_证据门"] is False
    assert midea["case"]["as_of"] == "2026-09-22"
    assert midea["case"]["financial_summary"]["reported_fy2025_weighted_shares_thousand"] == 7559265
    assert midea["case"]["financial_summary"]["year_end_issued_total_shares"] == 7597145346
    assert midea["case"]["financial_summary"]["attributable_ordinary_equity_cny"] == "223221305000"
    assert midea["case"]["financial_summary"]["minority_equity_cny"] == "13202918000"
    assert midea["case"]["financial_summary"]["attributable_ordinary_net_profit_cny"] == "43945411000"
    assert any(ref["id"] == "midea_share_basis" for ref in midea["case"]["evidence_refs"])
    assert any(ref["id"] == "midea_equity_scope" for ref in midea["case"]["evidence_refs"])
    assert any(ref["id"] == "midea_equity_return_history" for ref in midea["case"]["evidence_refs"])
    assert any(ref["id"] == "midea_finance_co" for ref in midea["case"]["evidence_refs"])
    assert any("会计加权普通股" in item["text"] for item in midea["case"]["counter_evidence"])
    assert any("金融业务" in item["text"] for item in midea["case"]["counter_evidence"])
    assert any("2014-2024历史权益" in item["text"] for item in midea["case"]["counter_evidence"])
    assert any("2014-2024年归母普通股权益" in item["text"] for item in midea["case"]["positives"])
    assert any("美的财务公司2025年经审计总资产" in item["text"] for item in midea["case"]["positives"])
    assert any("年末库存股股数" in blocker for blocker in midea["case"]["blockers"])
    assert any("剩余收益/权益价值路线" in blocker for blocker in midea["case"]["blockers"])
    assert any("2014-2024权益回报历史序列" in blocker for blocker in midea["case"]["blockers"])
    assert any("财务公司规模观察" in blocker for blocker in midea["case"]["blockers"])

    moutai = records["600519"]
    assert len(moutai["case"]["positives"]) >= 3
    assert len(moutai["case"]["counter_evidence"]) >= 3
    assert len(moutai["case"]["thesis_breakers"]) >= 3
    assert moutai["gate"]["results"]["G2_商业论点门"] is True
    assert moutai["gate"]["results"]["G3_估值门"] is False

    shenhua = records["601088"]
    assert len(shenhua["case"]["positives"]) >= 3
    assert len(shenhua["case"]["counter_evidence"]) >= 3
    assert len(shenhua["case"]["thesis_breakers"]) >= 3
    assert shenhua["gate"]["results"]["G0_证据门"] is True
    assert shenhua["gate"]["results"]["G1_财务门"] is False
    assert shenhua["gate"]["results"]["G2_商业论点门"] is True
    assert shenhua["gate"]["results"]["G3_估值门"] is False
    assert any(ref["id"] == "shenhua_candidates" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_ownership" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_ifrs_tax" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_operating_cycle" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_operating_cycle_audit" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_attributable_profit_series" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_price_cost_transport_bridge" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_external_index_provenance" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_public_index_history" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_bspi_reconciliation" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_bspi_point_in_time" for ref in shenhua["case"]["evidence_refs"])
    assert any("候选包" in item["text"] for item in shenhua["case"]["positives"])
    assert any("重要非全资子公司" in item["text"] for item in shenhua["case"]["positives"])
    assert any("HKEX 英文/IFRS 年报" in item["text"] for item in shenhua["case"]["positives"])
    assert any("2014-2025 运营周期证据包" in item["text"] for item in shenhua["case"]["positives"])
    assert any("运营序列审查已完成" in item["text"] for item in shenhua["case"]["positives"])
    assert any("2014-2025 归母经营利润与税负候选序列" in item["text"] for item in shenhua["case"]["positives"])
    assert any("外部煤价与 2025 内部成本/运输桥接" in item["text"] for item in shenhua["case"]["positives"])
    assert any("NCEI/BSPI/CCTD 运营方" in item["text"] for item in shenhua["case"]["positives"])
    assert any("五个历史图表端点" in item["text"] for item in shenhua["case"]["positives"])
    assert any("BSPI 公开端点逐年均值" in item["text"] for item in shenhua["case"]["positives"])
    assert any("BSPI 点时发布包已建立" in item["text"] for item in shenhua["case"]["positives"])
    assert any("中国能源网、易航网、CWESTC、国际煤炭网四处明确转载" in item["text"]
               for item in shenhua["case"]["positives"])
    assert any("分线路到港/到厂全成本披露审计" in item["text"] for item in shenhua["case"]["positives"])
    assert any("2018/2019/2020 运营方原文" in item["text"] for item in shenhua["case"]["next_events"])
    assert "重要非全资子公司逐户税前、税费与少数股东分配未披露" in shenhua["case"]["blockers"]
    assert "普通股分母已批准但尚未与估值日期和其余输入一并注册" in shenhua["case"]["blockers"]
    assert "运营、利润与现金税候选序列仅作研究边界，未批准为模型输入" in shenhua["case"]["blockers"]
    assert any("外部煤价与成本运输桥接存在指标断点" in item for item in shenhua["case"]["blockers"])
    assert any("2025 内部煤电销售 73.2 Mt 与发电耗用 77.7 Mt" in item for item in shenhua["case"]["blockers"])
    assert any(ref["id"] == "shenhua_internal_coal_power" for ref in shenhua["case"]["evidence_refs"])
    assert any(ref["id"] == "shenhua_route_delivered_cost" for ref in shenhua["case"]["evidence_refs"])
    assert any("2025 内部煤电 73.2 / 77.7 百万吨差异" in item["text"] for item in shenhua["case"]["positives"])
    assert any("BSPI 2018/2019/2020 仍缺运营方原文" in item
                   for item in shenhua["case"]["blockers"])
    assert any("分线路运输周转量、分线路成本或路线分配矩阵" in item
               for item in shenhua["case"]["blockers"])
    assert "多年运营序列尚未审核为正常化输入" not in shenhua["case"]["blockers"]


def test_stage_a_case_evidence_references_are_hash_addressable():
    payload = json.loads((ROOT / "runtime/excel-mvp-research-cases/evidence.json").read_text(encoding="utf-8"))
    for record in payload["records"]:
        for ref in record["case"]["evidence_refs"]:
            target = ROOT / ref["path"]
            assert target.exists(), ref
            assert hashlib.sha256(target.read_bytes()).hexdigest() == ref["sha256"]


def test_mvp_card_displays_price_bridge_margins_without_trade_promotion():
    workbook = Workbook()
    sheet = workbook.active
    record = {
        "case": {
            "name": "贵州茅台", "symbol": "600519", "industry": "高端白酒",
            "investment_path": "成熟优质复利", "research_status": "financial_scope_approved",
            "thesis": "测试论点", "return_driver": "测试回报", "mispricing_hypothesis": "未证明",
            "financial_summary": {"period_end": "2026-06-30"}, "financial_period": "2026-06-30",
            "positives": [], "counter_evidence": [], "thesis_breakers": [], "next_events": [],
            "evidence_status": "verified", "evidence_refs": [],
        },
        "gate": {
            "conclusion": "估值未就绪",
            "blockers": ["G3_估值模型通过"],
            "results": {
                "G0_证据门": True,
                "G1_财务门": True,
                "G2_商业论点门": True,
                "G3_估值门": False,
            },
        },
    }
    valuation = {
        "bear_value": "400", "base_value": "500", "bull_value": "600", "valuation_date": "2026-09-21",
        "confidence": "低",
        "blockers": ["低置信度不得升级为研究吸引力"], "status": "conditional_research_only",
        "assumptions": {"reverse_valuation": {"quote_date": "2026-09-21", "quote_price_cny": "550"}},
    }
    bridge = {"bridge_status": "READY", "current_price": "550", "margin_to_bear": "-0.375", "margin_to_base": "-0.1"}
    _mvp_research_card(sheet, 1, record, valuation, bridge)
    text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "价格桥接：550.00 元/股；相对熊/基准情景：-37.5% / -10.0%。" in text
    assert "条件研究（非交易）" in text
    assert "financial_scope_approved" not in text
    assert "财务范围已核验" in text
    assert "买入" not in text
    assert "订单" in text


def test_mvp_card_displays_unified_unready_valuation_without_values():
    workbook = Workbook()
    sheet = workbook.active
    record = {
        "case": {"name": "美的集团", "symbol": "000333", "industry": "家电",
                 "investment_path": "成长价值", "research_status": "financial_scope_blocked",
                 "thesis": "测试论点", "return_driver": "测试回报", "mispricing_hypothesis": "未证明",
                 "financial_summary": {}, "financial_period": "2025-09-30", "positives": [],
                 "counter_evidence": [], "thesis_breakers": [], "next_events": [],
                 "evidence_status": "partial", "evidence_refs": []},
        "gate": {"conclusion": "估值未就绪", "blockers": ["G3_估值模型通过"]},
    }
    valuation = {"symbol": "000333", "model_type": "FCFF", "valuation_date": "2025-09-30",
                 "bear_value": None, "base_value": None, "bull_value": None, "confidence": "低",
                 "blockers": ["ordinary_shares_not_verified"], "status": "not_ready"}

    _mvp_research_card(sheet, 1, record, valuation)

    text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "模型：FCFF；估值未就绪。" in text
    assert "ordinary_shares_not_verified" in text
    assert "买入" not in text


def test_mvp_card_displays_aggregated_current_status_without_trade_state():
    workbook = Workbook()
    sheet = workbook.active
    record = {
        "case": {
            "name": "贵州茅台", "symbol": "600519", "industry": "高端白酒",
            "investment_path": "成熟优质复利", "research_status": "financial_scope_approved",
            "thesis": "测试论点", "return_driver": "测试回报", "mispricing_hypothesis": "未证明",
            "financial_summary": {"period_end": "2026-06-30"}, "financial_period": "2026-06-30",
            "positives": [], "counter_evidence": [], "thesis_breakers": [], "next_events": [],
            "evidence_status": "verified", "evidence_refs": [],
        },
        "gate": {
            "conclusion": "估值未就绪",
            "blockers": ["G3_估值模型通过"],
            "results": {
                "G0_证据门": True,
                "G1_财务门": True,
                "G2_商业论点门": True,
                "G3_估值门": False,
            },
        },
    }
    valuation = {
        "symbol": "600519", "model_type": "归母权益剩余收益", "valuation_date": "2026-09-21",
        "bear_value": "400", "base_value": "500", "bull_value": "600", "confidence": "低",
        "assumptions": {
            "reverse_valuation": {"quote_date": "2026-09-21", "quote_price_cny": "550"}
        },
        "sensitivities": [], "evidence_refs": [{"id": "valuation"}], "blockers": [],
        "status": "conditional_research_only", "model_version": "fixture-v1",
    }
    bridge = {
        "symbol": "600519", "valuation_date": "2026-09-21", "quote_date": None,
        "current_price": None, "margin_to_bear": None, "margin_to_base": None,
        "model_validity_status": "VALID", "quote_status": "PENDING_EXTERNAL_DATA",
        "bridge_status": "PENDING_EXTERNAL_DATA",
        "evidence_refs": [{"id": "quote"}], "blockers": ["等待已验证收盘行情"],
    }
    current_status = current_research_status_from_payloads(
        record["gate"], valuation, bridge
    )
    _mvp_research_card(sheet, 1, record, valuation, bridge, current_status)
    text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)

    assert "统一当前状态" in text
    assert "工程链路：READY" in text
    assert "当前数据：PENDING_EXTERNAL_DATA" in text
    assert "估值结果已保留" in text
    assert "买入" not in text


def test_unready_fcff_card_uses_pending_status_without_price_or_orders():
    workbook = Workbook()
    sheet = workbook.active
    record = {
        "case": {
            "name": "美的集团", "symbol": "000333", "industry": "家电",
            "investment_path": "成长价值", "research_status": "financial_scope_blocked",
            "thesis": "测试论点", "return_driver": "测试回报", "mispricing_hypothesis": "未证明",
            "financial_summary": {}, "financial_period": "2025-12-31", "positives": [],
            "counter_evidence": [], "thesis_breakers": [], "next_events": [],
            "evidence_status": "partial", "evidence_refs": [],
        },
        "gate": {
            "conclusion": "数据不足",
            "blockers": ["G0_证据门", "G1_财务门", "G3_估值门"],
            "results": {
                "G0_证据门": False,
                "G1_财务门": False,
                "G2_商业论点门": True,
                "G3_估值门": False,
            },
        },
    }
    valuation = {
        "symbol": "000333", "model_type": "FCFF", "valuation_date": "2025-12-31",
        "bear_value": None, "base_value": None, "bull_value": None, "confidence": "低",
        "assumptions": {"scope": "input gate"}, "sensitivities": [],
        "evidence_refs": [{"id": "facts"}], "blockers": ["financial_facts_not_verified"],
        "status": "not_ready", "model_version": "fcff-input-gate-v1",
    }
    bridge = {
        "symbol": "000333", "valuation_date": "2025-12-31", "quote_date": None,
        "current_price": None, "margin_to_bear": None, "margin_to_base": None,
        "model_validity_status": "UNKNOWN", "quote_status": "PENDING_EXTERNAL_DATA",
        "bridge_status": "PENDING_EXTERNAL_DATA",
        "evidence_refs": [{"id": "facts"}],
        "blockers": ["正式估值未形成，价格桥接不启用"],
    }
    current_status = current_research_status_from_payloads(
        record["gate"], valuation, bridge
    )
    _mvp_research_card(sheet, 1, record, valuation, bridge, current_status)
    text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)

    assert current_status.research_conclusion == "数据不足"
    assert current_status.price_bridge_status == "PENDING_EXTERNAL_DATA"
    assert "正式估值未形成，价格桥接不启用" in current_status.current_data_status.blockers
    assert "统一当前状态" in text
    assert "已有估值对象已保留" in text
    assert "当前价格桥接" not in text
    assert "买入" not in text
