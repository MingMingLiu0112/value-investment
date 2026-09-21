import importlib.util
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "verify_decision_status_staged", ROOT / "scripts" / "verify_decision_status_staged.py")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def fixture(path: Path, action: bool) -> None:
    workbook = Workbook()
    research = workbook.active
    research.title = "09_公司研究"
    research.append(["placeholder"] * 16)
    research.append(["placeholder"] * 16)
    research.append(["代码", "公司"] + ["字段"] * 13 + ["研究状态"])
    research.append(["600519", "贵州茅台"] + [""] * 13 + ["原状态"])
    decisions = workbook.create_sheet("21_决策验证")
    decisions.append(["placeholder"] * 4)
    decisions.append(["placeholder"] * 4)
    decisions.append(["代码", "公司", "公式", "策略验证状态"])
    decisions.append(["600519", "贵州茅台", "=1+1", "原验证状态"])
    if action:
        research["P4"] = MODULE.ACTION_PREFIX + "\n" + MODULE.REASON + "\n原状态"
        decisions["D4"] = MODULE.ACTION_PREFIX + "\n" + MODULE.REASON + "\n原验证状态"
    workbook.save(path)


def test_verifier_accepts_only_the_two_status_cells(tmp_path):
    source, staged, receipt = tmp_path / "source.xlsx", tmp_path / "staged.xlsx", tmp_path / "receipt.json"
    fixture(source, action=False)
    fixture(staged, action=True)
    receipt_data = MODULE.verify(source, staged)
    assert {tuple(item.values()) for item in receipt_data["changed_cells"]} == {
        ("09_公司研究", "P4"),
        ("21_决策验证", "D4"),
    }


def test_verifier_rejects_an_unrelated_change(tmp_path):
    source, staged = tmp_path / "source.xlsx", tmp_path / "staged.xlsx"
    fixture(source, action=False)
    fixture(staged, action=True)
    workbook = MODULE.load_workbook(staged)
    workbook["21_决策验证"]["C4"] = "=2+2"
    workbook.save(staged)
    try:
        MODULE.verify(source, staged)
    except ValueError as error:
        assert "Unexpected workbook change" in str(error)
    else:
        raise AssertionError("Verifier accepted an unrelated formula change")
