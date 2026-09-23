from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import importlib.util
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from value_investment_agent.m2_discovery_engine import M2ScreeningPolicy, build_discovery_receipt
from value_investment_agent.m2_opportunity_discovery import (
    CANDIDATE_CLASS_LEAD,
    CHANNEL_DIVIDEND,
    CHANNEL_VALUE,
    DATA_PARTIAL,
    EvidenceReference,
)
from value_investment_agent.m2_research_report import (
    M2ResearchEvidence,
    M2ResearchReport,
    M2ResearchReportBatch,
    VERDICT_PENDING,
    VERDICT_REJECTED,
)


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "build_m2_original_workbook_candidate",
    ROOT / "scripts" / "build_m2_original_workbook_candidate.py",
)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def _sha(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _reference(ref_id: str) -> EvidenceReference:
    return EvidenceReference(
        id=ref_id,
        path=f"runtime/{ref_id}.json",
        sha256=_sha(ref_id),
        source_name=ref_id,
        source_url=f"https://example.test/{ref_id}",
        fetched_at=datetime(2026, 9, 23, 8, 0, tzinfo=timezone.utc),
    )


def _receipt_payload() -> dict:
    official = {
        "parser_version": "test",
        "complete": True,
        "scope": "沪深A股测试",
        "fetched_at": "2026-09-23T08:00:00+00:00",
        "records": [
            {"symbol": "600001", "name": "质量制造", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "制造业", "security_type": "A股"},
            {"symbol": "600002", "name": "测试银行", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "金融业", "security_type": "A股"},
            {"symbol": "600003", "name": "测试煤业", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "采矿业", "security_type": "A股"},
            {"symbol": "600004", "name": "测试高息", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "公用事业", "security_type": "A股"},
        ],
    }
    tencent = {
        "rows": [
            {"code": "sh600001", "name": "质量制造", "zxj": "20", "pe_ttm": "10", "pn": "2", "zsz": "200", "stock_type": "GP-A", "state": "", "zdf_y": "5"},
            {"code": "sh600002", "name": "测试银行", "zxj": "10", "pe_ttm": "6", "pn": "0.8", "zsz": "500", "stock_type": "GP-A", "state": "", "zdf_y": "1"},
            {"code": "sh600003", "name": "测试煤业", "zxj": "8", "pe_ttm": "8", "pn": "1", "zsz": "300", "stock_type": "GP-A", "state": "", "zdf_y": "-20"},
            {"code": "sh600004", "name": "测试高息", "zxj": "10", "pe_ttm": "12", "pn": "1.5", "zsz": "100", "stock_type": "GP-A", "state": "", "zdf_y": "3"},
        ],
        "total": 4,
    }
    sina = {
        "rows": [
            {"code": "600001", "name": "质量制造", "class": "机械设备", "trade": "20.01", "per": "10", "pb": "2", "mktcap": "2000000", "ticktime": "15:00:00"},
            {"code": "600002", "name": "测试银行", "class": "银行", "trade": "10.00", "per": "6", "pb": "0.8", "mktcap": "5000000", "ticktime": "15:00:00"},
            {"code": "600003", "name": "测试煤业", "class": "煤炭", "trade": "8.00", "per": "8", "pb": "1", "mktcap": "3000000", "ticktime": "15:00:00"},
            {"code": "600004", "name": "测试高息", "class": "公用事业", "trade": "10.00", "per": "12", "pb": "1.5", "mktcap": "1000000", "ticktime": "15:00:00"},
        ]
    }
    dividend_rows = []
    for code, name, amount, yield_value in (("600004", "测试高息", "1.00", "0.05"), ("600002", "测试银行", "0.50", "0.04")):
        row = [""] * 18
        row[0] = code
        row[1] = name
        row[5] = amount
        row[6] = yield_value
        row[13] = "2026-04-01"
        row[15] = "2026-07-01"
        row[16] = "实施中"
        dividend_rows.append(row)
    financial_points = [
        {
            "symbol": "600001",
            "field_name": field_name,
            "value": value,
            "unit": unit,
            "period_label": "2025-12-31",
            "validation_status": "verified",
            "human_reviewed": False,
            "metadata": {"automatic_cross_source_verification": True, "complete_debt_verified": True},
            "created_at": "2026-09-20T07:00:00+00:00",
            "fetched_at": "2026-09-20T07:00:00+00:00",
            "source_id": "financial-fixture",
        }
        for field_name, (value, unit) in {
            "roe": ("20", "percent"),
            "gross_margin": ("40", "percent"),
            "net_margin": ("15", "percent"),
            "operating_cash_flow_to_net_income": ("120", "percent"),
            "cash": ("100", "CNY"),
            "interest_bearing_debt": ("50", "CNY"),
            "debt_ratio": ("30", "percent"),
            "revenue_yoy": ("10", "percent"),
            "net_income_yoy": ("10", "percent"),
        }.items()
    ]
    refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    return build_discovery_receipt(
        run_id="m2-original-publisher-fixture",
        generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
        official_payload=official,
        tencent_payload=tencent,
        sina_payload=sina,
        dividend_payload={"fiscal_year": "2025", "rows": dividend_rows},
        financial_points=financial_points,
        quote_date="2026-09-23",
        run_refs=refs,
        policy=M2ScreeningPolicy(max_per_channel=10),
    ).as_policy()


def _source_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "保留页"
    sheet["A1"] = "用户手工内容"
    sheet["B1"] = "=1+2"
    workbook.save(path)


def _write_receipt(path: Path) -> None:
    path.write_text(json.dumps(_receipt_payload(), ensure_ascii=False), encoding="utf-8")


def _research_batch() -> M2ResearchReportBatch:
    evidence = M2ResearchEvidence(
        symbol="600001",
        field_name="free_cash_flow",
        period_label="2025-12-31",
        value="100",
        unit="CNY 100M",
        validation_status="verified",
        source_name="测试公告",
        source_url="https://example.test/disclosure.pdf",
        source_sha256="0" * 64,
        published_at="2026-04-01T00:00:00+00:00",
        fetched_at="2026-09-20T07:00:00+00:00",
    )
    def report(symbol, name, channel, verdict, evidence=()):
        return M2ResearchReport(
            report_id=f"m2-ac8-{symbol}-fixture",
            symbol=symbol,
            name=name,
            channels=(channel,),
            primary_channel=channel,
            verdict=verdict,
            channel_verdicts={channel: verdict},
            candidate_class=CANDIDATE_CLASS_LEAD,
            data_status=DATA_PARTIAL,
            title=f"{name}：M2 实质研究/否决",
            conclusion="固定输入可追溯，但尚需补充通道深研证据。",
            positives=("核心财务点可追溯",),
            counter_evidence=("缺失周期或派息长期证据",),
            missing_evidence=("补充研究证据",),
            next_events=("年度报告",),
            market_context={"pe_ttm": "10", "pb": "2"},
            evidence=evidence,
        )

    reports = (
        report("600001", "质量制造", CHANNEL_DIVIDEND, VERDICT_PENDING, (evidence,)),
        report("600002", "测试银行", CHANNEL_VALUE, VERDICT_REJECTED),
        report("600003", "测试煤业", CHANNEL_VALUE, VERDICT_REJECTED),
    )
    return M2ResearchReportBatch(
        schema_version="m2-ac8-research-report-v1",
        policy_version="fixture-v1",
        action="no_order",
        as_of=datetime(2026, 9, 23, tzinfo=timezone.utc).date(),
        minimum_substantive_reports=3,
        minimum_channels=2,
        source_binding={"m2_receipt_sha256": "0" * 64},
        reports=reports,
        machine_status="MACHINE_CHECKS_PASS",
        acceptance_status="AC8_REVIEW_PENDING",
    )


def test_candidate_prepends_m2_sheets_without_rewriting_original_parts(tmp_path):
    source = tmp_path / "canonical.xlsx"
    receipt = tmp_path / "receipt.json"
    output = tmp_path / "candidate.xlsx"
    _source_workbook(source)
    _write_receipt(receipt)
    expected = MODULE.digest(source)

    manifest = MODULE.build_candidate(
        source=source,
        receipt_path=receipt,
        output=output,
        expected_source_sha256=expected,
        project_root=tmp_path,
    )

    loaded = load_workbook(output, read_only=False)
    assert loaded.sheetnames[0] == "00_M2总览"
    assert loaded.sheetnames[-1] == "保留页"
    assert "10_逐通道覆盖" in loaded.sheetnames
    assert loaded["保留页"]["A1"].value == "用户手工内容"
    assert loaded["保留页"]["B1"].value == "=1+2"
    assert manifest["status"] == "candidate_verified_not_published"
    assert manifest["original_sheets_preserved"] == 1
    assert manifest["candidate_sha256"] == MODULE.digest(output)

    with ZipFile(source) as original, ZipFile(output) as candidate:
        original_sheet = original.read("xl/worksheets/sheet1.xml")
        retained_sheet = candidate.read("xl/worksheets/sheet1.xml")
        assert retained_sheet == original_sheet
    MODULE.STAGE_FRONTEND["validate_package_relationships"](output)


def test_candidate_includes_research_reports_and_evidence_links(tmp_path):
    source = tmp_path / "canonical.xlsx"
    receipt = tmp_path / "receipt.json"
    output = tmp_path / "candidate.xlsx"
    _source_workbook(source)
    _write_receipt(receipt)
    expected = MODULE.digest(source)
    batch = _research_batch()

    manifest = MODULE.build_candidate(
        source=source,
        receipt_path=receipt,
        output=output,
        expected_source_sha256=expected,
        research_batch=batch,
        project_root=tmp_path,
    )

    loaded = load_workbook(output, read_only=False)
    assert "11_研究报告" in loaded.sheetnames
    assert "12_研究证据" in loaded.sheetnames
    reports = loaded["11_研究报告"]
    assert reports["A1"].value == "AC8 实质研究与通道否决"
    assert reports["F5"].value == "待深研"
    evidence = loaded["12_研究证据"]
    assert evidence["H5"].value == "https://example.test/disclosure.pdf"
    assert evidence["H5"].hyperlink.target == "https://example.test/disclosure.pdf"
    overview = loaded["00_M2总览"]
    assert any(
        value == "研究/否决报告"
        for row in overview.iter_rows(values_only=True)
        for value in row
    )
    assert manifest["research_report_sha256"] is not None
    assert manifest["research_summary"]["report_count"] == 3
    assert manifest["research_machine_status"] == "MACHINE_CHECKS_PASS"

    forbidden = {"buy", "sell", "target_weight", "position_size", "order_quantity"}
    for worksheet in loaded.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            assert not forbidden & {str(value).lower() for value in row if value is not None}


def test_candidate_refuses_source_change_after_inspection(tmp_path):
    source = tmp_path / "canonical.xlsx"
    receipt = tmp_path / "receipt.json"
    output = tmp_path / "candidate.xlsx"
    _source_workbook(source)
    _write_receipt(receipt)
    expected = MODULE.digest(source)
    source.write_bytes(source.read_bytes() + b"changed")

    with pytest.raises(ValueError, match="changed since inspection"):
        MODULE.build_candidate(
            source=source,
            receipt_path=receipt,
            output=output,
            expected_source_sha256=expected,
        )


def test_candidate_refuses_existing_output(tmp_path):
    source = tmp_path / "canonical.xlsx"
    receipt = tmp_path / "receipt.json"
    output = tmp_path / "candidate.xlsx"
    _source_workbook(source)
    _write_receipt(receipt)
    output.write_text("occupied", encoding="utf-8")

    with pytest.raises(ValueError, match="already exists"):
        MODULE.build_candidate(
            source=source,
            receipt_path=receipt,
            output=output,
            expected_source_sha256=MODULE.digest(source),
        )
