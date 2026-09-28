from __future__ import annotations

from openpyxl import Workbook, load_workbook

from scripts.current.publish_product_workbench_to_canonical import (
    _assert_retained,
    _hide_legacy_sheets,
    _snapshot,
)
from value_investment_agent.presentation.excel.product_workbench import WORKBOOK_SHEETS


def test_legacy_tabs_are_hidden_without_changing_their_content(tmp_path) -> None:
    workbook = Workbook()
    workbook.active.title = WORKBOOK_SHEETS[0]
    for name in WORKBOOK_SHEETS[1:]:
        workbook.create_sheet(name)
    legacy = workbook.create_sheet("00_组合与股息")
    legacy["A21"] = "历史组合说明"
    legacy["D22"] = "历史 IPS / 持仓 / 现金字段"
    legacy["E23"] = "=1+1"
    legacy["F24"].hyperlink = "https://example.invalid/source"
    legacy["F24"] = "历史证据链接"
    legacy.protection.sheet = True

    original_path = tmp_path / "original.xlsx"
    published_path = tmp_path / "published.xlsx"
    workbook.save(original_path)
    workbook.close()

    before = _snapshot(original_path)
    candidate = load_workbook(original_path, data_only=False, keep_links=True)
    hidden = _hide_legacy_sheets(candidate)
    candidate.save(published_path)
    candidate.close()
    after = _snapshot(published_path)

    _assert_retained(before, after)
    assert hidden == ["00_组合与股息"]
    published = load_workbook(published_path, data_only=False, keep_links=True)
    try:
        assert published.active.title == WORKBOOK_SHEETS[0]
        assert [sheet.title for sheet in published if sheet.sheet_state == "visible"] == list(WORKBOOK_SHEETS)
        legacy = published["00_组合与股息"]
        assert legacy.sheet_state == "hidden"
        assert legacy["A21"].value == "历史组合说明"
        assert legacy["D22"].value == "历史 IPS / 持仓 / 现金字段"
        assert legacy["E23"].value == "=1+1"
        assert legacy["F24"].hyperlink.target == "https://example.invalid/source"
        assert legacy.protection.sheet is True
    finally:
        published.close()
