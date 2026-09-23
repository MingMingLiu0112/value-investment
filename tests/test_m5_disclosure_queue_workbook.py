from __future__ import annotations

from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from value_investment_agent.m5_disclosure_queue import (
    DISCLOSURE_QUEUE_SCHEMA,
    DisclosureReviewQueue,
    build_cninfo_event_scan,
)
from value_investment_agent.m5_disclosure_queue_workbook import (
    build_m5_disclosure_queue_workbook,
    write_m5_disclosure_queue_workbook,
)


RETRIEVED_AT = datetime(2026, 9, 24, 16, 0, tzinfo=timezone.utc)
SCAN_FROM = date(2026, 8, 27)
SCAN_TO = date(2026, 9, 24)


def _payload(rows: list[dict]) -> dict:
    return {
        "url": "https://www.cninfo.com.cn/new/hisAnnouncement/query",
        "total_announcements": len(rows),
        "announcements": rows,
    }


def _row(announcement_id: str, title: str) -> dict:
    return {
        "secCode": "600887",
        "announcementId": announcement_id,
        "announcementTitle": title,
        "announcementTime": int(
            datetime(2026, 9, 2, 1, 0, tzinfo=timezone.utc).timestamp() * 1000
        ),
        "adjunctUrl": f"finalpage/2026-09-02/{announcement_id}.PDF",
    }


def _downloader(source_url: str, target: Path) -> str:
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(b"%PDF-fixture")
    return "a" * 64


def _queue(tmp_path: Path) -> DisclosureReviewQueue:
    scan = build_cninfo_event_scan(
        _payload(
            [
                _row("1225542001", "关于回购公司股份的公告"),
                _row("1225542002", "董事会决议公告"),
            ]
        ),
        symbol="600887",
        scan_from=SCAN_FROM,
        scan_to=SCAN_TO,
        retrieved_at=RETRIEVED_AT,
        evidence_dir=tmp_path / "600887",
        root=tmp_path,
        pdf_downloader=_downloader,
    )
    return DisclosureReviewQueue(
        queue_id="fixture-disclosure-queue",
        schema_version=DISCLOSURE_QUEUE_SCHEMA,
        provider="cninfo",
        parser_version="m5-cninfo-announcement-title-v1",
        scan_from=SCAN_FROM,
        scan_to=SCAN_TO,
        retrieved_at=RETRIEVED_AT,
        scans=(scan,),
    )


def test_disclosure_queue_workbook_is_presentation_only(tmp_path):
    queue = _queue(tmp_path)
    workbook = build_m5_disclosure_queue_workbook(
        queue,
        security_names={"600887": "伊利股份"},
    )

    assert workbook.sheetnames == [
        "00_总览",
        "01_待复核公告",
        "02_来源覆盖",
        "03_输入与边界",
    ]
    assert workbook["00_总览"]["B10"].value == 1
    assert workbook["00_总览"]["B11"].value == 0
    assert workbook["00_总览"]["B14"].value == "no_order"
    assert workbook["01_待复核公告"].max_row == 5
    assert workbook["02_来源覆盖"].max_row == 5
    assert workbook["03_输入与边界"].max_row == 13

    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "自动材料性" not in text
    assert "自动交易" not in text
    assert "下单" not in text
    assert "目标仓位" not in text


def test_disclosure_queue_workbook_writer_pins_result_and_rejects_overwrite(tmp_path):
    queue = _queue(tmp_path / "queue")
    output = tmp_path / "disclosure-queue.xlsx"
    result = write_m5_disclosure_queue_workbook(
        queue,
        output=output,
        root=tmp_path,
        security_names={"600887": "伊利股份"},
    )
    loaded = load_workbook(output)

    assert result["sheet_count"] == 4
    assert result["company_count"] == 1
    assert result["pending_candidate_count"] == 1
    assert result["source_unavailable_count"] == 0
    assert result["coverage_complete_count"] == 1
    assert result["action"] == "no_order"
    assert loaded.sheetnames[0] == "00_总览"

    with pytest.raises(ValueError, match="already exists"):
        write_m5_disclosure_queue_workbook(
            queue,
            output=output,
            root=tmp_path,
            security_names={"600887": "伊利股份"},
        )
