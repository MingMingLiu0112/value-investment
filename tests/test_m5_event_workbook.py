from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.build_m5_event_infrastructure_candidate import (
    build_models,
    load_input,
)
from value_investment_agent.m5_event_workbook import (
    build_m5_event_workbook,
    write_m5_event_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "m5_event_infrastructure_demo.json"


def test_simulated_m5_workbook_has_expected_sheets_and_boundary(tmp_path):
    payload, _ = load_input(FIXTURE)
    receipt, watermark = build_models(payload)
    workbook = build_m5_event_workbook(
        receipt,
        watermark,
        security_names=payload["security_names"],
    )

    assert workbook.sheetnames == [
        "00_总览",
        "01_事件账",
        "02_水位与检查点",
        "03_依赖失效与重算",
        "04_Outbox",
        "05_输入与边界",
    ]
    overview = workbook["00_总览"]
    assert overview["B5"].value == "SIMULATED"
    assert overview["B6"].value == "需关注"
    assert overview["B7"].value == 7
    assert overview["B9"].value == 6
    assert overview["B17"].value == "no_order"
    assert workbook["01_事件账"].max_row == 11
    assert workbook["04_Outbox"].max_row == 10

    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "自动交易" not in text
    assert "下单" not in text
    assert "目标仓位" not in text


def test_writer_pins_hash_and_rejects_existing_output(tmp_path):
    payload, _ = load_input(FIXTURE)
    receipt, watermark = build_models(payload)
    output = tmp_path / "m5-candidate.xlsx"
    result = write_m5_event_workbook(
        receipt,
        watermark,
        output=output,
        root=tmp_path,
        security_names=payload["security_names"],
    )
    loaded = load_workbook(output)

    assert result["sheet_count"] == 6
    assert result["event_count"] == 7
    assert result["active_event_count"] == 6
    assert result["invalidation_count"] == 6
    assert result["alert_count"] == 6
    assert result["action"] == "no_order"
    assert loaded.sheetnames[0] == "00_总览"
    assert result["workbook_sha256"] == result["workbook_sha256"].lower()

    with pytest.raises(ValueError, match="already exists"):
        write_m5_event_workbook(
            receipt,
            watermark,
            output=tmp_path / "m5-candidate.xlsx",
            root=tmp_path,
            security_names=payload["security_names"],
        )
