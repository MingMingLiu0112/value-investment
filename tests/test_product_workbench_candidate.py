from __future__ import annotations

import copy
from datetime import datetime, timezone
from pathlib import Path
import re

from openpyxl import load_workbook
import pytest

from value_investment_agent.application.product.product_workbench_candidate import (
    ACTION_NO_ORDER,
    build_product_workbench_candidate_payload,
    validate_product_workbench_candidate_payload,
)
from value_investment_agent.presentation.excel.product_workbench import (
    SHEET_PORTFOLIO,
    SHEET_SYSTEM_AUDIT,
    USER_SHEETS,
    WORKBOOK_SHEETS,
    build_product_workbench_workbook,
)
from value_investment_agent.presentation.read_models.product_workbench import (
    product_workbench_from_payload,
)

from product_workbench_candidate_fixture import (
    materialize_synthetic_legacy_packet,
    synthetic_legacy_packet,
)


GENERATED_AT = datetime(2026, 9, 26, 3, 0, tzinfo=timezone.utc)


def _packet() -> dict:
    return synthetic_legacy_packet(GENERATED_AT.isoformat())


def _materialized_packet(root: Path) -> dict:
    packet = _packet()
    materialize_synthetic_legacy_packet(root, packet)
    return packet


def _user_text(workbook) -> str:
    return "\n".join(
        str(cell.value)
        for title in USER_SHEETS
        for row in workbook[title].iter_rows()
        for cell in row
        if cell.value is not None
    )


def test_synthetic_packet_builds_five_page_candidate_without_portfolio_metrics(
    tmp_path: Path,
) -> None:
    packet = _materialized_packet(tmp_path)
    payload = build_product_workbench_candidate_payload(packet, root=tmp_path)
    model = product_workbench_from_payload(payload)
    workbook = build_product_workbench_workbook(model)

    assert model.action == ACTION_NO_ORDER
    assert tuple(workbook.sheetnames) == WORKBOOK_SHEETS
    assert tuple(workbook.sheetnames[:5]) == USER_SHEETS
    assert workbook.sheetnames[-1] == SHEET_SYSTEM_AUDIT
    assert len(model.today_items) == 6
    assert len(model.opportunities) == 18
    assert len(model.companies) == 3
    assert len(model.events) == 1
    assert len(model.audit_evidence) == 13
    assert model.portfolio.real_data_available is False
    assert model.portfolio.summary == ()
    assert model.portfolio.positions == ()
    assert model.opportunities[0].research_status.user_label == "等待关键经营证据"
    assert model.companies[0].research_status.user_label == "等待更多证据"

    user_text = _user_text(workbook)
    assert not re.search(r"\bM[2-6]\b", user_text)
    for hidden in (
        "INSUFFICIENT_EVIDENCE",
        "INSUFFICIENT_RESEARCH",
        "ENGINEERING_DONE_OFFLINE",
        "PREFLIGHT_DONE",
        "PENDING_USER_PRIVATE_INPUT",
        "target_weight",
        "position_size",
        "trade_approved",
        "建议买入",
        "目标仓位",
        "下单",
    ):
        assert hidden not in user_text

    portfolio_text = "\n".join(
        str(cell.value)
        for row in workbook[SHEET_PORTFOLIO].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "尚未接入真实组合" in portfolio_text
    for simulated_metric in (
        "总资产",
        "股票仓位",
        "预计年股息",
        "正常化年股息",
        "单股集中度",
    ):
        assert simulated_metric not in portfolio_text
    assert not any(
        isinstance(cell.value, (int, float))
        for row in workbook[SHEET_PORTFOLIO].iter_rows()
        for cell in row
    )

    evidence_links = [
        cell
        for title in USER_SHEETS
        for row in workbook[title].iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("查看证据 (")
    ]
    assert evidence_links
    assert all(cell.hyperlink is not None for cell in evidence_links)

    audit_text = "\n".join(
        str(cell.value)
        for row in workbook[SHEET_SYSTEM_AUDIT].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "SHA-256" in audit_text
    assert packet["audit"]["artifacts"][0]["sha256"] in audit_text
    assert packet["audit"]["artifacts"][0]["path"] in audit_text
    assert ACTION_NO_ORDER in audit_text


def test_synthetic_candidate_can_be_written_and_reopened(tmp_path: Path) -> None:
    packet = _materialized_packet(tmp_path)
    payload = build_product_workbench_candidate_payload(packet, root=tmp_path)
    model = product_workbench_from_payload(payload)
    workbook = build_product_workbench_workbook(model)
    output = tmp_path / "m7-product-ux-candidate-v1-20260926.xlsx"
    workbook.save(output)

    reopened = load_workbook(output, data_only=False)

    assert tuple(reopened.sheetnames) == WORKBOOK_SHEETS
    assert reopened[SHEET_PORTFOLIO]["A4"].value == "尚未接入真实组合"


def test_projection_fails_closed_on_missing_evidence() -> None:
    packet = _packet()
    packet["audit"]["artifacts"] = []

    with pytest.raises(ValueError, match="Pinned audit evidence"):
        build_product_workbench_candidate_payload(packet, root=Path.cwd())


def test_projection_fails_closed_on_missing_evidence_file(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Missing evidence file"):
        build_product_workbench_candidate_payload(_packet(), root=tmp_path)


def test_projection_requires_an_evidence_root() -> None:
    with pytest.raises(TypeError, match="root"):
        build_product_workbench_candidate_payload(_packet())

    with pytest.raises(ValueError, match="root must be a pathlib.Path"):
        build_product_workbench_candidate_payload(_packet(), root=None)  # type: ignore[arg-type]


def test_projection_fails_closed_on_unsupported_status(tmp_path: Path) -> None:
    packet = _materialized_packet(tmp_path)
    packet["m2"]["rows"][0]["status"] = "M2_VERIFIED"

    with pytest.raises(ValueError, match="Unsupported m2 row status"):
        build_product_workbench_candidate_payload(packet, root=tmp_path)


def test_projection_fails_closed_on_execution_keys(tmp_path: Path) -> None:
    packet = _materialized_packet(tmp_path)
    packet["m2"]["rows"][0]["target_weight"] = 0.1

    with pytest.raises(ValueError, match="active execution key"):
        build_product_workbench_candidate_payload(packet, root=tmp_path)


def test_projection_fails_closed_on_simulated_portfolio_metrics(
    tmp_path: Path,
) -> None:
    packet = _materialized_packet(tmp_path)
    packet["m4"]["normalized_annual_dividend"] = "0.00"

    with pytest.raises(ValueError, match="Simulated portfolio metric"):
        build_product_workbench_candidate_payload(packet, root=tmp_path)


def test_payload_validator_rejects_fake_numeric_placeholders(tmp_path: Path) -> None:
    packet = _materialized_packet(tmp_path)
    payload = build_product_workbench_candidate_payload(packet, root=tmp_path)
    malformed = copy.deepcopy(payload)
    malformed["companies"][0]["valuation"]["value_text"] = "0.00"

    with pytest.raises(ValueError, match="numeric placeholder"):
        validate_product_workbench_candidate_payload(malformed)


def test_payload_validator_rejects_internal_stage_tokens_in_user_copy(
    tmp_path: Path,
) -> None:
    packet = _materialized_packet(tmp_path)
    payload = build_product_workbench_candidate_payload(packet, root=tmp_path)
    malformed = copy.deepcopy(payload)
    malformed["companies"][0]["valuation"]["unavailable_reason"] = (
        "M3 尚未完成。"
    )

    with pytest.raises(ValueError, match="internal stage token"):
        validate_product_workbench_candidate_payload(malformed)
