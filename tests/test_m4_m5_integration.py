from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
import json

import pytest
from openpyxl import load_workbook

from scripts.build_m4_m5_integrated_candidate import build_models, load_input
from value_investment_agent.m4_m5_integration import (
    M4M5ArtifactDefinition,
    STATUS_NEGATIVE,
    STATUS_PAUSED,
    STATUS_PARTIAL,
    STATUS_READY,
    build_m4_m5_integration,
    integration_artifact_definition_from_payload,
)
from value_investment_agent.m4_m5_integration_workbook import (
    build_m4_m5_integration_workbook,
    write_m4_m5_integration_workbook,
)
from value_investment_agent.m5_event_core import (
    CONFIDENCE_HIGH,
    EVENT_TYPE_NEW_FINANCIAL_REPORT,
    NAMESPACE_SIMULATED,
    SEVERITY_HIGH,
    ChangeEventInput,
)
from value_investment_agent.m5_event_dependencies import (
    KIND_FACTS,
    DependencyGraph,
    DependencyNode,
)
from value_investment_agent.m5_event_run import run_event_batch
from value_investment_agent.m5_event_watermark import (
    SOURCE_HEALTHY,
    WATERMARK_COVERAGE_COMPLETE,
    ScanWatermark,
)


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "m4_m5_integrated_demo.json"


@pytest.fixture
def models():
    payload, input_receipts = load_input(FIXTURE)
    result, risk, guidance, dividend, model_receipts = build_models(
        payload,
        base_dir=ROOT,
    )
    return payload, result, risk, guidance, dividend


def test_simulated_joint_contract_projects_precise_invalidations(models):
    payload, result, risk, guidance, dividend = models

    assert result.namespace == "SIMULATED"
    assert result.action == "no_order"
    assert len(result.receipt.active_events) == 4
    assert len(result.receipt.invalidations) == 4
    assert risk.status == "VIOLATION"
    assert guidance.status == "BUDGET_CONFLICT"
    assert dividend.status == "READY"

    portfolio_risk = result.artifact("portfolio_risk")
    guidance_state = result.artifact("position_guidance")
    distribution = result.artifact("distribution_history_000333")
    sustainability = result.artifact("dividend_sustainability_000333")
    finance_facts = result.artifact("financial_facts_601088")
    finance_inputs = result.artifact("valuation_inputs_601088")
    finance_validity = result.artifact("model_validity_601088")
    untouched_facts = result.artifact("financial_facts_600519")
    thesis = result.artifact("thesis_600519")
    decision = result.artifact("decision_review_600519")
    entry = result.artifact("entry_consistency_600519")

    assert portfolio_risk.status == STATUS_PAUSED
    assert portfolio_risk.triggering_event_ids == ("600519-portfolio-risk-001",)
    assert guidance_state.status == STATUS_PAUSED
    assert guidance_state.triggering_event_ids == (
        "600519-portfolio-risk-001",
        "000333-dividend-change-001",
    )
    assert distribution.status == STATUS_PAUSED
    assert distribution.triggering_event_ids == ("000333-dividend-change-001",)
    assert sustainability.status == STATUS_PAUSED
    assert finance_facts.status == STATUS_PAUSED
    assert finance_facts.triggering_event_ids == ("601088-financial-report-001",)
    assert finance_inputs.status == STATUS_PAUSED
    assert finance_validity.status == STATUS_PAUSED
    assert untouched_facts.status == STATUS_READY
    assert thesis.status == STATUS_NEGATIVE
    assert decision.status == STATUS_NEGATIVE
    assert entry.status == STATUS_NEGATIVE
    assert result.affected_artifact_count() == 10

    policy_text = json.dumps(result.as_policy(), ensure_ascii=False)
    assert "target_weight" not in policy_text
    assert "position_size" not in policy_text
    assert "buy" not in policy_text


def test_joint_contract_fails_closed_on_bad_definition_and_baseline(models):
    payload, result, risk, guidance, dividend = models
    receipt = result.receipt
    graph = result.graph
    definitions = tuple(
        integration_artifact_definition_from_payload(item)
        for item in payload["artifacts"]
    )
    baselines = {
        item.artifact_id: item.status
        for item in result.artifacts
    }

    missing = (
        replace(
            definitions[0],
            node_ids=("missing-node",),
        ),
        *definitions[1:],
    )
    with pytest.raises(ValueError, match="missing-node"):
        build_m4_m5_integration(
            receipt=receipt,
            graph=graph,
            artifacts=missing,
            baseline_statuses=baselines,
            generated_at=result.generated_at,
        )

    bad_baseline = dict(baselines)
    bad_baseline[definitions[0].artifact_id] = "INVALID"
    with pytest.raises(ValueError, match="Invalid baseline"):
        build_m4_m5_integration(
            receipt=receipt,
            graph=graph,
            artifacts=definitions,
            baseline_statuses=bad_baseline,
            generated_at=result.generated_at,
        )


def test_simulated_joint_workbook_has_expected_sheets_and_boundary(models, tmp_path):
    payload, result, risk, guidance, dividend = models
    workbook = build_m4_m5_integration_workbook(
        result,
        risk,
        guidance,
        dividend,
        security_names=payload["security_names"],
    )

    assert workbook.sheetnames == [
        "00_总览",
        "01_M4组合基线",
        "02_M5事件批",
        "03_依赖失效映射",
        "04_联合产品状态",
        "05_输入与边界",
    ]
    overview = workbook["00_总览"]
    assert overview["B5"].value == "SIMULATED"
    assert overview["B6"].value == STATUS_NEGATIVE
    assert overview["B7"].value == "VIOLATION"
    assert overview["B8"].value == "BUDGET_CONFLICT"
    assert overview["B9"].value == "READY"
    assert overview["B10"].value == 4
    assert overview["B11"].value == 10
    assert overview["B13"].value == "no_order"
    assert workbook["04_联合产品状态"].max_row == 15

    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "目标仓位" not in text
    assert "下单" not in text
    assert "position_size" not in text


def test_joint_workbook_writer_pins_hash_and_rejects_existing_output(models, tmp_path):
    payload, result, risk, guidance, dividend = models
    output = tmp_path / "m4m5-joint.xlsx"
    write_result = write_m4_m5_integration_workbook(
        result,
        risk,
        guidance,
        dividend,
        output=output,
        root=tmp_path,
        security_names=payload["security_names"],
    )
    loaded = load_workbook(output)

    assert write_result["sheet_count"] == 6
    assert write_result["artifact_count"] == 11
    assert write_result["affected_artifact_count"] == 10
    assert write_result["event_count"] == 4
    assert write_result["invalidation_count"] == 4
    assert write_result["namespace"] == "SIMULATED"
    assert write_result["action"] == "no_order"
    assert loaded.sheetnames[0] == "00_总览"
    assert write_result["workbook_sha256"] == write_result["workbook_sha256"].lower()

    with pytest.raises(ValueError, match="already exists"):
        write_m4_m5_integration_workbook(
            result,
            risk,
            guidance,
            dividend,
            output=tmp_path / "m4m5-joint.xlsx",
            root=tmp_path,
            security_names=payload["security_names"],
        )
    assert "target_weight" not in json.dumps(write_result, ensure_ascii=False)


def test_deferred_only_artifact_is_partial_instead_of_invalid_paused() -> None:
    observed = datetime(2026, 9, 24, 16, 0, tzinfo=timezone(timedelta(hours=8)))
    graph = DependencyGraph(
        tuple(
            DependencyNode(
                node_id=f"n{index:02d}",
                kind=KIND_FACTS,
                symbol="600519",
                inputs=() if index == 0 else (f"n{index - 1:02d}",),
                version="deferred-boundary-v1",
                evidence_refs=({"id": f"node-{index:02d}"},),
            )
            for index in range(26)
        )
    )
    event = ChangeEventInput(
        source_id="cninfo",
        source_event_id="600519-deferred-only",
        symbol="600519",
        event_type=EVENT_TYPE_NEW_FINANCIAL_REPORT,
        detected_at=observed - timedelta(minutes=5),
        available_at=observed - timedelta(minutes=5),
        effective_at=observed - timedelta(minutes=5),
        previous_state={},
        current_state={"report": "fixture"},
        severity=SEVERITY_HIGH,
        reason="synthetic deferred-only dependency boundary",
        evidence_refs=({"id": "deferred-only-source"},),
        confidence=CONFIDENCE_HIGH,
        requires_human_review=True,
        namespace=NAMESPACE_SIMULATED,
    )
    receipt = run_event_batch(
        events=(event,),
        observed_times=(observed,),
        watermark=ScanWatermark(
            watermark_id="deferred-only-scan",
            scope="ALL",
            source="cninfo",
            coverage_through=observed,
            retrieved_at=observed,
            parser_version="deferred-boundary-v1",
            coverage_status=WATERMARK_COVERAGE_COMPLETE,
            source_health=SOURCE_HEALTHY,
            evidence_refs=({"id": "deferred-only-scan"},),
        ),
        graph=graph,
        run_id="deferred-only-run",
        generated_at=observed,
    )
    definition = M4M5ArtifactDefinition(
        artifact_id="deferred-only",
        kind=KIND_FACTS,
        symbol="600519",
        node_ids=("n25",),
    )

    result = build_m4_m5_integration(
        receipt=receipt,
        graph=graph,
        artifacts=(definition,),
        baseline_statuses={"deferred-only": STATUS_READY},
        generated_at=observed,
    )

    state = result.artifact("deferred-only")
    assert state.status == STATUS_PARTIAL
    assert state.invalidated_node_ids == ()
    assert state.deferred_node_ids == ("n25",)
    assert state.triggering_event_ids == ("600519-deferred-only",)
    assert result.affected_artifact_count() == 1


def test_integration_rejects_cross_symbol_forgery_and_binds_result_identity() -> None:
    observed = datetime(2026, 9, 24, 16, 0, tzinfo=timezone(timedelta(hours=8)))
    graph = DependencyGraph(
        (
            DependencyNode(
                node_id="facts-600519",
                kind=KIND_FACTS,
                symbol="600519",
                inputs=(),
                version="cross-symbol-v1",
                evidence_refs=({"id": "facts-600519"},),
            ),
            DependencyNode(
                node_id="facts-601088",
                kind=KIND_FACTS,
                symbol="601088",
                inputs=(),
                version="cross-symbol-v1",
                evidence_refs=({"id": "facts-601088"},),
            ),
        )
    )
    event = ChangeEventInput(
        source_id="cninfo",
        source_event_id="600519-cross-symbol",
        symbol="600519",
        event_type=EVENT_TYPE_NEW_FINANCIAL_REPORT,
        detected_at=observed - timedelta(minutes=5),
        available_at=observed - timedelta(minutes=5),
        effective_at=observed - timedelta(minutes=5),
        previous_state={},
        current_state={"report": "fixture"},
        severity=SEVERITY_HIGH,
        reason="synthetic cross-symbol invalidation fixture",
        evidence_refs=({"id": "cross-symbol-source"},),
        confidence=CONFIDENCE_HIGH,
        requires_human_review=True,
        namespace=NAMESPACE_SIMULATED,
    )
    receipt = run_event_batch(
        events=(event,),
        observed_times=(observed,),
        watermark=ScanWatermark(
            watermark_id="cross-symbol-scan",
            scope="ALL",
            source="cninfo",
            coverage_through=observed,
            retrieved_at=observed,
            parser_version="cross-symbol-v1",
            coverage_status=WATERMARK_COVERAGE_COMPLETE,
            source_health=SOURCE_HEALTHY,
            evidence_refs=({"id": "cross-symbol-scan"},),
        ),
        graph=graph,
        run_id="cross-symbol-run",
        generated_at=observed,
    )
    definition = M4M5ArtifactDefinition(
        artifact_id="foreign-facts",
        kind=KIND_FACTS,
        symbol="601088",
        node_ids=("facts-601088",),
    )
    forged_item = replace(
        receipt.invalidations[0].affected_nodes[0],
        node_id="facts-601088",
        kind=KIND_FACTS,
        reason="forged cross-symbol invalidation",
    )
    forged_invalidation = replace(
        receipt.invalidations[0],
        affected_nodes=(forged_item,),
    )
    forged_receipt = replace(
        receipt,
        invalidations=(forged_invalidation,),
    )

    with pytest.raises(ValueError, match="crosses event and graph symbols"):
        build_m4_m5_integration(
            receipt=forged_receipt,
            graph=graph,
            artifacts=(definition,),
            baseline_statuses={"foreign-facts": STATUS_READY},
            generated_at=observed,
        )

    local_definition = M4M5ArtifactDefinition(
        artifact_id="local-facts",
        kind=KIND_FACTS,
        symbol="600519",
        node_ids=("facts-600519",),
    )
    ready = build_m4_m5_integration(
        receipt=receipt,
        graph=graph,
        artifacts=(local_definition,),
        baseline_statuses={"local-facts": STATUS_READY},
        generated_at=observed,
    )
    partial = build_m4_m5_integration(
        receipt=receipt,
        graph=graph,
        artifacts=(local_definition,),
        baseline_statuses={"local-facts": STATUS_PARTIAL},
        generated_at=observed,
    )
    assert ready.result_id != partial.result_id
