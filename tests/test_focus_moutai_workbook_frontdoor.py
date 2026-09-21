import importlib.util
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "focus_moutai_workbook_frontdoor.py"
SPEC = importlib.util.spec_from_file_location("focus_moutai_workbook_frontdoor", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MODULE)


def test_focus_hides_only_non_moutai_data_rows(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    research = workbook.active
    research.title = "09_公司研究"
    decision = workbook.create_sheet("21_决策验证")
    queue = workbook.create_sheet("00_待完成公司")
    for sheet in (research, decision):
        sheet.append(["header"])
        sheet.append(["note"])
        sheet.append(["代码"])
        sheet.append(["000001", "other"])
        sheet.append(["600519", "贵州茅台"])
        sheet.append(["600900", "other two"])
    for _ in range(3):
        queue.append(["header"])
    queue.append(["000001", "other", None, None, None, None, "研究链接"])
    queue["G4"].hyperlink = "#'09_公司研究'!N6"
    workbook.save(source)

    result = MODULE.build(source, output)
    assert result["scope"] == "row_visibility_and_internal_navigation_only"
    assert result["sheets"]["09_公司研究"]["symbol_row"] == 5
    assert result["sheets"]["21_决策验证"]["hidden_rows"] == 2

    focused = load_workbook(output)
    for sheet_name in MODULE.SHEETS:
        sheet = focused[sheet_name]
        assert sheet["A4"].value == "000001"
        assert sheet["A5"].value == "600519"
        assert sheet["A6"].value == "600900"
        assert sheet.row_dimensions[4].hidden is True
        assert sheet.row_dimensions[5].hidden is False
        assert sheet.row_dimensions[6].hidden is True
        assert sheet.freeze_panes == "A4"
    assert focused["00_待完成公司"]["G4"].hyperlink.location == "'09_公司研究'!N5"
