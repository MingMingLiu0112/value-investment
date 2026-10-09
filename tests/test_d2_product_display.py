"""Synthetic presentation checks; no valuation or human approval is created."""
from copy import deepcopy
from dataclasses import asdict
import hashlib
import json
import os
from pathlib import Path
import subprocess

from openpyxl import load_workbook
import pytest

from test_product_workbench_excel import _payload
from value_investment_agent.presentation.excel.product_workbench import (
    SHEET_TODAY, USER_SHEETS, WORKBOOK_SHEETS, _company_main_reason, _user_text,
    build_product_workbench_workbook,
)
from value_investment_agent.presentation.read_models.product_workbench import (
    DECISION_STEP_TITLES, product_workbench_from_payload,
)


ROOT = Path(__file__).resolve().parents[1]


def _synthetic_model():
    payload = _payload()
    company = payload["companies"][0]
    item = payload["today_items"][0]
    payload["companies"] = []
    payload["today_items"] = []
    # Same four review shapes as the case, with explicitly synthetic explanations.
    for symbol, name, label, reason in (
        ("600519", "合成茅台", "为什么未进入更高状态", "合成：商业质量尚未批准。"),
        ("000333", "合成美的", "尚缺证据", "合成：维持性资本开支与普通股口径尚缺。"),
        ("601088", "合成神华", "尚缺证据", "合成：跨周期利润和成本勾稽尚缺。"),
        ("600887", "合成伊利", "最强反证", "合成：利润下行，不能机械外推历史ROE。"),
    ):
        card = deepcopy(company)
        card.update(symbol=symbol, company_name=name, next_trigger="合成：等待下一次披露。")
        card["decision_review"] = [
            {"label": label, "value": reason},
            {"label": "当前决策状态", "value": "暂不进入人工买入复核"},
        ]
        card["valuation"].update(
            available=True, value_text="Bear 404.82 / Base 479.88 / Bull 572.74 元；conditional_research_only",
            unavailable_reason=None, needed_evidence=None, status="UNDER_REVIEW",
        )
        card["decision_process"] = [dict(
            key=key, status="BLOCKED", reason="合成：当前 PriceBridge 未通过。",
            next_action="合成：等待新证据。", evidence_refs=["evidence-1"],
        ) for key in DECISION_STEP_TITLES]
        card["decision_process"][3].update(status="CONDITIONAL", assessment_id="synthetic-valuation")
        payload["companies"].append(card)
        today = deepcopy(item)
        today.update(symbol=symbol, company=name, category="RESEARCH_CHANGE",
                     why_it_matters="合成：旧泛化描述。", next_step=card["next_trigger"])
        payload["today_items"].append(today)
    payload["overview"]["pending_count"] = len(payload["today_items"])
    return product_workbench_from_payload(payload)


@pytest.mark.parametrize("reviews, expected", [
    ([('最强反证', '反证'), ('尚缺证据', '缺证'), ('为什么未进入更高状态', '原因')], '原因'),
    ([('最强反证', '反证'), ('尚缺证据', '缺证')], '缺证'),
    ([('最强反证', '反证')], '反证'),
    ([], '决定步骤阻断'),
])
def test_main_reason_uses_recorded_review_before_trigger(reviews, expected):
    from dataclasses import replace
    company = _synthetic_model().companies[0]
    steps = tuple(replace(step, reason='决定步骤阻断' if step.key == 'decision_gate' else '其他阻断')
                  for step in company.decision_process)
    company = replace(company, decision_review=tuple(reviews), decision_process=steps)
    assert _company_main_reason(company) == expected
    assert _company_main_reason(company) != company.next_trigger
    passed = tuple(replace(step, status='PASS', assessment_id='synthetic-pass') for step in steps)
    assert _company_main_reason(replace(company, decision_review=(), decision_process=passed)) is None


def test_python_render_localizes_conditionals_without_mutating_decisions():
    model = _synthetic_model()
    before = asdict(model)
    workbook = build_product_workbench_workbook(model)
    try:
        today = workbook[SHEET_TODAY]
        for company in model.companies:
            row = next(row for row in today if row[0].value == company.company_name)
            assert row[2].value == _company_main_reason(company)
            assert row[4].value == company.next_trigger
        text = '\n'.join(str(cell.value) for name in USER_SHEETS for row in workbook[name]
                         for cell in row if cell.value is not None)
        assert 'conditional_research_only' not in text and 'PriceBridge' not in text
        assert '仅供条件压力情景研究（非中性合理价值）' in text
        assert '价格与估值口径桥接' in text
        for value in ('404.82', '479.88', '572.74', '暂不进入人工买入复核'):
            assert value in text
        assert asdict(model) == before
        assert model.action == 'no_order' and model.portfolio.positions == ()
        assert not model.portfolio.real_data_available
    finally:
        workbook.close()
    assert _user_text('中性合理价值 100.00 元') == '中性合理价值 100.00 元'
    assert _user_text('估值假设 base（非预测批准）') == '估值假设 中情景（非预测批准）'


@pytest.fixture(scope="module")
def artifact_export_case(tmp_path_factory):
    tmp_path = tmp_path_factory.mktemp('synthetic-artifact-display')
    bundle = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    node = bundle / 'bin/node.exe'
    modules = bundle / 'node_modules'
    pwsh = Path.home() / 'AppData/Local/Programs/PowerShell/7/pwsh.exe'
    if os.name != 'nt' or not node.is_file() or not (modules / '@oai/artifact-tool').is_dir():
        pytest.skip('Bundled artifact-tool runtime is unavailable')
    script = tmp_path / 'scripts/current/render_product_artifact_pages.mjs'
    script.parent.mkdir(parents=True)
    script.write_bytes((ROOT / 'scripts/current/render_product_artifact_pages.mjs').read_bytes())
    subprocess.run([str(pwsh), '-NoProfile', '-Command',
                    'New-Item -ItemType Junction -Path $env:D2_ARTIFACT_JUNCTION '
                    '-Target $env:D2_ARTIFACT_MODULES | Out-Null'], check=True, capture_output=True, timeout=30,
                   env=dict(os.environ, D2_ARTIFACT_JUNCTION=str(tmp_path / 'node_modules'),
                            D2_ARTIFACT_MODULES=str(modules)))
    runtime = tmp_path / 'runtime'
    runtime.mkdir()

    def write(name, value):
        target = runtime / name
        target.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
        return target

    def binding(target):
        return dict(path=target.relative_to(tmp_path).as_posix(),
                    sha256=hashlib.sha256(target.read_bytes()).hexdigest())

    model = _synthetic_model()
    snapshot = json.loads(json.dumps(asdict(model), default=lambda value: value.isoformat()))
    snapshot_before = deepcopy(snapshot)
    snapshot['companies'][0]['decision_review'][0][1] *= 90
    decision = write('synthetic-decision.json', dict(symbol='600519', valuation=dict(
        bear_value='404.82', base_value='479.88', bull_value='572.74'), action='no_order',
        recommendation='NO_ACTION', buy_price=None, position_guidance=None))
    read_model = write('synthetic-read-model.json', dict(decision_workbench_binding=binding(decision)))
    policy = runtime / 'synthetic-display-policy.py'
    policy.write_bytes((ROOT / 'src/value_investment_agent/presentation/excel/product_workbench.py').read_bytes())
    publication = write('synthetic-publication.json', dict(snapshot=snapshot, action='no_order',
        publication_approved=False, canonical_written=False, read_model_binding=binding(read_model),
        source_bindings=[binding(read_model), binding(decision)]))
    strings = set()

    def visit(value):
        if isinstance(value, str):
            strings.add(value)
        elif isinstance(value, dict):
            for item in value.values():
                visit(item)
        elif isinstance(value, list):
            for item in value:
                visit(item)

    visit(snapshot)
    display = write('synthetic-display.json', dict(input_sha256=binding(publication)['sha256'],
        policy_binding=binding(policy), strings={value: _user_text(value) for value in strings}))
    frozen_hashes = {target: binding(target)['sha256'] for target in (publication, display, decision, read_model)}
    output = runtime / 'synthetic-preview.xlsx'
    result = subprocess.run([str(node), str(script), '--publication-input', str(publication),
        '--sha256', binding(publication)['sha256'], '--display-input', str(display), '--output', str(output)],
        capture_output=True, text=True, encoding='utf-8', timeout=120)
    assert result.returncode == 0, result.stdout + result.stderr
    return dict(model=model, snapshot_before=snapshot_before, output=output, publication=publication,
                decision=decision, frozen_hashes=frozen_hashes, binding=binding)


def test_artifact_export_uses_shared_display_and_preserves_blocked_states(artifact_export_case):
    case = artifact_export_case
    model, output = case['model'], case['output']
    binding, publication = case['binding'], case['publication']
    workbook = load_workbook(output, data_only=False)
    try:
        assert tuple(workbook.sheetnames) == WORKBOOK_SHEETS
        today = workbook[SHEET_TODAY]
        for company in model.companies:
            row = next(row for row in today if company.symbol in str(row[0].value))
            assert _company_main_reason(company) in row[3].value
            assert row[4].value == company.next_trigger
            assert '03_公司' in row[0].value
        texts = [str(cell.value) for name in USER_SHEETS for row in workbook[name]
                 for cell in row if cell.value is not None]
        text = '\n'.join(texts)
        assert 'conditional_research_only' not in text and 'PriceBridge' not in text
        assert '仅供条件压力情景研究（非中性合理价值）' in text
        assert '暂不进入人工买入复核' in text
        decision_states = [row[2].value for row in workbook['决策过程'].iter_rows(min_row=6)]
        assert decision_states.count('未通过') == 28
        assert decision_states.count('条件性结果') == 4
        values = {cell.value for row in workbook['03_公司'] for cell in row if cell.data_type == 'n'}
        assert {404.82, 479.88, 572.74} <= values
        formulas = [cell.value for sheet in workbook for row in sheet for cell in row if cell.data_type == 'f']
        assert formulas and all(formula.startswith('=HYPERLINK("#\'') for formula in formulas)
        assert any('原始投资逻辑（证据）' in formula for formula in formulas)
        assert '个人仓位' in text and '尚未接入真实组合' in text
    finally:
        workbook.close()
    receipt = json.loads(output.with_name('synthetic-preview-receipt.json').read_text(encoding='utf-8'))
    assert receipt['action'] == 'no_order' and receipt['canonical_written'] is False
    assert receipt['input_sha256'] == binding(publication)['sha256']
    assert receipt['workbook_sha256'] == binding(output)['sha256']
    for target, digest in case['frozen_hashes'].items():
        assert binding(target)['sha256'] == digest
    assert asdict(model)['companies'][0]['decision_review'][0][1] == case['snapshot_before']['companies'][0]['decision_review'][0][1]


def test_pressure_table_keeps_three_values_and_fits_the_previous_three_row_space(artifact_export_case):
    case = artifact_export_case
    decision = json.loads(case['decision'].read_text(encoding='utf-8'))
    workbook = load_workbook(case['output'], data_only=False)
    try:
        sheet = workbook['03_公司']
        headers = [row for row in sheet if row[2].value == '条件情景值（元/股，非合理价）']
        assert len(headers) == 1
        start = headers[0][0].row
        assert [cell.value for cell in headers[0]] == [
            '公司代码', '压力情景', '条件情景值（元/股，非合理价）',
        ]
        labels = ['低值压力情景', '中值压力情景', '高值压力情景']
        expected = [float(decision['valuation'][f'{key}_value']) for key in ('bear', 'base', 'bull')]
        assert expected == [404.82, 479.88, 572.74]
        for offset, (label, value) in enumerate(zip(labels, expected), start=1):
            cells = list(sheet.iter_rows(min_row=start+offset, max_row=start+offset, max_col=3))[0]
            assert [cell.value for cell in cells] == [decision['symbol'], label, value]
            assert cells[2].data_type == 'n' and cells[2].number_format == '0.00'
        heights = [sheet.row_dimensions[row].height for row in range(start, start+4)]
        assert heights == [22.0] * 4
        assert sum(heights) <= 3 * 30
        for row in sheet.iter_rows(min_row=start, max_row=start+3, max_col=3):
            assert all(cell.font.sz == 11 for cell in row)
    finally:
        workbook.close()
