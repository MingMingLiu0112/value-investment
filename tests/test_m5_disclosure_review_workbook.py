from __future__ import annotations

from datetime import date, datetime, timezone
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest

from value_investment_agent.event_materiality import DECISION_REQUIRES_RECALCULATION
from value_investment_agent.m5_disclosure_queue import (
    DISCLOSURE_QUEUE_SCHEMA,
    PARSER_VERSION,
    PROVIDER_CNINFO,
    DisclosureReviewQueue,
    build_cninfo_event_scan,
)
from value_investment_agent.m5_disclosure_review import (
    ACTION_NO_ORDER,
    build_disclosure_materiality_reviews,
    disclosure_queue_sha256,
)
from value_investment_agent.m5_disclosure_review_workbook import (
    INPUT_SHEET,
    OVERVIEW_SHEET,
    QUEUE_HASH_LABEL,
    QUEUE_ID_LABEL,
    REVIEWED_AT_LABEL,
    build_m5_disclosure_review_workbook,
    read_m5_disclosure_review_workbook,
    write_m5_disclosure_review_workbook,
)


SCAN_FROM = date(2026, 8, 27)
SCAN_TO = date(2026, 9, 24)
RETRIEVED_AT = datetime(2026, 9, 24, 16, 0, tzinfo=timezone.utc)


def _timestamp(day: int = 2) -> int:
    return int(datetime(2026, 9, day, 1, 0, tzinfo=timezone.utc).timestamp() * 1000)


def _row(announcement_id: str, *, day: int = 2) -> dict:
    return {
        "secCode": "600887",
        "announcementId": announcement_id,
        "announcementTitle": "关于回购公司股份的公告",
        "announcementTime": _timestamp(day),
        "adjunctUrl": f"finalpage/2026-09-{day:02d}/{announcement_id}.PDF",
    }


def _downloader(source_url: str, target: Path) -> str:
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(b"%PDF-fixture")
    return "a" * 64


def _queue(tmp_path: Path, rows: list[dict]) -> DisclosureReviewQueue:
    scan = build_cninfo_event_scan(
        {
            "url": "https://www.cninfo.com.cn/new/hisAnnouncement/query",
            "total_announcements": len(rows),
            "announcements": rows,
        },
        symbol="600887",
        scan_from=SCAN_FROM,
        scan_to=SCAN_TO,
        retrieved_at=RETRIEVED_AT,
        evidence_dir=tmp_path / "600887",
        root=tmp_path,
        pdf_downloader=_downloader,
    )
    return DisclosureReviewQueue(
        queue_id="fixture-cninfo-review",
        schema_version=DISCLOSURE_QUEUE_SCHEMA,
        provider=PROVIDER_CNINFO,
        parser_version=PARSER_VERSION,
        scan_from=SCAN_FROM,
        scan_to=SCAN_TO,
        retrieved_at=RETRIEVED_AT,
        scans=(scan,),
        action=ACTION_NO_ORDER,
    )


def _metadata_row(ws, label: str) -> tuple[int, int]:
    for row in ws.iter_rows(min_col=1, max_col=2):
        if row[0].value == label:
            return row[1].row, row[1].column
    raise AssertionError(f"missing metadata: {label}")


def test_workbook_has_no_default_verdicts_and_reports_missing(tmp_path: Path):
    queue = _queue(tmp_path, [_row("1"), _row("2", day=3)])
    wb = build_m5_disclosure_review_workbook(queue)
    assert wb[INPUT_SHEET].max_row == 6
    assert all(
        wb[INPUT_SHEET].cell(row, 7).value in (None, "")
        for row in range(5, 7)
    )

    output = tmp_path / "review.xlsx"
    wb.save(output)
    wb = load_workbook(output)
    overview = wb[OVERVIEW_SHEET]
    row, column = _metadata_row(overview, REVIEWED_AT_LABEL)
    overview.cell(row, column, "2026-09-24")
    wb.save(output)
    with pytest.raises(ValueError, match="decision is missing.*1.*2"):
        read_m5_disclosure_review_workbook(output, queue)


def test_filled_workbook_round_trips_and_builds_reviews(tmp_path: Path):
    queue = _queue(tmp_path, [_row("1"), _row("2", day=3)])
    output = tmp_path / "review.xlsx"
    write_m5_disclosure_review_workbook(
        queue,
        output=output,
        root=tmp_path,
    )
    wb = load_workbook(output)
    input_sheet = wb[INPUT_SHEET]
    input_sheet.cell(5, 7, DECISION_REQUIRES_RECALCULATION)
    input_sheet.cell(5, 8, "balance_sheet_risk")
    input_sheet.cell(5, 12, "liability facts changed")
    input_sheet.cell(6, 7, "NOT_MATERIAL")
    input_sheet.cell(6, 12, "routine buyback disclosure")
    overview = wb[OVERVIEW_SHEET]
    row, column = _metadata_row(overview, REVIEWED_AT_LABEL)
    overview.cell(row, column, "2026-09-24 18:00")
    wb.save(output)

    intake = read_m5_disclosure_review_workbook(output, queue)
    assert intake.action == ACTION_NO_ORDER
    assert len(intake.decisions) == 2
    reviews = build_disclosure_materiality_reviews(queue, intake)
    assert len(reviews) == 1
    assert reviews[0].decisions[0].source_sha256 == "a" * 64
    assert reviews[0].decisions[0].human_decision == DECISION_REQUIRES_RECALCULATION


def test_missing_note_and_non_candidate_rows_fail(tmp_path: Path):
    queue = _queue(tmp_path, [_row("1"), _row("2", day=3)])
    output = tmp_path / "review.xlsx"
    write_m5_disclosure_review_workbook(
        queue,
        output=output,
        root=tmp_path,
    )
    wb = load_workbook(output)
    sheet = wb[INPUT_SHEET]
    sheet.cell(5, 7, "NOT_MATERIAL")
    sheet.cell(5, 12, "reviewed")
    sheet.cell(6, 7, "NOT_MATERIAL")
    sheet.cell(6, 12, "")
    overview = wb[OVERVIEW_SHEET]
    row, column = _metadata_row(overview, REVIEWED_AT_LABEL)
    overview.cell(row, column, "2026-09-24")
    wb.save(output)

    with pytest.raises(ValueError, match="note is missing.*2"):
        read_m5_disclosure_review_workbook(output, queue)

    sheet.cell(6, 12, "reviewed")
    sheet.cell(7, 1, "600887")
    sheet.cell(7, 2, "999999")
    sheet.cell(7, 7, "NOT_MATERIAL")
    sheet.cell(7, 12, "extra row")
    wb.save(output)
    with pytest.raises(ValueError, match="non-candidate"):
        read_m5_disclosure_review_workbook(output, queue)


def test_queue_fingerprint_change_rejects_workbook(tmp_path: Path):
    queue = _queue(tmp_path, [_row("1")])
    output = tmp_path / "review.xlsx"
    write_m5_disclosure_review_workbook(
        queue,
        output=output,
        root=tmp_path,
    )
    wb = load_workbook(output)
    row, column = _metadata_row(wb[OVERVIEW_SHEET], QUEUE_HASH_LABEL)
    wb[OVERVIEW_SHEET].cell(row, column, "c" * 64)
    wb.save(output)

    with pytest.raises(ValueError, match="queue hash"):
        read_m5_disclosure_review_workbook(output, queue)


def test_write_refuses_existing_output(tmp_path: Path):
    queue = _queue(tmp_path, [_row("1")])
    output = tmp_path / "review.xlsx"
    output.write_bytes(b"occupied")
    with pytest.raises(ValueError, match="already exists"):
        write_m5_disclosure_review_workbook(
            queue,
            output=output,
            root=tmp_path,
        )


def test_apply_script_writes_reviews_and_batches_without_events(tmp_path: Path):
    queue = _queue(tmp_path, [_row("1")])
    queue_path = tmp_path / "queue.json"
    queue_path.write_text(queue.to_json() + "\n", encoding="utf-8")
    output = tmp_path / "review.xlsx"
    write_m5_disclosure_review_workbook(
        queue,
        output=output,
        root=tmp_path,
    )
    wb = load_workbook(output)
    sheet = wb[INPUT_SHEET]
    sheet.cell(5, 7, "NOT_MATERIAL")
    sheet.cell(5, 12, "human reviewed")
    overview = wb[OVERVIEW_SHEET]
    row, column = _metadata_row(overview, REVIEWED_AT_LABEL)
    overview.cell(row, column, "2026-09-24 18:00")
    wb.save(output)

    runtime_root = tmp_path / "applied"
    result = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve().parents[1] / "scripts" / "apply_m5_disclosure_review.py"),
            "--queue",
            str(queue_path),
            "--workbook",
            str(output),
            "--runtime-root",
            str(runtime_root),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert result.returncode == 0, result.stderr
    manifest = json.loads(
        (runtime_root / "manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["action"] == ACTION_NO_ORDER
    assert manifest["events_not_applied"] is True
    assert manifest["review_count"] == 1
    assert manifest["decision_count"] == 1
    assert manifest["event_count"] == 0
    assert (runtime_root / "bridge_batches.json").is_file()
