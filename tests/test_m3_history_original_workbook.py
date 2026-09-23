from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook
import pytest

from value_investment_agent.m3_history_original_workbook_acceptance_audit import (
    M3HistoryOriginalWorkbookAcceptanceSpec,
    PENDING_HUMAN_REVIEW,
    audit,
    write_receipt,
)


ROOT = Path(__file__).resolve().parents[1]
BUILDER_SPEC = importlib.util.spec_from_file_location(
    "m3_history_original_workbook_builder",
    ROOT / "scripts" / "build_m3_history_original_workbook_candidate.py",
)
BUILDER = importlib.util.module_from_spec(BUILDER_SPEC)
BUILDER_SPEC.loader.exec_module(BUILDER)

STAGE_SPEC = importlib.util.spec_from_file_location(
    "stage_frontend_package",
    ROOT / "scripts" / "stage_frontend_package.py",
)
STAGE = importlib.util.module_from_spec(STAGE_SPEC)
STAGE_SPEC.loader.exec_module(STAGE)


GENERATED_AT = datetime(2026, 9, 24, 0, 0, tzinfo=timezone.utc)
HISTORY_SHEETS = (
    "00_历史链",
    "01_原Entry",
    "02_决策日志",
    "03_一致性复核",
    "04_来源哈希",
)


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _source_workbook(path: Path) -> None:
    workbook = Workbook()
    retained = workbook.active
    retained.title = "用户页"
    retained["A1"] = "保留的手工内容"
    decision = workbook.create_sheet("00_决策复核")
    decision["A1"] = "决策复核（M3 非个人化负向卡）"
    decision["A2"] = "action=no_order"
    decision["A3"] = "000651 600741 600887"
    decision["A4"] = "Checkpoint B：未完成"
    workbook.save(path)


def _addon_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    overview = workbook.create_sheet(HISTORY_SHEETS[0])
    overview["A1"] = "M3 论点连续性历史链"
    overview["A2"] = "模拟演示链路：只解释原始理由是否变化，不生成订单或真实成交。"
    header = ["证券代码", "公司", "命名空间", "Entry类型", "Entry日期", "原始论点摘要", "最新人工决定", "最新一致性", "动作"]
    data = ["600887", "伊利股份", "模拟", "模拟", "2026-09-01", "模拟理由", "确认减仓", "已破坏", "no_order"]
    for column, value in enumerate(header, start=1):
        overview.cell(row=4, column=column, value=value)
    for column, value in enumerate(data, start=1):
        overview.cell(row=5, column=column, value=value)
    for name in HISTORY_SHEETS[1:]:
        sheet = workbook.create_sheet(name)
        sheet["A1"] = name
    workbook.save(path)


def _spec(tmp_path: Path, *, tamper: bool = False) -> M3HistoryOriginalWorkbookAcceptanceSpec:
    source = tmp_path / "source.xlsx"
    addon = tmp_path / "addon.xlsx"
    history_input = tmp_path / "history.json"
    output = tmp_path / "candidate.xlsx"
    _source_workbook(source)
    _addon_workbook(addon)
    history_input.write_text("{}", encoding="utf-8")
    manifest = BUILDER.build_candidate(
        source=source,
        addon=addon,
        history_input=history_input,
        output=output,
        expected_source_sha256=_digest(source),
        expected_addon_sha256=_digest(addon),
        expected_history_input_sha256=_digest(history_input),
        generated_at=GENERATED_AT,
        project_root=tmp_path,
    )
    candidate_sha256 = _digest(output)
    if tamper:
        output.write_bytes(output.read_bytes() + b"tamper")
        candidate_sha256 = candidate_sha256

    manifest_path = output.with_name(output.stem + ".candidate.manifest.json")
    wps_dir = tmp_path / "WPSDrive"
    wps_dir.mkdir()
    wps_candidate = wps_dir / "candidate.xlsx"
    wps_candidate.write_bytes(output.read_bytes())
    canonical_wps = wps_dir / "canonical.xlsx"
    canonical_wps.write_bytes(source.read_bytes())
    wps_receipt = tmp_path / "wps-receipt.json"
    wps_receipt.write_text(
        json.dumps(
            {
                "status": "passed",
                "mode": "protected_history_overlay_candidate",
                "sha256": candidate_sha256,
                "sheets": manifest["sheet_count"],
                "history_sheets": list(HISTORY_SHEETS),
                "target_sheet": "00_决策复核",
                "source_sheets_preserved": manifest["original_sheets_preserved"],
                "new_sheets": 5,
                "namespace": "simulated",
                "action": "no_order",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return M3HistoryOriginalWorkbookAcceptanceSpec(
        source_path=source,
        source_sha256=_digest(source),
        addon_path=addon,
        addon_sha256=_digest(addon),
        history_input_path=history_input,
        history_input_sha256=_digest(history_input),
        candidate_path=output,
        candidate_sha256=candidate_sha256,
        manifest_path=manifest_path,
        manifest_sha256=_digest(manifest_path),
        generated_at=GENERATED_AT,
        wps_candidate_path=wps_candidate,
        wps_receipt_path=wps_receipt,
        wps_receipt_sha256=_digest(wps_receipt),
        canonical_wps_path=canonical_wps,
        canonical_sha256=_digest(source),
        expected_sheet_count=manifest["sheet_count"],
        expected_new_sheets=5,
        expected_source_sheets=manifest["original_sheets_preserved"],
        expected_unchanged_parts=manifest["original_parts_unchanged"],
    )


def test_builder_grafts_history_pages_deterministically(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    addon = tmp_path / "addon.xlsx"
    history_input = tmp_path / "history.json"
    output = tmp_path / "candidate.xlsx"
    second_output = tmp_path / "second.xlsx"
    _source_workbook(source)
    _addon_workbook(addon)
    history_input.write_text("{}", encoding="utf-8")

    manifest = BUILDER.build_candidate(
        source=source,
        addon=addon,
        history_input=history_input,
        output=output,
        expected_source_sha256=_digest(source),
        expected_addon_sha256=_digest(addon),
        expected_history_input_sha256=_digest(history_input),
        generated_at=GENERATED_AT,
        project_root=tmp_path,
    )
    second = STAGE.graft(source, addon, second_output, _digest(source))

    assert manifest["sheet_count"] == 7
    assert manifest["new_sheets"] == list(HISTORY_SHEETS)
    assert manifest["original_sheets_preserved"] == 2
    assert manifest["namespace"] == "simulated"
    assert manifest["action"] == "no_order"
    assert manifest["candidate_sha256"] == second["candidate_sha256"]
    assert _digest(output) == _digest(second_output)

    with pytest.raises(ValueError, match="already exists"):
        BUILDER.build_candidate(
            source=source,
            addon=addon,
            history_input=history_input,
            output=output,
            expected_source_sha256=_digest(source),
            expected_addon_sha256=_digest(addon),
            expected_history_input_sha256=_digest(history_input),
            generated_at=GENERATED_AT,
            project_root=tmp_path,
        )


def test_audit_verifies_overlay_without_publishing(tmp_path: Path):
    receipt = audit(
        tmp_path,
        _spec(tmp_path),
        ci_evidence={"status": "success"},
        test_evidence={"passed": True, "passed_count": 8, "skipped_count": 0},
    )

    assert receipt["status"] == PENDING_HUMAN_REVIEW
    assert all(
        receipt["criteria"][key]["status"] == "DONE"
        for key in (
            "hoc1_identity_and_manifest",
            "hoc2_history_simulated_boundary",
            "hoc3_original_workbook_structure",
            "hoc4_decision_sheet_fail_closed",
            "hoc5_wps_and_canonical_boundary",
            "hoc6_offline_regression_and_ci",
        )
    )
    assert receipt["criteria"]["hoc7_human_checkpoint_review"]["status"] == PENDING_HUMAN_REVIEW
    assert receipt["blockers"] == []
    assert receipt["action"] == "no_order"


def test_audit_fails_closed_when_candidate_changes(tmp_path: Path):
    with pytest.raises(ValueError, match="M3 history overlay candidate"):
        audit(
            tmp_path,
            _spec(tmp_path, tamper=True),
            ci_evidence={"status": "success"},
            test_evidence={"passed": True, "passed_count": 8, "skipped_count": 0},
        )


def test_receipt_writer_creates_versioned_pointer(tmp_path: Path):
    receipt = audit(
        tmp_path,
        _spec(tmp_path),
        ci_evidence={"status": "success"},
        test_evidence={"passed": True, "passed_count": 8, "skipped_count": 0},
    )
    output = write_receipt(receipt, root=tmp_path)
    receipt_path = tmp_path / output["receipt_path"]
    pointer_path = tmp_path / output["pointer_path"]
    assert receipt_path.is_file()
    assert pointer_path.is_file()
    assert _digest(receipt_path) == output["receipt_sha256"]
    assert json.loads(pointer_path.read_text(encoding="utf-8"))["sha256"] == output["receipt_sha256"]
