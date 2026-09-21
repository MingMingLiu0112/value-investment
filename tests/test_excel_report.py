from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
import pytest

from value_investment_agent.excel_report import build_report, choose_points


def fixture(tmp_path):
    w = Workbook()
    w.active.title = '00_首页Dashboard'
    for name in ['01_观察名单','02_质量评分','03_财务指标','04_估值跟踪','05_仓位管理',
                 '06_月度跟踪','07_月度复盘','08_交易记录','09_公司研究','10_年报跟踪','11_数据源审计']:
        w.create_sheet(name)
    w['09_公司研究'].append(['代码','公司'])
    w['09_公司研究']['A4']='000001'
    w['09_公司研究']['D4']='个人研究内容'
    w['05_仓位管理']['E10']=123
    w['08_交易记录']['A4']='个人历史交易'
    w['11_数据源审计']['A1501']='older-audit'
    w['06_月度跟踪']['A800']='2025-01'
    w['06_月度跟踪']['B800']='000001'
    w['06_月度跟踪']['D800']=15
    path=tmp_path/'template.xlsx'; w.save(path)
    today=datetime.now(timezone.utc)
    p={'market_candidates':[{'symbol':s,'name':name,'sector':sector,'source_id':'market',
        'screen_date':str(today.date()),'board':'主板','current_price':10,'pe':5,'pb':0.5}
        for s,name,sector in [('000001','平安银行','银行'),('600001','=danger','制造业')]],
        'generated_at':today.isoformat(),'points':[], 'monthly_snapshots':[]}
    return path,p


def test_company_research_preserves_personal_notes_and_has_live_link(tmp_path):
    path, payload = fixture(tmp_path)
    display = ['历史研究摘要', '反证', '估值及回测未完成', 'https://static.cninfo.com.cn/example.PDF']
    output = tmp_path / 'research.xlsx'
    build_report(path, payload, output, company_research={'000001': display})
    workbook = load_workbook(output)
    sheet = workbook['09_公司研究']
    row = next(row for row in sheet.iter_rows(min_row=4) if row[0].value == '000001')
    assert row[3].value == '个人研究内容'
    assert [cell.value for cell in row[13:17]] == display
    assert row[16].hyperlink.target == display[3]
    assert sheet.column_dimensions['N'].width == 56
    assert sheet.row_dimensions[row[0].row].height == 100


def test_unit_evidence_reaches_excel_evidence_sheet(tmp_path):
    path, payload = fixture(tmp_path)
    context = '单位依据（PDF第64页）：财务附注中报表的单位为：元\n数值原文：货币资金 100'
    payload['points'] = [{'symbol': '600001', 'field_name': 'cash', 'value': '100',
        'period_label': '2026-06-30', 'unit': 'CNY', 'validation_status': 'pending',
        'source_id': 'pdf', 'metadata': {'page_number': 66, 'candidate_excerpt': context}}]
    output = tmp_path / 'unit-evidence.xlsx'
    build_report(path, payload, output)
    with_open = load_workbook(output)
    evidence = with_open['18_指标证据']
    assert evidence.cell(3, 16).value == '原文与单位依据'
    row = next(r for r in evidence.iter_rows(min_row=4, values_only=True) if r[3] == 'cash')
    assert row[11] == 66
    assert row[15] == context
    with_open.close()


def test_legacy_provider_ratio_not_displayed_as_canonical_percent(tmp_path):
    path, payload = fixture(tmp_path)
    payload['points'] = [{'symbol': '600001', 'field_name': 'operating_cash_flow_to_net_income',
        'period_label': '2025-12-31', 'value': '0.947', 'unit': 'percent',
        'validation_status': 'pending', 'metadata': {}, 'source_id': 'legacy',
        'source_name': 'AkShare / Sina financial indicators',
        'source_url': 'https://example.test/report', 'sha256': 'a' * 64}]
    output = tmp_path / 'ratio.xlsx'
    build_report(path, payload, output)
    workbook = load_workbook(output)
    row = next(r for r in workbook['03_财务指标'].iter_rows(min_row=4) if r[0].value == '600001')
    assert row[13].value is None
    assert any(r[3] == 'operating_cash_flow_to_net_income' and r[5] == 0.947
               for r in workbook['18_指标证据'].iter_rows(min_row=4, values_only=True))
    workbook.close()


def test_all_main_views_include_all_candidates_and_preserve_user_data(tmp_path):
    path,p=fixture(tmp_path)
    out=tmp_path/'out.xlsx'
    build_report(path,p,out)
    w=load_workbook(out)
    for name in ['01_观察名单','02_质量评分','03_财务指标','04_估值跟踪','09_公司研究','10_年报跟踪']:
        assert {r[0] for r in w[name].iter_rows(min_row=4,values_only=True) if r[0]}=={'000001','600001'}
    assert w['09_公司研究']['D4'].value=='个人研究内容'
    assert w['05_仓位管理']['E10'].value==123
    assert w['08_交易记录']['A4'].value=='个人历史交易'
    assert w['11_数据源审计']['A1501'].value=='older-audit'
    assert w['01_观察名单']['B5'].data_type=='s'
    assert w['04_估值跟踪']['D4'].value is None
    assert w['12_提醒'].sheet_state=='hidden'
    assert w['00_公司总览'].sheet_state=='visible'
    assert {r[0] for r in w['00_待完成公司'].iter_rows(min_row=4,values_only=True)}=={'000001','600001'}
    assert w['00_待完成公司']['B5'].data_type=='s'
    assert w.sheetnames[1]=='00_公司总览'
    assert w['12_提醒'].max_row==5
    assert w['12_提醒']['O3'].value=='策略验证状态'
    assert w['12_提醒']['O4'].value=='历史回测未完成'
    assert w['12_提醒']['P4'].hyperlink.location=="'21_决策验证'!A4"
    assert w.sheetnames[2]=='00_待完成公司'
    assert w['21_决策验证'].sheet_state=='visible'
    assert w['21_决策验证'].max_row==5
    assert w['21_决策验证']['F4'].value==8
    assert w['21_决策验证']['I4'].value=='未通过'
    assert w['21_决策验证']['I4'].comment.text
    assert w['21_决策验证']['S4'].value is None
    assert w['21_决策验证']['X3'].value=='策略验证状态'
    assert w['21_决策验证']['X4'].value=='历史回测未完成'
    assert w['21_决策验证']['Y4'].value is None
    assert w['21_决策验证']['AA3'].value == '收盘交易日核验'
    assert '缺少带实际日期' in w['21_决策验证']['AA4'].value
    assert w['21_决策验证']['AB4'].value is None
    assert w['21_决策验证']['AC4'].value is None
    assert w['20_市场覆盖']['B6'].value=='尚未完成'


def test_broker_reference_does_not_promote_pending_evidence(tmp_path):
    path, payload = fixture(tmp_path)
    payload['market_candidates'][0].update(name='样例证券', sector='证券')
    payload['points'] = [{'symbol': '000001', 'field_name': 'capital_leverage',
        'period_label': '2026-06-30', 'unit': 'percent', 'value': '21.18',
        'source_id': 'official', 'validation_status': 'pending',
        'source_url': 'https://example.test/report', 'sha256': 'a' * 64,
        'metadata': {'statement_scope': '母公司（原文明确）', 'page_number': 12,
                     'automatic_cross_source_verification': False}}]
    output = tmp_path / 'broker.xlsx'
    build_report(path, payload, output)
    workbook = load_workbook(output)
    sheet = workbook['19_金融专用指标']
    row = next(r for r in sheet.iter_rows(min_row=4) if r[4].value == '资本杠杆率')
    assert row[5].value == 21.18
    assert row[7].value != '已交叉验证'
    assert row[12].value == 8 and row[13].value == 9.6
    assert row[14].value == '数值高于一般预警线'
    assert '不代表数据已验证' in row[14].comment.text
    assert workbook['21_决策验证']['X4'].value == '历史回测未完成'


def test_units_periods_and_evidence_are_not_conflated(tmp_path):
    path,p=fixture(tmp_path)
    base={'symbol':'600001','period_label':'2026-06-30','unit':'CNY','source_id':'test',
          'value':'100000000','validation_status':'pending','source_url':'https://example.test/report',
          'sha256':'a'*64,'metadata':{}}
    p['points']=[dict(base,field_name='revenue'),dict(base,field_name='net_income',period_label='2025-12-31'),
                 dict(base,field_name='eps_reported',unit='CNY/share',value='2')]
    out=tmp_path/'out.xlsx'; build_report(path,p,out); w=load_workbook(out)
    assert w['03_财务指标']['D5'].value==1
    assert w['03_财务指标']['F5'].value is None
    assert w['04_估值跟踪']['D5'].value is None  # Interim EPS is not TTM.
    assert '待交叉验证' in w['03_财务指标']['D5'].comment.text
    assert w['11_数据源审计']['A1502'].value=='test'


def test_snapshot_is_append_only_and_repeat_sync_is_idempotent(tmp_path):
    path,p=fixture(tmp_path)
    p['monthly_snapshots']=[{'symbol':'000001','snapshot_month':'2025-01','current_price':999},
                            {'symbol':'600001','snapshot_month':'2025-02','current_price':10}]
    out=tmp_path/'out.xlsx'; out2=tmp_path/'out2.xlsx'
    build_report(path,p,out); build_report(out,p,out2)
    w=load_workbook(out2)
    assert w['06_月度跟踪']['D800'].value==15
    assert w['06_月度跟踪']['D801'].value==10
    assert w['06_月度跟踪'].max_row==801


def test_reject_partial_count_and_never_prefer_old_verified_period(tmp_path):
    path,p=fixture(tmp_path)
    p['market_candidate_count']=3
    with pytest.raises(ValueError): build_report(path,p,tmp_path/'bad.xlsx')
    points=[{'symbol':'600001','field_name':'revenue','period_label':'2025-12-31','validation_status':'verified',
             'metadata':{'automatic_cross_source_verification':True}},
            {'symbol':'600001','field_name':'revenue','period_label':'2026-06-30','validation_status':'pending'}]
    assert choose_points(points)[('600001','revenue')]['period_label']=='2026-06-30'


def test_stale_export_blocks_build_signal(tmp_path):
    path,p=fixture(tmp_path)
    p['generated_at']='2000-01-01T00:00:00+00:00'
    p['valuations']=[{'symbol':'600001','build_signal':'建仓候选','target_weight':0.1}]
    out=tmp_path/'out.xlsx'; result=build_report(path,p,out)
    assert result['stale']
    assert load_workbook(out)['01_观察名单']['L5'].value is None


def test_monthly_screen_uses_daily_quotes_without_changing_entry_date(tmp_path):
    path, p = fixture(tmp_path)
    for c in p['market_candidates']:
        c['screen_date'] = '2020-01-01'
    p['candidate_tracking'] = [dict(symbol=c['symbol'], current_price=22,
        pe=30, pb=4, source_id='daily', quote_as_of=p['generated_at'],
        quote_status='matched', signal_blocked=False) for c in p['market_candidates']]
    out = tmp_path/'daily.xlsx'
    result = build_report(path, p, out)
    wb = load_workbook(out)
    assert not result['stale']
    assert wb['01_观察名单']['F4'].value == 22
    assert wb['01_观察名单']['R3'].value == '初筛日期'
    assert wb['01_观察名单']['R4'].value == '2020-01-01'
    assert wb['04_估值跟踪']['O4'].value == p['generated_at']
    dashboard = {r[0]:r[1] for r in wb['00_首页Dashboard'].iter_rows(min_row=4, values_only=True)}
    assert dashboard['初筛日期'] == '2020-01-01'
    assert dashboard['行情快照时间'] == p['generated_at']
    assert p['market_candidates'][0]['current_price'] == 10


def test_frontdoor_reapplication_keeps_columns_and_source_cells(tmp_path):
    from value_investment_agent.workbook_frontdoor import apply_frontdoor, HOME, PRIMARY
    from value_investment_agent.workbook_simple_overview import DERIVED
    path, payload = fixture(tmp_path)
    output = tmp_path / 'frontdoor.xlsx'
    build_report(path, payload, output)
    book = load_workbook(output)
    dimensions = {sheet.title: sheet.max_column for sheet in book if sheet.title not in DERIVED}
    records = {sheet.title: list(sheet.values) for sheet in book if sheet.title not in DERIVED}
    apply_frontdoor(book)
    apply_frontdoor(book)
    assert dimensions == {sheet.title: sheet.max_column for sheet in book if sheet.title not in DERIVED}
    assert records == {sheet.title: list(sheet.values) for sheet in book if sheet.title not in DERIVED}
    assert [sheet.title for sheet in book if sheet.sheet_state == 'visible'] == [s for s in PRIMARY if s in book]
    assert book.active.title == HOME
    book.close()


def test_missing_daily_run_does_not_display_initial_screen_price(tmp_path):
    path, p = fixture(tmp_path)
    p['candidate_tracking'] = []
    out = tmp_path/'missing.xlsx'
    result = build_report(path, p, out)
    wb = load_workbook(out)
    assert result['stale']
    assert wb['01_观察名单']['F4'].value is None
    assert wb['04_估值跟踪']['C4'].value is None


def test_historical_research_keeps_financial_evidence_without_reentering_pool(tmp_path):
    path, p = fixture(tmp_path)
    book = load_workbook(path)
    book['09_公司研究']['A5'] = '600900'
    book['09_公司研究']['B5'] = 'historical company'
    book.save(path)
    p['annual_points'] = [{'symbol':'600900','field_name':'long_term_borrowings',
        'period_label':'2025-12-31','value':'172310624303.31','unit':'CNY',
        'source_id':'official','metadata':{'page_number':81}}]
    out=tmp_path/'history-evidence.xlsx'
    build_report(path,p,out)
    book=load_workbook(out)
    evidence=[r for r in book['18_指标证据'].iter_rows(min_row=4,values_only=True) if r[0]=='600900']
    assert len(evidence)==1 and evidence[0][11]==81
    assert '600900' not in {r[0] for r in book['01_观察名单'].iter_rows(min_row=4,values_only=True)}


@pytest.mark.parametrize('debt_scope_verified', [True, False])
def test_annual_quality_is_separate_from_interim_tracking_with_retained_evidence(tmp_path, debt_scope_verified):
    from value_investment_agent.financial_quality import GENERAL_FIELDS
    path, p = fixture(tmp_path)
    base = {'symbol': '600001', 'value': 10, 'unit': 'percent', 'source_id': 'annual',
            'validation_status': 'verified', 'metadata': {'automatic_cross_source_verification': True}}
    p['annual_points'] = [dict(base, field_name=f, period_label='2025-12-31',
        metadata={**base['metadata'], **({'complete_debt_verified': True}
                  if f == 'interest_bearing_debt' and debt_scope_verified else {})}) for f in GENERAL_FIELDS]
    p['points'] = [dict(base, field_name='roe', period_label='2026-06-30', source_id='interim')]
    p['financial_quality'] = [{'symbol':'600001', 'quality_status':'已验证', 'total_score':80,
                                'calculation_details':{'period':'2025-12-31'}}]
    out = tmp_path / 'annual.xlsx'
    build_report(path, p, out)
    w = load_workbook(out)
    assert w['02_质量评分']['E5'].value == '2025-12-31'
    assert w['02_质量评分']['H5'].value == (80 if debt_scope_verified else None)
    assert w['02_质量评分']['N5'].value == 10
    assert 'annual' in w['02_质量评分']['N5'].comment.text
    assert w['03_财务指标']['C5'].value == '2026-06-30'
    assert w['01_观察名单']['K5'].value == '仅研究，不生成建仓信号'
    evidence = list(w['18_指标证据'].iter_rows(min_row=4, values_only=True))
    assert {r[4] for r in evidence if r[3] == 'roe'} == {'2025-12-31', '2026-06-30'}


def test_official_coverage_shows_evidence_and_differences(tmp_path):
    path,p=fixture(tmp_path)
    p['market_audit']={'coverage':{'official_universe_reconciled':False}}
    p['official_coverage']={'expected_count':3,'matched_count':2,'scope':'沪深北证券清单',
        'source_id':'official','sha256':'a'*64,'official_date':'2026-09-07','market_date':'2026-09-07',
        'missing_in_market':[{'symbol':'600003','name':'缺失样例'}],
        'sources':[{'label':'SZSE_A','sha256':'b'*64,'url':'https://www.szse.cn/api/report/ShowReport'}]}
    out=tmp_path/'official.xlsx'
    build_report(path,p,out)
    w=load_workbook(out)
    rows=list(w['20_市场覆盖'].iter_rows(min_row=4,values_only=True))
    assert any(r[0]=='逐代码匹配' and r[1]==2 for r in rows)
    assert any(r[0]=='行情缺失' and r[1]=='600003' for r in rows)
    assert any(r[0]=='SZSE_A' and r[1]=='b'*64 for r in rows)


def test_debt_research_links_do_not_change_signals_scores_or_create_tabs(tmp_path):
    template, payload = fixture(tmp_path)
    baseline, enriched = tmp_path / 'baseline.xlsx', tmp_path / 'enriched.xlsx'
    build_report(template, payload, baseline)
    research = {'companies': {'600001': '完整有息负债未验证'}, 'evidence': {'600001': [{
        'label': '负债附注研究：合计', 'field_name': 'research_current_maturity_total',
        'period': '2025-12-31', 'value': None, 'unit': 'CNY', 'status': '研究提取，非已验证指标',
        'source_id': 'official', 'source_url': 'https://example.test/report.pdf', 'sha256': 'a' * 64,
        'pages': '170,171', 'parser': 'fixture', 'fetched_at': None, 'excerpt': '本期合计为空，不填零'}]}}
    build_report(template, payload, enriched, debt_research=research)
    a, b = load_workbook(baseline), load_workbook(enriched)
    assert a.sheetnames == b.sheetnames
    for name in ('12_提醒', '02_质量评分', '04_估值跟踪', '05_仓位管理'):
        assert list(a[name].values) == list(b[name].values)
    for name in ('21_决策验证',):
        assert [r[:25] for r in list(a[name].values)[2:]] == [r[:25] for r in list(b[name].values)[2:]]
    rows = [(i, r) for i, r in enumerate(b['18_指标证据'].iter_rows(values_only=True), 1)
            if len(r) > 3 and r[3] == 'research_current_maturity_total']
    assert len(rows) == 1 and rows[0][1][5] is None
    decision_row = next(i for i, r in enumerate(b['21_决策验证'].iter_rows(values_only=True), 1) if r[0] == '600001')
    assert b['21_决策验证'].cell(decision_row, 26).hyperlink.location == f"'18_指标证据'!A{rows[0][0]}"
    assert b['21_决策验证'].cell(decision_row, 26).value == '完整有息负债未验证'
def test_lease_evidence_retains_raw_value_and_provenance_without_score_change(tmp_path):
    from value_investment_agent.excel_report import GENERAL
    path, payload = fixture(tmp_path)
    payload['annual_points'] = [{'symbol': '600001', 'field_name': 'lease_liabilities_noncurrent',
        'period_label': '2025-12-31', 'value': '2703584.10', 'unit': 'CNY',
        'validation_status': 'verified', 'source_id': 'lease-official',
        'source_url': 'https://example.test/official.pdf', 'sha256': 'a' * 64,
        'metadata': {'automatic_cross_source_verification': True, 'page_number': 85}}]
    output = tmp_path / 'lease.xlsx'
    build_report(path, payload, output)
    workbook = load_workbook(output)
    rows = list(workbook['18_指标证据'].iter_rows(min_row=4, values_only=True))
    lease = next(row for row in rows if row[3] == 'lease_liabilities_noncurrent')
    assert lease[2] == '非流动租赁负债'
    assert lease[4:7] == ('2025-12-31', 2703584.1, 'CNY')
    assert lease[8:12] == ('lease-official', 'https://example.test/official.pdf', 'a' * 64, 85)
    assert 'lease_liabilities_noncurrent' not in GENERAL
