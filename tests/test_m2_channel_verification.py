from datetime import date, datetime, timezone
from pathlib import Path

import pytest

from value_investment_agent.m2_channel_verification import (
    ACTION_NO_ORDER,
    CANDIDATE_CLASS_LEAD,
    CANDIDATE_CLASS_VERIFIED,
    CHANNELS,
    STATUS_INSUFFICIENT,
    STATUS_REJECTED,
    STATUS_VERIFIED,
    ChannelVerificationEvidence,
    ChannelVerificationResult,
    M2ChannelVerificationBatch,
    M2ChannelVerificationPolicy,
    build_m2_channel_verification,
)


ROOT = Path(__file__).resolve().parents[1]


def _policy() -> M2ChannelVerificationPolicy:
    return M2ChannelVerificationPolicy(
        schema_version="m2-channel-verification-policy-v1",
        policy_version="20260924-v1",
        as_of=date(2026, 9, 23),
        ac8_report_path="runtime/m2-ac8-research-reports-20260924-v1/report.json",
        ac8_report_sha256=(
            "dd55c02c75dec17ff766fa6b6ae529030d02a31376b305c02ddd0a8f5443c8a6"
        ),
        ac9_audit_report_path="runtime/m2-ac9-coverage-audit-20260923-v2/report.json",
        ac9_audit_report_sha256=(
            "3c30936a6e567bd55b8d03bf67163071c49d223ca10def66b93fcdc336a84695"
        ),
        m2_receipt_path="runtime/m2-live-20260923-v3/receipt.json",
        m2_receipt_sha256=(
            "869044a72c514be2d274308383c4479f7536bb393bfbf5ca10e492eee24bc220"
        ),
        minimum_second_stage_resolutions=3,
        minimum_channels=2,
    )


def _evidence() -> ChannelVerificationEvidence:
    return ChannelVerificationEvidence(
        id="600011-evidence-1",
        field_name="operating_cash_flow",
        period_label="2026-06-30",
        value="1000000",
        unit="CNY",
        validation_status="verified",
        source_name="CNINFO statutory disclosure",
        source_url="https://static.cninfo.com.cn/finalpage/example.PDF",
        source_sha256="a" * 64,
        published_at="2026-08-27T00:00:00+08:00",
        fetched_at="2026-09-09T00:00:00+00:00",
    )


def test_result_requires_evidence_before_verified_status():
    with pytest.raises(ValueError, match="source-bound evidence"):
        ChannelVerificationResult(
            verification_id="v1",
            report_id="report",
            symbol="600011",
            name="测试公司",
            channel="dividend_cash_return",
            source_verdict="PENDING_DEEP_RESEARCH",
            status=STATUS_VERIFIED,
            reason="test",
            evidence=(),
            missing_evidence=(),
            positives=(),
            counter_evidence=(),
            market_context={},
            rule_version="v1",
            source_report_sha256="a" * 64,
        )


def test_verified_result_changes_candidate_class_only():
    verified = ChannelVerificationResult(
        verification_id="v1",
        report_id="report",
        symbol="600011",
        name="测试公司",
        channel="dividend_cash_return",
        source_verdict="PENDING_DEEP_RESEARCH",
        status=STATUS_VERIFIED,
        reason="test",
        evidence=(_evidence(),),
        missing_evidence=(),
        positives=(),
        counter_evidence=(),
        market_context={},
        rule_version="v1",
        source_report_sha256="a" * 64,
    )
    rejected = ChannelVerificationResult(
        verification_id="v2",
        report_id="report",
        symbol="600011",
        name="测试公司",
        channel="value",
        source_verdict="REJECTED_FOR_CHANNEL",
        status=STATUS_REJECTED,
        reason="test",
        evidence=(_evidence(),),
        missing_evidence=(),
        positives=(),
        counter_evidence=(),
        market_context={},
        rule_version="v1",
        source_report_sha256="a" * 64,
    )

    assert verified.candidate_class == CANDIDATE_CLASS_VERIFIED
    assert rejected.candidate_class == CANDIDATE_CLASS_LEAD
    assert verified.action == rejected.action == ACTION_NO_ORDER


def test_real_pre_registered_ac8_build_is_a_formal_second_stage_packet():
    policy = _policy()
    batch = build_m2_channel_verification(
        policy,
        root=ROOT,
        generated_at=datetime(2026, 9, 24, 0, 0, tzinfo=timezone.utc),
    )

    counts = batch.counts()
    assert batch.machine_status == "MACHINE_CHECKS_PASS"
    assert batch.acceptance_status == "CHECKPOINT_A_READY_FOR_HUMAN_RESUBMISSION"
    assert counts[STATUS_VERIFIED] == 3
    assert counts[STATUS_REJECTED] == 13
    assert counts[STATUS_INSUFFICIENT] == 2
    assert batch.resolution_count() == 18
    assert len({item.channel for item in batch.results}) >= 2
    assert all(item.action == ACTION_NO_ORDER for item in batch.results)
    assert all(
        item.channel in CHANNELS
        for item in batch.results
    )


def test_real_batch_human_packet_contains_verified_symbols_and_no_order():
    policy = _policy()
    batch = build_m2_channel_verification(
        policy,
        root=ROOT,
        generated_at=datetime(2026, 9, 24, 0, 0, tzinfo=timezone.utc),
    )
    packet = batch.human_review_packet()

    assert packet["review_required"] is True
    assert packet["summary"]["lead_count"] == 18
    assert packet["summary"]["verified_for_deep_research"] == 3
    assert packet["summary"]["verified_symbols"] == ["002327", "002867", "600011"]
    assert packet["action"] == ACTION_NO_ORDER


def test_pinned_report_hash_change_fails_closed():
    policy = _policy()
    changed = M2ChannelVerificationPolicy(
        schema_version=policy.schema_version,
        policy_version=policy.policy_version,
        as_of=policy.as_of,
        ac8_report_path=policy.ac8_report_path,
        ac8_report_sha256="b" * 64,
        ac9_audit_report_path=policy.ac9_audit_report_path,
        ac9_audit_report_sha256=policy.ac9_audit_report_sha256,
        m2_receipt_path=policy.m2_receipt_path,
        m2_receipt_sha256=policy.m2_receipt_sha256,
        minimum_second_stage_resolutions=policy.minimum_second_stage_resolutions,
        minimum_channels=policy.minimum_channels,
    )

    with pytest.raises(ValueError, match="SHA-256 mismatch"):
        build_m2_channel_verification(
            changed,
            root=ROOT,
            generated_at=datetime(2026, 9, 24, 0, 0, tzinfo=timezone.utc),
        )


def test_human_review_workbook_is_presentation_only(tmp_path):
    from openpyxl import load_workbook
    from value_investment_agent.m2_channel_verification_workbook import (
        OVERVIEW_SHEET,
        RESOLUTION_SHEET,
        build_channel_verification_workbook,
    )

    batch = build_m2_channel_verification(
        _policy(),
        root=ROOT,
        generated_at=datetime(2026, 9, 24, 0, 0, tzinfo=timezone.utc),
    )
    workbook = build_channel_verification_workbook(batch)
    output = tmp_path / "m2-channel-verification.xlsx"
    workbook.save(output)

    loaded = load_workbook(output, read_only=True, data_only=False)
    assert OVERVIEW_SHEET in loaded.sheetnames
    assert RESOLUTION_SHEET in loaded.sheetnames
    overview = loaded[OVERVIEW_SHEET]
    assert any(
        "no_order" in str(cell.value)
        for row in overview.iter_rows()
        for cell in row
    )
