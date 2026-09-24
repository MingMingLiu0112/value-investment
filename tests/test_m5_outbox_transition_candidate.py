from __future__ import annotations

from copy import deepcopy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from scripts.build_m5_event_infrastructure_candidate import (
    build_models,
    load_input,
)
from scripts.build_m5_outbox_transition_candidate import (
    apply_transition_operations,
    build_candidate,
    load_transition_input,
)
from value_investment_agent.m5_event_core import EVENT_IDENTITY_LEGACY
from value_investment_agent.m5_event_outbox import (
    ALERT_ACKNOWLEDGED,
    ALERT_FAILED_TERMINAL,
    ALERT_PENDING,
    ALERT_SENT,
)
from value_investment_agent.m5_event_run_state import M5EventRunState
from value_investment_agent.m5_event_workbook import build_m5_event_workbook


ROOT = Path(__file__).resolve().parents[1]
BASE_FIXTURE = ROOT / "tests" / "fixtures" / "m5_event_infrastructure_demo.json"
TRANSITION_FIXTURE = (
    ROOT / "tests" / "fixtures" / "m5_outbox_transition_demo.json"
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sheet_rows(workbook, sheet_name: str) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[4]]
    return [
        dict(zip(headers, (cell.value for cell in row), strict=True))
        for row in sheet.iter_rows(min_row=5)
        if any(cell.value is not None for cell in row)
    ]


def _build(tmp_path: Path):
    output = tmp_path / "m5-outbox-transition.xlsx"
    manifest, state = build_candidate(
        base_input_path=BASE_FIXTURE,
        transition_input_path=TRANSITION_FIXTURE,
        output=output,
    )
    return output, manifest, state


def _expected_alert_ids() -> dict[tuple[str, str], str]:
    payload, _ = load_input(BASE_FIXTURE)
    receipt, _ = build_models(payload)
    expected: dict[tuple[str, str], str] = {}
    for source_event_id, alert_type in (
        ("600519-financial-report-001", "REVIEW_DUE"),
        ("000333-dividend-001", "DIVIDEND_ALERT"),
        ("600519-price-zone-001", "REVIEW_DUE"),
    ):
        event = next(
            item
            for item in receipt.state.event_ledger.active_events()
            if item.source_event_id == source_event_id
        )
        alert = next(
            item
            for item in receipt.state.outbox.alerts()
            if item.event_id == event.event_id and item.alert_type == alert_type
        )
        expected[(source_event_id, alert_type)] = alert.alert_id
    return expected


def test_candidate_is_v2_identity_workbook_with_explicit_transitions(tmp_path):
    output, manifest, state = _build(tmp_path)
    workbook = load_workbook(output, data_only=True)

    assert workbook.sheetnames == [
        "00_总览",
        "01_事件账",
        "02_水位与检查点",
        "03_依赖失效与重算",
        "04_Outbox",
        "05_Outbox迁移",
        "06_输入与边界",
    ]
    overview = workbook["00_总览"]
    assert overview["B5"].value == "SIMULATED"
    assert overview["B9"].value == 5
    assert overview["B17"].value == 7
    assert overview["B18"].value == 7
    assert overview["B19"].value == "no_order"

    event_sheet = workbook["01_事件账"]
    assert [cell.value for cell in event_sheet[4]][:7] == [
        "序号",
        "状态",
        "事件ID",
        "身份版本",
        "来源ID",
        "证券",
        "事件类型",
    ]
    event_rows = _sheet_rows(workbook, "01_事件账")
    assert len(event_rows) == 7
    assert {row["身份版本"] for row in event_rows} == {
        "历史兼容（无来源ID）"
    }
    assert {row["来源ID"] for row in event_rows} == {"unspecified-source"}
    assert {
        row["事件ID"]: row["状态"]
        for row in event_rows
        if row["状态"] == "已被替代"
    } == {"m5-1476e5fd3e68c2ca8706505cd8591d99": "已被替代"}

    outbox_sheet = workbook["04_Outbox"]
    assert outbox_sheet["K4"].value == "发送时间"
    outbox_rows = _sheet_rows(workbook, "04_Outbox")
    status_by_alert = {
        str(row["提醒ID"]): str(row["状态"])
        for row in outbox_rows
    }
    expected_alert_ids = _expected_alert_ids()
    assert {
        key: status_by_alert[alert_id]
        for key, alert_id in expected_alert_ids.items()
    } == {
        ("600519-financial-report-001", "REVIEW_DUE"): ALERT_ACKNOWLEDGED,
        ("000333-dividend-001", "DIVIDEND_ALERT"): ALERT_SENT,
        ("600519-price-zone-001", "REVIEW_DUE"): ALERT_FAILED_TERMINAL,
    }

    transition_rows = _sheet_rows(workbook, "05_Outbox迁移")
    assert len(transition_rows) == 7
    assert [row["修订"] for row in transition_rows] == list(range(1, 8))
    assert {row["动作"] for row in transition_rows} == {"no_order"}
    assert [cell.value for cell in workbook["05_Outbox迁移"][4]][:6] == [
        "修订",
        "迁移ID",
        "提醒ID",
        "来源事件ID",
        "来源ID",
        "身份版本",
    ]
    assert {row["来源ID"] for row in transition_rows} == {
        "unspecified-source"
    }
    assert {row["身份版本"] for row in transition_rows} == {
        "历史兼容（无来源ID）"
    }
    assert workbook["06_输入与边界"].max_row == 17

    assert state.outbox_revision == 7
    assert len(state.outbox_transitions) == 7
    assert manifest["active_event_count"] == 5
    assert manifest["state_sha256"] == state.state_sha256()


def test_manifest_binds_inputs_workbook_and_final_statuses(tmp_path):
    output, manifest, _ = _build(tmp_path)
    manifest_path = output.with_suffix(".manifest.json")

    assert manifest["base_input"]["sha256"] == _sha256(BASE_FIXTURE)
    assert manifest["transition_input"]["sha256"] == _sha256(TRANSITION_FIXTURE)
    assert manifest["workbook_sha256"] == _sha256(output)
    assert manifest["schema_version"] == "m5-outbox-transition-candidate-v2"
    assert manifest["generated_at"] == "2026-09-24T16:10:00+08:00"
    assert manifest["producer"]["script"]["relative_path"] == (
        "scripts/build_m5_outbox_transition_candidate.py"
    )
    assert manifest["producer"]["workbook_module"]["relative_path"] == (
        "src/value_investment_agent/m5_event_workbook.py"
    )
    assert manifest["producer"]["parent_v1_workbook"]["sha256"] == (
        "2b86953f793df46e199c40c614f3291e19b249cc0e53e2a670b0000506403dae"
    )
    assert manifest["base_state_sha256"] != manifest["state_sha256"]
    assert manifest["event_identity_counts"] == {
        EVENT_IDENTITY_LEGACY: 6
    }
    assert json.loads(manifest_path.read_text(encoding="utf-8")) == manifest
    assert len(manifest["final_outbox_statuses"]) == 6
    expected_alert_ids = _expected_alert_ids()
    statuses_by_alert = {
        item["alert_id"]: item["status"]
        for item in manifest["final_outbox_statuses"]
    }
    assert {
        statuses_by_alert[expected_alert_ids[key]]
        for key in expected_alert_ids
    } == {ALERT_ACKNOWLEDGED, ALERT_SENT, ALERT_FAILED_TERMINAL}
    assert list(statuses_by_alert.values()).count(ALERT_PENDING) == 3
    assert manifest["requested_operation_count"] == 7
    assert manifest["idempotent_operation_count"] == 0


def test_duplicate_operation_is_idempotent_without_extra_history():
    payload, _ = load_transition_input(
        TRANSITION_FIXTURE,
        base_input_path=BASE_FIXTURE,
    )
    base_payload, _ = load_input(BASE_FIXTURE)
    receipt, _ = build_models(base_payload)
    duplicate = deepcopy(payload)
    duplicate["operations"].insert(1, deepcopy(duplicate["operations"][0]))

    state, operations = apply_transition_operations(
        receipt=receipt,
        payload=duplicate,
    )

    assert len(operations) == 8
    assert sum(item["idempotent"] for item in operations) == 1
    assert len(state.outbox_transitions) == 7
    assert state.outbox_revision == 7


def test_transition_input_fails_closed_on_action_hash_and_unknown_fields(tmp_path):
    original = json.loads(TRANSITION_FIXTURE.read_text(encoding="utf-8"))

    bad_action = deepcopy(original)
    bad_action["action"] = "BUY"
    bad_action_path = tmp_path / "bad-action.json"
    bad_action_path.write_text(
        json.dumps(bad_action, ensure_ascii=False),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="no_order"):
        load_transition_input(bad_action_path, base_input_path=BASE_FIXTURE)

    bad_hash = deepcopy(original)
    bad_hash["base_input_sha256"] = "0" * 64
    bad_hash_path = tmp_path / "bad-hash.json"
    bad_hash_path.write_text(
        json.dumps(bad_hash, ensure_ascii=False),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="hash"):
        load_transition_input(bad_hash_path, base_input_path=BASE_FIXTURE)

    unknown = deepcopy(original)
    unknown["operations"][0]["trade_instruction"] = "BUY"
    unknown_path = tmp_path / "unknown-field.json"
    unknown_path.write_text(
        json.dumps(unknown, ensure_ascii=False),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="unknown fields"):
        load_transition_input(unknown_path, base_input_path=BASE_FIXTURE)


def test_transition_operations_reject_future_time_and_ambiguous_version():
    payload, _ = load_transition_input(
        TRANSITION_FIXTURE,
        base_input_path=BASE_FIXTURE,
    )
    base_payload, _ = load_input(BASE_FIXTURE)
    receipt, _ = build_models(base_payload)

    future = deepcopy(payload)
    future["operations"][0]["occurred_at"] = "2030-01-01T00:00:00+08:00"
    with pytest.raises(ValueError, match="later than"):
        apply_transition_operations(receipt=receipt, payload=future)

    ambiguous = deepcopy(payload)
    ambiguous["operations"] = [
        {
            "source_event_id": "601088-event-v1",
            "alert_type": "REVIEW_DUE",
            "from_status": "PENDING",
            "to_status": "SENT",
            "occurred_at": "2026-09-24T16:06:00+08:00",
            "error": None,
        }
    ]
    with pytest.raises(ValueError, match="exactly one alert"):
        apply_transition_operations(receipt=receipt, payload=ambiguous)


def test_current_state_must_share_receipt_payload_and_require_history_display():
    payload, _ = load_input(BASE_FIXTURE)
    receipt, watermark = build_models(payload)
    transition_payload, _ = load_transition_input(
        TRANSITION_FIXTURE,
        base_input_path=BASE_FIXTURE,
    )
    final_state, _ = apply_transition_operations(
        receipt=receipt,
        payload=transition_payload,
    )

    with pytest.raises(ValueError, match="transition history display"):
        build_m5_event_workbook(
            receipt,
            watermark,
            current_state=final_state,
            include_outbox_transitions=False,
        )

    other_state = M5EventRunState.empty(state_key="other-state")
    with pytest.raises(ValueError, match="non-outbox payload"):
        build_m5_event_workbook(
            receipt,
            watermark,
            current_state=other_state,
            include_outbox_transitions=True,
        )


def test_candidate_and_manifest_refuse_existing_outputs(tmp_path):
    output, manifest, _ = _build(tmp_path)

    with pytest.raises(ValueError, match="output already exists"):
        build_candidate(
            base_input_path=BASE_FIXTURE,
            transition_input_path=TRANSITION_FIXTURE,
            output=output,
        )

    alternate_output = tmp_path / "alternate.xlsx"
    existing_manifest = alternate_output.with_suffix(".manifest.json")
    existing_manifest.write_text("{}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="manifest already exists"):
        build_candidate(
            base_input_path=BASE_FIXTURE,
            transition_input_path=TRANSITION_FIXTURE,
            output=alternate_output,
        )
    assert manifest["action"] == "no_order"


def test_candidate_contains_no_trade_or_position_instruction(tmp_path):
    output, manifest, _ = _build(tmp_path)
    workbook = load_workbook(output, data_only=True)
    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    manifest_text = json.dumps(manifest, ensure_ascii=False)

    for forbidden in ("自动交易", "下单", "目标仓位", "position_size", "target_weight"):
        assert forbidden not in text
        assert forbidden not in manifest_text
    assert manifest["action"] == "no_order"
    assert workbook["06_输入与边界"]["B17"].value == "no_order"
