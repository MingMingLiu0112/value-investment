from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl import load_workbook
import pytest


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "stage_frontend_package",
    ROOT / "scripts" / "stage_frontend_package.py",
)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _source(path: Path) -> None:
    workbook = Workbook()
    retained = workbook.active
    retained.title = "保留页"
    retained["A1"] = "用户手工内容"
    retained["B1"] = "=1+2"
    review = workbook.create_sheet("00_决策复核")
    review["A1"] = "旧占位"
    review["A2"] = "旧状态"
    workbook.save(path)


def _addon(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "00_决策复核"
    sheet["A1"] = "M3 决策复核"
    sheet["A2"] = "研究证据不足"
    sheet["A3"] = "action=no_order"
    sheet.freeze_panes = "A4"
    workbook.save(path)


def test_replace_sheet_preserves_other_parts_and_sheet_order(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    addon = tmp_path / "addon.xlsx"
    output = tmp_path / "candidate.xlsx"
    _source(source)
    _addon(addon)

    receipt = MODULE.replace_sheet(
        source,
        addon,
        output,
        _digest(source),
        "00_决策复核",
    )

    assert receipt["status"] == "candidate_verified_not_published"
    assert receipt["original_sheets_preserved"] == 1
    assert receipt["derived_sheets_replaced"] == 1
    assert receipt["candidate_sha256"] == _digest(output)

    workbook = load_workbook(output, read_only=False)
    assert workbook.sheetnames == ["保留页", "00_决策复核"]
    assert workbook["保留页"]["A1"].value == "用户手工内容"
    assert workbook["保留页"]["B1"].value == "=1+2"
    assert workbook["00_决策复核"]["A1"].value == "M3 决策复核"
    assert workbook["00_决策复核"]["A2"].value == "研究证据不足"

    with ZipFile(source) as original, ZipFile(output) as candidate:
        assert candidate.read("xl/worksheets/sheet1.xml") == original.read(
            "xl/worksheets/sheet1.xml"
        )
    MODULE.validate_package_relationships(output)


def test_replace_sheet_refuses_source_change_after_inspection(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    addon = tmp_path / "addon.xlsx"
    output = tmp_path / "candidate.xlsx"
    _source(source)
    _addon(addon)
    expected = _digest(source)
    source.write_bytes(source.read_bytes() + b"changed")

    with pytest.raises(ValueError, match="changed since inspection"):
        MODULE.replace_sheet(
            source,
            addon,
            output,
            expected,
            "00_决策复核",
        )
