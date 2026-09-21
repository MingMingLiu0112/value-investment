import importlib.util
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "research_status_workbook", ROOT / "scripts" / "build_research_status_workbook_candidate.py")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def test_candidate_updates_only_the_intended_research_status_cells(tmp_path):
    source, output = tmp_path / "source.xlsx", tmp_path / "candidate.xlsx"
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "00_首页Dashboard"
    dashboard["C11"] = "old risk"
    pending = workbook.create_sheet("00_待完成公司")
    pending.append(["代码", "公司"])
    pending.append(["000651", "格力电器"])
    research = workbook.create_sheet("09_公司研究")
    research.append(["代码", "公司"])
    research.append(["000651", "格力电器"])
    workbook.save(source)

    result = MODULE.build(source, output)
    candidate = load_workbook(output)
    assert candidate["00_待完成公司"]["D2"].value == "R0一手事实已核验"
    assert candidate["09_公司研究"]["M2"].value.startswith("R0一手事实")
    assert candidate["09_公司研究"]["Q2"].hyperlink.target.startswith("https://static.cninfo.com.cn/")
    assert "next_trade_session_unavailable" not in candidate["00_首页Dashboard"]["C11"].value
    assert result["gree_card"]["sha256"] == MODULE.digest(
        ROOT / "runtime/company-research/000651-r0-research-card-20260917T080936Z/evidence.json")
