from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook
import pytest

from value_investment_agent.m3_decision_application import (
    build_nonpersonal_decision_card_collection,
)
from value_investment_agent.m3_decision_review_sheet import (
    TARGET_SHEET,
    build_m3_decision_review_sheet,
    write_m3_decision_review_addon,
)


GENERATED_AT = datetime(2026, 9, 23, 11, 42, 28, tzinfo=timezone.utc)
NAMES = {
    "000651": "格力电器",
    "600741": "华域汽车",
    "600887": "伊利股份",
}


def _run(symbol: str) -> dict:
    return {
        "run_id": f"{symbol}-m1-run",
        "symbol": symbol,
        "status": "COMPLETED_WITH_BLOCKERS",
        "action": "no_order",
        "pre_decision_eligibility": {
            "schema_version": "post-m1-predecision-eligibility-v1",
            "symbol": symbol,
            "decision_as_of": "2026-09-22",
            "status": "NOT_ELIGIBLE",
            "approval_status": "REJECTED_NEEDS_REWORK",
            "model_validity_status": "STALE",
            "price_bridge_status": "STALE_MODEL",
            "event_review_watermark": "2026-09-22",
            "blockers": ["research_gate_not_ready"],
            "evidence_refs": [
                {
                    "id": f"{symbol}-filing",
                    "kind": "official_issuer_filing",
                    "source_url": f"https://example.com/{symbol}.pdf",
                    "sha256": "a" * 64,
                }
            ],
            "action": "no_order",
        },
        "gate": {"conclusion": "估值未就绪", "blockers": ["research_gate_not_ready"]},
        "human_approval": {
            "symbol": symbol,
            "status": "REJECTED_NEEDS_REWORK",
            "blockers": ["research_gate_not_ready"],
            "action": "no_order",
        },
        "valuation": {
            "symbol": symbol,
            "model_type": "fixture",
            "valuation_date": "2026-09-22",
            "confidence": "低",
            "status": "conditional_research_only",
        },
        "model_validity": {"symbol": symbol, "status": "STALE", "blockers": []},
        "price_bridge": {
            "symbol": symbol,
            "bridge_status": "STALE_MODEL",
            "blockers": [],
        },
        "price_attractiveness": {
            "symbol": symbol,
            "status": "NOT_ASSESSABLE",
            "blockers": [],
        },
        "current_research_status": {
            "symbol": symbol,
            "research_conclusion": "估值未就绪",
            "blockers": [],
        },
    }


def _collection():
    return build_nonpersonal_decision_card_collection(
        [_run("000651"), _run("600741"), _run("600887")],
        generated_at=GENERATED_AT,
        source_run_id="fixture-m1-run",
    )


def _all_text(workbook: Workbook) -> str:
    return "\n".join(
        str(cell.value)
        for sheet in workbook
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


def test_review_sheet_is_negative_only_and_contains_no_formulas():
    workbook = build_m3_decision_review_sheet(
        _collection(),
        security_names=NAMES,
    )

    assert workbook.sheetnames == [TARGET_SHEET]
    sheet = workbook[TARGET_SHEET]
    assert sheet["A1"].value == "决策复核（M3 非个人化负向卡）"
    assert sheet["A6"].value == "000651"
    assert sheet["B6"].value == "格力电器"
    assert sheet["D6"].value == "研究证据不足"
    assert sheet["F6"].value == "未提交"
    assert sheet["G6"].value == "缺失"
    assert sheet["H6"].value == "本卡不需要"
    assert sheet["I6"].value == "无"
    assert sheet["A8"].value == "600887"

    text = _all_text(workbook)
    assert "action=no_order" in text
    assert "全部要求人工复核" in text or "三张卡全部要求人工复核" in text
    assert "Checkpoint B：未完成" in text
    assert "研究与估值前置条件未完成" in text
    for forbidden in ("买入", "加仓", "减仓", "目标仓位", "BUY", "ADD"):
        assert forbidden not in text

    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type != "f"


def test_addon_writer_is_no_overwrite_and_hash_pinned(tmp_path: Path):
    output = tmp_path / "m3-decision-review-addon.xlsx"
    result = write_m3_decision_review_addon(
        _collection(),
        output=output,
        root=tmp_path,
        security_names=NAMES,
    )

    target = tmp_path / result["addon_path"]
    assert target == output
    assert result["target_sheet"] == TARGET_SHEET
    assert result["card_count"] == 3
    assert result["positive_review_count"] == 0
    assert result["action"] == "no_order"
    loaded = load_workbook(target, read_only=False)
    assert loaded.sheetnames == [TARGET_SHEET]

    with pytest.raises(ValueError, match="already exists"):
        write_m3_decision_review_addon(
            _collection(),
            output=output,
            root=tmp_path,
            security_names=NAMES,
        )
