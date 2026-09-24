from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import importlib.util
import json
from pathlib import Path

import pytest

from value_investment_agent.m2_discovery_engine import (
    M2ScreeningPolicy,
    build_discovery_receipt,
)
from value_investment_agent.m2_market_data import (
    parse_eastmoney_dividends,
    parse_sina_industry,
    parse_tencent_board,
)
from value_investment_agent.m2_opportunity_discovery import (
    ACTION_NO_ORDER,
    CANDIDATE_CLASS_LEAD,
    CHANNEL_CYCLICAL,
    CHANNEL_DIVIDEND,
    CHANNEL_QUALITY,
    CHANNEL_VALUE,
    DATA_COMPLETE,
    DATA_PARTIAL,
    EVALUATION_BUDGET_EXCLUDED,
    EVALUATION_NOT_EVALUATED,
    EvidenceReference,
    PROFILE_UNSUPPORTED,
    _coverage_signature_v1,
    coverage_signature,
    discovery_receipt_from_payload,
)


ROOT = Path(__file__).resolve().parents[1]
RUN_SPEC = importlib.util.spec_from_file_location(
    "run_m2_opportunity_discovery",
    ROOT / "scripts" / "run_m2_opportunity_discovery.py",
)
RUN_M2 = importlib.util.module_from_spec(RUN_SPEC)
RUN_SPEC.loader.exec_module(RUN_M2)


def _sha(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _reference(ref_id: str) -> EvidenceReference:
    return EvidenceReference(
        id=ref_id,
        path=f"runtime/{ref_id}.json",
        sha256=_sha(ref_id),
        source_name=ref_id,
        source_url=f"https://example.test/{ref_id}",
        fetched_at=datetime(2026, 9, 23, 8, 0, tzinfo=timezone.utc),
    )


def _official_payload() -> dict:
    return {
        "parser_version": "test",
        "complete": True,
        "scope": "沪深A股测试",
        "fetched_at": "2026-09-23T08:00:00+00:00",
        "records": [
            {"symbol": "600001", "name": "质量制造", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "制造业", "security_type": "A股"},
            {"symbol": "600002", "name": "测试银行", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "金融业", "security_type": "A股"},
            {"symbol": "600003", "name": "测试煤业", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "采矿业", "security_type": "A股"},
            {"symbol": "600004", "name": "测试高息", "board": "主板", "exchange": "SSE", "listed_on": "2000-01-01", "official_industry": "公用事业", "security_type": "A股"},
        ],
    }


def _tencent_payload() -> dict:
    return {
        "rows": [
            {"code": "sh600001", "name": "质量制造", "zxj": "20", "pe_ttm": "10", "pn": "2", "zsz": "200", "stock_type": "GP-A", "state": "", "zdf_y": "5"},
            {"code": "sh600002", "name": "测试银行", "zxj": "10", "pe_ttm": "6", "pn": "0.8", "zsz": "500", "stock_type": "GP-A", "state": "", "zdf_y": "1"},
            {"code": "sh600003", "name": "测试煤业", "zxj": "8", "pe_ttm": "8", "pn": "1", "zsz": "300", "stock_type": "GP-A", "state": "", "zdf_y": "-20"},
            {"code": "sh600004", "name": "测试高息", "zxj": "10", "pe_ttm": "12", "pn": "1.5", "zsz": "100", "stock_type": "GP-A", "state": "", "zdf_y": "3"},
        ],
        "total": 4,
    }


def _sina_payload() -> dict:
    return {
        "rows": [
            {"code": "600001", "name": "质量制造", "class": "机械设备", "trade": "20.01", "per": "10", "pb": "2", "mktcap": "2000000", "ticktime": "15:00:00"},
            {"code": "600002", "name": "测试银行", "class": "银行", "trade": "10.00", "per": "6", "pb": "0.8", "mktcap": "5000000", "ticktime": "15:00:00"},
            {"code": "600003", "name": "测试煤业", "class": "煤炭", "trade": "8.00", "per": "8", "pb": "1", "mktcap": "3000000", "ticktime": "15:00:00"},
            {"code": "600004", "name": "测试高息", "class": "公用事业", "trade": "10.00", "per": "12", "pb": "1.5", "mktcap": "1000000", "ticktime": "15:00:00"},
        ]
    }


def _dividend_payload() -> dict:
    rows = []
    for code, amount, yield_value, status in (
        ("600004", "1.00", "0.05", "实施中"),
        ("600002", "0.50", "0.04", "实施中"),
    ):
        row = [""] * 18
        row[0] = code
        row[1] = "测试高息" if code == "600004" else "测试银行"
        row[5] = amount
        row[6] = yield_value
        row[13] = "2026-04-01"
        row[15] = "2026-07-01"
        row[16] = status
        rows.append(row)
    return {"fiscal_year": "2025", "rows": rows}


def _financial_points() -> list[dict]:
    common = {
        "period_label": "2025-12-31",
        "validation_status": "verified",
        "human_reviewed": False,
        "metadata": {
            "automatic_cross_source_verification": True,
            "complete_debt_verified": True,
        },
        "created_at": "2026-09-20T07:00:00+00:00",
        "fetched_at": "2026-09-20T07:00:00+00:00",
        "source_id": "financial-fixture",
    }
    values = {
        "roe": ("20", "percent"),
        "gross_margin": ("40", "percent"),
        "net_margin": ("15", "percent"),
        "operating_cash_flow_to_net_income": ("120", "percent"),
        "cash": ("100", "CNY"),
        "interest_bearing_debt": ("50", "CNY"),
        "debt_ratio": ("30", "percent"),
        "revenue_yoy": ("10", "percent"),
        "net_income_yoy": ("10", "percent"),
    }
    return [
        {
            **common,
            "symbol": "600001",
            "field_name": field_name,
            "value": value,
            "unit": unit,
        }
        for field_name, (value, unit) in values.items()
    ]


def _receipt():
    run_refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    return build_discovery_receipt(
        run_id="m2-fixture-run",
        generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
        official_payload=_official_payload(),
        tencent_payload=_tencent_payload(),
        sina_payload=_sina_payload(),
        dividend_payload=_dividend_payload(),
        financial_points=_financial_points(),
        quote_date="2026-09-23",
        run_refs=run_refs,
        policy=M2ScreeningPolicy(max_per_channel=10),
    )


def test_receipt_round_trip_preserves_no_order_and_candidate_signature():
    receipt = _receipt()
    restored = discovery_receipt_from_payload(receipt.as_policy())

    assert restored == receipt
    assert restored.action == ACTION_NO_ORDER
    assert restored.coverage_signature == receipt.coverage_signature
    assert restored.candidate_signature == receipt.candidate_signature


def test_legacy_coverage_receipt_without_extended_fields_still_decodes():
    receipt = _receipt()
    payload = receipt.as_policy()
    for channel_payload in payload["channel_results"].values():
        for evaluation in channel_payload["evaluations"]:
            for field in ("trigger_metrics", "trigger_reasons", "policy_version", "raw_rank_before_budget"):
                evaluation.pop(field, None)
    payload["coverage_signature"] = _coverage_signature_v1(receipt.channel_results)

    restored = discovery_receipt_from_payload(payload)

    assert restored.coverage_signature == payload["coverage_signature"]
    assert restored.candidate_signature == receipt.candidate_signature
    assert all(
        evaluation.trigger_metrics == {} and not evaluation.trigger_reasons
        for channel_result in restored.channel_results.values()
        for evaluation in channel_result.evaluations
    )


def test_legacy_coverage_signature_rejects_unhashed_extended_enrichment():
    receipt = _receipt()
    payload = receipt.as_policy()
    for channel_payload in payload["channel_results"].values():
        for evaluation in channel_payload["evaluations"]:
            for field in ("trigger_metrics", "trigger_reasons", "policy_version", "raw_rank_before_budget"):
                evaluation.pop(field, None)
    payload["coverage_signature"] = _coverage_signature_v1(receipt.channel_results)
    next(iter(payload["channel_results"].values()))["evaluations"][0][
        "trigger_metrics"
    ] = {"pe_ttm": "1.0"}

    with pytest.raises(ValueError, match="coverage signature changed"):
        discovery_receipt_from_payload(payload)


def test_bank_is_isolated_and_never_enters_general_channels():
    receipt = _receipt()
    all_candidates = receipt.candidate_pool()

    assert "600002" not in all_candidates
    dividend_result = receipt.channel_results[CHANNEL_DIVIDEND]
    assert any(item.symbol == "600002" for item in dividend_result.excluded)
    assert any(
        item.symbol == "600002" and item.profile_status == PROFILE_UNSUPPORTED
        for item in dividend_result.excluded
    )
    assert receipt.data_health.unsupported_financial_count >= 1


def test_missing_analytics_remain_null_not_zero():
    receipt = _receipt()
    value_candidates = receipt.channel_results[CHANNEL_VALUE].candidates
    quality_candidates = receipt.channel_results[CHANNEL_QUALITY].candidates

    assert any(candidate.metrics.get("fcf_yield") is None for candidate in value_candidates)
    assert any(candidate.metrics.get("ev_ebit") is None for candidate in value_candidates)
    assert any(candidate.metrics.get("normalized_earnings") is None for candidate in receipt.channel_results["cyclical"].candidates)
    assert any(candidate.symbol == "600001" for candidate in quality_candidates)


def test_future_dividend_declaration_cannot_enter_an_earlier_snapshot():
    dividend_payload = _dividend_payload()
    future_row = [""] * 18
    future_row[0] = "600001"
    future_row[1] = "质量制造"
    future_row[5] = "2.00"
    future_row[6] = "0.08"
    future_row[13] = "2026-10-01"
    future_row[15] = "2026-11-01"
    future_row[16] = "预案"
    dividend_payload["rows"].append(future_row)
    run_refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    receipt = build_discovery_receipt(
        run_id="m2-future-dividend-fixture",
        generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
        official_payload=_official_payload(),
        tencent_payload=_tencent_payload(),
        sina_payload=_sina_payload(),
        dividend_payload=dividend_payload,
        financial_points=_financial_points(),
        quote_date="2026-09-23",
        run_refs=run_refs,
        policy=M2ScreeningPolicy(max_per_channel=10),
    )

    dividend_result = receipt.channel_results[CHANNEL_DIVIDEND]
    assert not any(candidate.symbol == "600001" for candidate in dividend_result.candidates)
    assert next(item for item in dividend_result.evaluations if item.symbol == "600001").status == EVALUATION_NOT_EVALUATED


def test_source_external_quote_is_isolated_from_formal_pool_and_legacy_shadow():
    tencent_payload = _tencent_payload()
    tencent_payload["rows"].append({
        "code": "sh600099",
        "name": "非官方证券",
        "zxj": "5",
        "pe_ttm": "5",
        "pn": "1",
        "zsz": "100",
        "stock_type": "GP-A",
        "state": "",
        "zdf_y": "1",
    })
    sina_payload = _sina_payload()
    sina_payload["rows"].append({
        "code": "600099",
        "name": "非官方证券",
        "class": "机械设备",
        "trade": "5.00",
        "per": "5",
        "pb": "1",
        "mktcap": "1000000",
        "ticktime": "15:00:00",
    })
    run_refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    receipt = build_discovery_receipt(
        run_id="m2-extra-quote-fixture",
        generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
        official_payload=_official_payload(),
        tencent_payload=tencent_payload,
        sina_payload=sina_payload,
        dividend_payload=_dividend_payload(),
        financial_points=_financial_points(),
        quote_date="2026-09-23",
        run_refs=run_refs,
        policy=M2ScreeningPolicy(max_per_channel=10),
    )

    assert "600099" not in receipt.candidate_pool()
    assert "600099" not in receipt.legacy_comparison.legacy_candidates
    assert receipt.data_health.extra_quote_count == 1
    assert all(result.coverage_count == 4 for result in receipt.channel_results.values())


def test_cross_channel_candidate_reasons_are_preserved_in_pool():
    receipt = _receipt()
    reasons = receipt.candidate_pool().get("600001", ())

    assert len(reasons) == 2
    assert {item.channel for item in reasons} == {CHANNEL_QUALITY, CHANNEL_VALUE}
    assert all(reason.reasons for reason in reasons)


def test_screen_outputs_are_research_leads_until_a_later_verification_gate_passes():
    receipt = _receipt()
    all_reasons = [reason for reasons in receipt.candidate_pool().values() for reason in reasons]

    assert all(reason.candidate_class == CANDIDATE_CLASS_LEAD for reason in all_reasons)
    assert receipt.verified_candidate_pool() == {}


def test_budget_truncation_keeps_budget_excluded_accounting():
    run_refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    receipt = build_discovery_receipt(
        run_id="m2-budget-fixture",
        generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
        official_payload=_official_payload(),
        tencent_payload=_tencent_payload(),
        sina_payload=_sina_payload(),
        dividend_payload=_dividend_payload(),
        financial_points=_financial_points(),
        quote_date="2026-09-23",
        run_refs=run_refs,
        policy=M2ScreeningPolicy(max_per_channel=1),
    )
    value_result = receipt.channel_results[CHANNEL_VALUE]

    assert len(value_result.candidates) == 1
    assert value_result.pass_count == 1
    assert value_result.budget_excluded_count == 2
    assert value_result.coverage_count == 4


def test_empty_candidate_pool_is_a_legal_snapshot():
    from value_investment_agent.m2_discovery_engine import build_channel_results
    from value_investment_agent.m2_opportunity_discovery import candidate_signature

    policy = M2ScreeningPolicy()
    channels = build_channel_results(
        {}, {}, {}, run_refs={}, policy=policy, quote_date="2026-09-23"
    )

    assert all(not result.candidates for result in channels.values())
    assert candidate_signature(channels)
    assert coverage_signature(channels)


def test_adapter_parsers_preserve_units_and_missing_values():
    tencent = parse_tencent_board(_tencent_payload())
    sina = parse_sina_industry(_sina_payload())
    dividends = parse_eastmoney_dividends(_dividend_payload())

    assert tencent[0]["market_cap"] == Decimal("20000000000")
    assert tencent[0]["pe_ttm"] == Decimal("10")
    assert sina[0]["industry"] == "机械设备"
    assert sina[0]["market_cap"] == Decimal("20000000000")
    assert dividends[0]["cash_dps"] == Decimal("0.1")
    assert dividends[0]["fiscal_year"] == "2025"


def test_workbook_is_presentation_only_and_contains_channel_sheets(tmp_path):
    from openpyxl import load_workbook
    from value_investment_agent.m2_discovery_workbook import (
        COVERAGE_SHEET,
        HEALTH_SHEET,
        OVERVIEW_SHEET,
        POOL_SHEET,
        build_discovery_workbook,
    )

    receipt = _receipt()
    workbook = build_discovery_workbook(receipt, M2ScreeningPolicy(max_per_channel=10))
    output = tmp_path / "m2.xlsx"
    workbook.save(output)

    loaded = load_workbook(output, read_only=True, data_only=False)
    assert OVERVIEW_SHEET in loaded.sheetnames
    assert POOL_SHEET in loaded.sheetnames
    assert HEALTH_SHEET in loaded.sheetnames
    assert COVERAGE_SHEET in loaded.sheetnames
    overview = loaded[OVERVIEW_SHEET]
    assert any(
        "不生成估值、BUY、仓位或订单" in str(cell.value)
        for row in overview.iter_rows()
        for cell in row
    )
    assert any(
        "Quality Channel Coverage" in str(cell.value)
        for row in overview.iter_rows()
        for cell in row
    )
    assert any(
        "COVERAGE_LIMITED" in str(cell.value)
        for row in overview.iter_rows()
        for cell in row
    )
    assert any(
        "不能把 0 候选解释成市场没有高质量公司" in str(cell.value)
        for row in overview.iter_rows()
        for cell in row
    )


def test_cheap_screen_passes_remain_data_partial_until_deep_evidence_is_verified():
    receipt = _receipt()

    quality = receipt.channel_results[CHANNEL_QUALITY].candidates[0]
    assert quality.data_status == DATA_COMPLETE
    for channel in (CHANNEL_DIVIDEND, CHANNEL_VALUE, CHANNEL_CYCLICAL):
        candidates = receipt.channel_results[channel].candidates
        assert candidates
        assert all(candidate.data_status == DATA_PARTIAL for candidate in candidates)
        assert all(candidate.candidate_class == CANDIDATE_CLASS_LEAD for candidate in candidates)


def test_receipt_rejects_quote_date_newer_than_as_of():
    run_refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    with pytest.raises(ValueError, match="quote_date"):
        build_discovery_receipt(
            run_id="m2-future-quote-fixture",
            generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
            official_payload=_official_payload(),
            tencent_payload=_tencent_payload(),
            sina_payload=_sina_payload(),
            dividend_payload=_dividend_payload(),
            financial_points=_financial_points(),
            quote_date="2026-09-24",
            run_refs=run_refs,
        )


def test_receipt_rejects_evidence_fetched_after_generated_at():
    future = datetime(2026, 9, 23, 10, 0, tzinfo=timezone.utc)
    run_refs = {key: _reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    run_refs["tencent"] = EvidenceReference(
        id="tencent",
        path="runtime/tencent.json",
        sha256=_sha("tencent"),
        source_name="tencent",
        source_url="https://example.test/tencent",
        fetched_at=future,
    )

    with pytest.raises(ValueError, match="fetched_at"):
        build_discovery_receipt(
            run_id="m2-future-evidence-fixture",
            generated_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
            official_payload=_official_payload(),
            tencent_payload=_tencent_payload(),
            sina_payload=_sina_payload(),
            dividend_payload=_dividend_payload(),
            financial_points=_financial_points(),
            quote_date="2026-09-23",
            run_refs=run_refs,
        )


def test_future_financial_point_is_excluded_from_snapshot_evidence():
    from value_investment_agent.m2_discovery_engine import build_financial_evidence

    point = _financial_points()[0]
    point["fetched_at"] = "2026-09-23T09:00:00+00:00"
    evidence = build_financial_evidence(
        [point],
        {"600001": "质量制造"},
        {"600001": "机械设备"},
        evaluation_at=datetime(2026, 9, 23, 8, 5, tzinfo=timezone.utc),
    )

    assert evidence == {}


def test_reuse_inputs_preserves_original_clock_instead_of_relabeling_today(tmp_path):
    project = tmp_path / "project"
    output_dir = project / "run"
    output_dir.mkdir(parents=True)
    generated_at = datetime(2026, 9, 22, 8, 5, tzinfo=timezone.utc)

    def clocked_reference(ref_id: str) -> EvidenceReference:
        return EvidenceReference(
            id=ref_id,
            path=f"run/{ref_id}.json",
            sha256=_sha(ref_id),
            source_name=ref_id,
            source_url=f"https://example.test/{ref_id}",
            fetched_at=datetime(2026, 9, 22, 7, 0, tzinfo=timezone.utc),
        )

    official_payload = _official_payload()
    official_payload["fetched_at"] = "2026-09-22T07:00:00+00:00"
    tencent_payload = _tencent_payload()
    tencent_payload["fetched_at"] = "2026-09-22T07:00:00+00:00"
    sina_payload = _sina_payload()
    sina_payload["fetched_at"] = "2026-09-22T07:00:00+00:00"
    dividend_payload = _dividend_payload()
    dividend_payload["fetched_at"] = "2026-09-22T07:00:00+00:00"
    financial_points = _financial_points()

    run_refs = {key: clocked_reference(key) for key in ("official", "tencent", "sina", "dividend", "financial")}
    original = build_discovery_receipt(
        run_id="m2-original-clock-fixture",
        generated_at=generated_at,
        official_payload=official_payload,
        tencent_payload=tencent_payload,
        sina_payload=sina_payload,
        dividend_payload=dividend_payload,
        financial_points=financial_points,
        quote_date="2026-09-22",
        run_refs=run_refs,
        policy=M2ScreeningPolicy(max_per_channel=10),
    )
    (output_dir / "receipt.json").write_text(
        json.dumps(original.as_policy(), ensure_ascii=False),
        encoding="utf-8",
    )
    (output_dir / "official-universe.json").write_text(json.dumps(official_payload, ensure_ascii=False), encoding="utf-8")
    (output_dir / "tencent-market.json").write_text(json.dumps(tencent_payload, ensure_ascii=False), encoding="utf-8")
    (output_dir / "sina-industry-quotes.json").write_text(json.dumps(sina_payload, ensure_ascii=False), encoding="utf-8")
    (output_dir / "eastmoney-dividends.json").write_text(json.dumps(dividend_payload, ensure_ascii=False), encoding="utf-8")
    (output_dir / "retained-financial-points.json").write_text(
        json.dumps(
            {"points": financial_points, "retained_generated_at": "2026-09-20T07:00:00+00:00"},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    summary = RUN_M2.run_once(
        root=project,
        output_dir=output_dir,
        wps_dir=None,
        max_per_channel=10,
        skip_dividend=False,
        reuse_inputs=True,
    )

    assert summary["generated_at"] == generated_at.isoformat()
    assert summary["as_of"] == "2026-09-22"
    assert summary["quote_date"] == "2026-09-22"
    assert summary["replay_of_run_id"] == original.run_id
    assert summary["replay"]["receipt_bytes_match"] is True
    assert summary["replay"]["raw_inputs_match"] is True
