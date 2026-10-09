"""Synthetic authored-page integration; no real workbook or approval is touched.

Openpyxl writes test fixtures only. Production integration runs the real OOXML
adapter, with byte checks and the existing preservation/source verifiers.
"""
from copy import deepcopy
from dataclasses import asdict, replace
import hashlib
import json
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
import pytest

from scripts import stage_frontend_package as stage
from scripts.current import prepare_research_handoff_excel as handoff_cli
from scripts.current import publish_product_workbench_to_canonical as publisher
from test_canonical_workbook_preservation import _workbook
from test_product_workbench_excel import _payload
from value_investment_agent.application.product.research_publication_input import (
    prepare_research_publication_input,
)
from value_investment_agent.presentation.excel.product_workbench import WORKBOOK_SHEETS
from value_investment_agent.presentation.read_models.product_workbench import (
    product_workbench_from_payload,
)


SYNTHETIC = "SYNTHETIC_AUTHORED_INTEGRATION_TEST_ONLY"


def _hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _rewrite_fixture_zip(path, replacements):
    with ZipFile(path) as archive:
        original = [(item, archive.read(item.filename)) for item in archive.infolist()]
    rewritten = path.with_name("rewritten-" + path.name)
    with ZipFile(rewritten, "w", ZIP_DEFLATED) as archive:
        for item, data in original:
            archive.writestr(item, replacements.get(item.filename, data))
        for name, data in replacements.items():
            if name not in {item.filename for item, _ in original}:
                archive.writestr(name, data)
    rewritten.replace(path)


def _inline_fixture_to_shared_strings(path):
    shared = ET.Element(f"{{{stage.M}}}sst")
    changes = {}
    with ZipFile(path) as archive:
        _, entries = stage.sheets(archive)
        for _, part in entries:
            sheet = stage.xml(archive.read(part))
            for cell in sheet.iter(f"{{{stage.M}}}c"):
                if cell.get("t") != "inlineStr":
                    continue
                text = cell.find(f"{{{stage.M}}}is")
                item = deepcopy(text)
                item.tag = f"{{{stage.M}}}si"
                cell.remove(text)
                cell.set("t", "s")
                ET.SubElement(cell, f"{{{stage.M}}}v").text = str(len(shared))
                shared.append(item)
            changes[part] = stage.serialized(sheet)
        relations = stage.xml(archive.read("xl/_rels/workbook.xml.rels"))
        ET.SubElement(relations, f"{{{stage.P}}}Relationship", {
            "Id": "syntheticSharedStrings", "Type": stage.R + "/sharedStrings",
            "Target": "sharedStrings.xml",
        })
        content_types = stage.xml(archive.read("[Content_Types].xml"))
        ET.SubElement(content_types, f"{{{stage.C}}}Override", {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        })
    shared.set("count", str(len(shared)))
    shared.set("uniqueCount", str(len(shared)))
    changes.update({"xl/sharedStrings.xml": stage.serialized(shared),
                    "xl/_rels/workbook.xml.rels": stage.serialized(relations),
                    "[Content_Types].xml": stage.serialized(content_types)})
    _rewrite_fixture_zip(path, changes)
    stage.validate_package_relationships(path)


@pytest.fixture
def authored_case(tmp_path):
    canonical = tmp_path / "synthetic-canonical.xlsx"
    _workbook(canonical)
    workbook = load_workbook(canonical)
    for name in WORKBOOK_SHEETS:
        workbook[name]["A1"] = SYNTHETIC + " OLD " + name
    workbook.worksheets[-1].sheet_state = "hidden"
    for index in range(54):
        sheet = workbook.create_sheet(f"synthetic_history_{index:02d}")
        sheet["A1"] = SYNTHETIC + " preserved historical user record"
        sheet["B2"] = index
        sheet["C2"] = "=B2+1"
        sheet.sheet_state = "veryHidden" if index == 0 else "hidden"
    workbook.save(canonical)
    workbook.close()

    runtime = tmp_path / "runtime"
    runtime.mkdir()
    payload = _payload()
    payload["audit"]["evidence"].append({
        **payload["audit"]["evidence"][0], "evidence_id": "synthetic-evidence-2",
        "title": SYNTHETIC + " second source",
    })
    for index, evidence in enumerate(payload["audit"]["evidence"]):
        source = runtime / f"synthetic-evidence-{index}.json"
        source_url = f"https://example.test/synthetic/source-{index}.pdf"
        source.write_text(json.dumps({"scope": SYNTHETIC, "evidence_id": evidence["evidence_id"],
                                      "source_url": source_url}), encoding="utf-8")
        evidence.update(path=source.relative_to(tmp_path).as_posix(), sha256=_hash(source),
                        source_url=source_url)
    model = product_workbench_from_payload(payload)
    authored = runtime / "synthetic-authored.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, name in enumerate(WORKBOOK_SHEETS):
        sheet = workbook.create_sheet(name)
        sheet["A1"] = SYNTHETIC + " NEW " + name
        sheet["B2"] = index + 1
        sheet["C2"] = f'=HYPERLINK("#\'{WORKBOOK_SHEETS[0]}\'!A1","synthetic navigation")'
        sheet["B2"].number_format = "0.000"
        sheet["A1"].font = Font(name="Calibri", bold=True, color="12704F")
        sheet["A1"].fill = PatternFill("solid", fgColor="EDF5EE")
        sheet.freeze_panes = "B4"
        sheet.column_dimensions["A"].width = 33
        sheet.row_dimensions[1].height = 30
        if name == "03_\u516c\u53f8":
            for row, company in enumerate(model.companies, start=5):
                sheet.cell(row, 1, company.symbol)
    workbook.save(authored)
    workbook.close()
    _inline_fixture_to_shared_strings(authored)
    output = runtime / "integration" / "canonical-integration-historical-preview.xlsx"
    return dict(root=tmp_path, canonical=canonical, authored=authored, model=model,
                output=output, source_hash=_hash(canonical))


def _integrate(case, digest=None):
    return publisher.build_protected_research_preview(
        case["root"], case["canonical"], case["output"], case["model"],
        authored_pages=case["authored"],
        authored_pages_sha256=_hash(case["authored"]) if digest is None else digest,
    )


def _author_receipt(case, handoff):
    runtime = case["root"] / "runtime"
    policy = runtime / "synthetic-display-policy.json"
    policy.write_text(json.dumps({"scope": SYNTHETIC, "action": "no_order"}), encoding="utf-8")
    display = runtime / "synthetic-display.json"
    display.write_text(json.dumps({
        "scope": SYNTHETIC, "input_sha256": _hash(handoff),
        "policy_binding": {"path": policy.relative_to(case["root"]).as_posix(), "sha256": _hash(policy)},
    }), encoding="utf-8")
    receipt = case["authored"].with_name(case["authored"].stem + "-receipt.json")
    receipt.write_text(json.dumps({
        "schema_version": "d2-product-engineering-preview-v1", "scope": SYNTHETIC,
        "input_sha256": _hash(handoff), "workbook_sha256": _hash(case["authored"]),
        "source_count": len(json.loads(handoff.read_text(encoding="utf-8"))["source_bindings"]),
        "canonical_written": False, "action": "no_order",
        "display_binding": {"path": display.relative_to(case["root"]).as_posix(), "sha256": _hash(display)},
    }), encoding="utf-8")


def _set_authored_fixture_cells(case, name, values):
    with ZipFile(case["authored"]) as archive:
        _, entries = stage.sheets(archive)
        part = next(part for sheet, part in entries if sheet.get("name") == name)
        sheet = stage.xml(archive.read(part))
    data = sheet.find(f"{{{stage.M}}}sheetData")
    dimension = sheet.find(f"{{{stage.M}}}dimension")
    _, _, max_column, max_row = range_boundaries(dimension.get("ref"))
    for coordinate, value in values.items():
        row_number, column = coordinate_to_tuple(coordinate)
        max_row, max_column = max(max_row, row_number), max(max_column, column)
        row = next((row for row in data if row.get("r") == str(row_number)), None)
        if row is None:
            row = ET.SubElement(data, f"{{{stage.M}}}row", {"r": str(row_number)})
        cell = next((cell for cell in row if cell.get("r") == coordinate), None)
        if cell is None:
            cell = ET.SubElement(row, f"{{{stage.M}}}c", {"r": coordinate})
        for child in list(cell):
            cell.remove(child)
        if value.startswith("="):
            cell.attrib.pop("t", None)
            ET.SubElement(cell, f"{{{stage.M}}}f").text = value[1:]
        else:
            cell.set("t", "inlineStr")
            ET.SubElement(ET.SubElement(cell, f"{{{stage.M}}}is"), f"{{{stage.M}}}t").text = value
        row[:] = sorted(row, key=lambda cell: coordinate_to_tuple(cell.get("r"))[1])
    data[:] = sorted(data, key=lambda row: int(row.get("r")))
    dimension.set("ref", f"A1:{get_column_letter(max_column)}{max_row}")
    _rewrite_fixture_zip(case["authored"], {part: stage.serialized(sheet)})


def test_seven_authored_pages_preserve_55_history_pages_and_all_other_parts(
    authored_case, monkeypatch,
):
    case = authored_case
    before = publisher._snapshot(case["canonical"])
    authored_hash = _hash(case["authored"])

    def forbid_save(*args, **kwargs):
        raise AssertionError("production authored integration must not call openpyxl save")

    monkeypatch.setattr(Workbook, "save", forbid_save)
    receipt = _integrate(case)
    after = publisher._snapshot(case["output"])
    publisher._assert_retained(before, after)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert _hash(case["authored"]) == authored_hash
    assert receipt["canonical_written"] is False
    assert receipt["status"] == "PREVIEW_PENDING_NATIVE_REVIEW"
    proof = json.loads((case["output"].parent / "canonical-preservation.json").read_text(encoding="utf-8"))
    assert proof["preservation"] == "PASS" and proof["preserved_sheet_count"] == 55
    assert proof["source_sha256"] == case["source_hash"]
    assert proof["authored_pages_sha256"] == authored_hash
    assert proof["strict_pit"] == "NOT_PROVEN"
    assert proof["current_price_bridge"] == "NOT_ADMITTED"
    assert proof["canonical_touched"] is False and proof["action"] == "no_order"
    with ZipFile(case["canonical"]) as source, ZipFile(case["output"]) as result:
        _, entries = stage.sheets(source)
        managed_parts = {part for sheet, part in entries if sheet.get("name") in WORKBOOK_SHEETS}
        assert set(source.namelist()) == set(result.namelist())
        for part in source.namelist():
            if part not in managed_parts | {"xl/styles.xml"}:
                assert result.read(part) == source.read(part), part
        for part in managed_parts:
            assert result.read(part) != source.read(part)
            sheet = stage.xml(result.read(part))
            pane = sheet.find("m:sheetViews/m:sheetView/m:pane", stage.NS)
            assert pane.get("topLeftCell") == "B4"
            assert pane.get("xSplit") == "1" and pane.get("ySplit") == "3"
            assert pane.get("activePane") == "bottomRight"
        original_styles = stage.xml(source.read("xl/styles.xml"))
        integrated_styles = stage.xml(result.read("xl/styles.xml"))
        for section in ("fonts", "fills", "borders", "cellStyleXfs", "cellXfs"):
            original_items = list(original_styles.find(f"{{{stage.M}}}{section}"))
            integrated_items = list(integrated_styles.find(f"{{{stage.M}}}{section}"))
            assert [stage.serialized(item) for item in integrated_items[:len(original_items)]] == [
                stage.serialized(item) for item in original_items
            ], section
        assert result.testzip() is None
    workbook = load_workbook(case["output"], data_only=False)
    try:
        assert len(workbook.sheetnames) == 62
        assert tuple(workbook.sheetnames[:7]) == WORKBOOK_SHEETS
        for index, name in enumerate(WORKBOOK_SHEETS):
            assert workbook[name]["A1"].value == SYNTHETIC + " NEW " + name
            assert workbook[name]["B2"].value == index + 1
            assert workbook[name]["B2"].number_format == "0.000"
            assert workbook[name]["C2"].value == (
                f'=HYPERLINK("#\'{WORKBOOK_SHEETS[0]}\'!A1","synthetic navigation")'
            )
            assert workbook[name].freeze_panes == "B4"
            assert workbook[name].page_setup.orientation == "landscape"
            assert str(workbook[name].page_setup.paperSize) == "8"
            assert workbook[name].page_setup.fitToWidth == 1
            assert workbook[name].page_setup.fitToHeight == 0
            assert workbook[name].sheet_properties.pageSetUpPr.fitToPage is True
        manual = workbook["人工持仓"]
        assert manual["A1"].value == "\u7528\u6237\u8f93\u5165"
        assert manual["B2"].value == "=1+1"
        assert manual["D4"].comment.author == "user"
        assert len(manual.data_validations.dataValidation) == 1
        assert manual.freeze_panes == "B3"
        assert workbook["synthetic_history_00"].sheet_state == "veryHidden"
    finally:
        workbook.close()
    stage.validate_package_relationships(case["output"])


@pytest.mark.parametrize("missing", [False, True])
def test_authored_hash_is_required_and_must_match(authored_case, missing):
    case = authored_case
    with pytest.raises(ValueError, match="authored workbook path or hash"):
        publisher.build_protected_research_preview(
            case["root"], case["canonical"], case["output"], case["model"],
            authored_pages=case["authored"], authored_pages_sha256=None if missing else "0" * 64,
        )
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "extra", "reordered"])
def test_authored_pages_must_match_the_exact_managed_navigation(authored_case, mutation):
    case = authored_case
    with ZipFile(case["authored"]) as archive:
        workbook = stage.xml(archive.read("xl/workbook.xml"))
    sheets = workbook.find(f"{{{stage.M}}}sheets")
    if mutation == "missing":
        sheets.remove(sheets[-1])
    elif mutation == "duplicate":
        sheets[1].set("name", sheets[0].get("name"))
    elif mutation == "extra":
        extra = deepcopy(sheets[-1])
        extra.set("name", "synthetic_unmanaged_extra")
        extra.set("sheetId", "999")
        sheets.append(extra)
    else:
        sheets[:] = [sheets[1], sheets[0], *list(sheets)[2:]]
    _rewrite_fixture_zip(case["authored"], {"xl/workbook.xml": stage.serialized(workbook)})
    with pytest.raises(ValueError, match="exactly the managed pages"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.mark.parametrize("index", [0, 6])
def test_external_worksheet_dependencies_are_rejected_even_on_the_last_page(authored_case, index):
    case = authored_case
    with ZipFile(case["authored"]) as archive:
        _, entries = stage.sheets(archive)
    part = entries[index][1]
    rel_part = part.replace("xl/worksheets/", "xl/worksheets/_rels/") + ".rels"
    relations = ET.Element(f"{{{stage.P}}}Relationships")
    ET.SubElement(relations, f"{{{stage.P}}}Relationship", {
        "Id": "syntheticExternal", "Type": stage.R + "/hyperlink",
        "Target": "https://example.test/synthetic", "TargetMode": "External",
    })
    _rewrite_fixture_zip(case["authored"], {rel_part: stage.serialized(relations)})
    with pytest.raises(ValueError, match="worksheet dependencies"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


def test_authored_bytes_changing_between_replacements_cannot_get_a_preservation_proof(
    authored_case, monkeypatch,
):
    case = authored_case
    actual_replace = stage.replace_sheet

    def change_after_first(*args, **kwargs):
        result = actual_replace(*args, **kwargs)
        case["authored"].write_bytes(case["authored"].read_bytes() + b"SYNTHETIC_DRIFT")
        return result

    monkeypatch.setattr(stage, "replace_sheet", change_after_first)
    with pytest.raises(ValueError, match="changed during integration"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.mark.parametrize("local", [False, True], ids=["workbook-name", "sheet-local-name"])
def test_authored_named_dependencies_are_rejected(authored_case, local):
    case = authored_case
    with ZipFile(case["authored"]) as archive:
        workbook = stage.xml(archive.read("xl/workbook.xml"))
    names = workbook.find(f"{{{stage.M}}}definedNames")
    if names is None:
        names = ET.SubElement(workbook, f"{{{stage.M}}}definedNames")
    attributes = {"name": "SyntheticDependency"}
    if local:
        attributes["localSheetId"] = "0"
    ET.SubElement(names, f"{{{stage.M}}}definedName", attributes).text = (
        f"'{WORKBOOK_SHEETS[0]}'!$B$2"
    )
    _rewrite_fixture_zip(case["authored"], {"xl/workbook.xml": stage.serialized(workbook)})
    with pytest.raises(ValueError, match="named dependencies"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.mark.parametrize("formula", [
    "B2*2",
    "SyntheticDependency",
    'HYPERLINK("#\'synthetic_history_00\'!A1","retained page")',
    'HYPERLINK("https://example.test/synthetic","external target")',
], ids=["calculation", "defined-name", "retained-page", "external-formula"])
def test_authored_formulas_cannot_introduce_unimported_dependencies(authored_case, formula):
    case = authored_case
    with ZipFile(case["authored"]) as archive:
        _, entries = stage.sheets(archive)
        part = entries[-1][1]
        sheet = stage.xml(archive.read(part))
    cell = next(cell for cell in sheet.iter(f"{{{stage.M}}}c") if cell.get("r") == "C2")
    cell.find(f"{{{stage.M}}}f").text = formula
    _rewrite_fixture_zip(case["authored"], {part: stage.serialized(sheet)})
    with pytest.raises(ValueError, match="only internal navigation formulas"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.mark.parametrize("mutation", ["missing", "extra"])
def test_authored_company_symbols_must_equal_the_restored_model(authored_case, mutation):
    case = authored_case
    with ZipFile(case["authored"]) as archive:
        _, entries = stage.sheets(archive)
        part = next(part for sheet, part in entries if sheet.get("name") == "03_\u516c\u53f8")
        sheet = stage.xml(archive.read(part))
    cell = next(cell for cell in sheet.iter(f"{{{stage.M}}}c") if cell.get("r") == "A5")
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    assert "999999" not in {company.symbol for company in case["model"].companies}
    value = "" if mutation == "missing" else case["model"].companies[0].symbol + " 999999"
    ET.SubElement(ET.SubElement(cell, f"{{{stage.M}}}is"), f"{{{stage.M}}}t").text = value
    _rewrite_fixture_zip(case["authored"], {part: stage.serialized(sheet)})
    with pytest.raises(ValueError, match="company pages do not match"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.fixture
def prepared_authored_case(authored_case):
    return _prepare_authored_case(authored_case)


def _prepare_authored_case(case):
    envelope = case["root"] / "runtime" / "synthetic-read-model.json"
    envelope.write_text(json.dumps({
        "schema_version": "historical-company-read-model-preview-v1",
        "snapshot": asdict(case["model"]), "canonical_written": False,
        "historical_preview": True, "action": "no_order",
    }, default=lambda value: value.isoformat()), encoding="utf-8")
    handoff = envelope.with_name("synthetic-publication-input.json")
    prepare_research_publication_input(root=case["root"], read_model_path=envelope,
                                      expected_sha256=_hash(envelope), output_path=handoff)
    _author_receipt(case, handoff)
    receipt = handoff_cli.prepare(
        root=case["root"], canonical=case["canonical"], handoff=handoff,
        digest=_hash(handoff), folder=case["output"].parent,
        authored_pages=case["authored"], authored_pages_sha256=_hash(case["authored"]),
    )
    assert receipt["canonical_written"] is False
    proof = case["output"].parent / "canonical-preservation.json"
    sidecar = case["output"].with_name(case["output"].stem + ".source-bindings.json")
    declaration = json.loads(sidecar.read_text(encoding="utf-8"))
    source_bindings = json.loads(handoff.read_text(encoding="utf-8"))["source_bindings"]
    assert declaration["source_bindings"][:len(source_bindings)] == source_bindings
    assert len(declaration["source_bindings"]) == len(source_bindings) + 4
    assert {"path": case["authored"].relative_to(case["root"]).as_posix(),
            "sha256": _hash(case["authored"])} in declaration["source_bindings"]
    assert declaration["authored_pages_binding"] == {
        "path": case["authored"].relative_to(case["root"]).as_posix(),
        "sha256": _hash(case["authored"]),
    }
    return dict(case, handoff=handoff, proof=proof, sidecar=sidecar)


def test_prepared_authored_candidate_source_bindings_pass_the_formal_verifier(prepared_authored_case):
    case = prepared_authored_case
    binding_result = publisher._verify_reviewed_research_source_bindings(
        case["root"], case["output"], case["proof"], case["sidecar"],
    )
    assert binding_result["verified_source_count"] >= 3
    assert _hash(case["canonical"]) == case["source_hash"]


def test_exact_audit_source_url_formulas_pass_protected_preview_and_provenance(authored_case):
    case = authored_case
    before = publisher._snapshot(case["canonical"])
    values = {}
    for row, evidence in enumerate(case["model"].audit_evidence, start=6):
        values.update({
            f"A{row}": evidence.evidence_id,
            f"E{row}": f'=HYPERLINK("{evidence.source_url}","{evidence.source_url}")',
            f"F{row}": evidence.sha256,
        })
    _set_authored_fixture_cells(case, WORKBOOK_SHEETS[-1], values)
    prepared = _prepare_authored_case(case)
    verification = publisher._verify_reviewed_research_source_bindings(
        case["root"], case["output"], prepared["proof"], prepared["sidecar"],
    )
    assert verification["verified_source_count"] >= len(case["model"].audit_evidence) + 4
    publisher._assert_retained(before, publisher._snapshot(case["output"]))
    proof = json.loads(prepared["proof"].read_text(encoding="utf-8"))
    assert proof["preserved_sheet_count"] == 55
    handoff = json.loads(prepared["handoff"].read_text(encoding="utf-8"))
    sources = {entry["path"]: entry["sha256"] for entry in handoff["source_bindings"]}
    snapshot_evidence = {entry["evidence_id"]: entry for entry in handoff["snapshot"]["audit_evidence"]}
    workbook = load_workbook(case["output"], data_only=False)
    try:
        for row, evidence in enumerate(case["model"].audit_evidence, start=6):
            assert sources[evidence.path] == evidence.sha256
            declared = snapshot_evidence[evidence.evidence_id]
            assert declared["source_url"] == evidence.source_url
            assert declared["sha256"] == evidence.sha256
            assert workbook[WORKBOOK_SHEETS[-1]][f"A{row}"].value == evidence.evidence_id
            assert workbook[WORKBOOK_SHEETS[-1]][f"E{row}"].value == values[f"E{row}"]
    finally:
        workbook.close()
    with ZipFile(case["canonical"]) as original, ZipFile(case["output"]) as integrated:
        _, entries = stage.sheets(original)
        for sheet, part in entries:
            if sheet.get("name") not in WORKBOOK_SHEETS:
                assert original.read(part) == integrated.read(part), part
    assert _hash(case["canonical"]) == case["source_hash"]


@pytest.mark.parametrize("mutation", [
    "wrong-page", "wrong-column", "header-row", "unknown-evidence", "wrong-evidence",
    "unbound-url", "wrong-display", "http", "credentials", "arbitrary-formula",
])
def test_audit_external_formula_exception_is_exact_and_cannot_escape(authored_case, mutation):
    case = authored_case
    first, second = case["model"].audit_evidence
    name, coordinate, row, evidence_id = WORKBOOK_SHEETS[-1], "E6", 6, first.evidence_id
    url, display = first.source_url, first.source_url
    if mutation == "wrong-page":
        name = WORKBOOK_SHEETS[3]
    elif mutation == "wrong-column":
        coordinate = "F6"
    elif mutation == "header-row":
        coordinate, row = "E5", 5
    elif mutation == "unknown-evidence":
        evidence_id = "synthetic-not-in-model"
    elif mutation == "wrong-evidence":
        evidence_id = second.evidence_id
    elif mutation == "unbound-url":
        url = display = "https://example.test/synthetic/not-an-approved-source.pdf"
    elif mutation == "wrong-display":
        display = second.source_url
    elif mutation == "http":
        url = display = first.source_url.replace("https:", "http:")
    elif mutation == "credentials":
        url = display = first.source_url.replace("https://", "https://synthetic-user:synthetic-only@")
        # Even an exact model URL cannot admit embedded credentials.
        case["model"] = replace(case["model"], audit_evidence=(replace(first, source_url=url), second))
    formula = "=1+1" if mutation == "arbitrary-formula" else f'=HYPERLINK("{url}","{display}")'
    _set_authored_fixture_cells(case, name, {f"A{row}": evidence_id, coordinate: formula})
    with pytest.raises(ValueError, match="only internal navigation formulas"):
        _integrate(case)
    assert _hash(case["canonical"]) == case["source_hash"]
    assert not case["output"].exists()
    assert not (case["output"].parent / "canonical-preservation.json").exists()


@pytest.mark.parametrize("mutation, message", [
    ("receipt-wrong-input", "authored product receipt does not bind"),
    ("display-drift", "authored display snapshot changed"),
    ("policy-drift", "authored display policy does not bind"),
    ("extra-source", "research handoff source bindings differ"),
    ("missing-original-source", "research handoff source bindings differ"),
    ("proof-wrong-authored-hash", "authored workbook is not bound to preservation proof"),
])
def test_formal_verifier_rejects_authored_provenance_drift(
    prepared_authored_case, mutation, message,
):
    case = prepared_authored_case
    verifier_args = (case["root"], case["output"], case["proof"], case["sidecar"])
    # Start with an admitted synthetic candidate; then isolate each source-contract violation.
    publisher._verify_reviewed_research_source_bindings(*verifier_args)
    candidate_hash = _hash(case["output"])
    declaration = json.loads(case["sidecar"].read_text(encoding="utf-8"))
    author_receipt = case["authored"].with_name(case["authored"].stem + "-receipt.json")
    receipt = json.loads(author_receipt.read_text(encoding="utf-8"))
    display_path = case["root"] / receipt["display_binding"]["path"]
    display = json.loads(display_path.read_text(encoding="utf-8"))
    policy_path = case["root"] / display["policy_binding"]["path"]

    if mutation == "receipt-wrong-input":
        receipt["input_sha256"] = "0" * 64
        author_receipt.write_text(json.dumps(receipt), encoding="utf-8")
        # Re-pin the outer hash so only the incorrect handoff binding causes rejection.
        receipt_relative = author_receipt.relative_to(case["root"]).as_posix()
        for binding in declaration["source_bindings"]:
            if binding["path"] == receipt_relative:
                binding["sha256"] = _hash(author_receipt)
    elif mutation == "display-drift":
        display["scope"] = SYNTHETIC + " DRIFT"
        display_path.write_text(json.dumps(display), encoding="utf-8")
    elif mutation == "policy-drift":
        policy = json.loads(policy_path.read_text(encoding="utf-8"))
        policy["scope"] = SYNTHETIC + " DRIFT"
        policy_path.write_text(json.dumps(policy), encoding="utf-8")
    elif mutation == "extra-source":
        extra = case["root"] / "runtime" / "synthetic-unapproved-source.json"
        extra.write_text(json.dumps({"scope": SYNTHETIC}), encoding="utf-8")
        declaration["source_bindings"].append({
            "path": extra.relative_to(case["root"]).as_posix(), "sha256": _hash(extra),
        })
    elif mutation == "missing-original-source":
        handoff = json.loads(case["handoff"].read_text(encoding="utf-8"))
        assert handoff["source_bindings"]
        assert declaration["source_bindings"][0] == handoff["source_bindings"][0]
        del declaration["source_bindings"][0]
    else:
        proof = json.loads(case["proof"].read_text(encoding="utf-8"))
        proof["authored_pages_sha256"] = "0" * 64
        case["proof"].write_text(json.dumps(proof), encoding="utf-8")
        # Preserve the outer manifest binding; the inner authored hash must still be checked.
        declaration["output_manifest_sha256"] = _hash(case["proof"])
    case["sidecar"].write_text(json.dumps(declaration), encoding="utf-8")

    with pytest.raises(ValueError, match=message):
        publisher._verify_reviewed_research_source_bindings(*verifier_args)
    assert _hash(case["output"]) == candidate_hash
    assert _hash(case["canonical"]) == case["source_hash"]
