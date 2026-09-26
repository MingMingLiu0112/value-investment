from __future__ import annotations

from collections import deque
import hashlib
import json
from pathlib import Path
import re
import tempfile

from openpyxl import load_workbook

from value_investment_agent.presentation.excel.product_workbench import (
    SHEET_COMPANIES,
    SHEET_EVENTS,
    SHEET_OPPORTUNITIES,
    SHEET_PORTFOLIO,
    SHEET_SYSTEM_AUDIT,
    SHEET_TODAY,
    USER_SHEETS,
    WORKBOOK_SHEETS,
    build_product_workbench_workbook,
    write_product_workbench_candidate,
)
from value_investment_agent.presentation.read_models.product_workbench import (
    ACTION_NO_ORDER,
    PRODUCT_WORKBENCH_SCHEMA_VERSION,
    product_workbench_from_payload,
)


ROOT = Path(__file__).resolve().parents[1]
SHA = "b" * 64


def _assessment(
    *,
    status: str,
    available: bool,
    value_text: str | None = None,
    reason: str | None = None,
    needed: str | None = None,
) -> dict:
    return {
        "status": status,
        "available": available,
        "value_text": value_text,
        "unavailable_reason": reason,
        "needed_evidence": needed,
        "evidence_refs": ["evidence-1"],
    }


def _payload() -> dict:
    sections = [
        {
            "key": key,
            "status": "PARTIAL",
            "summary": f"{key} 的上游结论。",
            "evidence_refs": ["evidence-1"],
        }
        for key in (
            "business_quality",
            "financial_quality",
            "capital_allocation",
            "valuation",
            "dividend",
            "risks_counterevidence",
        )
    ]
    scenarios = [
        {
            "key": key,
            "assessment": _assessment(
                status="NOT_READY",
                available=False,
                reason=f"{key} 情景输入未复核。",
                needed="同口径情景假设。",
            ),
        }
        for key in ("bear", "base", "bull")
    ]
    return {
        "schema_version": PRODUCT_WORKBENCH_SCHEMA_VERSION,
        "generated_at": "2026-09-26T03:00:00+00:00",
        "as_of": "2026-09-25",
        "action": ACTION_NO_ORDER,
        "overview": {"pending_count": 1},
        "system_health": {
            "status": "EVIDENCE_INSUFFICIENT",
            "message": "部分输入仍待人工复核。",
        },
        "stages": {
            "m2": {"status": "DONE", "detail": "初筛已人工通过。"},
            "m3": {"status": "PARTIAL", "detail": "研究资料仍在补齐。"},
            "m4": {
                "status": "PENDING_USER_PRIVATE_INPUT",
                "detail": "尚未收到真实组合输入。",
            },
            "m5": {"status": "PENDING_HUMAN_REVIEW", "detail": "有事件待复核。"},
            "m6": {
                "status": "PREFLIGHT_DONE",
                "detail": "真实监控尚未开始。",
            },
        },
        "audit": {
            "evidence": [
                {
                    "evidence_id": "evidence-1",
                    "title": "固定测试证据",
                    "artifact_type": "research_payload",
                    "path": "runtime/product-fixture.json",
                    "sha256": SHA,
                    "available_at": "2026-09-25",
                    "action": ACTION_NO_ORDER,
                }
            ]
        },
        "today_items": [
            {
                "category": "PRICE_WATCH",
                "company": "测试公司",
                "symbol": "600000",
                "what_happened": "价格进入关注范围。",
                "why_it_matters": "需要复核上游研究结论。",
                "current_status": "等待人工复核。",
                "next_step": "核对价格与研究时点。",
                "evidence_refs": ["evidence-1"],
            }
        ],
        "opportunities": [
            {
                "symbol": "600000",
                "company_name": "测试公司",
                "why_now": "价格状态已变化。",
                "research_status": "WAIT_FOR_PRICE",
                "valuation_status": "NOT_READY",
                "dividend_status": "INSUFFICIENT_EVIDENCE",
                "price_status": "IN_WATCH_RANGE",
                "main_risk": "价格变化不等于投资结论。",
                "next_trigger": "等待人工复核。",
                "evidence_refs": ["evidence-1"],
            }
        ],
        "companies": [
            {
                "symbol": "600000",
                "company_name": "测试公司",
                "research_status": "PARTIAL",
                "price": _assessment(
                    status="AVAILABLE",
                    available=True,
                    value_text="上游价格 10 元",
                ),
                "valuation": _assessment(
                    status="NOT_READY",
                    available=False,
                    reason="估值输入尚未复核。",
                    needed="同一时点估值输入。",
                ),
                "margin_of_safety": _assessment(
                    status="NOT_READY",
                    available=False,
                    reason="估值输入尚未复核。",
                    needed="同一时点估值输入。",
                ),
                "dividend": _assessment(
                    status="INSUFFICIENT_EVIDENCE",
                    available=False,
                    reason="股息证据不足。",
                    needed="多财年分配与现金流证据。",
                ),
                "sections": sections,
                "scenarios": scenarios,
                "latest_change": "价格进入关注范围。",
                "next_trigger": "等待人工复核。",
                "original_thesis": "高质量经营与现金回报。",
                "thesis_change": "UNKNOWN",
                "evidence_refs": ["evidence-1"],
            }
        ],
        "portfolio": {
            "real_data_available": False,
            "status": "PENDING_USER_PRIVATE_INPUT",
            "connection_hint": "通过受保护的私密入口接入真实组合。",
            "summary": [],
            "positions": [],
            "action": ACTION_NO_ORDER,
        },
        "events": [
            {
                "event_id": "event-1",
                "event_type": "MATERIAL_REQUIRES_RECALCULATION",
                "company_name": "测试公司",
                "what_happened": "发生需要复核的重大事项。",
                "impact_area": "投资逻辑。",
                "current_conclusion": "等待人工判断。",
                "research_action": "REOPEN_RESEARCH",
                "next_step": "补充事件后证据。",
                "evidence_refs": ["evidence-1"],
            }
        ],
    }


def _model():
    return product_workbench_from_payload(_payload())


def _all_text(workbook) -> str:
    values: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))
    return "\n".join(values)


def _user_text(workbook) -> str:
    values: list[str] = []
    for title in USER_SHEETS:
        sheet = workbook[title]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))
    return "\n".join(values)


def _link_sheet(cell) -> str | None:
    hyperlink = cell.hyperlink
    if hyperlink is None:
        return None
    raw = hyperlink.location or hyperlink.target or ""
    match = re.match(r"^#'?([^'!]+)'?!", raw)
    return match.group(1) if match else None


def _click_distance(workbook, start: str, target: str) -> int | None:
    queue = deque([(start, 0)])
    seen = {start}
    while queue:
        title, distance = queue.popleft()
        if title == target:
            return distance
        sheet = workbook[title]
        for row in sheet.iter_rows():
            for cell in row:
                linked = _link_sheet(cell)
                if linked and linked not in seen and linked in workbook.sheetnames:
                    seen.add(linked)
                    queue.append((linked, distance + 1))
    return None


def _empty_payload() -> dict:
    payload = _payload()
    payload["overview"]["pending_count"] = 0
    payload["today_items"] = []
    payload["opportunities"] = []
    payload["companies"] = []
    payload["events"] = []
    payload["audit"]["evidence"] = []
    payload["stages"] = {
        key: {"status": "WAIT", "detail": f"{key} 等待输入。"}
        for key in ("m2", "m3", "m4", "m5", "m6")
    }
    payload["portfolio"] = {
        "real_data_available": False,
        "status": "NOT_STARTED",
        "connection_hint": "尚未接入真实组合。",
        "summary": [],
        "positions": [],
        "action": ACTION_NO_ORDER,
    }
    return payload


def test_candidate_has_five_user_sheets_and_secondary_audit() -> None:
    workbook = build_product_workbench_workbook(_model())

    assert tuple(workbook.sheetnames) == WORKBOOK_SHEETS
    assert tuple(workbook.sheetnames[:5]) == USER_SHEETS
    assert workbook.sheetnames[-1] == SHEET_SYSTEM_AUDIT
    assert workbook[USER_SHEETS[0]].sheet_view.showGridLines is False


def test_company_detail_is_reachable_within_three_clicks() -> None:
    workbook = build_product_workbench_workbook(_model())

    distance = _click_distance(workbook, SHEET_TODAY, SHEET_COMPANIES)

    assert distance is not None
    assert distance <= 3
    assert _click_distance(workbook, SHEET_OPPORTUNITIES, SHEET_COMPANIES) == 1


def test_user_pages_hide_internal_codes_hashes_and_execution_semantics() -> None:
    workbook = build_product_workbench_workbook(_model())
    text = _user_text(workbook)

    for hidden in (
        "PENDING_HUMAN_REVIEW",
        "PENDING_USER_PRIVATE_INPUT",
        "PREFLIGHT_DONE",
        "SHA-256",
        SHA,
        "target_weight",
        "position_size",
        "trade_approved",
        "建议买入",
        "目标仓位",
        "下单",
    ):
        assert hidden not in text
    assert workbook[SHEET_OPPORTUNITIES]["C5"].value == "研究完成，等待价格"


def test_absent_portfolio_renders_only_unconnected_state() -> None:
    workbook = build_product_workbench_workbook(_model())
    sheet = workbook[SHEET_PORTFOLIO]
    text = "\n".join(
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )

    assert "尚未接入真实组合" in text
    assert "查看接入说明" in text
    for simulated_metric in (
        "总资产",
        "股票仓位",
        "现金",
        "预计年股息",
        "正常化年股息",
    ):
        assert simulated_metric not in text


def test_evidence_link_reaches_audit_and_audit_contains_hash() -> None:
    workbook = build_product_workbench_workbook(_model())
    evidence_links = [
        cell
        for row in workbook[SHEET_OPPORTUNITIES].iter_rows()
        for cell in row
        if cell.value == "查看证据 (1)"
    ]

    assert evidence_links
    assert _link_sheet(evidence_links[0]) == SHEET_SYSTEM_AUDIT
    audit_text = "\n".join(
        str(cell.value)
        for row in workbook[SHEET_SYSTEM_AUDIT].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert SHA in audit_text
    assert "runtime/product-fixture.json" in audit_text
    assert "no_order" in audit_text


def test_unavailable_assessment_has_reason_and_no_empty_numeric_placeholder() -> None:
    workbook = build_product_workbench_workbook(_model())
    company_text = "\n".join(
        str(cell.value)
        for row in workbook[SHEET_COMPANIES].iter_rows()
        for cell in row
        if cell.value is not None
    )

    assert "暂不可评估" in company_text
    assert "原因：估值输入尚未复核。" in company_text
    assert "需要：同一时点估值输入。" in company_text
    assert "0.00" not in company_text


def test_all_wait_empty_candidate_renders_legal_empty_states() -> None:
    model = product_workbench_from_payload(_empty_payload())
    workbook = build_product_workbench_workbook(model)

    assert tuple(workbook.sheetnames) == WORKBOOK_SHEETS
    today_text = "\n".join(
        str(cell.value)
        for row in workbook[SHEET_TODAY].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "当前没有已分类的今日事项" in today_text
    assert "尚未接入个人组合" in today_text


def test_manifest_is_no_order_and_does_not_touch_canonical_pointer() -> None:
    pointer = ROOT / "config" / "current-trial-workbook.json"
    pointer_before = hashlib.sha256(pointer.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as temporary:
        temporary_root = Path(temporary)
        output = temporary_root / "m7-product-ux-candidate.xlsx"

        receipt = write_product_workbench_candidate(
            _model(),
            output=output,
            root=temporary_root,
        )

        assert output.exists()
        assert receipt["action"] == ACTION_NO_ORDER
        assert receipt["workbook_kind"] == "M7_PRODUCT_UX_CANDIDATE"
        assert receipt["final_user_acceptance"] == "NOT_PASSED"
        assert receipt["candidate_not_formal_workbook"] is True
        assert receipt["canonical_pointer_modified"] is False
        manifest_path = output.with_name(
            output.stem + ".m7-product-workbench-candidate-manifest.json"
        )
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assert manifest["action"] == ACTION_NO_ORDER
    assert hashlib.sha256(pointer.read_bytes()).hexdigest() == pointer_before
